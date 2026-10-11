"""Diapositivas con errores puestos a propósito, para las revisiones que no hace el verificador solo
(PLAN.md §5.8). La cantidad de diapositivas y las notas las siembra tests/unidad/test_verificador.py.

- desborde: un texto que no cabe en su recuadro. Lo mide PowerPoint (tests/unidad/test_office.py).
- fuera de orden, bloque inventado y bloque sin diapositivas: los revisa el programa antes del PowerPoint
  (tests/unidad/test_diapositivas.py).
- contenido que no está en la lectura: lo busca la revisión del contenido (tests/ia/test_diapositivas_sembradas.py).
"""

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Cm, Pt

TEXTO_QUE_NO_CABE = "El equipo revisa el tablero al final de cada Sprint y decide qué cambiar. " * 6
TEXTO_QUE_CABE = "Los pilares de Scrum son tres."


def con_desborde(destino: Path) -> Path:
    """Dos diapositivas: la primera con un texto que cabe y la segunda con uno que se sale de su recuadro."""
    presentacion = Presentation()
    presentacion.slide_width, presentacion.slide_height = Cm(33.867), Cm(19.05)
    for texto in (TEXTO_QUE_CABE, TEXTO_QUE_NO_CABE):
        diapositiva = presentacion.slides.add_slide(presentacion.slide_layouts[6])
        caja = diapositiva.shapes.add_textbox(Cm(2), Cm(3), Cm(14), Cm(4))
        caja.text_frame.word_wrap = True
        caja.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        caja.text_frame.paragraphs[0].text = texto
        caja.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    presentacion.save(destino)
    return destino
