"""Crea por programa una fuente de cada formato, con texto conocido."""

from datetime import date
from pathlib import Path

import openpyxl
import pymupdf
from docx import Document
from docx.enum.text import WD_BREAK
from pptx import Presentation
from pptx.util import Inches


def pdf_con_texto(ruta: Path) -> Path:
    """Dos páginas. La segunda tiene un número impreso, como las láminas de un curso."""
    documento = pymupdf.open()
    pagina = documento.new_page()
    pagina.insert_text((72, 72), "Scrum tiene tres pilares.", fontsize=12)
    pagina = documento.new_page()
    pagina.insert_text((72, 60), "43", fontsize=10)
    pagina.insert_text((72, 120), "Scrum es gratuito.", fontsize=12)
    documento.save(ruta)
    return ruta


def _imagen_de_texto(texto: str) -> bytes:
    documento = pymupdf.open()
    pagina = documento.new_page(width=400, height=120)
    pagina.insert_text((20, 60), texto, fontsize=16)
    return pagina.get_pixmap(dpi=100).tobytes("png")


def pdf_escaneado(ruta: Path) -> Path:
    """Una página con texto y una página que es solo una imagen."""
    documento = pymupdf.open()
    documento.new_page().insert_text((72, 72), "Página con texto.", fontsize=12)
    pagina = documento.new_page()
    pagina.insert_image(pymupdf.Rect(72, 72, 472, 192), stream=_imagen_de_texto("Texto escaneado"))
    documento.save(ruta)
    return ruta


def word_dos_paginas(ruta: Path) -> Path:
    documento = Document()
    documento.add_paragraph("Primera página del sílabo.")
    documento.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    documento.add_paragraph("Segunda página del sílabo.")
    tabla = documento.add_table(rows=1, cols=2)
    tabla.cell(0, 0).text = "Sesión 1"
    tabla.cell(0, 1).text = "Fundamentos"
    documento.save(ruta)
    return ruta


def powerpoint(ruta: Path) -> Path:
    """Lámina 1 con título y notas. Lámina 2 con una tabla. Lámina 3 solo con una imagen."""
    presentacion = Presentation()
    lamina = presentacion.slides.add_slide(presentacion.slide_layouts[1])
    lamina.shapes.title.text = "Valores de Scrum"
    lamina.placeholders[1].text = "Compromiso"
    lamina.notes_slide.notes_text_frame.text = "Pregunta al grupo por un ejemplo."
    lamina = presentacion.slides.add_slide(presentacion.slide_layouts[6])
    tabla = lamina.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    tabla.cell(0, 0).text = "Pilar"
    tabla.cell(0, 1).text = "Definición"
    tabla.cell(1, 0).text = "Inspección"
    tabla.cell(1, 1).text = "Revisar el avance."
    ruta_imagen = ruta.with_suffix(".png")
    ruta_imagen.write_bytes(_imagen_de_texto("Diagrama del Sprint"))
    lamina = presentacion.slides.add_slide(presentacion.slide_layouts[6])
    lamina.shapes.add_picture(str(ruta_imagen), Inches(1), Inches(1))
    ruta_imagen.unlink()
    presentacion.save(ruta)
    return ruta


def excel(ruta: Path) -> Path:
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Registro"
    hoja["A1"] = "Fecha de referencia"
    hoja["B1"] = date(2026, 10, 5)
    hoja["A2"] = "Pedidos listos"
    hoja["B2"] = 12
    otra = libro.create_sheet("Datos")
    otra["C3"] = "Almacén central"
    libro.save(ruta)
    return ruta


def excel_con_formula_sin_valor(ruta: Path) -> Path:
    """openpyxl guarda la fórmula sin su resultado, como un archivo nunca abierto en Excel."""
    libro = openpyxl.Workbook()
    libro.active["A1"] = 2
    libro.active["A2"] = "=A1*3"
    libro.save(ruta)
    return ruta


def imagen(ruta: Path, texto: str = "Fecha de referencia: 5 de octubre de 2026") -> Path:
    ruta.write_bytes(_imagen_de_texto(texto))
    return ruta
