"""Decisión 28: horas y número de sesiones, botones «Completar», Guardar que confirma y carátula del modelo."""

import asyncio

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.fichas import almacen, flujo, plantillas
from app.materiales import catalogo, estilo, modelos
from app.materiales import docx as generador
from app.materiales.contenido import Lectura
from tests.conftest import consulta_simulada
from tests.fixtures import fichas as datos
from tests.unidad.test_lectura import LIMPIA, sesion  # noqa: F401  (fixture)


def correr(coro):
    return asyncio.run(coro)


@pytest.fixture
def curso(entorno):
    entorno["env"].write_text("ANTHROPIC_API_KEY=sk-ant-de-prueba\n", encoding="utf-8")
    carpeta = entorno["cursos"] / "excel"
    (carpeta / "fuentes").mkdir(parents=True)
    return carpeta


def _ficha(curso, valores):
    ficha = almacen.cargar(curso, "curso")
    almacen.aplicar_formulario(ficha, valores)
    almacen.guardar(curso, "curso", None, ficha)


def _respuesta(**campos):
    return consulta_simulada(structured_output={"campos": [{"id": c, "valor": v} for c, v in campos.items()]})


# ---------- Horas y número de sesiones ----------

def test_la_ficha_del_curso_tiene_horas_y_numero_de_sesiones():
    texto = plantillas.renderizar("curso", {"curso.horas": {"valor": "24"}, "curso.numero_sesiones": {"valor": "6"}})
    assert "- Horas del curso: 24\n- Número de sesiones: 6\n- Sesiones (obligatorio)." in texto


def test_las_horas_no_llegan_al_redactor(sesion):  # noqa: F811
    from app.materiales import lectura
    _ficha(sesion, {"curso.horas": "24", "curso.numero_sesiones": "2"})
    assert "Horas del curso" in almacen.ruta_ficha(sesion, "curso", extension=".md").read_text(encoding="utf-8")
    pedido = lectura.pedido_de_redaccion(sesion, 1)
    assert "Horas del curso" not in pedido and "24" not in pedido
    assert "Número de sesiones: 2" in pedido


def test_la_propuesta_no_inventa_horas_ni_numero_de_sesiones(curso):
    ficha = almacen.cargar(curso, "curso")
    propuesta = {"campos": [{"id": "curso.horas", "valor": "30", "origen": "deducido", "literal": False},
                            {"id": "curso.numero_sesiones", "valor": "5", "origen": "deducido", "literal": False}],
                 "conflictos": []}
    resumen = flujo.aplicar_propuesta(curso, ficha, propuesta)
    assert resumen["descartados"] == 2
    assert almacen.valor(ficha, "curso.horas") == "" and almacen.valor(ficha, "curso.numero_sesiones") == ""


# ---------- Botones «Completar» ----------

def _sesiones(*sesiones):
    return consulta_simulada(structured_output={"sesiones": [
        {"numero": n, "titulo": t, "alcance": a, "paginas": p} for n, t, a, p in sesiones]})


def test_repartir_las_sesiones_usa_horas_numero_y_publico_y_reemplaza_lo_que_habia(curso):
    _ficha(curso, {"curso.nombre": "Excel para análisis", "curso.horas": "12", "curso.numero_sesiones": "3",
                   "publico.quienes": "Asistentes administrativos", "curso.sesiones": "1. Lo que había"})
    consulta = _sesiones((2, "Fórmulas", "SUMAR.SI.", []), (1, "Tablas", "Orden y filtros.", []))
    mensaje = correr(flujo.completar(curso, "excel", "sesiones", consulta=consulta))
    ficha = almacen.cargar(curso, "curso")
    assert ficha["campos"]["curso.sesiones"] == {"valor": "1. Tablas: Orden y filtros.\n2. Fórmulas: SUMAR.SI.",
                                                 "origen": "deducido", "estado": "propuesto"}
    assert "Ojo: pediste 3 sesiones y la propuesta trae 2." in mensaje
    assert "Sin fuentes convertidas no se calcula la cobertura." in mensaje
    assert "cobertura_sesiones" not in ficha
    pedido = consulta.llamadas[0]["prompt"]
    assert "- Horas del curso: 12" in pedido and "Asistentes administrativos" in pedido
    assert "- Horas de cada sesión: 4" in pedido and "exactamente 3 sesiones" in pedido
    assert "Lo que había" not in pedido
    assert not consulta.llamadas[0]["options"].allowed_tools   # sin fuentes, sin herramientas


def test_repartir_las_sesiones_calcula_la_cobertura_con_las_paginas_de_las_fuentes(curso, monkeypatch):
    _ficha(curso, {"curso.nombre": "Agilidad", "curso.horas": "6", "curso.numero_sesiones": "2"})
    monkeypatch.setattr(flujo.verificacion, "fuentes_utilizables", lambda c: ["Guia.pdf", "Manifiesto.pdf"])
    monkeypatch.setattr(flujo, "paginas_de_las_fuentes",
                        lambda c, f: {"Guia.pdf": list(range(1, 11)), "Manifiesto.pdf": [1, 2, 3, 4, 5, 7, 8, 9, 10, 11]})
    consulta = _sesiones(
        (1, "Dominios", "Los dominios.", [{"fuente": "Guia.pdf.jsonl", "desde": 1, "hasta": 8},
                                          {"fuente": "Inventada.pdf", "desde": 1, "hasta": 99}]),
        (2, "Valores", "Los cuatro valores.", [{"fuente": "Manifiesto.pdf", "desde": 9, "hasta": 3}]))
    mensaje = correr(flujo.completar(curso, "agilidad", "sesiones", consulta=consulta))
    ficha = almacen.cargar(curso, "curso")
    c = flujo.cobertura_vigente(ficha)
    # Guía: 8 de 10. Manifiesto: 3 a 9 son 3, 4, 5, 7, 8 y 9. Una fuente que no es del curso no cuenta.
    assert (c["cubiertas"], c["total"], c["porcentaje"], c["vigente"]) == (14, 20, 70, True)
    assert c["fuentes"][0] == {"fuente": "Guia.pdf", "cubiertas": 8, "total": 10, "fuera": "9 y 10"}
    assert c["fuentes"][1]["fuera"] == "1, 2, 10 y 11"
    assert c["sesiones"] == {"1": ["Guia.pdf: páginas 1 a 8"], "2": ["Manifiesto.pdf: páginas 3 a 5 y 7 a 9"]}
    assert "Cubren el 70 % de las páginas de las fuentes (14 de 20)." in mensaje
    pedido = consulta.llamadas[0]["prompt"]
    assert "Guia.pdf (páginas 1 a 10)" in pedido and "Si no cabe, elige los temas más importantes" in pedido
    # Si el profesor cambia las sesiones, la cobertura deja de valer.
    almacen.aplicar_formulario(ficha, {"curso.sesiones": "1. Otra cosa"})
    assert flujo.cobertura_vigente(ficha)["vigente"] is False


def test_los_tramos_de_paginas_se_leen_como_una_frase():
    assert flujo._tramos([1, 2, 3, 7, 9, 10]) == "1 a 3, 7, 9 y 10"
    assert flujo._tramos([4]) == "4" and flujo._tramos([]) == ""


def test_la_pagina_muestra_la_cobertura_y_avisa_si_ya_no_vale(entorno):
    from fastapi.testclient import TestClient
    from app import servidor
    cliente = TestClient(servidor.app)
    cliente.post("/cursos", data={"nombre": "Excel"})
    curso = entorno["cursos"] / "excel"
    ficha = almacen.cargar(curso, "curso")
    ficha["campos"]["curso.sesiones"] = {"valor": "1. Dominios: los dominios.", "origen": "deducido", "estado": "propuesto"}
    ficha["cobertura_sesiones"] = flujo.cobertura(
        [{"numero": 1, "paginas": [{"fuente": "Guia.pdf", "desde": 1, "hasta": 3}]}], {"Guia.pdf": [1, 2, 3, 4]},
        "1. Dominios: los dominios.")
    almacen.guardar(curso, "curso", None, ficha)
    pagina = cliente.get("/cursos/excel/ficha").text
    assert "Las sesiones cubren el 75 % del material" in pagina and "Quedan fuera las páginas 4." in pagina
    assert "Sesión 1: Guia.pdf: páginas 1 a 3." in pagina
    almacen.aplicar_formulario(ficha, {"curso.sesiones": "1. Otra"})
    almacen.guardar(curso, "curso", None, ficha)
    assert "la cobertura ya no vale" in cliente.get("/cursos/excel/ficha").text


def test_la_pregunta_de_adjuntos_solo_se_muestra_con_ia():
    from app.fichas import controles
    assert controles.vista("herramientas.adjuntos", "")["depende"] == ("herramientas.ejercicios", "IA|Excel e IA")
    assert not controles.se_muestra("herramientas.adjuntos", {"herramientas.ejercicios": "ninguna"})
    assert not controles.se_muestra("herramientas.adjuntos", {})
    assert controles.se_muestra("herramientas.adjuntos", {"herramientas.ejercicios": "Excel e IA"})
    assert controles.se_muestra("herramientas.neutral", {})


def test_completar_pide_primero_lo_que_necesita_y_no_gasta_tokens(curso):
    _ficha(curso, {"curso.nombre": "Excel para análisis"})
    consulta = _respuesta()
    assert correr(flujo.completar(curso, "excel", "sesiones", consulta=consulta)) == \
        "Para repartir las sesiones, llena primero «Horas del curso» y «Número de sesiones»."
    assert correr(flujo.completar(curso, "excel", "datos", consulta=consulta)) == \
        "Para completar los datos fijos, llena primero «Empresa o institución ficticia y rubro»."
    assert consulta.llamadas == []


def test_completar_el_caso_llena_solo_lo_vacio(curso):
    _ficha(curso, {"curso.nombre": "Excel para análisis", "publico.quienes": "Contadores",
                   "caso.empresa": "Distribuidora Andina, venta de abarrotes"})
    consulta = _respuesta(**{"caso.empresa": "Otra", "caso.areas": "Ventas, Compras, Almacén",
                             "caso.moneda": "Soles (PEN), Perú"})
    correr(flujo.completar(curso, "excel", "caso", consulta=consulta))
    ficha = almacen.cargar(curso, "curso")
    assert almacen.valor(ficha, "caso.empresa") == "Distribuidora Andina, venta de abarrotes"
    assert ficha["campos"]["caso.empresa"]["estado"] == ""
    assert ficha["campos"]["caso.areas"]["estado"] == "propuesto"
    pedido = consulta.llamadas[0]["prompt"]
    assert "Contadores" in pedido and "- caso.areas:" in pedido and "- caso.empresa:" not in pedido
    otra = _respuesta()
    assert "ya está llena" in correr(flujo.completar(curso, "excel", "caso", consulta=otra))
    assert otra.llamadas == []


def test_completar_los_datos_fijos_usa_el_caso_y_avisa_lo_que_no_cuadra(curso):
    _ficha(curso, {"curso.nombre": "Excel para análisis", "caso.empresa": "Distribuidora Andina",
                   "caso.moneda": "Soles (PEN), Perú"})
    consulta = _respuesta(**{"datos.fecha": "15 de marzo de 2026",
                             "datos.fijos": "Gerente de ventas: Rosa Quispe\nEl cierre dura dos horas"})
    mensaje = correr(flujo.completar(curso, "excel", "datos", consulta=consulta))
    assert "Distribuidora Andina" in consulta.llamadas[0]["prompt"]
    assert "Ojo: 1 líneas de los datos fijos no tienen la forma «Etiqueta: valor»." in mensaje
    assert "la propuesta menciona un tiempo" in mensaje
    ficha = almacen.cargar(curso, "curso")
    assert ficha["confirmada"] is False and ficha["campos"]["datos.fecha"]["estado"] == "propuesto"


# ---------- La carátula del modelo da la disposición de la portada ----------

def _png(ruta):
    pixeles = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 30), 0)
    pixeles.clear_with(90)
    pixeles.save(ruta)
    return ruta


def _modelo(carpeta, *, centrada, franja, logo_x=None):
    (carpeta / "modelos").mkdir(exist_ok=True)
    documento = pymupdf.open()
    pagina = documento.new_page()   # 595 × 842
    if franja:
        pagina.draw_rect(pymupdf.Rect(0, 0, 595, 60), color=None, fill=(0.12, 0.31, 0.62))
    if logo_x is not None:
        pagina.insert_image(pymupdf.Rect(logo_x, 120, logo_x + 80, 160), filename=str(_png(carpeta / "l.png")))
    titulo = "Curso de prueba"
    ancho = pymupdf.get_text_length(titulo, fontsize=28)
    pagina.insert_text(((595 - ancho) / 2 if centrada else 72, 260), titulo, fontsize=28)
    documento.new_page().insert_text((72, 80), "Segunda página", fontsize=11)
    documento.save(carpeta / "modelos" / "modelo.pdf")
    _ficha(carpeta, {"modelos.aprobados": "modelo.pdf"})


def test_la_caratula_centrada_con_franja_y_logo_a_la_derecha(sesion):  # noqa: F811
    _modelo(sesion, centrada=True, franja=True, logo_x=470)
    assert modelos.disposicion_de_portada(sesion) == ("modelo.pdf", {"centrada": True, "franja": True, "logo": "derecha"})


def test_la_caratula_alineada_a_la_izquierda_sin_franja_ni_logo(sesion):  # noqa: F811
    _modelo(sesion, centrada=False, franja=False)
    assert modelos.disposicion_de_portada(sesion)[1] == {"centrada": False, "franja": False, "logo": ""}


def test_con_modelo_la_portada_sigue_la_caratula_y_la_ficha_manda_en_colores(sesion, tmp_path):  # noqa: F811
    _modelo(sesion, centrada=True, franja=True, logo_x=470)
    _ficha(sesion, {"visual.principal": "#6D1A36", "visual.estilo": "minimalista"})
    identidad = estilo.desde_ficha(sesion)
    assert identidad.principal == "6D1A36"                       # la identidad visual de la ficha manda
    assert identidad.estilo.nombre == "minimalista" and identidad.estilo.borde_recuadro == 12
    assert (identidad.estilo.portada_centrada, identidad.estilo.franja, identidad.estilo.logo) == (True, True, "derecha")
    assert any("La portada sigue la carátula del modelo (modelo.pdf)" in a for a in identidad.avisos)


def test_el_word_pone_el_logo_donde_lo_pone_la_caratula(tmp_path):
    from dataclasses import replace
    diseno = replace(catalogo.estilo("profesional"), portada_centrada=True, logo="derecha")
    identidad = estilo.Identidad(**estilo.PALETA_SOBRIA, tipografia="Calibri", logo=_png(tmp_path / "logo.png"),
                                 estilo=diseno)
    portada = generador.Portada("Scrum Master con IA", "Sesión 1 · Fundamentos", "Lectura")
    documento = Document(generador.generar_lectura(Lectura.model_validate(LIMPIA), identidad, portada,
                                                   tmp_path / "S1_Lectura.docx"))
    con_imagen = next(p for p in documento.paragraphs if p._p.xpath(".//pic:pic"))
    assert con_imagen.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    titulo = next(p for p in documento.paragraphs if p.style.name == "Title")
    assert titulo.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_las_diapositivas_ponen_el_logo_donde_lo_pone_la_caratula(tmp_path):
    from dataclasses import replace

    from pptx import Presentation

    from app.materiales import pptx as diapositivas
    from app.materiales.contenido import Diapositivas
    from tests.unidad.test_diapositivas import LIMPIAS
    diseno = replace(catalogo.estilo("profesional"), logo="derecha")
    identidad = estilo.Identidad(**estilo.PALETA_SOBRIA, tipografia="Calibri", logo=_png(tmp_path / "logo.png"),
                                 estilo=diseno)
    portada = generador.Portada("Scrum Master con IA", "Sesión 1 · Fundamentos", "Diapositivas")
    ruta = diapositivas.generar_diapositivas(Diapositivas.model_validate(LIMPIAS), identidad, portada,
                                             tmp_path / "S1_Diapositivas.pptx")
    presentacion = Presentation(ruta)
    imagen = next(f for f in presentacion.slides[0].shapes if f.shape_type == 13)   # 13: imagen
    assert imagen.left + imagen.width > presentacion.slide_width / 2


def test_guardar_completa_confirma_y_sin_cambios_no_vuelve_a_confirmar(curso):
    almacen.guardar(curso, "curso", None, datos.ficha("curso", datos.CURSO))
    from app import servidor
    assert correr(servidor._guardar_y_confirmar(curso, "excel", "curso", None)) == "Ficha del curso guardada y confirmada."
    assert correr(servidor._guardar_y_confirmar(curso, "excel", "curso", None)) == "Ficha guardada. Sigue confirmada."
