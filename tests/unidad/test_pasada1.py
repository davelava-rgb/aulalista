"""Etapa 5b: primera pasada, veracidad oración por oración (PLAN.md §5.2)."""

import asyncio
import copy
import json

import openpyxl
import pytest

from app import herramientas
from app.fichas import almacen
from app.materiales import lectura
from app.validacion import pasada1
from app.validacion.pasajes import Corpus, Fuente, _lineas_de_ficha
from tests.conftest import SIN_COMPARAR, consulta_en_secuencia, pasada, segunda
from tests.unidad.test_lectura import LIMPIA, generar, sesion  # noqa: F401  (fixture)

NORMA = Fuente("guia.pdf", [
    {"texto": "Scrum tiene tres pilares: transparencia, inspección y adaptación.", "ubicacion": "página 35 · número impreso 54"},
    {"texto": "Los artefactos deben inspeccionarse con frecuencia.", "ubicacion": "página 37 · número impreso 56"},
])
FICHA = Fuente("ficha del curso", [{"texto": "Duración de cada Sprint: dos semanas", "ubicacion": "Cifras, reglas y nombres"}])


def oracion(n, texto, seccion="S1_Lectura.docx · Pilares"):
    return {"n": n, "huella": f"h{n}-{hash(texto)}", "seccion": seccion, "texto": texto}


def respuesta(n, tipo, veredicto="coincide", fuente="", pasaje="", **comparacion):
    return {"n": n, "tipo": tipo, "fuente": fuente, "pasaje": pasaje, "veredicto": veredicto,
            "motivo": "Motivo de prueba.", "comparacion": {**SIN_COMPARAR, **comparacion}}


def correr(oraciones, *salidas, anteriores=None, titulos=frozenset(), corpus=None):
    consulta = consulta_en_secuencia(*salidas)
    resultado = asyncio.run(pasada1.ejecutar(
        oraciones, titulos=set(titulos), corpus=corpus or Corpus([NORMA, FICHA]), anteriores=anteriores or {},
        curso="_sistema", sesion=1, material="lectura",
        instruccion_sin_ejecucion=lectura.INSTRUCCION_SIN_EJECUCION, consulta=consulta))
    return resultado, consulta


@pytest.fixture(autouse=True)
def clave(entorno):
    entorno["env"].write_text("ANTHROPIC_API_KEY=sk-ant-de-prueba\n", encoding="utf-8")


# ---------- Comprobaciones del programa sobre lo que dice la IA ----------

def test_un_pasaje_que_existe_queda_con_su_ubicacion_real():
    o = oracion(1, "Scrum tiene tres pilares.")
    resultado, _ = correr([o], {"oraciones": [respuesta(1, "norma", fuente="guia.pdf", pasaje="Scrum tiene tres pilares",
                                                        numero="igual", termino="igual", cantidades="igual")]})
    fila = resultado.filas[o["huella"]]
    assert fila["veredicto"] == "coincide"
    assert fila["fuente"] == "guia.pdf, página 35 · número impreso 54"


def test_un_pasaje_inventado_queda_sin_fuente():
    o = oracion(1, "Scrum tiene cuatro pilares.")
    resultado, _ = correr([o], {"oraciones": [respuesta(1, "norma", fuente="guia.pdf", pasaje="Scrum tiene cuatro pilares")]})
    fila = resultado.filas[o["huella"]]
    assert fila["veredicto"] == "sin fuente"
    assert "no aparece tal cual" in fila["motivo"]


def test_el_pasaje_se_busca_en_otra_fuente_si_la_ia_se_equivoca_de_nombre():
    o = oracion(1, "Cada Sprint dura dos semanas.")
    resultado, _ = correr([o], {"oraciones": [respuesta(1, "dato del caso", fuente="guia.pdf",
                                                        pasaje="Duración de cada Sprint: dos semanas")]})
    assert resultado.filas[o["huella"]]["fuente"] == "ficha del curso, Cifras, reglas y nombres"


def test_una_norma_con_una_diferencia_no_coincide_aunque_la_ia_diga_que_si():
    o = oracion(1, "Los artefactos pueden inspeccionarse con frecuencia.")
    resultado, _ = correr([o], {"oraciones": [respuesta(1, "norma", fuente="guia.pdf",
                                                        pasaje="Los artefactos deben inspeccionarse con frecuencia",
                                                        obligacion="distinto")]})
    fila = resultado.filas[o["huella"]]
    assert fila["veredicto"] == "no coincide"
    assert "obligacion" in fila["motivo"]


def test_regla_del_curso_necesita_en_este_curso():
    con = oracion(1, "En este curso, una situación contradice un aspecto solo si valora más el segundo elemento.")
    sin = oracion(2, "Una situación contradice un aspecto solo si valora más el segundo elemento.")
    resultado, _ = correr([con, sin], {"oraciones": [respuesta(1, "regla del curso"), respuesta(2, "regla del curso")]})
    assert resultado.filas[con["huella"]]["veredicto"] == "coincide"
    assert resultado.filas[sin["huella"]]["veredicto"] == "no coincide"


def test_instruccion_de_la_lectura_sin_archivos_que_ejecutar():
    o = oracion(1, "Elige un trabajo de tu área.")
    resultado, _ = correr([o], {"oraciones": [respuesta(1, "instrucción")]})
    fila = resultado.filas[o["huella"]]
    assert (fila["veredicto"], fila["pasaje"]) == ("coincide", lectura.INSTRUCCION_SIN_EJECUCION)


def test_una_oracion_que_la_ia_no_devolvio_no_queda_aprobada():
    o = oracion(1, "Scrum tiene tres pilares.")
    resultado, _ = correr([o], {"oraciones": []})
    assert resultado.filas[o["huella"]]["veredicto"] == "sin fuente"


# ---------- Ahorro: lo que no va a la IA ----------

def test_titulos_portada_y_encabezado_no_van_a_la_ia():
    oraciones = [oracion(1, "Pilares de Scrum"), oracion(2, "Scrum Master con IA", "S1_Lectura.docx · inicio"),
                 oracion(3, "Scrum Master con IA", "S1_Lectura.docx · encabezado y pie")]
    resultado, consulta = correr(oraciones, titulos={"Pilares de Scrum"})
    assert consulta.llamadas == []
    assert {f["tipo"] for f in resultado.filas.values()} == {"sin afirmación"}
    assert {f["pasaje"] for f in resultado.filas.values()} == {"no aplica"}


def test_las_oraciones_sin_cambios_conservan_su_resultado():
    o = oracion(1, "Scrum tiene tres pilares.")
    guardada = {"tipo": "norma", "pasaje": "Scrum tiene tres pilares", "fuente": "guia.pdf, página 35",
                "veredicto": "coincide", "motivo": "ok"}
    resultado, consulta = correr([o], anteriores={o["huella"]: guardada})
    assert consulta.llamadas == [] and resultado.filas[o["huella"]] == guardada


def test_las_oraciones_van_en_grupos_de_cuarenta_con_sus_candidatos():
    oraciones = [oracion(n, f"Scrum tiene tres pilares, oración {n}.") for n in range(1, 46)]
    resultado, consulta = correr(oraciones, pasada(), pasada())
    assert len(consulta.llamadas) == 2
    assert resultado.enviadas_a_la_ia == 45
    primero = consulta.llamadas[0]["prompt"]
    assert "(guia.pdf · página 35 · número impreso 54) Scrum tiene tres pilares" in primero
    assert "### Primera pasada · Veracidad" in primero
    opciones = consulta.llamadas[0]["options"]
    assert opciones.allowed_tools == [herramientas.BUSCAR] and opciones.tools == []
    assert opciones.model == "claude-sonnet-5-5"
    assert opciones.effort == "low"
    assert primero.index("### Primera pasada") < primero.index("ORACIONES:")  # la parte fija va primero


def test_la_herramienta_de_busqueda_devuelve_pasajes_con_su_ubicacion():
    servidor, uso = herramientas.servidor_de_busqueda(Corpus([NORMA]))
    assert servidor["name"] == "aulalista"
    texto = herramientas.formatear(Corpus([NORMA]).buscar("pilares de Scrum", 1))
    assert texto == "- (guia.pdf · página 35 · número impreso 54) Scrum tiene tres pilares: transparencia, inspección y adaptación."


# ---------- Ficha como fuente y datos nuevos del caso ----------

def test_cada_valor_de_la_ficha_es_un_pasaje_con_su_campo(sesion):  # noqa: F811
    pasajes = _lineas_de_ficha(almacen.ruta_ficha(sesion, "curso", extension=".md"))
    assert {"texto": "Duración de cada Sprint: dos semanas",
            "ubicacion": "Cifras, reglas y nombres que deben repetirse igual en todas las sesiones"} in pasajes
    sesion_md = _lineas_de_ficha(almacen.ruta_ficha(sesion, "sesion", 1, ".md"))
    assert {"texto": "Clasificar situaciones", "ubicacion": "Ejercicio 2 · Acción que realiza el alumno"} in sesion_md
    assert not any(p["texto"].endswith(":") for p in sesion_md)


def test_agregar_datos_fijos_no_repite_y_deja_la_ficha_confirmada(sesion):  # noqa: F811
    nuevos = almacen.agregar_datos_fijos(sesion, ["Plan de inducción: seis meses", "Duración de cada Sprint: dos semanas",
                                                  "sin formato"], "Lectura S1")
    ficha = almacen.cargar(sesion, "curso")
    assert nuevos == ["Plan de inducción: seis meses"]
    assert ficha["confirmada"] is True
    assert ficha["datos_agregados"][0]["origen"] == "Lectura S1"
    assert almacen.valor(ficha, "datos.fijos").endswith("Plan de inducción: seis meses")


def test_lectura_sin_fuente_se_corrige_con_un_dato_nuevo_del_caso(sesion):  # noqa: F811
    inventada = "La inducción dura seis meses."
    con_dato = copy.deepcopy(LIMPIA)
    con_dato["bloques"][1]["parrafos"].append(inventada)
    material, consulta = generar(
        sesion, con_dato,
        pasada({inventada: {"tipo": "dato del caso", "veredicto": "sin fuente", "pasaje": ""}}),
        segunda(),
        {"cambios": [], "explicaciones": [], "datos_nuevos": ["Plan de inducción del ejemplo: seis meses"]},
        pasada({inventada: {"tipo": "dato del caso", "fuente": "ficha del curso",
                            "pasaje": "Plan de inducción del ejemplo: seis meses"}}),
    )
    assert material["estado"] == "verificada"
    assert material["correcciones"] == {"verificador": 0, "primera pasada": 1, "segunda pasada": 0}
    assert material["datos_agregados"] == ["Plan de inducción del ejemplo: seis meses"]
    assert "VERACIDAD · sin fuente · «La inducción dura seis meses.»" in consulta.llamadas[3]["prompt"]
    assert consulta.llamadas[4]["prompt"].count("Oración:") == 1  # la segunda pasada solo revisa lo que no coincidía
    base = json.loads((almacen.carpeta_sesion(sesion, 1) / "verificacion.json").read_text(encoding="utf-8"))
    assert {"etiqueta": "Plan de inducción del ejemplo", "valor": "seis meses"} in base["datos_fijos"]

    libro = openpyxl.load_workbook(lectura.archivos(sesion, 1)["excel"])
    filas = list(libro["Oraciones"].iter_rows(min_row=2, values_only=True))
    assert all(f[4] and f[5] and f[7] for f in filas)  # ninguna fila sin tipo, pasaje y veredicto
    dato = next(f for f in filas if f[3] == inventada)
    assert dato[4:8] == ("dato del caso", "Plan de inducción del ejemplo: seis meses",
                         "ficha del curso, Cifras, reglas y nombres que deben repetirse igual en todas las sesiones", "coincide")
    historial = [f for f in libro["Hallazgos"].iter_rows(min_row=2, values_only=True)]
    assert len(historial) == 1
    assert historial[0][:4] == ("FALLA", "S1_Lectura.docx · Pilares de Scrum", inventada, "veracidad: sin fuente")
    assert historial[0][5] == "Enviado a corrección en la vuelta 1. Resuelto."


def test_una_norma_no_se_apoya_en_la_ficha():
    o = oracion(1, "Cada Sprint dura dos semanas según la guía.")
    resultado, _ = correr([o], {"oraciones": [respuesta(1, "norma", fuente="ficha del curso",
                                                        pasaje="Duración de cada Sprint: dos semanas")]})
    fila = resultado.filas[o["huella"]]
    assert fila["veredicto"] == "sin fuente" and "no con una ficha" in fila["motivo"]


def test_el_buscador_une_el_titulo_con_su_definicion_y_prefiere_palabras_raras():
    corpus = Corpus([Fuente("guia.pdf", [
        {"texto": "Transparencia", "ubicacion": "página 36"},
        {"texto": "El proceso y el trabajo emergentes deben ser visibles para quienes realizan el trabajo.", "ubicacion": "página 36"},
        {"texto": "Scrum es gratuito.", "ubicacion": "página 24"},
        {"texto": "Transparencia en Scrum.", "ubicacion": "página 40"},
    ])])
    primero = corpus.buscar("La transparencia pide que el proceso y el trabajo sean visibles.")[0]
    assert primero["ubicacion"] == "página 36"
    assert "deben ser visibles" in primero["texto"]


def test_las_notas_de_la_ficha_no_entran_en_el_pasaje(sesion):  # noqa: F811
    ficha = almacen.cargar(sesion, "curso")
    ficha["campos"]["caso.moneda"] = {"valor": "Soles", "origen": "Material (V01) SP.pdf", "estado": "propuesto"}
    almacen.guardar(sesion, "curso", None, ficha)
    pasajes = _lineas_de_ficha(almacen.ruta_ficha(sesion, "curso", extension=".md"))
    assert {"texto": "Soles", "ubicacion": "Moneda y país de los casos"} in pasajes



def test_una_oracion_escrita_tal_cual_en_la_fuente_se_aprueba_sin_ia():
    norma = oracion(1, "Los artefactos deben inspeccionarse con frecuencia.")
    dato = oracion(2, "Duración de cada Sprint: dos semanas")
    corta = oracion(3, "Scrum tiene tres pilares")  # está tal cual, pero tiene menos de cinco palabras
    resultado, consulta = correr([norma, dato, corta], pasada())
    assert resultado.aprobadas_por_programa == 2
    assert resultado.filas[norma["huella"]]["tipo"] == "norma"
    assert resultado.filas[norma["huella"]]["fuente"] == "guia.pdf, página 37 · número impreso 56"
    assert resultado.filas[dato["huella"]]["tipo"] == "dato del caso"
    assert consulta.llamadas[0]["prompt"].count("Oración:") == 1


def test_una_norma_sin_su_comparacion_no_queda_aprobada():
    o = oracion(1, "Scrum tiene tres pilares según la guía.")
    sin_comparacion = respuesta(1, "norma", fuente="guia.pdf", pasaje="Scrum tiene tres pilares")
    del sin_comparacion["comparacion"]
    resultado, _ = correr([o], {"oraciones": [sin_comparacion]})
    assert resultado.filas[o["huella"]]["veredicto"] == "no coincide"


def test_el_pasaje_del_redactor_es_el_primer_candidato():
    o = oracion(1, "Los artefactos se inspeccionan con frecuencia, como pide la guía.")
    anclas = [{"oracion": "Los artefactos se inspeccionan con frecuencia, como pide la guía.", "fuente": "guia",
               "ubicacion": "?", "texto": "deben inspeccionarse con frecuencia"}]
    consulta = consulta_en_secuencia(pasada())
    asyncio.run(pasada1.ejecutar([o], titulos=set(), corpus=Corpus([NORMA, FICHA]), anteriores={}, curso="_sistema",
                                 sesion=1, material="lectura", anclas=anclas, consulta=consulta))
    candidatos = consulta.llamadas[0]["prompt"].split("Candidatos:")[1]
    primero = candidatos.strip().splitlines()[0]
    assert primero == "- (guia.pdf · página 37 · número impreso 56) deben inspeccionarse con frecuencia"


def test_un_pasaje_del_redactor_que_no_existe_no_se_usa():
    o = oracion(1, "Los artefactos se inspeccionan cada hora.")
    anclas = [{"oracion": o["texto"], "fuente": "guia.pdf", "ubicacion": "?", "texto": "se inspeccionan cada hora"}]
    assert pasada1._candidato_de_ancla(o, anclas, Corpus([NORMA])) is None


# ---------- Reglas aflojadas (PLAN.md §0, decisión 10) ----------

def test_la_norma_se_compara_en_cuatro_puntos():
    assert pasada1.COMPARACIONES == ["numero", "termino", "cantidades", "obligacion"]
    esquema = pasada1.ESQUEMA["properties"]["oraciones"]["items"]["properties"]["comparacion"]
    assert esquema["required"] == pasada1.COMPARACIONES


def test_el_orden_o_quien_distintos_ya_no_bastan_para_no_coincidir():
    o = oracion(1, "Inspección, transparencia y adaptación son los tres pilares de Scrum.")
    resultado, _ = correr([o], {"oraciones": [respuesta(
        1, "norma", fuente="guia.pdf", pasaje="Scrum tiene tres pilares: transparencia, inspección y adaptación",
        numero="no aplica", termino="igual", cantidades="igual", obligacion="no aplica", orden="distinto")]})
    assert resultado.filas[o["huella"]]["veredicto"] == "coincide"


def test_un_resumen_con_en_este_curso_no_se_compara_punto_por_punto():
    o = oracion(1, "En este curso, resumimos la norma: los artefactos se revisan a menudo.")
    sin_comparacion = respuesta(1, "norma", fuente="guia.pdf", pasaje="Los artefactos deben inspeccionarse con frecuencia",
                                obligacion="distinto")
    del sin_comparacion["comparacion"]["numero"]
    resultado, _ = correr([o], {"oraciones": [sin_comparacion]})
    assert resultado.filas[o["huella"]]["veredicto"] == "coincide"


def test_un_resumen_que_contradice_sigue_sin_coincidir_y_necesita_su_pasaje():
    contradice = oracion(1, "En este curso, resumimos: los artefactos no se revisan.")
    inventado = oracion(2, "En este curso, resumimos: Scrum tiene cinco pilares.")
    resultado, _ = correr([contradice, inventado], {"oraciones": [
        respuesta(1, "norma", "no coincide", fuente="guia.pdf", pasaje="Los artefactos deben inspeccionarse con frecuencia"),
        respuesta(2, "norma", fuente="guia.pdf", pasaje="Scrum tiene cinco pilares")]})
    assert resultado.filas[contradice["huella"]]["veredicto"] == "no coincide"
    assert resultado.filas[inventado["huella"]]["veredicto"] == "sin fuente"


def test_la_ia_recibe_la_regla_de_cuatro_puntos_y_no_la_de_siete():
    o = oracion(1, "Los artefactos se revisan.")
    _, consulta = correr([o], {"oraciones": [respuesta(1, "sin afirmación")]})
    pedido = consulta.llamadas[0]["prompt"]
    assert pasada1.REGLA_NORMA_DE_AULALISTA in pedido
    assert "estas siete cosas" not in pedido and "las siete cosas" not in pedido
