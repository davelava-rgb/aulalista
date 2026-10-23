"""Agente diseñador: el programa pptxgenjs se revisa, corre con Node y su PowerPoint se controla (decisión 32)."""

import pytest
from pptx import Presentation

from app.materiales import disenador, estilo
from app.materiales.contenido import Diapositivas
from app.materiales.docx import Portada
from tests.unidad.test_diapositivas import LIMPIAS

pytestmark = pytest.mark.skipif(not disenador.disponible(), reason="Node o los paquetes del diseñador no están instalados")

IDENTIDAD = estilo.Identidad(**estilo.PALETA_SOBRIA, tipografia="Calibri", logo=None)
PORTADA = Portada("Scrum Master con IA", "Sesión 1 · Fundamentos", "Diapositivas", docente="José Fernando Davelouis Mendía",
                  contacto="www.linkedin.com/in/jdavelouis; info@futurasoft.net")

# Un diseño simple, escrito a mano, como el que devolvería el agente: una diapositiva por parte, con íconos.
DISENO = r"""
module.exports = async (h) => {
  const C = h.colores, c = h.contenido;
  const texto = (s, t, x, y, w, alto, o = {}) => s.addText(t, { x, y, w, h: alto, fontFace: h.fuente, fontSize: 16,
    color: C.texto, isTextBox: true, valign: "top", ...o });
  const icono = await h.icono("FaDatabase", C.blanco);
  const visibles = (d) => [...d.frases, ...d.columnas.flatMap((x) => [x.titulo, ...x.frases]),
    ...(d.tabla ? [d.tabla.titulo, ...d.tabla.encabezados, ...d.tabla.filas.flat()] : []),
    ...(d.grafico ? [d.grafico.centro, ...d.grafico.ejes,
                     ...d.grafico.elementos.flatMap((e) => [e.texto, e.detalle, ...e.hijos])] : [])].filter(Boolean);

  let s = h.pres.addSlide();
  s.background = { color: C.principal };
  texto(s, h.portada.curso, 0.5, 1.5, 6, 1, { fontSize: 36, bold: true, color: C.blanco });
  texto(s, h.portada.sesion, 0.5, 2.6, 6, 0.6, { color: C.blanco });
  texto(s, [h.portada.docente, ...h.portada.contacto].map((p) => ({ text: p, options: { breakLine: true } })), 0.5, 3.4, 6, 1.2, { color: C.blanco });
  s.addImage({ data: icono, x: 7.5, y: 1.8, w: 1.2, h: 1.2 });

  s = h.pres.addSlide();
  texto(s, "Idea central", 0.5, 0.4, 9, 0.6, { fontSize: 24, bold: true, color: C.principal });
  texto(s, c.idea_central.frases.join(" "), 0.5, 1.4, 9, 2);

  for (const [i, d] of c.diapositivas.entries()) {
    s = h.pres.addSlide();
    s.addShape(h.pres.shapes.ROUNDED_RECTANGLE, { x: 0.5, y: 0.2, w: 1.4, h: 0.3, rectRadius: 0.1, fill: { color: C.acento } });
    texto(s, `${i + 1} · Bloque`, 0.5, 0.2, 1.4, 0.3, { fontSize: 12, color: C.blanco, margin: 0 });
    texto(s, d.titulo, 0.5, 0.6, 9, 0.7, { fontSize: 24, bold: true, color: C.principal });
    const partes = visibles(d);
    texto(s, partes.map((p) => ({ text: p, options: { breakLine: true } })), 0.5, 1.5, 9, 3.5);
    // pie
  }
  for (const [titulo, parte, lineas] of [["Aplícalo así", c.aplicalo, [...c.aplicalo.frases, ...c.aplicalo.plantilla]],
                                          ["Cuidado con", c.cuidado, c.cuidado.riesgos]]) {
    s = h.pres.addSlide();
    texto(s, titulo, 0.5, 0.4, 9, 0.7, { fontSize: 24, bold: true, color: C.principal });
    texto(s, lineas.map((p) => ({ text: p, options: { breakLine: true } })), 0.5, 1.4, 9, 3);
  }
  s = h.pres.addSlide();
  texto(s, c.cierre.frase, 0.5, 1.5, 9, 1.5, { fontSize: 28, bold: true, color: C.principal });
};
"""


# El mismo diseño, por partes, como lo devuelve el agente (ESQUEMA_DISENO): una función común que dibuja la
# diapositiva n y una entrada por diapositiva que la llama.
COMUN = r"""
  const C = h.colores, c = h.contenido;
  const texto = (s, t, x, y, w, alto, o = {}) => s.addText(t, { x, y, w, h: alto, fontFace: h.fuente, fontSize: 16,
    color: C.texto, isTextBox: true, valign: "top", ...o });
  const visibles = (d) => [...d.frases, ...d.columnas.flatMap((x) => [x.titulo, ...x.frases]),
    ...(d.tabla ? [d.tabla.titulo, ...d.tabla.encabezados, ...d.tabla.filas.flat()] : []),
    ...(d.grafico ? [d.grafico.centro, ...d.grafico.ejes,
                     ...d.grafico.elementos.flatMap((e) => [e.texto, e.detalle, ...e.hijos])] : [])].filter(Boolean);
  const lista = (s, lineas, y) => texto(s, lineas.map((p) => ({ text: p, options: { breakLine: true } })), 0.5, y, 9, 3);
  const dibujar = async (n) => {
    const s = h.pres.addSlide();
    const k = c.diapositivas.length;
    if (n === 1) {
      s.background = { color: C.principal };
      texto(s, h.portada.curso, 0.5, 1.5, 6, 1, { fontSize: 36, bold: true, color: C.blanco });
      texto(s, h.portada.sesion, 0.5, 2.6, 6, 0.6, { color: C.blanco });
      texto(s, [h.portada.docente, ...h.portada.contacto].map((p) => ({ text: p, options: { breakLine: true } })), 0.5, 3.4, 6, 1.2, { color: C.blanco });
      s.addImage({ data: await h.icono("FaDatabase", C.blanco), x: 7.5, y: 1.8, w: 1.2, h: 1.2 });
      return;
    }
    if (n === 2) {
      texto(s, "Idea central", 0.5, 0.4, 9, 0.6, { fontSize: 24, bold: true, color: C.principal });
      texto(s, c.idea_central.frases.join(" "), 0.5, 1.4, 9, 2);
      return;
    }
    if (n <= k + 2) {
      const d = c.diapositivas[n - 3];
      texto(s, d.titulo, 0.5, 0.6, 9, 0.7, { fontSize: 24, bold: true, color: C.principal });
      lista(s, visibles(d), 1.5);
      return;
    }
    if (n === k + 3 || n === k + 4) {
      const [titulo, parte, lineas] = n === k + 3 ? ["Aplícalo así", c.aplicalo, [...c.aplicalo.frases, ...c.aplicalo.plantilla]]
                                                   : ["Cuidado con", c.cuidado, c.cuidado.riesgos];
      texto(s, titulo, 0.5, 0.4, 9, 0.7, { fontSize: 24, bold: true, color: C.principal });
      lista(s, lineas, 1.4);
      return;
    }
    texto(s, c.cierre.frase, 0.5, 1.5, 9, 1.5, { fontSize: 28, bold: true, color: C.principal });
  };
"""


def respuesta_por_partes(total: int, comun: str = COMUN, notas: str = "Tarjetas y una portada sobre el color principal.") -> dict:
    """Lo que devuelve el diseñador: solo las diapositivas (decisión 68 d). El código de prueba que comparten va
    repetido dentro de cada una, porque el esquema ya no tiene «comun»."""
    return {"diapositivas": [{"numero": n, "codigo": f"{comun}\nawait dibujar({n});"} for n in range(1, total + 1)],
            "notas": notas}


@pytest.fixture
def contenido():
    return Diapositivas.model_validate(LIMPIAS)


def correr(tmp_path, contenido, codigo=DISENO):
    datos = disenador.datos(contenido, IDENTIDAD, PORTADA, tmp_path)
    salida = tmp_path / "S1_Diapositivas.pptx"
    return disenador.correr(codigo, datos, tmp_path / "diseno", salida), salida


def test_el_diseno_corre_con_node_y_cumple_los_controles(tmp_path, contenido):
    corrida, salida = correr(tmp_path, contenido)
    assert corrida.ok, corrida.error
    assert corrida.iconos == ("FaDatabase",)
    assert len(Presentation(salida).slides) == contenido.total
    assert disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD) == []


def test_un_programa_con_require_o_process_se_rechaza_antes_de_correr():
    problemas = disenador.revisar_codigo('const fs = require("fs");\nmodule.exports = async (h) => { process.exit(0); };')
    assert any("require" in p and "process" in p for p in problemas)
    assert disenador.revisar_codigo("async function f(h) {}") != []


def test_node_no_deja_escribir_fuera_de_la_carpeta(tmp_path, contenido):
    """Aunque un programa se saltara la revisión, Node le niega los archivos fuera de su carpeta."""
    fuera = tmp_path.parent / "fuera_del_material.txt"
    codigo = ('module.exports = async (h) => { const m = h.pres.constructor.constructor("return this")().process;'
              f' m.mainModule.require("fs").writeFileSync({str(fuera)!r}, "x"); }};')
    corrida, _ = correr(tmp_path, contenido, codigo)
    assert not corrida.ok and not fuera.exists()
    assert "ERR_ACCESS_DENIED" in corrida.error or "restricted" in corrida.error, corrida.error


def test_un_error_del_programa_vuelve_como_problema(tmp_path, contenido):
    corrida, _ = correr(tmp_path, contenido, 'module.exports = async (h) => { await h.icono("FaNoExiste", "000000"); };')
    assert not corrida.ok and "No existe el ícono «FaNoExiste»" in corrida.error


def test_los_controles_detectan_texto_ajeno_letra_chica_colores_y_cantidad(tmp_path, contenido):
    malo = (DISENO
            .replace('texto(s, h.portada.sesion, 0.5, 2.6, 6, 0.6, { color: C.blanco });',
                     'texto(s, "Scrum nació en Japón en 1986 y lo usan todas las empresas.", 0.5, 2.6, 6, 0.6, '
                     '{ color: "00FF00", fontSize: 9 });')
            .replace("  s = h.pres.addSlide();\n  texto(s, c.cierre.frase", "  s = h.pres.addSlide();\n  s = h.pres.addSlide();\n  texto(s, c.cierre.frase"))
    corrida, salida = correr(tmp_path, contenido, malo)
    assert corrida.ok, corrida.error
    problemas = " ".join(disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD))
    assert "no está tal cual en el contenido aprobado" in problemas
    assert "letra de 9" in problemas
    assert "00FF00" in problemas
    assert f"deben ser {contenido.total}" in problemas


# ---------- El programa por partes: un arreglo devuelve solo lo que cambia ----------

def test_el_programa_armado_por_partes_corre_y_cumple_los_controles(tmp_path, contenido):
    partes, faltan = disenador.fusionar(None, respuesta_por_partes(contenido.total), contenido.total)
    assert faltan == []
    corrida, salida = correr(tmp_path, contenido, disenador.armar(partes))
    assert corrida.ok, corrida.error
    assert len(Presentation(salida).slides) == contenido.total
    assert disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD) == []


def test_un_arreglo_cambia_solo_las_diapositivas_que_devuelve():
    partes, _ = disenador.fusionar(None, respuesta_por_partes(5, comun="const a = 1;"), 5)
    arreglo = {"comun": "const b = 2;", "diapositivas": [{"numero": 3, "codigo": "await otra(3);"},
                                                         {"numero": 9, "codigo": "fuera de rango"}], "notas": ""}
    nuevas, faltan = disenador.fusionar(partes, arreglo, 5)
    assert faltan == []
    assert not hasattr(nuevas, "comun")                          # el común es fijo: el diseñador no lo escribe
    assert nuevas.diapositivas[3] == "await otra(3);"
    assert nuevas.diapositivas[2] == "const a = 1;\nawait dibujar(2);" and 9 not in nuevas.diapositivas
    assert partes.diapositivas[3] == "const a = 1;\nawait dibujar(3);"   # las partes anteriores no se tocan


def test_un_diseno_al_que_le_faltan_diapositivas_vuelve_como_problema():
    respuesta = respuesta_por_partes(5)
    del respuesta["diapositivas"][1:3]
    _, faltan = disenador.fusionar(None, respuesta, 5)
    assert faltan and "diapositiva 2, 3" in faltan[0]


def test_el_arreglo_muestra_el_programa_por_partes_y_pide_solo_lo_que_cambia(contenido):
    partes, _ = disenador.fusionar(None, respuesta_por_partes(3, comun="const a = 1;"), 3)
    pedido = disenador.pedido_de_arreglo(partes, ["FALLA · diseño · diapositiva 2: letra de 9 puntos."])
    assert "// diapositiva 2\nconst a = 1;\nawait dibujar(2);" in pedido and "// comun" not in pedido
    assert "el código común «A» de AulaLista no cambia" in pedido
    assert "devuelve SOLO las diapositivas que cambias" in pedido
    assert "- FALLA · diseño · diapositiva 2: letra de 9 puntos." in pedido


# ---------- Etapa A de la optimización (decisión 41) ----------

def test_la_guia_y_los_datos_van_en_el_mensaje_de_sistema_y_no_en_los_pedidos(tmp_path, contenido):
    """Lo que no cambia entre el diseño y sus arreglos va en un archivo de sistema, que el SDK guarda en caché."""
    datos = disenador.datos(contenido, IDENTIDAD, PORTADA, tmp_path)
    ruta = disenador.escribir_sistema(datos, tmp_path / "diseno")
    sistema = ruta.read_text(encoding="utf-8")
    assert sistema.startswith("ERES EL DISEÑADOR") and contenido.cierre.frase in sistema
    partes, _ = disenador.fusionar(None, respuesta_por_partes(3), 3)
    for pedido in (disenador.pedido_de_diseno(3), disenador.pedido_de_arreglo(partes, ["un problema"])):
        assert "ERES EL DISEÑADOR" not in pedido and contenido.cierre.frase not in pedido


def test_letra_chica_y_colores_ajenos_se_arreglan_sin_ia(tmp_path, contenido):
    malo = DISENO.replace('texto(s, h.portada.sesion, 0.5, 2.6, 6, 0.6, { color: C.blanco });',
                          'texto(s, h.portada.sesion, 0.5, 2.6, 6, 0.6, { color: "00FF00", fontSize: 9 });')
    corrida, salida = correr(tmp_path, contenido, malo)
    assert corrida.ok, corrida.error
    assert disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD) != []
    ajustes = " ".join(disenador.ajustar_pptx(salida, contenido, IDENTIDAD))
    assert "subí a 12 puntos 1 trozos" in ajustes
    assert "00FF00 por" in ajustes
    assert disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD) == []
    assert disenador.ajustar_pptx(salida, contenido, IDENTIDAD) == []     # un PowerPoint correcto no se toca


def test_el_texto_ajeno_no_se_arregla_sin_ia(tmp_path, contenido):
    """Solo se arregla sin IA lo que tiene una sola respuesta correcta. Un texto ajeno vuelve al diseñador."""
    malo = DISENO.replace('texto(s, h.portada.sesion, 0.5, 2.6, 6, 0.6, { color: C.blanco });',
                          'texto(s, "Scrum nació en Japón en 1986 y lo usan todas las empresas.", 0.5, 2.6, 6, 0.6);')
    corrida, salida = correr(tmp_path, contenido, malo)
    assert corrida.ok, corrida.error
    disenador.ajustar_pptx(salida, contenido, IDENTIDAD)
    assert "no está tal cual" in " ".join(disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD))


# ---------- Contradicciones entre la guía y los controles ----------

LISTA = "    texto(s, partes.map((p) => ({ text: p, options: { breakLine: true } })), 0.5, 1.5, 9, 3.5);"


def con_lista(trozos: str) -> str:
    """Las frases de cada diapositiva como viñetas, con la primera palabra en un trozo y el resto en otro."""
    assert LISTA in DISENO
    return DISENO.replace(LISTA, (
        "    texto(s, partes.flatMap((p) => { const [a, ...b] = p.split(' ');\n"
        "      if (!b.length) return [{ text: p, options: { bullet: true, breakLine: true } }];\n"
        f"      return {trozos}; }}), 0.5, 1.5, 9, 3.5);"))


def test_las_etiquetas_y_el_pie_unidos_con_punto_medio_se_aceptan(tmp_path, contenido):
    """La guía pide «1 · Qué es el gobierno de datos» y el pie con el curso y la sesión: el número y el separador no
    cuentan como palabras, y cada tramo que sale del contenido o de la portada tampoco."""
    bien = (DISENO.replace("`${i + 1} · Bloque`", "`${i + 1} · ${d.bloque}`")
            .replace("    // pie",
                     "    texto(s, `${h.portada.curso}  ·  ${h.portada.sesion}`, 0.5, 5.1, 9, 0.3, { fontSize: 12 });"))
    corrida, salida = correr(tmp_path, contenido, bien)
    assert corrida.ok, corrida.error
    assert disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD) == []


def test_un_texto_ajeno_unido_con_punto_medio_sigue_siendo_ajeno(tmp_path, contenido):
    malo = DISENO.replace("`${i + 1} · Bloque`",
                          "`${i + 1} · Scrum · nació en Japón en 1986 y hoy lo usan todas las empresas`")
    corrida, salida = correr(tmp_path, contenido, malo)
    assert corrida.ok, corrida.error
    assert "nació en Japón" in " ".join(disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD))


def test_una_frase_partida_en_varias_vinetas_vuelve_al_disenador(tmp_path, contenido):
    """Con bullet en cada trozo, pptxgenjs hace una viñeta por trozo y la frase queda partida (Agilidad S1)."""
    malo = con_lista("[{ text: a + ' ', options: { bullet: true } }, "
                     "{ text: b.join(' '), options: { bullet: true, bold: true, breakLine: true } }]")
    corrida, salida = correr(tmp_path, contenido, malo)
    assert corrida.ok, corrida.error
    problemas = " ".join(disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD))
    assert "quedó partida en varias viñetas" in problemas and "bullet: true solo en su primer trozo" in problemas


def test_una_frase_resaltada_con_bullet_solo_en_su_primer_trozo_queda_entera(tmp_path, contenido):
    bien = con_lista("[{ text: a + ' ', options: { bullet: true } }, "
                     "{ text: b.join(' '), options: { bold: true, breakLine: true } }]")
    corrida, salida = correr(tmp_path, contenido, bien)
    assert corrida.ok, corrida.error
    assert disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD) == []
    lista = max(Presentation(salida).slides[2].shapes, key=lambda f: len(f.text_frame.paragraphs))
    assert len(lista.text_frame.paragraphs) == len(contenido.diapositivas[0].textos_visibles())


def test_el_texto_que_el_diseno_quita_o_pasa_a_las_notas_vuelve_al_disenador(tmp_path, contenido):
    """Las notas se copian del contenido sin IA: un texto que el diseño pasa a las notas se perdería."""
    malo = DISENO.replace("[...c.aplicalo.frases, ...c.aplicalo.plantilla]", "c.aplicalo.plantilla")
    corrida, salida = correr(tmp_path, contenido, malo)
    assert corrida.ok, corrida.error
    problemas = " ".join(disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD))
    assert f"falta el texto «{contenido.aplicalo.frases[0][:120]}»" in problemas


# ---------- Decisión 56 ----------

def test_una_etiqueta_de_hasta_ocho_palabras_se_acepta(tmp_path, contenido):
    bien = DISENO.replace("`${i + 1} · Bloque`", "`${i + 1} · Lo que el equipo mira cada día aquí`")
    corrida, salida = correr(tmp_path, contenido, bien)
    assert corrida.ok, corrida.error
    assert disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD) == []


def test_el_diseno_tiene_el_tamano_de_la_version_base(tmp_path, contenido):
    """Así una diapositiva base puede reemplazar a una diseñada sin cambiar de escala."""
    from app.materiales import pptx as generador
    corrida, salida = correr(tmp_path, contenido)
    assert corrida.ok, corrida.error
    diseno = Presentation(salida)
    assert abs(diseno.slide_width - generador.ANCHO) <= disenador.MEDIO_MILIMETRO
    assert abs(diseno.slide_height - generador.ALTO) <= disenador.MEDIO_MILIMETRO


def test_diapositivas_de_lee_los_numeros_de_los_problemas():
    problemas = ["FALLA · diseño · diapositiva 9: el texto «x» no está tal cual.",
                 "FALLA · desborde · S1_Diapositivas.pptx · diapositiva 15: Diapositiva 15 («Y»): no cabe."]
    assert disenador.diapositivas_de(problemas) == [9, 15]
    assert disenador.diapositivas_de([*problemas, "FALLA · diseño · hay 14 diapositivas; deben ser 15."]) == []
    assert disenador.diapositivas_de([*problemas, "FALLA · diseño · el programa no corrió"], todas=False) == [9, 15]


# ---------- Decisión 68 d: el código común es de AulaLista; el diseñador solo escribe las diapositivas ----------

CON_AYUDAS = {
    1: "const s = h.pres.addSlide(); await A.portada(s);",
    2: 'const s = h.pres.addSlide(); A.titulo(s, "Idea central", "", 2); A.cita(s, h.contenido.idea_central.frases.join(" "));',
    "concepto": (
        "const s = h.pres.addSlide(); const d = h.contenido.diapositivas[{i}];\n"
        "const y = A.titulo(s, d.titulo, d.bloque, {n});\n"
        "const partes = [...d.frases, ...d.columnas.flatMap((x) => [x.titulo, ...x.frases]),\n"
        "  ...(d.tabla ? [] : []), ...(d.grafico ? [d.grafico.centro, ...d.grafico.ejes,\n"
        "  ...d.grafico.elementos.flatMap((e) => [e.texto, e.detalle, ...e.hijos])] : [])].filter(Boolean);\n"
        "if (d.tabla) {{ if (partes.length) A.frases(s, partes, y, 0.8, 18); A.tabla(s, d.tabla, y + (partes.length ? 0.9 : 0), 4); }}\n"
        "else if (d.tipo === 'columnas') A.columnas(s, d.columnas[0], d.columnas[1], y);\n"
        "else if (d.tipo === 'pasos' || d.tipo === 'proceso') A.pasos(s, partes, y);\n"
        "else if (d.tipo === 'tarjetas') await A.tarjetas(s, [...d.frases.map((t) => ({{ texto: t }})), "
        "...d.grafico.elementos.map((e) => ({{ texto: e.texto, detalle: e.detalle }}))], y, 4.4, ['FaDatabase']);\n"
        "else A.frases(s, partes, y);"),
    "aplicalo": ('const s = h.pres.addSlide(); const y = A.titulo(s, "Aplícalo así", "", {n}); '
                 "A.frases(s, [...h.contenido.aplicalo.frases, ...h.contenido.aplicalo.plantilla], y);"),
    "cuidado": ('const s = h.pres.addSlide(); const y = A.titulo(s, "Cuidado con", "", {n}); '
                "A.frases(s, h.contenido.cuidado.riesgos, y);"),
    "cierre": "const s = h.pres.addSlide(); A.cierre(s);",
}


def respuesta_con_ayudas(total: int) -> dict:
    conceptos = total - 5
    codigos = {1: CON_AYUDAS[1], 2: CON_AYUDAS[2]}
    for i in range(conceptos):
        codigos[i + 3] = CON_AYUDAS["concepto"].format(i=i, n=i + 3)
    codigos[conceptos + 3] = CON_AYUDAS["aplicalo"].format(n=conceptos + 3)
    codigos[conceptos + 4] = CON_AYUDAS["cuidado"].format(n=conceptos + 4)
    codigos[conceptos + 5] = CON_AYUDAS["cierre"]
    return {"diapositivas": [{"numero": n, "codigo": c} for n, c in codigos.items()], "notas": "Con las ayudas de A."}


def test_el_esquema_del_diseno_no_pide_el_codigo_comun():
    assert "comun" not in disenador.ESQUEMA_DISENO["properties"]
    assert "A.titulo(" in disenador.GUIA and "No definas funciones comunes" in disenador.GUIA
    assert disenador.revisar_codigo(disenador.armar(disenador.Partes({1: "// nada"}))) == []   # el común pasa el control


def test_las_diapositivas_escritas_con_las_ayudas_corren_y_cumplen_los_controles(tmp_path, contenido):
    partes, faltan = disenador.fusionar(None, respuesta_con_ayudas(contenido.total), contenido.total)
    assert faltan == []
    codigo = disenador.armar(partes)
    assert "const A = (() => {" in codigo and codigo.index("const A") < codigo.index("// diapositiva 1")
    corrida, salida = correr(tmp_path, contenido, codigo)
    assert corrida.ok, corrida.error
    assert "FaDatabase" in corrida.iconos
    assert len(Presentation(salida).slides) == contenido.total
    assert disenador.revisar_pptx(salida, contenido, PORTADA, IDENTIDAD) == []


def test_el_disenador_no_puede_redefinir_las_ayudas(tmp_path, contenido):
    respuesta = respuesta_con_ayudas(contenido.total)
    respuesta["diapositivas"][0]["codigo"] = "A = null; const s = h.pres.addSlide();"
    partes, _ = disenador.fusionar(None, respuesta, contenido.total)
    corrida, _ = correr(tmp_path, contenido, disenador.armar(partes))
    assert not corrida.ok and "constant" in corrida.error.lower()
