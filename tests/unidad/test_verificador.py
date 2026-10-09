"""Etapa 3: programa verificador y verificacion.json (PLAN.md §5.1 y §6)."""

import importlib.util
import json
import sys
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from pptx import Presentation

from tests.sembrados.casos import CASOS_POR_ORACION, ORACIONES_LIMPIAS

RAIZ = Path(__file__).resolve().parents[2]
RUTA_VERIFICADOR = RAIZ / ".claude" / "skills" / "material-de-clase" / "scripts" / "verificar.py"
MODELO = RAIZ / "S1_Laboratorio_Verificacion.xlsx"

especificacion = importlib.util.spec_from_file_location("verificar", RUTA_VERIFICADOR)
verificar = importlib.util.module_from_spec(especificacion)
sys.modules["verificar"] = verificar
especificacion.loader.exec_module(verificar)


def sin_word(_ruta):
    return 1


# ---------- Material y configuración de prueba ----------

@pytest.fixture
def sesion(tmp_path):
    """Carpeta de una sesión con su fuente convertida, su carpeta de práctica y verificacion.json."""
    fuente = tmp_path / "guia.pdf.jsonl"
    pasajes = [
        {"texto": "Definición de Scrum", "ubicacion": "página 24", "pagina": 24, "numero_impreso": 43},
        {"texto": "Scrum es gratuito y se aplica completo en cada equipo.", "ubicacion": "página 24",
         "pagina": 24, "numero_impreso": 43},
        {"texto": "Tres pilares de Scrum", "ubicacion": "página 35", "pagina": 35, "numero_impreso": 54},
    ]
    fuente.write_text("\n".join(json.dumps(p, ensure_ascii=False) for p in pasajes), encoding="utf-8")
    practica = tmp_path / "practica"
    practica.mkdir()
    configuracion = {
        "sesion": 1,
        "material": "Laboratorio",
        "fuentes": [{
            "nombre": "Guía de Scrum", "texto": "guia.pdf.jsonl", "norma": True,
            "referencias": {"patron": r"lámina (\d+)", "indice": "numero_impreso"},
        }],
        "datos_fijos": [{"etiqueta": "Fecha de referencia", "valor": "lunes 5 de octubre de 2026"}],
        "vocabulario": [{"concepto": "Manifiesto Ágil", "variantes": ["Agile"]}],
        "nunca_se_incluye": ["objetivos de aprendizaje", "requisitos previos", "glosario"],
        "permitidos": ["notas del profesor", "puntos de historia"],
        "carpeta_practica": "practica",
    }
    ruta = tmp_path / "verificacion.json"
    ruta.write_text(json.dumps(configuracion, ensure_ascii=False), encoding="utf-8")
    return {"carpeta": tmp_path, "config": ruta, "datos": configuracion}


def con_archivo_de_practica(sesion):
    (sesion["carpeta"] / "practica" / "S1_E1_registro-almacen.xlsx").write_bytes(b"")


def guardar_config(sesion, **cambios):
    sesion["datos"].update(cambios)
    sesion["config"].write_text(json.dumps(sesion["datos"], ensure_ascii=False), encoding="utf-8")


def word(ruta: Path, parrafos: list[str], titulo: str = "Ejercicio 1 · Revisa el registro") -> Path:
    documento = Document()
    documento.add_paragraph("Laboratorio", style="Title")
    documento.add_heading(titulo, level=1)
    for parrafo in parrafos:
        documento.add_paragraph(parrafo)
    documento.save(ruta)
    return ruta


def correr(sesion, archivos, **opciones):
    return verificar.verificar(sesion["config"], archivos, contar_paginas=opciones.get("contar_paginas", sin_word))


# ---------- Formato del Excel ----------

def test_excel_tiene_las_hojas_columnas_y_estilos_del_modelo(sesion, tmp_path):
    material = word(tmp_path / "S1_Laboratorio.docx", ["Esta pregunta vale 5 puntos."])
    salida = verificar.escribir_excel(correr(sesion, [material]), tmp_path / "S1_Laboratorio_Verificacion.xlsx")
    modelo = openpyxl.load_workbook(MODELO)
    nuevo = openpyxl.load_workbook(salida)
    assert nuevo.sheetnames == modelo.sheetnames
    for nombre in modelo.sheetnames:
        m, n = modelo[nombre], nuevo[nombre]
        assert [c.value for c in n[1]] == [c.value for c in m[1]], nombre
        assert n.freeze_panes == m.freeze_panes
        for letra, dimension in m.column_dimensions.items():
            assert n.column_dimensions[letra].width == dimension.width, (nombre, letra)
        for celda_m, celda_n in zip(m[1], n[1]):
            assert (celda_n.font.name, celda_n.font.b, celda_n.font.color.rgb, celda_n.fill.fgColor.rgb) == \
                   (celda_m.font.name, celda_m.font.b, celda_m.font.color.rgb, celda_m.fill.fgColor.rgb)
    for nombre in ("Oraciones", "Hallazgos"):
        celda_m, celda_n = modelo[nombre]["D2" if nombre == "Oraciones" else "C2"], nuevo[nombre]["D2" if nombre == "Oraciones" else "C2"]
        assert (celda_n.font.name, celda_n.font.sz, celda_n.alignment.wrap_text, celda_n.alignment.vertical) == \
               (celda_m.font.name, celda_m.font.sz, celda_m.alignment.wrap_text, celda_m.alignment.vertical)


def test_excel_no_copia_los_defectos_del_modelo(sesion, tmp_path):
    seccion_larga = "Ejercicio 2 · Encuentra el pilar que falla en el registro del almacén"
    material = word(tmp_path / "S1_Laboratorio.docx", ["El Manifiesto Ágil tiene cuatro aspectos."], titulo=seccion_larga)
    salida = verificar.escribir_excel(correr(sesion, [material]), tmp_path / "v.xlsx")
    libro = openpyxl.load_workbook(salida)
    secciones = [c.value for c in libro["Oraciones"]["B"][1:]]
    assert f"S1_Laboratorio.docx · {seccion_larga}" in secciones  # nombre completo, sin cortar
    repetidos = {f[0].value: f[2].value for f in libro["Datos repetidos"].iter_rows(min_row=2)}
    assert repetidos["Manifiesto Ágil"].startswith("S1_Laboratorio.docx")  # sin la «L» de más


def test_cada_fila_de_oraciones_tiene_numero_seccion_parte_y_texto(sesion, tmp_path):
    material = word(tmp_path / "S1_Laboratorio.docx", ["Scrum tiene tres pilares. El tablero está en la pared."])
    resultado = correr(sesion, [material])
    assert [(o.n, o.seccion, o.parte, o.texto) for o in resultado.oraciones] == [
        (1, "S1_Laboratorio.docx · inicio", "texto", "Laboratorio"),
        (2, "S1_Laboratorio.docx · Ejercicio 1 · Revisa el registro", "texto", "Ejercicio 1 · Revisa el registro"),
        (3, "S1_Laboratorio.docx · Ejercicio 1 · Revisa el registro", "texto", "Scrum tiene tres pilares."),
        (4, "S1_Laboratorio.docx · Ejercicio 1 · Revisa el registro", "texto", "El tablero está en la pared."),
    ]


def test_la_huella_cambia_si_cambia_una_letra():
    a = verificar.Oracion(1, "a.docx", "a.docx · inicio", "texto", "Scrum tiene tres pilares.")
    b = verificar.Oracion(9, "a.docx", "a.docx · inicio", "texto", "Scrum tiene tres pilares.")
    c = verificar.Oracion(1, "a.docx", "a.docx · inicio", "texto", "Scrum tiene tres pilares!")
    assert a.huella == b.huella != c.huella


# ---------- Separador de oraciones ----------

def oraciones_del_modelo() -> list[tuple[str, str, str]]:
    hoja = openpyxl.load_workbook(MODELO)["Oraciones"]
    return [(f[1], f[2], f[3]) for f in hoja.iter_rows(min_row=2, values_only=True)]


def test_no_corta_ninguna_oracion_del_modelo():
    mal_cortadas = [t for _, _, t in oraciones_del_modelo() if verificar.separar_oraciones(t) != [t]]
    assert mal_cortadas == []


def test_separa_oraciones_seguidas_del_modelo():
    filas = oraciones_del_modelo()
    pares = [
        (a[2], b[2]) for a, b in zip(filas, filas[1:])
        if a[0] == b[0] and a[1] == b[1] == "texto" and a[2].endswith((".", "?", ":")) is not False
        and a[2][-1] in ".?" and b[2][:1].isupper()
    ]
    assert len(pares) > 100
    mal = [(a, b) for a, b in pares if verificar.separar_oraciones(f"{a} {b}") != [a, b]]
    assert mal == []


def test_no_corta_abreviaturas_ni_decimales():
    assert verificar.separar_oraciones("Usa p. ej. el registro. El valor es 3,5 y no 4.") == [
        "Usa p. ej. el registro.", "El valor es 3,5 y no 4."]


# ---------- Errores sembrados: uno por revisión ----------

@pytest.mark.parametrize("ident, oracion, nivel, regla", CASOS_POR_ORACION, ids=[c[0] for c in CASOS_POR_ORACION])
def test_cada_error_sembrado_queda_marcado_en_su_oracion(sesion, tmp_path, ident, oracion, nivel, regla):
    material = word(tmp_path / "S1_Laboratorio.docx", ["Scrum tiene tres pilares.", oracion])
    resultado = correr(sesion, [material])
    fila = next(o for o in resultado.oraciones if o.texto == oracion)
    marcados = [(h.nivel, h.regla) for h in resultado.hallazgos if h.n == fila.n]
    assert (nivel, regla) in marcados
    assert not [h for h in resultado.hallazgos if h.n != fila.n]  # la oración correcta queda limpia


def test_control_limpio_no_marca_nada(sesion, tmp_path):
    con_archivo_de_practica(sesion)
    material = word(tmp_path / "S1_Laboratorio.docx", ORACIONES_LIMPIAS)
    resultado = correr(sesion, [material])
    assert [(h.regla, h.oracion, h.detalle) for h in resultado.hallazgos] == []
    assert resultado.resumen() == f"Oraciones: {len(ORACIONES_LIMPIAS) + 2} · Fallas: 0 · Avisos: 0"


def test_lenguaje_de_ia_dentro_de_una_cita_es_aviso(sesion, tmp_path):
    material = word(tmp_path / "S1_Laboratorio.docx", ["El borrador decía «es importante destacar»."])
    resultado = correr(sesion, [material])
    assert [(h.nivel, h.regla) for h in resultado.hallazgos] == [("AVISO", "lenguaje de IA")]


def test_la_clausula_citada_anota_su_titulo(sesion, tmp_path):
    material = word(tmp_path / "S1_Laboratorio.docx", ["Lo explica la lámina 43 del capítulo."])
    fila = correr(sesion, [material]).oraciones[-1]
    assert fila.fuente == "Guía de Scrum, lámina 43: Definición de Scrum"


def test_cita_presente_en_la_ficha_es_valida(sesion, tmp_path):
    (tmp_path / "ficha.md").write_text("El problema sembrado es que el tablero no se actualiza nunca.", encoding="utf-8")
    guardar_config(sesion, textos_permitidos=["ficha.md"])
    material = word(tmp_path / "S1_Laboratorio.docx", ["La ficha dice «el tablero no se actualiza nunca»."])
    assert correr(sesion, [material]).hallazgos == []


def test_indice_por_patron_de_clausulas(sesion, tmp_path):
    norma = tmp_path / "norma.jsonl"
    norma.write_text(json.dumps({"texto": "5.1 Liderazgo y compromiso", "ubicacion": "página 9"}), encoding="utf-8")
    guardar_config(sesion, fuentes=[{
        "nombre": "Norma", "texto": "norma.jsonl", "norma": True,
        "referencias": {"patron": r"cláusula (\d+(?:\.\d+)*)", "indice": "patron",
                        "patron_indice": r"^(\d+(?:\.\d+)*)\s+(.+)$"},
    }])
    material = word(tmp_path / "S1_Laboratorio.docx", ["La cláusula 5.1 pide liderazgo.", "La cláusula 7.9 pide recursos."])
    resultado = correr(sesion, [material])
    assert resultado.oraciones[-2].fuente == "Norma, cláusula 5.1: Liderazgo y compromiso"
    assert [(h.n, h.regla) for h in resultado.hallazgos] == [(resultado.oraciones[-1].n, "referencia inexistente")]


# ---------- Revisiones del material completo ----------

def test_archivo_de_practica_sin_mencionar_y_con_nombre_mal_formado(sesion, tmp_path):
    con_archivo_de_practica(sesion)
    (tmp_path / "practica" / "registro final.xlsx").write_bytes(b"")
    material = word(tmp_path / "S1_Laboratorio.docx", ["Scrum tiene tres pilares."])
    reglas = {(h.oracion, h.regla) for h in correr(sesion, [material]).hallazgos}
    assert ("S1_E1_registro-almacen.xlsx", "archivo sin mencionar") in reglas
    assert ("registro final.xlsx", "nombre de archivo") in reglas


def test_material_sin_carpeta_de_practica_no_puede_mencionar_archivos(sesion, tmp_path):
    guardar_config(sesion, carpeta_practica=None)
    material = word(tmp_path / "S1_Lectura.docx", ["Abre el archivo S1_E1_registro-almacen.xlsx."])
    assert [h.regla for h in correr(sesion, [material]).hallazgos] == ["archivo inexistente"]


def test_mas_paginas_que_el_maximo_es_falla(sesion, tmp_path):
    guardar_config(sesion, limites={"paginas_max": 6})
    material = word(tmp_path / "S1_Lectura.docx", ["Scrum tiene tres pilares."])
    hallazgos = correr(sesion, [material], contar_paginas=lambda _: 7).hallazgos
    assert [(h.nivel, h.regla, h.detalle) for h in hallazgos] == [("FALLA", "páginas", "tiene 7 páginas; el máximo es 6")]


def test_si_word_no_cuenta_las_paginas_queda_aviso(sesion, tmp_path):
    guardar_config(sesion, limites={"paginas_max": 6})
    def falla(_):
        raise RuntimeError("Word no está")
    material = word(tmp_path / "S1_Lectura.docx", ["Scrum tiene tres pilares."])
    assert [(h.nivel, h.regla) for h in correr(sesion, [material], contar_paginas=falla).hallazgos] == [("AVISO", "páginas")]


def test_conteo_de_ejercicios_y_titulo_opcional(sesion, tmp_path):
    guardar_config(sesion, limites={
        "conteos": [{"nombre": "ejercicios", "patron": r"^Ejercicio \d+", "min": 3, "max": 5}],
        "titulos_con": [{"patron": "^Mejora el resultado", "debe_contener": "(opcional)"}],
    })
    documento = Document()
    documento.add_heading("Ejercicio 1 · Clasifica", level=1)
    documento.add_heading("Mejora el resultado", level=2)
    documento.add_heading("Ejercicio 2 · Revisa", level=1)
    documento.add_heading("Mejora el resultado (opcional)", level=2)
    ruta = tmp_path / "S1_Laboratorio.docx"
    documento.save(ruta)
    hallazgos = {(h.regla, h.detalle) for h in correr(sesion, [ruta]).hallazgos}
    assert ("ejercicios", "tiene 2; deben ser de 3 a 5") in hallazgos
    assert ("título", "debe contener «(opcional)»") in hallazgos
    assert len(hallazgos) == 2


def test_diapositivas_cuenta_y_notas(sesion, tmp_path):
    guardar_config(sesion, limites={"diapositivas_min": 14, "diapositivas_max": 18, "notas_en_todas": True})
    presentacion = Presentation()
    for numero in range(1, 3):
        lamina = presentacion.slides.add_slide(presentacion.slide_layouts[1])
        lamina.shapes.title.text = f"Concepto {numero}"
        if numero == 1:
            lamina.notes_slide.notes_text_frame.text = "Pregunta al grupo por un ejemplo."
    ruta = tmp_path / "S1_Diapositivas.pptx"
    presentacion.save(ruta)
    resultado = correr(sesion, [ruta])
    assert {(h.regla, h.detalle) for h in resultado.hallazgos} == {
        ("diapositivas", "tiene 2 diapositivas; deben ser de 14 a 18"),
        ("notas del profesor", "la diapositiva 2 no tiene notas"),
    }
    assert ("S1_Diapositivas.pptx · diapositiva 1", "notas", "Pregunta al grupo por un ejemplo.") in \
        [(o.seccion, o.parte, o.texto) for o in resultado.oraciones]


def test_excel_de_practica_da_una_fila_por_celda(sesion, tmp_path):
    libro = openpyxl.Workbook()
    libro.active.title = "Registro"
    libro.active["A1"] = "Fecha de referencia: lunes 5 de octubre de 2026"
    libro.active["B2"] = "Esta pregunta vale 5 puntos."
    ruta = tmp_path / "practica" / "S1_E1_registro-almacen.xlsx"
    libro.save(ruta)
    resultado = correr(sesion, [ruta])
    assert [(o.seccion, o.parte) for o in resultado.oraciones] == [
        ("S1_E1_registro-almacen.xlsx · Registro", "celda")] * 2
    assert [h.regla for h in resultado.hallazgos] == ["puntaje", "archivo sin mencionar"]


def test_html_lee_texto_visible_y_datos_de_los_casos(sesion, tmp_path):
    ruta = tmp_path / "S1_Practica.html"
    ruta.write_text(
        "<html><head><style>p{}</style><script>var x = 'Activa el cronómetro';</script>"
        '<script type="application/json" id="datos">{"casos": [{"enunciado": "Resuelve en 10 minutos."}]}</script>'
        "</head><body><h1>Estación 1</h1><p>Scrum tiene tres pilares.</p></body></html>",
        encoding="utf-8",
    )
    resultado = correr(sesion, [ruta])
    textos = [o.texto for o in resultado.oraciones]
    assert "Scrum tiene tres pilares." in textos
    assert "Resuelve en 10 minutos." in textos
    assert "Activa el cronómetro" not in " ".join(textos)  # el código no es texto del alumno
    assert [h.regla for h in resultado.hallazgos] == ["tiempo"]


def test_datos_repetidos_lista_conceptos_variantes_y_cifras(sesion, tmp_path):
    material = word(tmp_path / "S1_Laboratorio.docx", ["El Manifiesto Ágil tiene cuatro aspectos.", "Son cuatro."])
    filas = {t: (n, o) for t, n, o in correr(sesion, [material]).datos_repetidos}
    assert filas["Agile"][0] == 0
    assert filas["Manifiesto Ágil"][0] == 1
    assert filas["cuatro"][0] == 2


# ---------- Programa desde la línea de comandos ----------

def test_linea_de_comandos_imprime_resumen_y_codigo(sesion, tmp_path, capsys, monkeypatch):
    monkeypatch.setattr(verificar, "_contar_paginas_con_word", sin_word)
    material = word(tmp_path / "S1_Laboratorio.docx", ["Esta pregunta vale 5 puntos."])
    salida = tmp_path / "S1_Laboratorio_Verificacion.xlsx"
    codigo = verificar.main(["--config", str(sesion["config"]), "--salida", str(salida),
                             "--json", str(tmp_path / "v.json"), str(material)])
    ultima = capsys.readouterr().out.strip().splitlines()[-1]
    assert codigo == 1
    assert ultima == "Oraciones: 3 · Fallas: 1 · Avisos: 0"
    assert salida.exists()
    datos = json.loads((tmp_path / "v.json").read_text(encoding="utf-8"))
    assert len(datos["oraciones"][0]["huella"]) == 16


def test_linea_de_comandos_sin_fallas_termina_en_cero(sesion, tmp_path, capsys):
    con_archivo_de_practica(sesion)
    material = word(tmp_path / "S1_Laboratorio.docx", ["Adjunta el archivo S1_E1_registro-almacen.xlsx."])
    codigo = verificar.main(["--config", str(sesion["config"]), "--salida", str(tmp_path / "v.xlsx"), str(material)])
    assert codigo == 0


@pytest.mark.parametrize("contenido, mensaje", [
    ("{no es json", "No se pudo leer"),
    ('{"sesion": 1}', "Falta «material»"),
    ('{"sesion": 1, "material": "L", "fuentes": [{"nombre": "X", "texto": "no-existe.jsonl"}]}', "No existe el texto"),
    ('{"sesion": 1, "material": "L", "datos_fijos": [{"etiqueta": "Fecha"}]}', "«etiqueta» y «valor»"),
])
def test_configuracion_con_errores_termina_en_dos(tmp_path, capsys, contenido, mensaje):
    config = tmp_path / "verificacion.json"
    config.write_text(contenido, encoding="utf-8")
    material = word(tmp_path / "S1_Lectura.docx", ["Scrum tiene tres pilares."])
    codigo = verificar.main(["--config", str(config), "--salida", str(tmp_path / "v.xlsx"), str(material)])
    assert codigo == 2
    assert mensaje in capsys.readouterr().out


def test_word_real_cuenta_las_paginas(tmp_path):
    documento = Document()
    documento.add_paragraph("Página uno.")
    documento.add_page_break()
    documento.add_paragraph("Página dos.")
    ruta = tmp_path / "dos.docx"
    documento.save(ruta)
    assert verificar._contar_paginas_con_word(ruta) == 2


def test_ejemplo_listo_para_probar_marca_sus_seis_errores(tmp_path):
    ejemplo = RAIZ / "tests" / "fixtures" / "verificador_ejemplo"
    resultado = verificar.verificar(ejemplo / "verificacion.json", [ejemplo / "S1_Lectura.docx"], contar_paginas=sin_word)
    assert [h.regla for h in resultado.hallazgos] == [
        "cita sin fuente", "tiempo", "dato fijo", "variante de un concepto", "operación", "lenguaje de IA"]
    assert resultado.oraciones[3].fuente == "Guía de Scrum, lámina 54: Tres Pilares de Scrum • Transparencia • Inspección • Adaptación"


def test_datos_repetidos_lista_fechas_y_nombres_propios(sesion, tmp_path):
    material = word(tmp_path / "S1_Laboratorio.docx", [
        "El equipo de la Sucursal Centro empieza el lunes 5 de octubre de 2026.",
        "Desde el 5 de octubre, la Sucursal Centro usa el tablero."])
    filas = {t: n for t, n, _ in correr(sesion, [material]).datos_repetidos}
    assert filas["Sucursal Centro"] == 2
    assert filas["lunes 5 de octubre de 2026"] == 1
    assert filas["5 de octubre"] == 1


def test_unidades_con_dos_cifras_distintas():
    unidades = verificar.unidades_con_cifras([("a", "El Sprint dura dos semanas."), ("b", "El Sprint dura tres semanas."),
                                              ("c", "Hay 12 pedidos y una tienda.")])
    assert set(unidades["semanas"]) == {"dos", "tres"}
    assert "tienda" not in unidades  # «una» es artículo, no cifra
