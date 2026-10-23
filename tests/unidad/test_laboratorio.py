"""Etapa 7a: guía del laboratorio, archivos de práctica reales, figuras sacadas de ellos y el .zip."""

import asyncio
import zipfile
from urllib.parse import unquote

import openpyxl
from claude_agent_sdk import AssistantMessage, TextBlock
import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app import servidor
from app.materiales import ejecucion, figuras, laboratorio, practica
from app.materiales.contenido import Laboratorio
from tests.conftest import consulta_en_secuencia, resultado_simulado
from tests.sembrados import laboratorio as sembrados
from tests.unidad.test_lectura import SIN_ERRORES, sesion, una_pagina  # noqa: F401

cliente = TestClient(servidor.app)
LLAMADAS_A_EXCEL = []


def recalcular_sin_excel(ruta):
    """En las pruebas no se abre Excel: se escribe el valor de la fórmula del laboratorio limpio."""
    LLAMADAS_A_EXCEL.append(ruta.name)
    libro = openpyxl.load_workbook(ruta)
    for hoja in libro.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if isinstance(celda.value, str) and celda.value.startswith("=SUM"):
                    celda.value = 10
    libro.save(ruta)
    return ruta


async def sin_ejecucion(contenido, **_):
    """La ejecución tiene sus propias pruebas (abajo): aquí cada ejercicio coincide sin llamar a Claude."""
    return [ejecucion.Ejecucion(e.numero, e.titulo, respuesta_ficha="La situación 4 contradice",
                                conversacion=[ejecucion.Turno("alumno", "prompt"), ejecucion.Turno("herramienta", "ok")])
            for e in contenido.ejercicios]


def generar(carpeta, *salidas, ejecutar=sin_ejecucion):
    consulta = consulta_en_secuencia(*salidas, SIN_ERRORES)
    material = asyncio.run(laboratorio.generar(carpeta, "scrum", 1, consulta=consulta, recalcular=recalcular_sin_excel,
                                               contar_paginas=una_pagina, ejecutar=ejecutar))
    return material, consulta


# ---------- Generación completa ----------

def test_genera_la_guia_los_archivos_las_figuras_y_el_zip(sesion):  # noqa: F811
    material, consulta = generar(sesion, sembrados.limpio())
    assert material["estado"] == "verificada", material.get("problemas")
    rutas = laboratorio.archivos(sesion, 1)
    nombres = ["S1_E1_pedidos-semana.xlsx", "S1_E2_aviso-almacen.docx", "S1_E3_capacitacion.xlsx"]
    assert sorted(p.name for p in rutas["practica"].iterdir()) == nombres
    assert sorted(zipfile.ZipFile(rutas["zip"]).namelist()) == nombres
    assert "S1_E3_capacitacion.xlsx" in LLAMADAS_A_EXCEL            # tiene una fórmula: Excel la calcula
    guia = Document(rutas["word"])
    textos = [p.text for p in guia.paragraphs]
    assert laboratorio.AVISO_ADJUNTOS in textos                       # la ficha no dice que acepta adjuntos
    tabla = guia.tables[0]
    assert [c.text for c in tabla.rows[0].cells] == ["Archivo", "Ejercicio", "Contenido"]
    assert [tabla.cell(i, 0).text for i in range(1, 4)] == nombres
    titulos = [p.text for p in guia.paragraphs if p.style.name.startswith("Heading")]
    assert titulos[:3] == ["Antes de empezar", "Archivos de práctica", "Ejercicio 1 · Clasifica los pedidos de la semana"]
    assert not any(t.startswith("Mejora el resultado") for t in titulos)   # decisión 66 b
    assert "Cada ejercicio funciona solo. Puedes hacerlos en cualquier orden." in textos
    assert "Los 3 archivos están en el archivo comprimido S1_Laboratorio_Archivos.zip." in textos
    assert "Área: Ventas" in textos and "Situación" in textos
    assert "S1_E1_pedidos-semana.xlsx. Tiene una hoja: Pedidos." in textos
    assert "2.	Revisa la columna Entrega de la hoja Pedidos. Mira la figura 1." in textos
    assert "Figura 1. La columna Entrega de la hoja Pedidos." in textos
    assert "Figura 2. La oración del aviso que habla del Sprint." in textos
    # Ejercicio 3: fórmulas en español en la guía, recuadro «Antes de copiar una fórmula» con su tabla.
    tablas = [[c.text for c in fila.cells] for tabla in guia.tables for fila in tabla.rows]
    assert ["Función en español", "Función en inglés"] in tablas and ["SI", "IF"] in tablas and ["PROMEDIO", "AVERAGE"] in tablas
    celdas = " ".join(c.text for tabla in guia.tables for fila in tabla.rows for c in fila.cells)
    assert '=SI(B2<5,"REVISAR","OK")' in celdas and "Fórmula 2 · Promedio de horas" in celdas
    pie = guia.sections[0].footer.paragraphs[0].text
    assert pie.startswith("Sesión 1 · Laboratorio · Página")
    assert len(guia.inline_shapes) == 3
    redaccion = consulta.llamadas[0]["prompt"]
    assert "## Material 3 · Laboratorio y archivos de práctica" in redaccion
    assert "Mira la figura N" in redaccion and laboratorio.AVISO_ADJUNTOS in redaccion
    entrega = {s["clave"]: s["lineas"] for s in material["entrega"]["secciones"]}
    assert entrega["archivos"][1].startswith("S1_Laboratorio_Archivos.zip: los 3 archivos de práctica")
    assert entrega["no_pude"][0].startswith("Ejecuté los prompts con Claude")
    assert any(l.startswith("Ejecuté 3 ejercicios como un alumno") for l in entrega["probe"])
    assert entrega["archivos"][-1].startswith("S1_Laboratorio_Ejecucion.md")
    assert rutas["registro"].exists()
    assert any("Problema sembrado en S1_E1_pedidos-semana.xlsx" in l for l in entrega["decidi"])


def test_los_archivos_de_practica_son_reales_y_sobrios(sesion):  # noqa: F811
    generar(sesion, sembrados.limpio())
    carpeta = laboratorio.archivos(sesion, 1)["practica"]
    libro = openpyxl.load_workbook(carpeta / "S1_E1_pedidos-semana.xlsx")
    hoja = libro["Pedidos"]
    assert [c.value for c in hoja[1]] == ["Pedido", "Cliente", "Entrega"] and hoja["C4"].value == "Envío"
    assert hoja["A1"].font.bold and hoja.freeze_panes == "A2"
    horas = openpyxl.load_workbook(carpeta / "S1_E3_capacitacion.xlsx")["Horas"]
    assert horas["B2"].value == 4                                     # un número queda como número
    aviso = Document(carpeta / "S1_E2_aviso-almacen.docx")
    assert [p.text for p in aviso.paragraphs][:2] == ["Aviso al equipo", "El Sprint dura dos semanas."]


def test_las_figuras_salen_del_archivo_real_con_marcas_ambar(sesion):  # noqa: F811
    generar(sesion, sembrados.limpio())
    rutas = laboratorio.archivos(sesion, 1)
    imagen = Image.open(rutas["figuras"] / "E1_paso2.png").convert("RGB")
    colores = {c for _, c in imagen.getcolors(maxcolors=1 << 20)}
    assert figuras.AMBAR in colores
    contenido = Laboratorio.model_validate(sembrados.limpio())
    figura = contenido.ejercicios[0].pasos[1].figura
    otra = figuras.hacer(rutas["practica"] / figura.archivo, figura, rutas["figuras"] / "prueba.png")
    assert otra.letra_pt >= figuras.LETRA_MINIMA_PT and otra.ancho_cm <= figuras.ANCHO_MAXIMO_CM
    word = Image.open(rutas["figuras"] / "E2_paso2.png").convert("RGB")
    assert figuras.AMBAR in {c for _, c in word.getcolors(maxcolors=1 << 20)}


def test_una_formula_en_espanol_se_guarda_en_ingles_dentro_del_archivo(tmp_path):
    from app.materiales import estilo
    from app.materiales.contenido import ArchivoPractica
    archivo = ArchivoPractica.model_validate({
        "nombre": "S1_E1_horas.xlsx", "descripcion": "Horas.", "parrafos": [], "tabla": None, "problema": "",
        "hojas": [{"nombre": "Horas", "encabezados": ["Persona", "Horas", "Revisión"],
                   "filas": [["Analista 1", 4, '=SI(Y(B2<5;B2>0);"REVISAR";"OK")'], ["Total", "=SUMA(B2:B2)", None]]}]})
    identidad = estilo.Identidad(**estilo.PALETA_SOBRIA, tipografia="Calibri", logo=None)
    hoja = openpyxl.load_workbook(practica.crear_excel(archivo, identidad, tmp_path / archivo.nombre))["Horas"]
    assert hoja["C2"].value == '=IF(AND(B2<5,B2>0),"REVISAR","OK")'
    assert hoja["B3"].value == "=SUM(B2:B2)"
    assert practica.funciones_usadas(['=SI(Y(A1>2,B1),1,0)', "=PROMEDIO(B2:B3)"]) == [
        ("SI", "IF"), ("Y", "AND"), ("PROMEDIO", "AVERAGE")]


# ---------- Errores sembrados: los detecta un programa antes de crear los archivos ----------

@pytest.mark.parametrize("nombre", list(sembrados.SEMBRADOS))
def test_cada_error_sembrado_se_detecta(sesion, nombre):  # noqa: F811
    cambio, esperado = sembrados.SEMBRADOS[nombre]
    contenido, problemas = laboratorio._validar(sembrados.con(cambio), sesion, 1)
    assert contenido is None
    assert any(esperado in p for p in problemas), problemas


def test_el_laboratorio_limpio_no_tiene_problemas(sesion):  # noqa: F811
    contenido, problemas = laboratorio._validar(sembrados.limpio(), sesion, 1)
    assert problemas == [] and contenido is not None


def test_la_respuesta_a_la_pregunta_de_la_herramienta_va_en_el_ejercicio():
    contenido = Laboratorio.model_validate(sembrados.limpio())
    problemas = practica.revisar(contenido, 1, 3, respuestas_a_la_herramienta={2: "Sí, usa el aviso de ayer"})
    assert any("ejercicio 2: la ficha dice que la herramienta pregunta algo" in p for p in problemas)


def test_un_problema_de_un_ejercicio_antes_de_crear_los_archivos_se_corrige_solo_en_ese_ejercicio(sesion):  # noqa: F811
    roto = sembrados.con(sembrados.SEMBRADOS["espacio por llenar"][0])
    corregido = {"ejercicios": [sembrados.limpio()["ejercicios"][0]], "archivos": [], "pasajes": [],
                 "explicaciones": [], "datos_nuevos": []}
    material, consulta = generar(sesion, roto, corregido)
    assert material["estado"] == "verificada"
    llamada = consulta.llamadas[1]
    assert "«[nombre del cliente]» en lo que el alumno copia" in llamada["prompt"]
    assert "Los problemas son solo de los ejercicios 1" in llamada["prompt"]
    assert llamada["options"].output_format["schema"] == laboratorio.ESQUEMA_CORRECCIONES
    assert "Llena solo «por_ejercicios»." in llamada["prompt"]
    assert llamada["options"].model == "claude-sonnet-5-5"


def test_un_error_del_esquema_se_corrige_con_el_laboratorio_completo(sesion):  # noqa: F811
    roto = sembrados.limpio()
    del roto["ejercicios"][0]["pasos"]
    corregido = {"laboratorio": sembrados.limpio(), "explicaciones": [], "datos_nuevos": []}
    material, consulta = generar(sesion, roto, corregido)
    assert material["estado"] == "verificada"
    assert consulta.llamadas[1]["options"].output_format["schema"] == laboratorio.ESQUEMA_CORRECCIONES
    assert "Llena solo «completa»." in consulta.llamadas[1]["prompt"]


# ---------- Página ----------

def test_la_pagina_de_la_sesion_genera_y_aprueba_el_laboratorio(sesion, monkeypatch):  # noqa: F811
    pagina = cliente.get("/cursos/scrum/sesiones/1/materiales/laboratorio").text
    assert '<h1 id="t-laboratorio">Laboratorio y archivos de práctica</h1>' in pagina and "Generar laboratorio" in pagina
    generar(sesion, sembrados.limpio())
    pagina = cliente.get("/cursos/scrum/sesiones/1/materiales/laboratorio").text
    assert "S1_Laboratorio.docx" in pagina and "S1_Laboratorio_Archivos.zip" in pagina
    r = cliente.get("/cursos/scrum/sesiones/1/descargar/S1_Laboratorio_Archivos.zip")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    r = cliente.post("/cursos/scrum/sesiones/1/materiales/laboratorio/aprobar", follow_redirects=False)
    assert "Laboratorio aprobado." in unquote(r.headers["location"])
    assert laboratorio.estado(sesion, 1)["estado"] == "aprobada"


def test_no_aplica_en_la_ficha_no_exige_una_respuesta(sesion):  # noqa: F811
    from app.fichas import almacen
    ficha = almacen.cargar(sesion, "sesion", 1)
    ficha["campos"]["ejercicio.1.pregunta_herramienta"] = {
        "valor": "No aplica: la herramienta no debe preguntar nada. (propuesto)", "origen": "x", "estado": ""}
    ficha["campos"]["ejercicio.2.pregunta_herramienta"] = {"valor": "Sí, usa el aviso de ayer", "origen": "x", "estado": ""}
    almacen.guardar(sesion, "sesion", 1, ficha)
    assert laboratorio.respuestas_a_la_herramienta(sesion, 1) == {2: "Sí, usa el aviso de ayer"}


def test_el_pedido_pide_la_figura_donde_se_ubica_algo_y_sin_comillas():
    texto = laboratorio.instrucciones_de_formato(1, "")
    assert "Ponla en el paso que pide ubicar algo en el archivo" in texto
    assert "la figura es obligatoria" in texto
    assert "va en **negrita**, sin comillas" in texto


# ---------- Ejecución de cada ejercicio como un alumno (PLAN.md §0, decisión 35) ----------

VALORES_DE_EXCEL = {"C2": "REVISAR", "C3": "OK", "B5": 5}


def excel_simulado(valores=VALORES_DE_EXCEL):
    """Escribe el valor que daría Excel en cada celda con fórmula que copió el programa."""
    def recalcular(ruta):
        libro = openpyxl.load_workbook(ruta)
        for hoja in libro.worksheets:
            for fila in hoja.iter_rows():
                for celda in fila:
                    if isinstance(celda.value, str) and celda.value.startswith("="):
                        celda.value = valores.get(celda.coordinate, 10)
        libro.save(ruta)
        return ruta
    return recalcular


def herramienta_y_comparacion(diferencias=()):
    """La herramienta responde con texto; la comparación devuelve las diferencias que se le pasan."""
    llamadas = []
    pendientes = list(diferencias)

    async def consulta(*, prompt, options):
        llamadas.append({"prompt": prompt, "options": options})
        if options.output_format is None:
            yield AssistantMessage(content=[TextBlock(text="Pedido P-103: Envío. La situación 4 contradice.")],
                                   model="claude-sonnet-5-5")
            yield resultado_simulado()
        else:
            yield resultado_simulado(structured_output={"diferencias": pendientes.pop(0) if pendientes else []})

    consulta.llamadas = llamadas
    return consulta


def ejecutar_limpio(sesion, consulta, valores=VALORES_DE_EXCEL):
    generar(sesion, sembrados.limpio())                       # crea los archivos de práctica reales
    rutas = laboratorio.archivos(sesion, 1)
    contenido = Laboratorio.model_validate(sembrados.limpio())
    return asyncio.run(ejecucion.ejecutar(
        contenido, carpeta_practica=rutas["practica"], carpeta=rutas["ejecucion"],
        respuestas_ficha={1: "La situación 4 contradice", 2: "La situación 4 contradice", 3: "C2 dice REVISAR"},
        respuestas_alumno={}, recalcular=excel_simulado(valores),
        llamada={"curso": "scrum", "sesion": "S1", "material": "laboratorio", "consulta": consulta}))


def test_las_formulas_se_copian_donde_dice_el_paso_y_las_calcula_excel(sesion):  # noqa: F811
    resultados = ejecutar_limpio(sesion, herramienta_y_comparacion())
    tercero = resultados[2]
    celdas = {c["celda"]: c for c in tercero.celdas}
    assert celdas["C2"]["formula"] == '=IF(B2<5,"REVISAR","OK")'       # en inglés dentro del archivo
    assert celdas["C3"]["formula"] == '=IF(B3<5,"REVISAR","OK")'       # «hacia abajo, hasta la fila 3»
    assert celdas["B5"]["formula"] == "=AVERAGE(B2:B3)" and celdas["B5"]["valor"] == 5   # la fórmula 2
    copia = openpyxl.load_workbook(laboratorio.archivos(sesion, 1)["ejecucion"] / "E3" / "S1_E3_capacitacion.xlsx")
    assert copia["Horas"]["A5"].value == "Promedio"                      # «Escribe **Promedio** en la celda A5»
    assert tercero.conversacion == [] and tercero.coincide


def test_cada_prompt_va_en_una_conversacion_nueva_con_el_texto_del_archivo(sesion):  # noqa: F811
    consulta = herramienta_y_comparacion()
    resultados = ejecutar_limpio(sesion, consulta)
    conversaciones = [l for l in consulta.llamadas if l["options"].output_format is None]
    assert len(conversaciones) == 2                                       # ejercicios 1 y 2; el 3 solo tiene fórmulas
    primero = conversaciones[0]["prompt"]
    assert "[Archivo adjunto: S1_E1_pedidos-semana.xlsx]" in primero and "P-103 | Cliente C | Envío" in primero
    assert primero.endswith("Restricciones: " + sembrados.NOTA)
    assert conversaciones[0]["options"].model == "claude-sonnet-5-5" and conversaciones[0]["options"].tools == []
    assert all(r.coincide for r in resultados)


def test_una_formula_con_error_es_una_falla_de_ejecucion(sesion):  # noqa: F811
    resultados = ejecutar_limpio(sesion, herramienta_y_comparacion(), valores={**VALORES_DE_EXCEL, "C2": "#NAME?"})
    assert any("la fórmula de la celda C2 (hoja Horas) da #NAME?" in f for f in resultados[2].fallas)
    assert not resultados[2].coincide


def test_una_diferencia_sin_prueba_se_descarta(sesion):  # noqa: F811
    consulta = herramienta_y_comparacion([
        [{"esperado": "La situación 4 contradice", "salio": "", "explicacion": "No marcó la situación 4."},
         {"esperado": "algo que la ficha no dice", "salio": "", "explicacion": "Inventada."}],
    ])
    resultados = ejecutar_limpio(sesion, consulta)
    assert [d.esperado for d in resultados[0].diferencias] == ["La situación 4 contradice"]
    assert resultados[0].descartadas == 1
    assert resultados[0].pendientes() == [
        "Ejecución · ejercicio 1: la ficha dice «La situación 4 contradice» y eso no salió. No marcó la situación 4."]


def test_una_diferencia_con_la_ficha_queda_como_pendiente_en_la_entrega(sesion):  # noqa: F811
    async def con_diferencia(contenido, **_):
        resultados = await sin_ejecucion(contenido)
        resultados[1].diferencias = [ejecucion.Diferencia("La situación 4 contradice", "", "No salió.")]
        return resultados

    material, _ = generar(sesion, sembrados.limpio(), ejecutar=con_diferencia)
    assert material["estado"] == "con pendientes"
    assert any(p.startswith("Ejecución · ejercicio 2") for p in material["problemas"])
    assert material["ejecucion"]["coinciden"] == [1, 3]
    registro = laboratorio.archivos(sesion, 1)["registro"].read_text(encoding="utf-8")
    assert "Resultado: no coincide." in registro and "## Ejercicio 2" in registro


def test_un_paso_que_no_dice_donde_va_la_formula_se_detecta_antes():
    datos = sembrados.limpio()
    datos["ejercicios"][2]["pasos"][1]["texto"] = "Copia la fórmula 1."
    datos["ejercicios"][2]["pasos"][2]["texto"] = "Revisa el resultado."
    problemas = practica.revisar(Laboratorio.model_validate(datos), 1, 3)
    assert any("ningún paso dice dónde se copia la fórmula 1" in p for p in problemas)


# ---------- Menos costo: corrección por ejercicio y páginas de referencia en el pedido ----------

def test_los_problemas_se_asignan_a_su_ejercicio():
    nombrados = laboratorio.ejercicios_nombrados
    assert nombrados(["S1_Laboratorio.docx · Ejercicio 2 · Revisa el aviso FALLA · tiempo"]) == {2}
    assert nombrados(["S1_E3_capacitacion.xlsx · Horas FALLA", "FALLA · ejecución · ejercicio 1: da #NAME?"]) == {1, 3}
    # Uno que no es de un ejercicio obliga a corregir el laboratorio completo.
    assert nombrados(["S1_Laboratorio.docx · Antes de empezar FALLA · tiempo", "FALLA · ejecución · ejercicio 1"]) is None


def test_la_correccion_por_ejercicio_no_toca_los_demas():
    datos = sembrados.limpio()
    nuevo_e2 = {**datos["ejercicios"][1], "titulo": "Revisa el aviso de Almacén", "archivos": ["S1_E2_aviso-nuevo.docx"]}
    nuevo_archivo = {**datos["archivos"][1], "nombre": "S1_E2_aviso-nuevo.docx"}
    respuesta = {"ejercicios": [nuevo_e2, {**datos["ejercicios"][0], "titulo": "No pedido: se ignora"}],
                 "archivos": [nuevo_archivo], "pasajes": [], "explicaciones": [], "datos_nuevos": []}
    nuevo, no_volvieron = laboratorio.aplicar_por_ejercicios(datos, respuesta, {2})
    assert no_volvieron == []
    assert nuevo["ejercicios"][1]["titulo"] == "Revisa el aviso de Almacén"
    assert nuevo["ejercicios"][0] == datos["ejercicios"][0] and nuevo["ejercicios"][2] == datos["ejercicios"][2]
    assert [a["nombre"] for a in nuevo["archivos"]] == ["S1_E1_pedidos-semana.xlsx", "S1_E3_capacitacion.xlsx",
                                                        "S1_E2_aviso-nuevo.docx"]   # el archivo viejo del E2 se quita


def test_una_falla_en_el_archivo_de_un_ejercicio_corrige_solo_ese_ejercicio(sesion):  # noqa: F811
    con_falla = sembrados.limpio()
    con_falla["archivos"][1]["parrafos"][2]["texto"] = "Los pedidos se recogen en tienda en 10 minutos."
    correccion = {"ejercicios": [sembrados.limpio()["ejercicios"][1]], "archivos": [sembrados.limpio()["archivos"][1]],
                  "pasajes": [], "explicaciones": [], "datos_nuevos": []}
    material, consulta = generar(sesion, con_falla, correccion)
    assert material["estado"] == "verificada", material.get("problemas")
    llamada = consulta.llamadas[1]
    assert llamada["options"].output_format["schema"] == laboratorio.ESQUEMA_CORRECCIONES
    assert "Llena solo «por_ejercicios»." in llamada["prompt"]
    assert "Los problemas son solo de los ejercicios 2" in llamada["prompt"]


def test_con_referencia_las_paginas_van_en_el_pedido_y_la_redaccion_no_busca(sesion):  # noqa: F811
    from app.fichas import almacen
    ficha = almacen.cargar(sesion, "sesion", 1)
    almacen.aplicar_formulario(ficha, {"sesion.referencia": "silabo.pdf: página 1"})
    ficha["confirmada"] = True
    almacen.guardar(sesion, "sesion", 1, ficha)
    _, consulta = generar(sesion, sembrados.limpio())
    redaccion = consulta.llamadas[0]
    assert "[silabo.pdf · página 1] Scrum Master con IA" in redaccion["prompt"]
    assert "busca con Grep" not in redaccion["prompt"]
    assert redaccion["options"].tools == [] and not redaccion["options"].allowed_tools


# ---------- Herramienta de los ejercicios (PLAN.md §0, decisión 36) ----------

def poner_herramienta(carpeta, valor):
    """Cambia «Herramienta de los ejercicios» en la ficha del curso sin tocar su confirmación."""
    from app.fichas import almacen
    ficha = almacen.cargar(carpeta, "curso")
    ficha["campos"]["herramientas.ejercicios"] = {"valor": valor, "origen": "profesor", "estado": ""}
    almacen.guardar(carpeta, "curso", None, ficha)


@pytest.mark.parametrize("valor, esperado", [
    ("", set()), ("ninguna", set()), ("Excel", {"excel"}), ("IA", {"ia"}), ("Excel e IA", {"excel", "ia"}),
    ("Ninguna: el alumno trabaja a mano", set()),
])
def test_la_herramienta_sale_de_la_ficha_del_curso(sesion, valor, esperado):  # noqa: F811
    poner_herramienta(sesion, valor)
    assert laboratorio.herramienta(sesion) == esperado


def test_un_curso_sin_herramienta_no_pide_prompts_ni_formulas():
    texto = laboratorio.instrucciones_de_formato(1, "", frozenset())
    assert "Herramienta de los ejercicios, según la ficha del curso: ninguna" in texto
    assert "En este curso no hay prompts ni fórmulas" in texto
    assert "«prompt»:" not in texto and "«formula»:" not in texto and "adjunta el archivo" not in texto
    assert "Antes de copiar una fórmula" not in texto
    assert "lo hace Claude a mano" in texto
    completo = laboratorio.instrucciones_de_formato(1, "", frozenset({"excel", "ia"}))
    assert "«prompt»:" in completo and "«formula»:" in completo and "adjunta el archivo" in completo
    assert "En este curso no hay" not in completo


def test_sin_ia_no_hay_aviso_de_adjuntos_ni_respuesta_a_la_herramienta(sesion):  # noqa: F811
    from app.fichas import almacen
    ficha = almacen.cargar(sesion, "sesion", 1)
    ficha["campos"]["ejercicio.2.pregunta_herramienta"] = {"valor": "Sí, usa el aviso de ayer", "origen": "x", "estado": ""}
    almacen.guardar(sesion, "sesion", 1, ficha)
    poner_herramienta(sesion, "ninguna")
    assert laboratorio.respuestas_a_la_herramienta(sesion, 1) == {}
    assert laboratorio.aviso_de_adjuntos(sesion) == ""
    pagina = cliente.get("/cursos/scrum/sesiones/1/ficha").text
    assert 'name="ejercicio.2.pregunta_herramienta"' not in pagina and 'name="ejercicio.2.respuesta"' in pagina
    poner_herramienta(sesion, "IA")
    assert 'name="ejercicio.2.pregunta_herramienta"' in cliente.get("/cursos/scrum/sesiones/1/ficha").text


def test_el_laboratorio_a_mano_no_tiene_problemas_en_un_curso_sin_herramienta(sesion):  # noqa: F811
    poner_herramienta(sesion, "ninguna")
    contenido, problemas = laboratorio._validar(sembrados.a_mano(), sesion, 1)
    assert problemas == [] and contenido is not None
    contenido, problemas = laboratorio._validar(sembrados.limpio(), sesion, 1)   # con prompts y fórmulas
    assert any("trae un recuadro de prompt y el curso no usa IA" in p for p in problemas)
    assert any("trae un recuadro de fórmula y el curso no usa Excel" in p for p in problemas)


@pytest.mark.parametrize("nombre", list(sembrados.SEMBRADOS_SIN_HERRAMIENTA))
def test_cada_error_de_herramienta_sembrado_se_detecta(sesion, nombre):  # noqa: F811
    poner_herramienta(sesion, "ninguna")
    cambio, esperado = sembrados.SEMBRADOS_SIN_HERRAMIENTA[nombre]
    contenido, problemas = laboratorio._validar(sembrados.con_a_mano(cambio), sesion, 1)
    assert contenido is None
    assert any(esperado in p for p in problemas), problemas


def test_en_un_curso_de_excel_vale_la_formula_y_no_el_prompt(sesion):  # noqa: F811
    poner_herramienta(sesion, "Excel")
    _, problemas = laboratorio._validar(sembrados.limpio(), sesion, 1)
    assert any("ejercicio 1: trae un recuadro de prompt" in p for p in problemas)
    assert not any("recuadro de fórmula" in p for p in problemas)


def test_un_ejercicio_a_mano_lo_hace_claude_como_alumno(sesion):  # noqa: F811
    poner_herramienta(sesion, "ninguna")
    generar(sesion, sembrados.a_mano())                       # crea los archivos de práctica reales
    rutas = laboratorio.archivos(sesion, 1)
    consulta = herramienta_y_comparacion()
    resultados = asyncio.run(ejecucion.ejecutar(
        Laboratorio.model_validate(sembrados.a_mano()), carpeta_practica=rutas["practica"], carpeta=rutas["ejecucion"],
        respuestas_ficha={1: "La situación 4 contradice", 2: "La situación 4 contradice", 3: "La situación 4 contradice"},
        respuestas_alumno={}, recalcular=excel_simulado(),
        llamada={"curso": "scrum", "sesion": "S1", "material": "laboratorio", "consulta": consulta}))
    conversaciones = [l for l in consulta.llamadas if l["options"].output_format is None]
    assert len(conversaciones) == 3                                      # uno por ejercicio, todos a mano
    primero = conversaciones[0]["prompt"]
    assert primero.startswith("Eres un alumno de un curso. Haz este ejercicio de laboratorio a mano")
    assert "[Archivo de práctica: S1_E1_pedidos-semana.xlsx]" in primero and "P-103 | Cliente C | Envío" in primero
    assert "3. Escribe Recoge o Envía en la columna Revisión de cada pedido." in primero
    assert "La situación 4 contradice" not in primero                    # Claude no ve la respuesta de la ficha
    assert conversaciones[0]["options"].tools == []
    assert all(r.a_mano and r.coincide and not r.no_pude for r in resultados)
    assert resultados[0].salida() == "Pedido P-103: Envío. La situación 4 contradice."
    registro = ejecucion.a_markdown(resultados, "Ejecución", "claude-sonnet-5-5")
    assert "**Lo que hizo Claude a mano:**" in registro


def test_la_entrega_de_un_laboratorio_a_mano_no_habla_de_prompts(sesion):  # noqa: F811
    poner_herramienta(sesion, "ninguna")

    async def a_mano(contenido, **_):
        return [ejecucion.Ejecucion(e.numero, e.titulo, respuesta_ficha="La situación 4 contradice", a_mano=True,
                                    conversacion=[ejecucion.Turno("guía", "pasos"), ejecucion.Turno("resultado", "ok")])
                for e in contenido.ejercicios]

    material, consulta = generar(sesion, sembrados.a_mano(), ejecutar=a_mano)
    assert material["estado"] == "verificada", material.get("problemas")
    entrega = {s["clave"]: s["lineas"] for s in material["entrega"]["secciones"]}
    assert "Ejecuté 3 ejercicios como un alumno: 3 ejercicios hechos a mano por Claude como alumno." in entrega["probe"]
    assert not any("prompt" in l for l in entrega["no_pude"])
    redaccion = consulta.llamadas[0]["prompt"]
    assert "En este curso no hay prompts ni fórmulas" in redaccion and laboratorio.AVISO_ADJUNTOS not in redaccion
    textos = [p.text for p in Document(laboratorio.archivos(sesion, 1)["word"]).paragraphs]
    assert laboratorio.AVISO_ADJUNTOS not in textos


def test_el_laboratorio_se_reanuda_desde_lo_guardado(sesion):  # noqa: F811
    generar(sesion, sembrados.limpio())
    laboratorio._actualizar(sesion, 1, estado="error", error="se cortó")
    assert laboratorio.se_puede_reanudar(sesion, 1)
    consulta = consulta_en_secuencia(SIN_ERRORES)
    material = asyncio.run(laboratorio.reanudar(sesion, "scrum", 1, consulta=consulta, recalcular=recalcular_sin_excel,
                                                contar_paginas=una_pagina, ejecutar=sin_ejecucion))
    assert material["estado"] in ("verificada", "con pendientes"), material.get("problemas")
    assert len(consulta.llamadas) == 1
    assert any("Reanudo desde lo generado el" in a for a in material["avance"])
    pagina = TestClient(servidor.app).get("/cursos/scrum/sesiones/1/materiales/laboratorio").text
    assert "Reanudar desde lo guardado" not in pagina        # ya terminó bien


# ---------- Listas de opciones en los Excel de práctica (PLAN.md §0, decisión 60) ----------

def _identidad():
    from app.materiales import estilo
    return estilo.Identidad(**estilo.PALETA_SOBRIA, tipografia="Calibri", logo=None)


def _archivo_con_listas():
    return Laboratorio.model_validate(sembrados.con_listas()).archivos[0]


def test_el_excel_trae_la_hoja_opciones_y_la_lista_en_su_columna(tmp_path):
    archivo = _archivo_con_listas()
    libro = openpyxl.load_workbook(practica.crear_excel(archivo, _identidad(), tmp_path / archivo.nombre))
    assert libro.sheetnames == ["Pedidos", practica.HOJA_DE_OPCIONES]
    opciones = libro[practica.HOJA_DE_OPCIONES]
    assert opciones.sheet_state == "visible"
    assert [c.value for c in opciones["A"] if c.value is not None] == ["¿Recoge en tienda?", "Sí", "No"]
    assert [c.value for c in opciones["B"]] == ["Concepto", *sembrados.CONCEPTOS]
    validaciones = {str(v.sqref): v for v in libro["Pedidos"].data_validations.dataValidation}
    assert set(validaciones) == {"D2:D5", "E2:E5"}
    si_no, conceptos = validaciones["D2:D5"], validaciones["E2:E5"]
    assert si_no.type == "list" and si_no.formula1 == "'Opciones'!$A$2:$A$3"
    assert conceptos.formula1 == "'Opciones'!$B$2:$B$4"
    assert si_no.allow_blank and si_no.showErrorMessage and si_no.errorStyle == "stop"
    assert si_no.error == practica.MENSAJE_DE_LA_LISTA == "Elige una opción de la lista."
    assert all(libro["Pedidos"].cell(row=f, column=4).value is None for f in range(2, 6))


def test_un_excel_sin_listas_no_tiene_validaciones_ni_hoja_opciones(tmp_path):
    archivo = Laboratorio.model_validate(sembrados.limpio()).archivos[0]
    libro = openpyxl.load_workbook(practica.crear_excel(archivo, _identidad(), tmp_path / archivo.nombre))
    assert libro.sheetnames == ["Pedidos"]
    assert not libro["Pedidos"].data_validations.dataValidation


def test_las_respuestas_de_la_lista_no_llegan_al_archivo_del_alumno(tmp_path):
    archivo = _archivo_con_listas()
    texto = ejecucion.texto_del_archivo(practica.crear_excel(archivo, _identidad(), tmp_path / archivo.nombre))
    # «Opciones» trae cada opción una vez; la columna del alumno va vacía.
    assert texto.count("Pilares de Scrum") == 1 and texto.count("Sí") == 1


@pytest.mark.parametrize("cambio, esperado", [
    (lambda h: h["listas"][0].update(columna="Estado"), "no es de una columna de la hoja"),
    (lambda h: h["listas"][0].update(opciones=["Sí", "sí"]), "opciones repetidas"),
    (lambda h: h["listas"][0].update(respuestas=["Sí"]), "va una por fila"),
    (lambda h: h["filas"][0].__setitem__(3, "Sí"), "va vacía, para que el alumno elija"),
    (lambda h: h["listas"].append(dict(h["listas"][0])), "tiene dos listas"),
])
def test_el_modelo_rechaza_una_lista_mal_formada(cambio, esperado):
    datos = sembrados.con_listas()
    cambio(datos["archivos"][0]["hojas"][0])
    with pytest.raises(ValueError, match=esperado):
        Laboratorio.model_validate(datos)


def test_el_origen_de_las_opciones_sale_de_las_fichas(sesion):  # noqa: F811
    origen = laboratorio.origen_de_opciones(sesion, 1)
    assert [c for c, _ in origen.vocabulario] == sembrados.CONCEPTOS
    assert dict(origen.vocabulario)["Manifiesto Ágil"] == ("Agile", "manifiesto agile")
    assert "Sucursal Centro" in origen.texto_de_las_fichas and "Recursos Humanos" in origen.texto_de_las_fichas
    assert origen.ensena_listas is False


@pytest.mark.parametrize("tema, ensena", [
    ("3 · Validación de datos en Excel", True), ("Listas desplegables para registrar pedidos", True),
    ("Crear una lista de opciones", True), ("Pilares y valores de Scrum", False),
])
def test_reconoce_una_sesion_que_ensena_listas(tema, ensena):
    assert bool(practica.ENSENA_LISTAS.search(tema)) is ensena


@pytest.mark.parametrize("nombre", list(sembrados.SEMBRADOS_LISTAS))
def test_cada_error_sembrado_en_una_lista_se_detecta(sesion, nombre):  # noqa: F811
    cambio, esperado = sembrados.SEMBRADOS_LISTAS[nombre]
    contenido, problemas = laboratorio._validar(sembrados.con_listas_y(cambio), sesion, 1)
    assert contenido is None
    assert any(esperado in p for p in problemas), problemas
    assert all(p.startswith("FALLA") for p in problemas)


def test_una_lista_en_una_sesion_que_ensena_validacion_de_datos_es_falla(sesion):  # noqa: F811
    import dataclasses
    origen = dataclasses.replace(laboratorio.origen_de_opciones(sesion, 1), ensena_listas=True)
    contenido = Laboratorio.model_validate(sembrados.con_listas())
    problemas = practica.revisar(contenido, 1, 3, permitidos=frozenset({"prompt", "formula", "plantilla"}),
                                 origen_de_opciones=origen)
    assert any(sembrados.LISTA_EN_SESION_DE_VALIDACION in p and p.startswith("FALLA") for p in problemas), problemas


def test_un_paso_que_no_dice_que_se_elige_de_la_lista_es_un_aviso(sesion):  # noqa: F811
    cambio, esperado = sembrados.SIN_PASO_DE_LA_LISTA
    hallazgos = laboratorio.avisos_de_listas(Laboratorio.model_validate(sembrados.con_listas_y(cambio)),
                                             laboratorio.origen_de_opciones(sesion, 1))
    assert len(hallazgos) == 1
    assert hallazgos[0].nivel == "AVISO" and hallazgos[0].regla == "lista de opciones"
    assert esperado in hallazgos[0].detalle and hallazgos[0].seccion.startswith("Ejercicio 1")


def test_el_laboratorio_con_listas_bien_puestas_no_tiene_problemas(sesion):  # noqa: F811
    contenido, problemas = laboratorio._validar(sembrados.con_listas(), sesion, 1)
    assert problemas == [] and contenido is not None
    assert laboratorio.avisos_de_listas(contenido, laboratorio.origen_de_opciones(sesion, 1)) == []


def test_el_laboratorio_con_listas_se_genera_y_el_verificador_no_marca_fallas(sesion):  # noqa: F811
    material, _ = generar(sesion, sembrados.con_listas())
    assert material["estado"] == "verificada", material.get("problemas")
    ruta = laboratorio.archivos(sesion, 1)["practica"] / "S1_E1_pedidos-semana.xlsx"
    assert practica.HOJA_DE_OPCIONES in openpyxl.load_workbook(ruta).sheetnames


def test_la_revision_del_contenido_ve_la_lista_y_sabe_que_no_da_la_respuesta():
    bloques = laboratorio.bloques_de(Laboratorio.model_validate(sembrados.con_listas()))
    texto = " ".join(o for b in bloques for o in b.oraciones)
    assert "lista de opciones de la columna Concepto" in texto
    assert "Valores de Scrum" in texto and "respuestas" not in texto
    assert "no da la respuesta" in laboratorio.INDICACIONES_DE_REVISION


# ---------- Opciones con menos restricciones (PLAN.md §0, decisión 61) ----------

@pytest.mark.parametrize("nombre", list(sembrados.OPCIONES_CON_AVISO))
def test_una_opcion_que_no_aparece_escrita_es_un_aviso_y_no_una_falla(sesion, nombre):  # noqa: F811
    cambio, opcion = sembrados.OPCIONES_CON_AVISO[nombre]
    contenido, problemas = laboratorio._validar(sembrados.con_listas_y(cambio), sesion, 1)
    assert problemas == [] and contenido is not None
    avisos = laboratorio.avisos_de_listas(contenido, laboratorio.origen_de_opciones(sesion, 1))
    assert [h.nivel for h in avisos] == ["AVISO"]
    assert f"la opción «{opcion}» no aparece escrita" in avisos[0].detalle
    assert avisos[0].seccion.startswith("Ejercicio 1")


@pytest.mark.parametrize("nombre", list(sembrados.OPCIONES_SIN_AVISO))
def test_una_opcion_escrita_en_la_ficha_antes_de_empezar_o_el_archivo_vale(sesion, nombre):  # noqa: F811
    contenido, problemas = laboratorio._validar(sembrados.con_listas_y(sembrados.OPCIONES_SIN_AVISO[nombre]), sesion, 1)
    assert problemas == [] and contenido is not None
    assert laboratorio.avisos_de_listas(contenido, laboratorio.origen_de_opciones(sesion, 1)) == []


def test_una_opcion_se_busca_como_palabra_entera_y_sin_mayusculas():
    assert practica._aparece("Alto", "nivel alto de riesgo")
    assert not practica._aparece("Alto", "el altoparlante")


# ---------- El orden de las opciones solo en los procesos (PLAN.md §0, decisión 62) ----------

def _con_proceso_en_la_ficha(carpeta):
    from app.fichas import almacen
    from tests.fixtures import fichas as datos
    sesion_ = {**datos.SESION, "vocabulario.conceptos": sembrados.VOCABULARIO_CON_PROCESO}
    almacen.guardar(carpeta, "sesion", 1, datos.ficha("sesion", sesion_))


def test_la_ficha_lee_la_marca_proceso_y_la_quita_del_nombre():
    from app.fichas import almacen
    from tests.fixtures import fichas as datos
    conceptos = almacen.vocabulario(datos.ficha("sesion", {"vocabulario.conceptos": sembrados.VOCABULARIO_CON_PROCESO}))
    assert [(c["concepto"], c["proceso"]) for c in conceptos] == [
        ("Manifiesto Ágil", True), ("Pilares de Scrum", True), ("Valores de Scrum", False)]
    assert conceptos[0]["variantes"] == ["Agile", "manifiesto agile"]
    assert practica.procesos_del_vocabulario(conceptos) == [("Manifiesto Ágil", "Pilares de Scrum")]


def test_conceptos_sin_proceso_en_otro_orden_pasan(sesion):  # noqa: F811
    contenido, problemas = laboratorio._validar(sembrados.con_listas_y(sembrados.OTRO_ORDEN), sesion, 1)
    assert problemas == [] and contenido is not None


def test_pasos_de_un_proceso_de_la_ficha_en_otro_orden_son_falla(sesion):  # noqa: F811
    _con_proceso_en_la_ficha(sesion)
    contenido, problemas = laboratorio._validar(sembrados.con_listas_y(sembrados.OTRO_ORDEN), sesion, 1)
    assert contenido is None
    assert any(sembrados.FUERA_DEL_PROCESO in p and p.startswith("FALLA") for p in problemas), problemas
    contenido, problemas = laboratorio._validar(sembrados.con_listas(), sesion, 1)
    assert problemas == [], problemas


def test_los_procesos_salen_de_las_diapositivas_y_de_aplicalo_asi():
    diapositivas_ = {"diapositivas": [
        {"tipo": "proceso", "grafico": {"elementos": [{"texto": "Planificar"}, {"texto": "Revisar"}]}},
        {"tipo": "pasos", "frases": ["Abrir", "Cerrar"]},
        {"tipo": "mapa", "grafico": {"elementos": [{"texto": "A"}, {"texto": "B"}]}},
    ]}
    lectura_ = {"aplicalo": {"plantilla": ["1. Define la meta.", "2) Elige las tareas", "Nota sin número"]}}
    assert practica.procesos_del_material(lectura_, diapositivas_) == [
        ("Planificar", "Revisar"), ("Abrir", "Cerrar"), ("Define la meta", "Elige las tareas")]
    assert practica.procesos_del_material(None, None) == []


def test_pasos_de_una_diapositiva_de_proceso_aprobada_en_otro_orden_son_falla(sesion, monkeypatch):  # noqa: F811
    from app.materiales import comun, diapositivas
    rutas = diapositivas.archivos(sesion, 1)
    rutas["contenido"].parent.mkdir(parents=True, exist_ok=True)
    comun.guardar_json(rutas["contenido"], {"diapositivas": [
        {"tipo": "proceso", "grafico": {"elementos": [{"texto": c} for c in sembrados.CONCEPTOS]}}]})
    monkeypatch.setattr(diapositivas, "estado", lambda *_: {"estado": "aprobada"})
    contenido, problemas = laboratorio._validar(sembrados.con_listas_y(sembrados.OTRO_ORDEN), sesion, 1)
    assert contenido is None
    assert any("son pasos de un proceso" in p for p in problemas), problemas
    monkeypatch.setattr(diapositivas, "estado", lambda *_: {"estado": "verificada"})
    assert laboratorio._validar(sembrados.con_listas_y(sembrados.OTRO_ORDEN), sesion, 1)[1] == []
