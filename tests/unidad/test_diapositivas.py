"""Las diapositivas: PowerPoint 16:9 desde la lectura aprobada, validado en dos pasadas (SKILL.md, Material 2)."""

import asyncio
import copy
import re
import shutil
import zipfile
from pathlib import Path
from urllib.parse import unquote

import openpyxl
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Cm, Pt

from app import office, servidor
from app.fichas import almacen
from app.materiales import comun, diapositivas, estilo, lectura
from app.materiales import pptx as generador
from app.materiales.contenido import Diapositivas
from app.materiales.docx import Portada
from tests.conftest import consulta_en_secuencia
from tests.unidad.test_lectura import LIMPIA, SIN_ERRORES, error, revision, sesion, una_pagina  # noqa: F401

cliente = TestClient(servidor.app)

def diapositiva(bloque, titulo, tipo="frases", frases=None, **otros):
    return {"bloque": bloque, "titulo": titulo, "tipo": tipo, "frases": frases or [], "columnas": [],
            "tabla": None, "grafico": None, **otros}


def elemento(texto, detalle="", hijos=None, valor=None):
    return {"texto": texto, "detalle": detalle, "hijos": hijos or [], "valor": valor}


def grafico(*elementos, centro="", ejes=None):
    return {"centro": centro, "elementos": list(elementos), "ejes": ejes or []}


MANIFIESTO, PILARES = "El Manifiesto Ágil", "Pilares de Scrum"
LIMPIAS = {
    "idea_central": {"frases": ["Scrum ayuda a un equipo a trabajar con lo que observa."]},
    "diapositivas": [
        diapositiva(MANIFIESTO, "El Manifiesto Ágil", "tarjetas", frases=["El Manifiesto Ágil tiene cuatro aspectos."],
                    grafico=grafico(elemento("Personas"), elemento("Producto"))),
        diapositiva(MANIFIESTO, "Qué valora más", "tabla", frases=["Cada aspecto valora más un elemento."],
                    tabla={"titulo": "", "encabezados": ["Aspecto", "Qué valora más"],
                           "filas": [["Primero", "Las personas"], ["Segundo", "El producto que funciona"]]}),
        diapositiva(MANIFIESTO, "Quién valora qué", "organigrama",
                    grafico=grafico(elemento("Personas"), elemento("Producto"), centro="El Manifiesto Ágil")),
        diapositiva(MANIFIESTO, "Personas y producto", "columnas",
                    columnas=[{"titulo": "Primero", "frases": ["Las personas."]},
                              {"titulo": "Segundo", "frases": ["El producto que funciona."]}]),
        diapositiva(MANIFIESTO, "Para recordar", "tarjetas",
                    grafico=grafico(elemento("Personas", "Las personas."), elemento("Producto", "El producto que funciona."))),
        diapositiva(PILARES, "Pilares de Scrum", "proceso", frases=["Los pilares de Scrum son tres."],
                    grafico=grafico(elemento("Ver"), elemento("Revisar"), elemento("Ajustar"))),
        diapositiva(PILARES, "Cómo se usan", "ciclo",
                    grafico=grafico(elemento("Observa"), elemento("Decide"), elemento("Cambia"))),
        diapositiva(PILARES, "Ejemplo: el tablero", "ejemplo", frases=["El equipo muestra su tablero a todos."]),
        diapositiva(PILARES, "Todos ven el trabajo", "mapa", frases=["El equipo muestra su tablero a todos."],
                    grafico=grafico(elemento("Transparencia"), elemento("Inspección"), elemento("Adaptación"),
                                    centro="Los tres pilares")),
        diapositiva(PILARES, "Para recordar", "tarjetas", frases=["Los pilares de Scrum son tres."],
                    grafico=grafico(elemento("Transparencia"), elemento("Inspección"), elemento("Adaptación"))),
    ],
    "aplicalo": {"frases": ["Usa estas dos preguntas al cerrar cada semana."],
                 "plantilla": ["Qué observamos esta semana.", "Qué cambiamos la próxima semana."]},
    "cuidado": {"riesgos": ["Revisar el avance sin cambiar nada después."]},
    "cierre": {"frase": "Scrum ayuda a un equipo a trabajar con lo que observa."},
    "pasajes": [{"oracion": "Los pilares de Scrum son tres.", "fuente": diapositivas.FUENTE_LECTURA,
                 "ubicacion": PILARES, "texto": "Los pilares de Scrum son tres."}],
    "decisiones": [{"texto": "Dividí el ejemplo del tablero en dos diapositivas.", "tipo": "decisión"}],
}


def con(**cambios) -> dict:
    datos = copy.deepcopy(LIMPIAS)
    datos.update(cambios)
    return datos


def medir_sin_desbordes(ruta, _carpeta):
    return office.Medicion(diapositivas=len(Presentation(ruta).slides), desbordes=[], imagenes=[])


@pytest.fixture
def con_lectura(sesion):  # noqa: F811
    """Curso con las dos fichas confirmadas y la lectura generada y aprobada."""
    asyncio.run(lectura.generar(sesion, "scrum", 1, consulta=consulta_en_secuencia(LIMPIA, SIN_ERRORES),
                                contar_paginas=una_pagina))
    lectura.aprobar(sesion, 1)
    return sesion


def generar(carpeta, *salidas, medir=medir_sin_desbordes, revisar=True, disenar=False, despues=()):
    """`despues`: respuestas que llegan después de la revisión del contenido (el diseñador y la revisión visual)."""
    consulta = consulta_en_secuencia(*salidas, *([SIN_ERRORES] if revisar else []), *despues)
    material = asyncio.run(diapositivas.generar(carpeta, "scrum", 1, consulta=consulta, medir=medir, disenar=disenar))
    return material, consulta


# ---------- PowerPoint con el diseño de la skill ----------

@pytest.fixture
def presentacion(tmp_path):
    identidad = estilo.Identidad(**estilo.PALETA_SOBRIA, tipografia="Calibri", logo=None)
    ruta = generador.generar_diapositivas(Diapositivas.model_validate(LIMPIAS), identidad,
                                          Portada("Scrum Master con IA", "Sesión 1 · Fundamentos", "Diapositivas"),
                                          tmp_path / "S1_Diapositivas.pptx")
    return ruta, identidad


def test_16_9_con_portada_idea_central_conceptos_y_cierre(presentacion):
    p = Presentation(presentacion[0])
    assert round(p.slide_width / p.slide_height, 2) == round(16 / 9, 2)
    assert len(p.slides) == 15
    ultima = " ".join(f.text_frame.text for f in p.slides[-1].shapes if f.has_text_frame)
    assert "La idea que te llevas" in ultima and LIMPIAS["cierre"]["frase"] in ultima
    titulos = [s.shapes.title.text if s.shapes.title is not None else "" for s in p.slides]
    assert titulos[1] == "Idea central" and titulos[-3:-1] == ["Aplícalo así", "Cuidado con"]
    portada = " ".join(f.text_frame.text for f in p.slides[0].shapes if f.has_text_frame)
    assert "Scrum Master con IA" in portada and "Sesión 1 · Fundamentos" in portada
    assert not any(t in ("Preguntas", "Gracias") for t in titulos)


def test_sin_notas_del_profesor(presentacion):
    """Decisión 66 a: las diapositivas no llevan notas del profesor."""
    for s in Presentation(presentacion[0]).slides:
        assert not s.has_notes_slide or not s.notes_slide.notes_text_frame.text.strip()


def test_el_esquema_no_pide_notas_ni_portada():
    import json
    from app.materiales.contenido import ESQUEMA_DIAPOSITIVAS
    assert "notas" not in json.dumps(ESQUEMA_DIAPOSITIVAS)
    assert "portada" not in ESQUEMA_DIAPOSITIVAS["properties"]
    assert "- notas:" not in diapositivas.instrucciones_de_formato(True)
    assert "Sin notas del profesor." in diapositivas.instrucciones_de_formato(True)


def test_letra_legible_a_distancia(presentacion):
    tamanos = [r.font.size for s in Presentation(presentacion[0]).slides for f in s.shapes if f.has_text_frame
               for par in f.text_frame.paragraphs for r in par.runs]
    assert tamanos and min(tamanos) >= Pt(generador.TAMANO_MINIMO)


def _xml_de(ruta, titulo: str) -> str:
    """El XML de la diapositiva que tiene ese título."""
    p = Presentation(ruta)
    numero = next(i for i, s in enumerate(p.slides, start=1) if s.shapes.title is not None and s.shapes.title.text == titulo)
    return zipfile.ZipFile(ruta).read(f"ppt/slides/slide{numero}.xml").decode("utf-8")


def test_tarjetas_con_encabezado_de_color_y_tabla_del_diseno_comun(presentacion):
    """Decisión 30: en las diapositivas, tarjetas con encabezado de color, sin el borde izquierdo de los Word."""
    ruta, identidad = presentacion
    ejemplo = _xml_de(ruta, "Ejemplo: el tablero")
    assert f'<a:srgbClr val="{identidad.acento}"/>' in ejemplo and f'<a:srgbClr val="{identidad.fondo}"/>' in ejemplo
    assert ">Ejemplo<" in ejemplo
    tabla = _xml_de(ruta, "Qué valora más")
    assert generador.SIN_ESTILO_NI_CUADRICULA in tabla          # sin bordes de la plantilla
    assert f'<a:srgbClr val="{estilo.GRIS_FILA_ALTERNA}"/>' in tabla
    assert f'<a:srgbClr val="{identidad.advertencia}"/>' in _xml_de(ruta, "Cuidado con")


def test_cada_diapositiva_de_concepto_lleva_su_bloque_y_su_avance(presentacion):
    p = Presentation(presentacion[0])
    textos = [f.text_frame.text for f in p.slides[2].shapes if f.has_text_frame]
    assert MANIFIESTO.upper() in textos
    puntos = [f for f in p.slides[2].shapes if f.shape_type == 1 and f.width == f.height and f.top < Cm(1.5)]
    assert len(puntos) == 2                                     # dos bloques


def test_diseno_uniforme_solo_usa_los_colores_de_la_identidad(presentacion):
    ruta, identidad = presentacion
    archivo = zipfile.ZipFile(ruta)
    xml = "".join(archivo.read(n).decode("utf-8") for n in archivo.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n))
    usados = set(re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml))
    assert usados and usados <= identidad.colores_permitidos()


def test_un_texto_largo_baja_de_tamano_sin_pasar_del_minimo():
    corto = generador.tamano_que_cabe(["Una frase corta."], generador.CUERPO_ANCHO, generador.CUERPO_ALTO)
    largo = generador.tamano_que_cabe(["Una frase larga que se repite. " * 12] * 6, generador.CUERPO_ANCHO,
                                      generador.CUERPO_ALTO)
    assert corto == generador.TAMANO_CUERPO and largo == generador.TAMANO_MINIMO


# ---------- Estructura: cantidad y orden de la lectura ----------

def lectura_limpia():
    from app.materiales.contenido import Lectura
    return Lectura.model_validate(LIMPIA)


def test_la_cantidad_la_da_el_contenido():
    """Decisión 66 c: sin cantidad fija. Nueve conceptos (14 con las fijas) o veinte valen igual."""
    contenido, problemas = diapositivas._validar(con(diapositivas=LIMPIAS["diapositivas"][:9]), lectura_limpia())
    assert problemas == [] and contenido.total == 14
    muchas = [dict(d, titulo=f"{d['titulo']} {i}") for i in range(2) for d in LIMPIAS["diapositivas"]]
    contenido, problemas = diapositivas._validar(con(diapositivas=muchas), lectura_limpia())
    assert not any("diapositivas:" in p for p in problemas)


def test_sigue_el_orden_de_la_lectura():
    desordenadas = LIMPIAS["diapositivas"][5:] + LIMPIAS["diapositivas"][:5]
    _, problemas = diapositivas._validar(con(diapositivas=desordenadas), lectura_limpia())
    assert any("no sigue el orden de la lectura" in p for p in problemas)


def test_cada_bloque_de_la_lectura_tiene_diapositivas_y_ninguna_inventa_un_bloque():
    sin_pilares = [d for d in LIMPIAS["diapositivas"] if d["bloque"] == MANIFIESTO] * 2
    _, problemas = diapositivas._validar(con(diapositivas=sin_pilares), lectura_limpia())
    assert any("«Pilares de Scrum» de la lectura no tiene diapositivas" in p for p in problemas)
    otra = copy.deepcopy(LIMPIAS["diapositivas"])
    otra[0]["bloque"] = "Historia de Scrum"
    _, problemas = diapositivas._validar(con(diapositivas=otra), lectura_limpia())
    assert any("«Historia de Scrum» no es un bloque de la lectura" in p for p in problemas)


def test_cada_tipo_trae_lo_que_necesita():
    rota = copy.deepcopy(LIMPIAS["diapositivas"])
    rota[3]["columnas"] = rota[3]["columnas"][:1]
    contenido, problemas = diapositivas._validar(con(diapositivas=rota), lectura_limpia())
    assert contenido is None and any("dos columnas" in p for p in problemas)


# ---------- Flujo: la lectura aprobada primero ----------

def test_sin_la_lectura_aprobada_no_empieza(sesion):  # noqa: F811
    assert diapositivas.faltan(sesion, 1)["otros"] == ["Aprueba primero la lectura: las diapositivas salen de ella."]
    with pytest.raises(diapositivas.NoSePuedeEmpezar):
        asyncio.run(diapositivas.generar(sesion, "scrum", 1, consulta=consulta_en_secuencia(), medir=medir_sin_desbordes))
    pagina = cliente.get("/cursos/scrum/sesiones/1/materiales/diapositivas").text
    assert '<h1 id="t-diapositivas">Diapositivas</h1>' in pagina and "Aprueba primero la lectura" in pagina
    boton = pagina.split("Generar diapositivas")[0].rsplit("<button", 1)[1]
    assert "disabled" in boton
    r = cliente.post("/cursos/scrum/sesiones/1/materiales/diapositivas", follow_redirects=False)
    assert "Aprueba primero la lectura" in unquote(r.headers["location"])


def test_genera_verifica_y_entrega(con_lectura):
    material, consulta = generar(con_lectura, LIMPIAS)
    assert material["estado"] == "verificada", material.get("problemas")
    rutas = diapositivas.archivos(con_lectura, 1)
    assert len(Presentation(rutas["pptx"]).slides) == 15
    redaccion = consulta.llamadas[0]["prompt"]
    assert "## Material 2 · Diapositivas" in redaccion and "LECTURA APROBADA" in redaccion
    assert "Los pilares de Scrum son tres." in redaccion         # la lectura va en el pedido
    revision_ = consulta.llamadas[1]["prompt"]
    assert diapositivas.FUENTE_LECTURA in revision_ and "no agregan contenido" in revision_
    assert "[Diapositivas de «Pilares de Scrum»]" in revision_
    entrega = {s["clave"]: s["lineas"] for s in material["entrega"]["secciones"]}
    assert entrega["archivos"][0].startswith("S1_Diapositivas.pptx: las diapositivas, 16:9, 15 diapositivas")
    assert any("Medí con PowerPoint cada texto de las 15 diapositivas: ninguno se sale" in l for l in entrega["probe"])
    assert any("contra la lectura aprobada" in l for l in entrega["valide"])
    assert sorted(p.name for p in rutas["pptx"].parent.glob("S1_Diapositivas*")) == [    # decisión 42
        "S1_Diapositivas.pptx", "S1_Diapositivas_Verificacion.xlsx"]
    assert rutas["verificacion_json"].exists() and comun.TRABAJO in rutas["contenido"].parts
    assert material["lectura_fecha"] == lectura.estado(con_lectura, 1)["fecha"]


def solo(numero: int, *nuevas) -> dict:
    """Respuesta de una corrección por diapositivas que reemplaza una de concepto."""
    return {"reemplazos": [{"numero": numero, "diapositivas": list(nuevas)}],
            "fijas": {k: None for k in diapositivas.FIJAS_CORREGIBLES}, "cambios": [], "explicaciones": [],
            "datos_nuevos": []}


def test_un_desborde_se_corrige_solo_en_su_diapositiva(con_lectura):
    medidas = []

    def medir(ruta, carpeta):
        medidas.append(ruta)
        if len(medidas) == 1:
            return office.Medicion(15, [office.Desborde(8, "Los pilares de Scrum son tres.", 1.4)], [])
        return medir_sin_desbordes(ruta, carpeta)

    corta = {**LIMPIAS["diapositivas"][5], "frases": ["Scrum tiene tres pilares."]}
    material, consulta = generar(con_lectura, LIMPIAS, solo(8, corta), medir=medir)
    assert material["estado"] == "verificada"
    correccion = consulta.llamadas[1]["prompt"]
    assert "FALLA · desborde" in correccion and "Diapositiva 8 («Pilares de Scrum»)" in correccion
    assert "sobran 1.4 cm" in correccion and "Corrige solo las diapositivas 8." in correccion
    assert consulta.llamadas[1]["options"].output_format["schema"] == diapositivas.ESQUEMA_CORRECCIONES
    assert "Llena solo «por_diapositivas»." in correccion
    final = comun.leer_json(diapositivas.archivos(con_lectura, 1)["contenido"])
    assert final["diapositivas"][5]["frases"] == ["Scrum tiene tres pilares."]
    assert final["diapositivas"][:5] == LIMPIAS["diapositivas"][:5]
    hoja = openpyxl.load_workbook(diapositivas.archivos(con_lectura, 1)["excel"])["Hallazgos"]
    reglas = [(f[3].value, f[5].value) for f in hoja.iter_rows(min_row=2)]
    assert any(r == "desborde" and "Resuelto" in (c or "") for r, c in reglas)


def test_lo_que_no_esta_en_la_lectura_se_corrige_sin_tocar_el_orden(con_lectura):
    inventada = "Scrum nació en Japón en 1986."
    datos = copy.deepcopy(LIMPIAS)
    datos["diapositivas"][5]["frases"].append(inventada)
    revision_con_error = revision(error(inventada, "dato inventado", bloque=f"Diapositivas de «{PILARES}»"))
    cambios = {"cambios": [{"oracion": inventada, "nueva": ""}], "explicaciones": [], "datos_nuevos": []}
    material, consulta = generar(con_lectura, datos, revision_con_error, cambios)
    assert material["estado"] == "verificada" and material["revision"]["corregidos"] == 1
    final = comun.leer_json(diapositivas.archivos(con_lectura, 1)["contenido"])
    assert inventada not in str(final)
    assert [d["bloque"] for d in final["diapositivas"]] == [d["bloque"] for d in LIMPIAS["diapositivas"]]


def test_una_correccion_por_frases_no_cambia_el_bloque_ni_el_tipo():
    datos, no_encontradas = comun.aplicar_cambios(
        LIMPIAS, [{"oracion": "Pilares de Scrum", "nueva": "Otra cosa"}], diapositivas.CAMPOS_QUE_NO_SE_CORRIGEN)
    assert not no_encontradas
    assert [d["bloque"] for d in datos["diapositivas"]] == [d["bloque"] for d in LIMPIAS["diapositivas"]]
    assert datos["diapositivas"][5]["titulo"] == "Otra cosa"


def test_si_la_lectura_se_genera_de_nuevo_quedan_desactualizadas(con_lectura):
    generar(con_lectura, LIMPIAS)
    assert not diapositivas.estado(con_lectura, 1).get("desactualizado")
    lectura._actualizar(con_lectura, 1, fecha="2099-01-01T00:00:00")
    assert diapositivas.estado(con_lectura, 1)["desactualizado"] is True
    with pytest.raises(diapositivas.NoSePuedeAprobar):
        diapositivas.aprobar(con_lectura, 1)
    pagina = cliente.get("/cursos/scrum/sesiones/1/materiales/diapositivas").text
    assert "Desactualizadas</span>" in pagina and '<span class="paso-nota">La lectura o las fichas cambiaron</span>' in pagina


def test_se_aprueban_desde_la_pagina(con_lectura):
    generar(con_lectura, LIMPIAS)
    pagina = cliente.get("/cursos/scrum/sesiones/1/materiales/diapositivas").text
    assert "Aprobar las diapositivas" in pagina and "S1_Diapositivas.pptx" in pagina
    r = cliente.post("/cursos/scrum/sesiones/1/materiales/diapositivas/aprobar", follow_redirects=False)
    assert "Diapositivas aprobadas." in unquote(r.headers["location"])
    assert diapositivas.estado(con_lectura, 1)["estado"] == "aprobada"
    assert "Aprobadas</span>" in cliente.get("/cursos/scrum/sesiones/1/materiales/diapositivas").text


def test_detener_y_material_desconocido(con_lectura):
    r = cliente.post("/cursos/scrum/sesiones/1/materiales/diapositivas/detener", follow_redirects=False)
    assert "Las diapositivas no estaban trabajando." in unquote(r.headers["location"])
    assert cliente.post("/cursos/scrum/sesiones/1/materiales/otro").status_code == 404


def test_revisor_independiente_deja_pendientes_y_se_pueden_aprobar(con_lectura):
    from tests.unidad.test_revisor import hallazgo
    generar(con_lectura, LIMPIAS)
    frase = "El equipo muestra su tablero a todos."
    consulta = consulta_en_secuencia({"hallazgos": [hallazgo(frase)], "bloques": [], "lista": []})
    material = asyncio.run(diapositivas.revisar_con_revisor(con_lectura, "scrum", 1, consulta=consulta))
    assert material["estado"] == "con pendientes" and material["revisor"]["hallazgos"] == 1
    assert "nombre_material" not in consulta.llamadas[0]["prompt"] and "Diapositivas" in consulta.llamadas[0]["prompt"]
    pendientes = next(s["lineas"] for s in material["entrega"]["secciones"] if s["clave"] == "pendientes")
    assert any(frase in l for l in pendientes)
    diapositivas.aprobar(con_lectura, 1)
    assert diapositivas.estado(con_lectura, 1)["estado"] == "aprobada"


# ---------- Sin lectura en la sesión (PLAN.md §0, decisión 26) ----------

@pytest.fixture
def sin_lectura(entorno):
    """Curso con una fuente y una sesión que produce las diapositivas, pero no la lectura."""
    from app.fichas import almacen, flujo
    from tests.fixtures import fichas as datos
    from tests.unidad.test_fichas import con_fuente
    entorno["env"].write_text("ANTHROPIC_API_KEY=sk-ant-de-prueba\n", encoding="utf-8")
    carpeta = entorno["cursos"] / "scrum"
    (carpeta / "fuentes").mkdir(parents=True)
    (carpeta / "curso.json").write_text('{"nombre": "Scrum"}', encoding="utf-8")
    con_fuente(carpeta)
    almacen.guardar(carpeta, "curso", None, datos.ficha("curso", datos.CURSO))
    flujo.confirmar_curso(carpeta)
    almacen.guardar(carpeta, "sesion", 1, datos.ficha("sesion", {**datos.SESION, "materiales.cuales": "diapositivas"}))
    revisor_de_ejercicios = consulta_en_secuencia({"ejercicios": [
        {"numero": n, "funciona_solo": True, "sin_llenar_a_mano": True, "motivo": "Cumple"} for n in (1, 2, 3)]})
    asyncio.run(flujo.confirmar_sesion(carpeta, "scrum", 1, consulta=revisor_de_ejercicios))
    return carpeta


def test_sin_lectura_salen_de_las_fuentes_y_las_fichas(sin_lectura):
    assert diapositivas.usa_lectura(sin_lectura, 1) is False
    assert diapositivas.faltan(sin_lectura, 1) == {"ficha": [], "otros": []}
    material, consulta = generar(sin_lectura, LIMPIAS)
    assert material["estado"] == "verificada", material.get("problemas")
    redaccion = consulta.llamadas[0]["prompt"]
    assert "ESTA SESIÓN NO TIENE LECTURA" in redaccion and "LECTURA APROBADA" not in redaccion
    assert "PASAJES DE LAS FUENTES DEL CURSO" in redaccion and "Grep" not in redaccion   # decisión 68 c
    revision_ = consulta.llamadas[1]["prompt"]
    assert diapositivas.INDICACIONES_SIN_LECTURA in revision_ and diapositivas.FUENTE_LECTURA not in revision_
    entrega = {s["clave"]: s["lineas"] for s in material["entrega"]["secciones"]}
    assert any("contra las fuentes y las fichas" in l for l in entrega["valide"])
    assert any("La sesión no tiene lectura" in l and "decisión 26" in l for l in entrega["decidi"])
    assert not material.get("desactualizado") and material["lectura_fecha"] == ""


def test_sin_lectura_las_paginas_de_referencia_van_en_el_pedido_y_la_redaccion_no_busca(sin_lectura):
    from app.fichas import almacen
    ficha = almacen.cargar(sin_lectura, "sesion", 1)
    almacen.aplicar_formulario(ficha, {"sesion.referencia": "silabo.pdf: página 1"})
    ficha["confirmada"] = True
    almacen.guardar(sin_lectura, "sesion", 1, ficha)
    assert diapositivas.paginas_de_referencia(sin_lectura, 1) == [
        "[silabo.pdf · página 1] Scrum Master con IA", "[silabo.pdf · página 1] Docente: jefe de capacitación"]
    _, consulta = generar(sin_lectura, LIMPIAS)
    redaccion = consulta.llamadas[0]
    assert "PÁGINAS DEL MATERIAL DE REFERENCIA" in redaccion["prompt"]
    assert "[silabo.pdf · página 1] Scrum Master con IA" in redaccion["prompt"]
    assert "busca con Grep" not in redaccion["prompt"]
    assert redaccion["options"].tools == [] and not redaccion["options"].allowed_tools
    # Las correcciones y la revisión siguen sin las páginas: solo la redacción las necesita.
    assert all("PÁGINAS DEL MATERIAL DE REFERENCIA" not in c["prompt"] for c in consulta.llamadas[1:])


def test_una_referencia_demasiado_larga_se_recorta_y_no_se_busca_con_herramientas(sin_lectura, monkeypatch):
    """Decisión 68 c: si las páginas pasan del tope, quedan las que nombran el vocabulario o los temas."""
    from app.fichas import almacen, referencia
    ficha = almacen.cargar(sin_lectura, "sesion", 1)
    almacen.aplicar_formulario(ficha, {"sesion.referencia": "silabo.pdf", "sesion.alcance": "Docente y capacitación"})
    ficha["confirmada"] = True
    almacen.guardar(sin_lectura, "sesion", 1, ficha)
    monkeypatch.setattr(referencia, "CARACTERES_EN_EL_PEDIDO", 60)
    assert diapositivas.paginas_de_referencia(sin_lectura, 1) == ["[silabo.pdf · página 1] Docente: jefe de capacitación"]
    _, consulta = generar(sin_lectura, LIMPIAS)
    redaccion = consulta.llamadas[0]
    assert "Docente: jefe de capacitación" in redaccion["prompt"] and "Scrum Master con IA" not in redaccion["prompt"].split(
        "PÁGINAS DEL MATERIAL DE REFERENCIA")[1]
    assert redaccion["options"].tools == [] and "Grep" not in redaccion["prompt"]


def test_sin_lectura_las_diapositivas_de_un_bloque_van_juntas():
    separadas = copy.deepcopy(LIMPIAS["diapositivas"])
    separadas[2], separadas[6] = separadas[6], separadas[2]
    _, problemas = diapositivas._validar(con(diapositivas=separadas), None)
    assert any("del bloque «El Manifiesto Ágil» no van juntas" in p for p in problemas)
    contenido, problemas = diapositivas._validar(LIMPIAS, None)
    assert contenido is not None and problemas == []


def test_sin_lectura_la_ficha_pide_temas_o_referencia_y_riesgos():
    from app.fichas import almacen
    from tests.fixtures import fichas as datos
    vacia = {**datos.SESION, "materiales.cuales": "diapositivas", "lectura.bloque.1": "", "lectura.bloque.2": "",
             "lectura.cuidado": ""}
    errores = almacen.errores_del_material(datos.ficha("sesion", vacia), "diapositivas")
    assert "Sin lectura, las diapositivas necesitan de 2 a 4 bloques, o un material de referencia." in errores
    assert "«Cuidado con» necesita de 1 a 3 riesgos, uno por línea." in errores
    con_lectura_ = {**vacia, "materiales.cuales": "lectura y diapositivas"}
    assert almacen.errores_del_material(datos.ficha("sesion", con_lectura_), "diapositivas") == []


# ---------- Recursos gráficos y cierre (PLAN.md §0, decisión 29) ----------

def test_cada_recurso_trae_sus_elementos():
    rota = copy.deepcopy(LIMPIAS["diapositivas"])
    rota[4]["grafico"]["elementos"] = rota[4]["grafico"]["elementos"][:1]
    contenido, problemas = diapositivas._validar(con(diapositivas=rota), lectura_limpia())
    assert contenido is None and any("de 2 a 6 elementos" in p for p in problemas)
    sin_centro = copy.deepcopy(LIMPIAS["diapositivas"])
    sin_centro[8]["grafico"]["centro"] = ""
    _, problemas = diapositivas._validar(con(diapositivas=sin_centro), lectura_limpia())
    assert any("necesita «centro»" in p for p in problemas)


def test_la_grafica_solo_usa_cifras_de_la_lectura():
    con_grafica = copy.deepcopy(LIMPIAS["diapositivas"])
    con_grafica[0] = diapositiva(MANIFIESTO, "Cuántos aspectos", "grafica", grafico=grafico(
        elemento("Aspectos", valor=4), elemento("Pilares", valor=3), centro="Cantidad"))
    lectura_ = lectura_limpia()
    texto = diapositivas._texto_de_la_lectura(lectura_)
    _, problemas = diapositivas._validar(con(diapositivas=con_grafica), lectura_)
    faltan = [p for p in problemas if "usa cifras que no están" in p]
    assert ("4" in texto and "3" in texto) == (not faltan)
    con_grafica[0]["grafico"]["elementos"][0]["valor"] = 1986
    _, problemas = diapositivas._validar(con(diapositivas=con_grafica), lectura_)
    assert any("usa cifras que no están en la lectura: 1986" in p for p in problemas)


def test_la_frase_de_cierre_sale_de_la_idea_central():
    _, problemas = diapositivas._validar(con(cierre={"frase": "¡Gracias por su atención!"}),
                                         lectura_limpia())
    assert any("una de las frases de la idea central" in p for p in problemas)


def test_cada_recurso_se_dibuja_con_formas_y_colores_de_la_ficha(tmp_path):
    tipos = {
        "proceso": grafico(elemento("Planificar", "Se elige el trabajo."), elemento("Hacer"), elemento("Revisar")),
        "ciclo": grafico(elemento("Observar"), elemento("Decidir"), elemento("Cambiar"), centro="Mejora"),
        "organigrama": grafico(elemento("Dueño de producto", hijos=["Prioriza"]), elemento("Equipo"),
                               centro="Comité"),
        "mapa": grafico(elemento("Transparencia"), elemento("Inspección"), elemento("Adaptación"), centro="Pilares"),
        "matriz": grafico(elemento("A"), elemento("B"), elemento("C"), elemento("D"), ejes=["Impacto", "Esfuerzo"]),
        "tarjetas": grafico(elemento("Personas", "Las personas."), elemento("Producto")),
        "grafica": grafico(elemento("Aspectos", valor=4), elemento("Pilares", valor=3), centro="Cantidad"),
    }
    datos = con(diapositivas=[diapositiva(MANIFIESTO, f"Recurso {t}", t, grafico=g) for t, g in tipos.items()]
                + LIMPIAS["diapositivas"][:3])
    identidad = estilo.Identidad(**estilo.PALETA_SOBRIA, tipografia="Calibri", logo=None)
    ruta = generador.generar_diapositivas(Diapositivas.model_validate(datos), identidad,
                                          Portada("Scrum Master con IA", "Sesión 1 · Fundamentos", "Diapositivas"),
                                          tmp_path / "recursos.pptx")
    p = Presentation(ruta)
    for numero, tipo in enumerate(tipos, start=2):
        formas = p.slides[numero].shapes
        textos = " ".join(f.text_frame.text for f in formas if f.has_text_frame)
        if tipo == "grafica":
            assert any(f.has_chart for f in formas)
        else:
            assert len(formas) >= 4, tipo
            for e in tipos[tipo]["elementos"]:
                assert e["texto"] in textos, (tipo, e["texto"])
    archivo = zipfile.ZipFile(ruta)
    xml = "".join(archivo.read(n).decode("utf-8") for n in archivo.namelist()
                  if re.match(r"ppt/(slides/slide|charts/chart)\d+\.xml$", n))
    assert set(re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml)) <= identidad.colores_permitidos()
    tamanos = [r.font.size for s in p.slides for f in s.shapes if f.has_text_frame
               for par in f.text_frame.paragraphs for r in par.runs]
    assert min(tamanos) >= Pt(generador.TAMANO_MINIMO)


# ---------- Revisión visual (PLAN.md §0, decisión 30) ----------

def medir_con_imagenes(ruta, carpeta):
    """Como PowerPoint: una imagen por diapositiva, sin desbordes."""
    carpeta.mkdir(parents=True, exist_ok=True)
    total = len(Presentation(ruta).slides)
    imagenes = []
    for n in range(1, total + 1):
        imagen = carpeta / f"diapositiva_{n:02d}.png"
        imagen.write_bytes(b"\x89PNG-de-prueba")
        imagenes.append(imagen)
    return office.Medicion(diapositivas=total, desbordes=[], imagenes=imagenes)


# ---------- Corrección por diapositivas (decisión 31) ----------

def test_solo_lo_que_nombra_una_diapositiva_se_corrige_por_diapositivas():
    nombradas = diapositivas.diapositivas_nombradas
    desborde = "FALLA · desborde · S1_Diapositivas.pptx · diapositiva 8: Diapositiva 8 («X»): no cabe; sobran 1.0 cm."
    assert nombradas([desborde], LIMPIAS) == {8}
    assert nombradas(["FALLA · estructura · idea central: tiene 40 palabras; el máximo es 30."], LIMPIAS) == {2}
    assert nombradas(["FALLA · estructura · diapositivas.2.grafico: falta"], LIMPIAS) == {5}
    assert nombradas(["DISEÑO · diapositiva 4 («Y»): cargada. Propuesta: tarjetas."], LIMPIAS) == {4}
    frase = "FALLA · lenguaje de IA · «Es importante destacar que…»: dice «es importante destacar»"
    assert nombradas([desborde, frase], LIMPIAS) == {8}
    # Del documento entero: orden, tablas y recursos en total.
    for problema in ["FALLA · estructura · tablas: hay 3 (diapositivas 4, 5, 6); el máximo es 2.",
                     "FALLA · estructura · diapositiva 9 («Z»): no sigue el orden de la lectura.",
                     "FALLA · estructura · recursos gráficos: 3 de 10 diapositivas de concepto son un diagrama."]:
        assert nombradas([desborde, problema], LIMPIAS) is None, problema


def test_la_correccion_por_diapositivas_no_toca_las_demas():
    nueva = diapositiva(PILARES, "Primera mitad", "frases", frases=["Scrum tiene tres pilares."])
    otra = diapositiva(PILARES, "Segunda mitad", "frases", frases=["Cada uno se usa cada día."])
    respuesta = solo(8, nueva, otra)
    respuesta["reemplazos"].append({"numero": 3, "diapositivas": [nueva]})        # no está en la lista: se ignora
    respuesta["fijas"]["cierre"] = {"frase": "Otra"}                             # no está en la lista: se ignora
    respuesta["fijas"]["idea_central"] = {"frases": ["Idea nueva."]}
    datos, no_encontradas = diapositivas.aplicar_por_diapositivas(LIMPIAS, respuesta, {2, 8})
    assert not no_encontradas
    assert [d["titulo"] for d in datos["diapositivas"]][5:7] == ["Primera mitad", "Segunda mitad"]
    assert datos["diapositivas"][:5] == LIMPIAS["diapositivas"][:5]
    assert datos["diapositivas"][7:] == LIMPIAS["diapositivas"][6:]
    assert datos["idea_central"]["frases"] == ["Idea nueva."] and datos["cierre"] == LIMPIAS["cierre"]


def test_el_orden_sigue_con_la_correccion_completa(con_lectura):
    """Un problema del documento entero (el orden de la lectura) va a la corrección completa. La cantidad ya no es
    un problema (decisión 66 c)."""
    desordenadas = con(diapositivas=LIMPIAS["diapositivas"][5:] + LIMPIAS["diapositivas"][:5])
    _, consulta = generar(con_lectura, desordenadas, {"diapositivas": LIMPIAS, "explicaciones": [], "datos_nuevos": []})
    assert consulta.llamadas[1]["options"].output_format["schema"] == diapositivas.ESQUEMA_CORRECCIONES
    assert "Llena solo «completa»." in consulta.llamadas[1]["prompt"]


# ---------- Agente diseñador (PLAN.md §0, decisión 32) ----------

def _con_disenador():
    from app.materiales import disenador
    if not disenador.disponible():
        pytest.skip("Node o los paquetes del diseñador no están instalados")
    from tests.unidad.test_disenador import respuesta_por_partes
    return respuesta_por_partes(Diapositivas.model_validate(LIMPIAS).total)


def _textos_del_pedido(prompt) -> tuple[str, list[str]]:
    """El texto de un pedido con imágenes (consultar_con_imagenes) y sus etiquetas de imagen."""
    async def leer():
        return [m async for m in prompt]
    contenido = asyncio.run(leer())[0]["message"]["content"]
    textos = [c["text"] for c in contenido if c["type"] == "text"]
    con_imagen = [contenido[i - 1]["text"] for i, c in enumerate(contenido) if c["type"] == "image"]
    return textos[-1], con_imagen


VISUAL_VACIA = {"cambios": []}


def test_con_el_contenido_aprobado_el_disenador_dibuja_y_la_revision_visual_pide_un_arreglo(con_lectura):
    diseno = _con_disenador()
    visual = {"cambios": [{"numero": 4, "problema": "Una tabla con mucho texto.", "propuesta": "Pasar a tarjetas."},
                          {"numero": 1, "problema": "La portada sin fuerza.", "propuesta": "Título más grande."}]}
    material, consulta = generar(con_lectura, LIMPIAS, medir=medir_con_imagenes, disenar=True,
                                 despues=(diseno, visual, diseno))
    assert material["estado"] == "verificada", material.get("problemas")
    assert material["diseno"]["estado"] == "diseñada" and material["diseno"]["arreglos"] == 1
    assert consulta.llamadas[2]["options"].output_format["schema"] == disenador_esquema()
    sistema = consulta.llamadas[2]["options"].system_prompt
    assert "module.exports = async (h)" in Path(sistema["path"]).read_text(encoding="utf-8")
    assert consulta.llamadas[4]["options"].system_prompt == sistema           # el mismo archivo: se lee de caché
    # Decisión 56: el diseño, la revisión visual y los arreglos van con Opus. Decisión 71: una sola revisión visual.
    assert {consulta.llamadas[n]["options"].model for n in (2, 3, 4)} == {"claude-opus-5-5"}
    assert len(consulta.llamadas) == 5
    rondas = [r["ronda"] for r in material["diseno"]["rondas"]]
    assert rondas == ["diseño", "revisión visual 1", "arreglo 1"]
    arreglo, imagenes = _textos_del_pedido(consulta.llamadas[4]["prompt"])
    assert "DISEÑO · diapositiva 4 («Qué valora más»): Una tabla con mucho texto. Propuesta: Pasar a tarjetas." in arreglo
    assert "DISEÑO · diapositiva 1" in arreglo and "PROGRAMA ACTUAL, POR DIAPOSITIVAS" in arreglo   # la portada también
    assert "Míralas para encontrar la causa" in arreglo
    assert imagenes == ["Diapositiva 1 («Portada»), como se ve ahora", "Diapositiva 4 («Qué valora más»), como se ve ahora"]
    assert material["revision_visual"] == visual["cambios"]
    entrega = {s["clave"]: s["lineas"] for s in material["entrega"]["secciones"]}
    assert any("lo dibujó el agente diseñador con 1 ícono" in l for l in entrega["valide"])
    assert any("Revisión visual de las imágenes" in l and "Pasar a tarjetas" in l for l in entrega["valide"])
    assert len(Presentation(diapositivas.archivos(con_lectura, 1)["pptx"]).slides) == 15


def _diseno_con_texto_ajeno(numero: int) -> dict:
    """El diseño de prueba con una frase inventada en la diapositiva `numero`: un control la rechaza siempre."""
    diseno = _con_disenador()
    return {**diseno, "diapositivas": [
        {**d, "codigo": d["codigo"] + ("\nh.pres.slides[h.pres.slides.length - 1].addText('Scrum nació en Japón en "
                                       "1986 y hoy lo usan todas las empresas del mundo.', { x: 0.5, y: 6.5, w: 9, "
                                       "h: 0.5, fontSize: 14 });" if d["numero"] == numero else "")}
        for d in diseno["diapositivas"]]}


def test_la_revision_visual_corre_aunque_el_primer_diseno_tenga_fallas_y_el_arreglo_ve_las_imagenes(con_lectura):
    """Decisión 56: la primera revisión visual no espera un diseño sin fallas, y el arreglo recibe la imagen de cada
    diapositiva con problemas."""
    material, consulta = generar(con_lectura, LIMPIAS, medir=medir_con_imagenes, disenar=True,
                                 despues=(_diseno_con_texto_ajeno(3), VISUAL_VACIA, _con_disenador()))
    assert material["diseno"]["estado"] == "diseñada", material.get("problemas")
    rondas = [r["ronda"] for r in material["diseno"]["rondas"]]
    assert rondas == ["diseño", "revisión visual 1", "arreglo 1"]
    arreglo, imagenes = _textos_del_pedido(consulta.llamadas[4]["prompt"])
    assert "Scrum nació en Japón" in arreglo
    assert imagenes == ["Diapositiva 3 («El Manifiesto Ágil»), como se ve ahora"]


def test_las_diapositivas_que_siguen_con_fallas_van_en_su_version_base_y_el_resto_queda_disenado(con_lectura):
    """Decisión 56: ya no es todo o nada."""
    malo = _diseno_con_texto_ajeno(3)
    material, consulta = generar(con_lectura, LIMPIAS, medir=medir_con_imagenes, disenar=True,
                                 despues=(malo, VISUAL_VACIA, *(malo,) * diapositivas.ARREGLOS_DE_DISENO))
    assert material["diseno"]["estado"] == "mixta" and material["diseno"]["base"] == [3]
    assert material["estado"] == "con pendientes"
    assert any("Diseño: FALLA · diseño · diapositiva 3" in p for p in material["problemas"])
    presentacion = Presentation(diapositivas.archivos(con_lectura, 1)["pptx"])
    textos = [" ".join(f.text_frame.text for f in s.shapes if f.has_text_frame) for s in presentacion.slides]
    assert "Scrum nació en Japón" not in textos[2] and "EL MANIFIESTO ÁGIL" in textos[2]   # la versión base
    assert "EL MANIFIESTO ÁGIL" not in textos[3] and "Qué valora más" in textos[3]           # el diseño
    entrega = {s["clave"]: s["lineas"] for s in material["entrega"]["secciones"]}
    assert any("salvo las diapositivas 3" in l for l in entrega["valide"])


def test_si_el_diseno_no_queda_bien_se_entrega_la_version_base(con_lectura):
    roto = _con_disenador()
    roto["diapositivas"][0]["codigo"] = 'throw new Error("se rompió");'       # el común es fijo: el error va en una diapositiva
    material, consulta = generar(con_lectura, LIMPIAS, medir=medir_con_imagenes, disenar=True,
                                 despues=(roto,) * (diapositivas.ARREGLOS_DE_DISENO + 1))
    assert material["estado"] == "con pendientes"
    assert material["diseno"]["estado"] == "base"
    assert any("Diseño: FALLA · diseño · el programa no corrió" in p and "se rompió" in p for p in material["problemas"])
    assert len(Presentation(diapositivas.archivos(con_lectura, 1)["pptx"]).slides) == 15   # la versión base


def test_un_programa_que_usa_require_no_se_corre(con_lectura):
    malo = _con_disenador()
    malo["diapositivas"][0]["codigo"] = 'const fs = require("fs");' 
    material, consulta = generar(con_lectura, LIMPIAS, medir=medir_con_imagenes, disenar=True,
                                 despues=(malo,) * (diapositivas.ARREGLOS_DE_DISENO + 1))
    assert material["diseno"]["estado"] == "base"
    assert "el programa usa lo que no está permitido: require" in consulta.llamadas[3]["prompt"]


def disenador_esquema():
    from app.materiales import disenador
    return disenador.ESQUEMA_DISENO


# ---------- Desbordes pequeños sin IA (PLAN.md §0, decisión 39) ----------

def _primer_texto(ruta, numero: int):
    """El primer recuadro con texto de una diapositiva: su texto como lo mide PowerPoint, el tamaño de su letra y
    la suma de sus márgenes interiores de arriba y abajo, en cm."""
    for figura in Presentation(ruta).slides[numero - 1].shapes:
        if figura.has_text_frame and figura.text_frame.text.strip():
            corridas = [c for p in figura.text_frame.paragraphs for c in p.runs]
            marco = figura.text_frame
            return (" ".join(marco.text.split())[:200], corridas[0].font.size,
                    round((marco.margin_top + marco.margin_bottom) / diapositivas.EMU_POR_CM, 2))
    raise AssertionError("sin texto")


def test_un_desborde_pequeno_recorta_el_margen_sin_ia_y_sin_bajar_la_letra(con_lectura):
    vistos = {}

    def medir(ruta, carpeta):
        if not vistos:
            vistos["texto"], vistos["letra"], vistos["margen"] = _primer_texto(ruta, 8)
            return office.Medicion(15, [office.Desborde(8, vistos["texto"], 0.1)], [])
        return medir_sin_desbordes(ruta, carpeta)

    material, consulta = generar(con_lectura, LIMPIAS, medir=medir)
    assert material["estado"] == "verificada", material.get("problemas")
    assert len(consulta.llamadas) == 2                       # redacción y revisión: el ajuste no llamó a la IA
    texto, letra, margen = _primer_texto(diapositivas.archivos(con_lectura, 1)["pptx"], 8)
    assert (texto, letra) == (vistos["texto"], vistos["letra"])          # la letra no cambia
    assert margen == round(vistos["margen"] - 0.2, 2)                    # 0.1 cm que sobra más 0.1 de holgura
    assert any("Ajuste sin IA: recorté el margen interior de 1 recuadro" in a for a in material["avance"])
    decidi = {s["clave"]: s["lineas"] for s in material["entrega"]["secciones"]}["decidi"]
    assert any(l.startswith("Diapositiva 8: recorté 0.2 cm del margen interior") and "decisión 39" in l for l in decidi)


def test_un_desborde_grande_una_tabla_o_sin_margen_no_se_ajustan(presentacion):
    ruta, _ = presentacion
    texto, letra, margen = _primer_texto(ruta, 2)
    copia = ruta.with_name("copia.pptx")
    shutil.copyfile(ruta, copia)
    assert diapositivas.ajustar_desbordes_pequenos(copia, [office.Desborde(2, texto, 0.6)]) == []
    assert diapositivas.ajustar_desbordes_pequenos(
        copia, [office.Desborde(2, texto, 0.1), office.Desborde(3, "tabla", 0.1, es_tabla=True)]) == []
    sin_margen = round(margen, 1)            # lo que sobra más la holgura supera el margen que hay
    assert diapositivas.ajustar_desbordes_pequenos(copia, [office.Desborde(2, texto, sin_margen)]) == []
    assert _primer_texto(copia, 2) == (texto, letra, margen)             # no tocó nada


# ---------- Reanudar desde lo guardado (PLAN.md §0, decisión 40) ----------

def _con_fallas(carpeta):
    """Una corrida que termina con fallas: el desborde de 1.4 cm no se va en ninguna de las tres correcciones."""
    def medir(ruta, carpeta_imagenes):
        return office.Medicion(15, [office.Desborde(8, "Los pilares de Scrum son tres.", 1.4)], [])
    corta = {**LIMPIAS["diapositivas"][5], "frases": ["Scrum tiene tres pilares."]}
    material, _ = generar(carpeta, LIMPIAS, *[solo(8, corta)] * diapositivas.VUELTAS_DEL_VERIFICADOR,
                          medir=medir, revisar=False)
    assert material["estado"] == "con fallas"
    return comun.leer_json(diapositivas.archivos(carpeta, 1)["contenido"])


def test_reanudar_sigue_desde_lo_guardado_sin_redactar_de_nuevo(con_lectura):
    guardado = _con_fallas(con_lectura)
    assert diapositivas.se_puede_reanudar(con_lectura, 1)
    pagina = cliente.get("/cursos/scrum/sesiones/1/materiales/diapositivas").text
    assert "Reanudar desde lo guardado</button>" in pagina
    consulta = consulta_en_secuencia(SIN_ERRORES)
    material = asyncio.run(diapositivas.reanudar(con_lectura, "scrum", 1, consulta=consulta,
                                                 medir=medir_sin_desbordes, disenar=False))
    assert material["estado"] == "verificada", material.get("problemas")
    assert len(consulta.llamadas) == 1                         # solo la revisión del contenido: no hubo redacción
    assert comun.leer_json(diapositivas.archivos(con_lectura, 1)["contenido"]) == guardado
    assert any("Reanudo desde lo generado el" in a for a in material["avance"])
    no_pude = {s["clave"]: s["lineas"] for s in material["entrega"]["secciones"]}["no_pude"]
    assert any("Reanudé desde lo generado el" in l and "decisión 40" in l for l in no_pude)
    assert not diapositivas.se_puede_reanudar(con_lectura, 1)   # ya quedó bien: no hay nada que reanudar


def test_no_se_reanuda_si_las_fichas_cambiaron(con_lectura):
    _con_fallas(con_lectura)
    ficha = almacen.cargar(con_lectura, "sesion", 1)   # un cambio real en la ficha: los materiales quedan desactualizados
    ficha["campos"]["lectura.cuidado"] = {"valor": "Otro riesgo", "origen": "profesor", "estado": ""}
    almacen.guardar(con_lectura, "sesion", 1, ficha)
    almacen.marcar_desactualizados(con_lectura, 1)
    assert not diapositivas.se_puede_reanudar(con_lectura, 1)
    assert "Reanudar desde lo guardado" not in cliente.get("/cursos/scrum/sesiones/1/materiales/diapositivas").text
    r = cliente.post("/cursos/scrum/sesiones/1/materiales/diapositivas/reanudar", follow_redirects=False)
    assert "No hay un trabajo guardado que se pueda reanudar" in unquote(r.headers["location"])
    with pytest.raises(diapositivas.NoSePuedeEmpezar):
        asyncio.run(diapositivas.reanudar(con_lectura, "scrum", 1, consulta=consulta_en_secuencia(),
                                          medir=medir_sin_desbordes, disenar=False))


def test_una_diapositiva_de_concepto_muestra_40_palabras_como_maximo():
    """Decisión 56: el diseñador no puede acortar el texto, así que el contenido llega con el largo de una diapositiva."""
    largas = copy.deepcopy(LIMPIAS)
    largas["diapositivas"][0]["frases"] = [" ".join(["palabra"] * 41)]
    largas["idea_central"]["frases"] = [" ".join(["idea"] * 41)]
    _, problemas = diapositivas._validar(largas, None)
    assert any("diapositiva 3" in p and "muestra 4" in p and "el máximo es 40" in p for p in problemas), problemas
    assert any("idea central: tiene 41 palabras; el máximo es 40" in p for p in problemas)
    justas = copy.deepcopy(LIMPIAS)
    justas["diapositivas"][0]["frases"] = [" ".join(["palabra"] * 30)]
    assert not [p for p in diapositivas._validar(justas, None)[1] if "palabras" in p]


def test_la_version_base_muestra_la_introduccion_de_las_columnas_y_el_titulo_de_la_tabla(tmp_path):
    """La versión base es la que se entrega cuando falla el diseño: tampoco puede omitir texto."""
    datos = copy.deepcopy(LIMPIAS)
    datos["diapositivas"][0] = diapositiva(MANIFIESTO, "Dos miradas", "columnas", frases=["Dos formas de verlo."],
                                           columnas=[{"titulo": "Antes", "frases": ["Plan fijo."]},
                                                     {"titulo": "Ahora", "frases": ["Plan que cambia."]}])
    datos["diapositivas"][1] = diapositiva(MANIFIESTO, "Los valores", "tabla", frases=[],
                                           tabla={"titulo": "Valores del manifiesto", "encabezados": ["Valor", "Frente a"],
                                                  "filas": [["Personas", "Procesos"]]})
    contenido = Diapositivas.model_validate(datos)
    identidad = estilo.Identidad(**estilo.PALETA_SOBRIA, tipografia="Calibri", logo=None)
    ruta = generador.generar_diapositivas(contenido, identidad, Portada("Scrum Master con IA", "Sesión 1", "Diapositivas"),
                                          tmp_path / "S1_Diapositivas.pptx")
    p = Presentation(ruta)
    textos = [" ".join(f.text_frame.text for f in s.shapes if f.has_text_frame) for s in p.slides]
    assert "Dos formas de verlo." in textos[2]
    assert "Valores del manifiesto" in textos[3]


def test_la_revision_visual_no_propone_quitar_texto_ni_pasarlo_a_las_notas():
    """Las notas se copian del contenido sin IA: un texto que el diseño pasara a las notas se perdería."""
    assert "no propongas quitarlo, acortarlo ni pasarlo a las notas" in diapositivas.PEDIDO_VISUAL
    assert "pasar el resto a las notas" not in diapositivas.PEDIDO_VISUAL


def test_la_orientacion_de_cantidad_no_es_tope():
    """Decisión 69: el pedido orienta a entre 12 y 20 de concepto; el programa no da falla por pasarse."""
    formato = diapositivas.instrucciones_de_formato(True)
    assert "apunta a entre 12 y 20 de concepto" in formato and "Nunca dejes fuera un concepto" in formato
    muchas = [dict(d, titulo=f"{d['titulo']} {i}") for i in range(3) for d in LIMPIAS["diapositivas"]]
    _, problemas = diapositivas._validar(con(diapositivas=muchas), lectura_limpia())
    assert not any("diapositivas:" in p for p in problemas)
