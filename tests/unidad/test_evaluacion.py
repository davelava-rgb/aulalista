"""Evaluación práctica (PBQ) y su clave, fase 9a (PLAN.md §0, decisión 64): el modelo de datos, lo que hace falta
para generarla, la revisión de sus preguntas en la ficha y el plan de la sesión. Sin IA."""

import asyncio
import json
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError

from app import servidor
from app.fichas import almacen, flujo
from claude_agent_sdk import AssistantMessage, TextBlock

from app.materiales import comun, ejecucion, evaluacion, practica
from app.materiales.contenido import ESQUEMA_EVALUACION, Evaluacion
from tests.conftest import consulta_en_secuencia, resultado_simulado
from tests.sembrados import evaluacion as sembrados
from tests.unidad.test_lectura import SIN_ERRORES, sesion  # noqa: F401

cliente = TestClient(servidor.app)


# ---------- Modelo de datos ----------

def test_la_evaluacion_limpia_cumple_el_modelo():
    contenido = Evaluacion.model_validate(sembrados.limpio())
    assert len(contenido.preguntas) == 3
    assert contenido.preguntas[0].tabla_de_respuesta.por_llenar == 3   # cada celda vacía es una decisión (64 c)
    assert set(ESQUEMA_EVALUACION["properties"]) == {"indicaciones", "preguntas", "archivos", "pasajes", "decisiones"}


@pytest.mark.parametrize("campo", ["puntos", "puntaje", "minutos", "tiempo", "nota"])
def test_el_modelo_no_acepta_puntos_ni_tiempos(campo):
    """SKILL.md, «Sin tiempos ni puntajes»: la evaluación no tiene dónde ponerlos."""
    datos = sembrados.limpio()
    datos["preguntas"][0][campo] = 5
    with pytest.raises(ValidationError):
        Evaluacion.model_validate(datos)
    datos = sembrados.limpio()
    datos[campo] = "30 minutos"
    with pytest.raises(ValidationError):
        Evaluacion.model_validate(datos)


@pytest.mark.parametrize("cambio, mensaje", [
    (lambda d: d["preguntas"][0]["clave"]["respuesta"][0].__setitem__(1, ""), "sin llenar"),
    (lambda d: d["preguntas"][0]["clave"]["respuesta"].pop(), "forma de la tabla"),
    (lambda d: d["preguntas"][0]["clave"]["respuesta"][0].__setitem__(0, "Hecho 9"), "cambia una celda"),
    (lambda d: d["preguntas"][0]["tabla_de_respuesta"].__setitem__("filas", [["Hecho 1", "Transparencia"]]),
     "forma de la tabla"),
    (lambda d: d["preguntas"][0]["clave"].__setitem__("partes", []), "partes"),
    (lambda d: d["preguntas"][0].__setitem__("tarea", "redactar"), "tarea"),
    (lambda d: d["preguntas"][1].__setitem__("numero", 3), "numeradas"),
    (lambda d: d["archivos"][0].__setitem__("nombre", "S1_E1_hechos.xlsx"), "no sigue la forma"),
    (lambda d: d["archivos"][0].__setitem__("nombre", "S1_P2_hechos-ventas.xlsx"), "es de la pregunta 2"),
    (lambda d: d["preguntas"][0]["archivos"].append("S1_P1_otro.xlsx"), "no está en los archivos"),
])
def test_el_modelo_rechaza_una_pregunta_sin_respuesta_clara(cambio, mensaje):
    datos = sembrados.limpio()
    cambio(datos)
    with pytest.raises(ValidationError, match=mensaje):
        Evaluacion.model_validate(datos)


def test_una_tabla_sin_celdas_para_el_alumno_no_es_una_pregunta():
    datos = sembrados.limpio()
    completa = datos["preguntas"][0]["clave"]["respuesta"]
    datos["preguntas"][0]["tabla_de_respuesta"]["filas"] = [list(f) for f in completa]
    with pytest.raises(ValidationError, match="ninguna celda para el alumno"):
        Evaluacion.model_validate(datos)


# ---------- Lo que hace falta para generarla (decisión 64 a) ----------

def _aprobar(carpeta, clave):
    ahora = datetime.now().isoformat(timespec="seconds")
    comun.actualizar(carpeta, 1, clave, estado="aprobada", fecha=ahora, aprobada=ahora, desactualizado=False)


def test_la_evaluacion_espera_la_lectura_y_el_laboratorio_aprobados_y_al_dia(sesion):  # noqa: F811
    assert evaluacion.faltan(sesion, 1) == {"ficha": [], "otros": [
        "Aprueba primero la lectura: la evaluación no evalúa nada que no se haya enseñado.",
        "Aprueba primero el laboratorio: la evaluación no evalúa nada que no se haya enseñado."]}
    _aprobar(sesion, "lectura")
    _aprobar(sesion, "laboratorio")
    assert evaluacion.faltan(sesion, 1) == {"ficha": [], "otros": []}
    comun.actualizar(sesion, 1, "laboratorio", desactualizado=True)
    assert evaluacion.faltan(sesion, 1)["otros"] == [
        "El laboratorio está desactualizado: genéralo y apruébalo de nuevo."]


def test_sin_laboratorio_en_la_sesion_la_evaluacion_no_lo_espera(sesion):  # noqa: F811
    ficha = almacen.cargar(sesion, "sesion", 1)
    ficha["campos"]["materiales.cuales"] = {"valor": "lectura, evaluación", "origen": "profesor", "estado": ""}
    almacen.guardar(sesion, "sesion", 1, ficha, "editada por el profesor")
    _aprobar(sesion, "lectura")
    assert evaluacion.faltan(sesion, 1) == {"ficha": [], "otros": []}


def test_sin_preguntas_en_la_ficha_lo_que_falta_es_la_ficha(sesion):  # noqa: F811
    ficha = almacen.cargar(sesion, "sesion", 1)
    for campo in [c for c in ficha["campos"] if c.startswith("pregunta.")]:
        ficha["campos"][campo] = {"valor": "", "origen": "", "estado": ""}
    almacen.guardar(sesion, "sesion", 1, ficha, "editada por el profesor")
    faltan = evaluacion.faltan(sesion, 1)
    assert faltan["otros"] == [] and any("al menos una pregunta" in f for f in faltan["ficha"])


def test_la_evaluacion_tiene_sus_propios_archivos(sesion):  # noqa: F811
    rutas = evaluacion.archivos(sesion, 1)
    nombres = {k: r.name for k, r in rutas.items() if k in ("word", "clave", "zip", "excel")}
    assert nombres == {"word": "S1_Evaluacion.docx", "clave": "S1_Evaluacion_Clave.docx",
                       "zip": "S1_Evaluacion_Archivos.zip", "excel": "S1_Evaluacion_Verificacion.xlsx"}
    assert rutas["zip"] != servidor.laboratorio.archivos(sesion, 1)["zip"]     # aparte del .zip del laboratorio


# ---------- La ficha y el plan de la sesión (decisión 64 h e i) ----------

def _revision(ejercicios, preguntas):
    def fila(n, solo=True, datos=True):
        return {"numero": n, "funciona_solo": solo, "sin_llenar_a_mano": datos, "motivo": "Cumple"}
    return {"ejercicios": [fila(n) for n in ejercicios], "preguntas": [fila(n, *p) for n, p in preguntas]}


def test_la_revision_de_la_ficha_revisa_tambien_las_preguntas(sesion):  # noqa: F811
    ficha = almacen.cargar(sesion, "sesion", 1)
    ficha.pop("revision_ejercicios", None)          # sin una revisión guardada de antes
    texto, preguntas = almacen.texto_de_preguntas(ficha)
    assert preguntas and "Pregunta 1 · Concepto que evalúa: Pilares de Scrum" in texto
    _, ejercicios = almacen.texto_de_ejercicios(ficha)
    consulta = consulta_en_secuencia(_revision(ejercicios, [(1, (True, True)), (2, (False, True))]))
    problemas = asyncio.run(flujo.revisar_ejercicios(ficha, "scrum", 1, consulta=consulta))
    assert problemas == ["Pregunta 2 de la evaluación: no funciona sola. Cumple"]
    assert "Revisa también cada pregunta de la evaluación práctica" in consulta.llamadas[0]["prompt"]


def test_sin_evaluacion_en_la_sesion_sus_preguntas_no_se_revisan(sesion):  # noqa: F811
    ficha = almacen.cargar(sesion, "sesion", 1)
    huella_con = almacen.huella_de_ejercicios(ficha)
    ficha["campos"]["materiales.cuales"] = {"valor": "lectura, laboratorio", "origen": "profesor", "estado": ""}
    assert almacen.texto_de_preguntas(ficha) == ("", [])
    assert almacen.huella_de_ejercicios(ficha) != huella_con   # cambia lo que se revisa: se revisa de nuevo
    _, ejercicios = almacen.texto_de_ejercicios(ficha)
    consulta = consulta_en_secuencia(_revision(ejercicios, []))
    asyncio.run(flujo.revisar_ejercicios(ficha, "scrum", 1, consulta=consulta))
    assert "No hay preguntas de evaluación" in consulta.llamadas[0]["prompt"]


def test_el_plan_de_la_sesion_cuenta_la_evaluacion_solo_si_la_produce(sesion):  # noqa: F811
    ficha = almacen.cargar(sesion, "sesion", 1)
    assert not flujo.plan_por_preparar(ficha)
    ficha["campos"]["pregunta.1.respuesta"] = {"valor": "", "origen": "", "estado": ""}
    assert flujo.plan_por_preparar(ficha)                         # falta algo que la evaluación pide
    ficha["campos"]["materiales.cuales"] = {"valor": "lectura, laboratorio", "origen": "profesor", "estado": ""}
    assert not flujo.plan_por_preparar(ficha)                     # la sesión no la produce: no cuenta


def test_la_ficha_ya_no_dice_que_la_evaluacion_no_se_usa(sesion):  # noqa: F811
    pagina = cliente.get("/cursos/scrum/sesiones/1/ficha").text
    assert "estado-no_se_usa" not in pagina and "AulaLista no genera la evaluación" not in pagina
    assert "decisión 37" not in pagina and "apagada" not in pagina
    assert "Las preguntas se revisan al guardar, junto con los ejercicios." in pagina
    # El material todavía no se genera desde la interfaz (decisión 46): llega en la fase 9d.
    assert "/materiales/evaluacion" not in cliente.get("/cursos/scrum").text


# ---------- Archivos y verificador (fase 9b) ----------

def _hallazgos(resultado):
    return [(h.nivel, h.regla, h.detalle) for h in resultado.hallazgos]


def _fallas(resultado):
    return [h for h in resultado.hallazgos if h.nivel == "FALLA"]


def _texto_docx(ruta):
    from docx import Document
    documento = Document(ruta)
    celdas = [c.text for t in documento.tables for f in t.rows for c in f.cells]
    return "\n".join([p.text for p in documento.paragraphs] + celdas)


def test_la_evaluacion_limpia_da_sus_dos_word_su_zip_y_su_excel_sin_fallas(sesion):  # noqa: F811
    contenido = Evaluacion.model_validate(sembrados.limpio())
    problemas, resultado = evaluacion.armar_y_verificar(sesion, 1, contenido)
    assert problemas == [] and _fallas(resultado) == [], _hallazgos(resultado)
    rutas = evaluacion.archivos(sesion, 1)
    for clave in ("word", "clave", "zip", "excel"):
        assert rutas[clave].exists(), clave
    assert practica.nombres_del_zip(rutas["zip"]) == sorted(a.nombre for a in contenido.archivos)
    hoja, clave = _texto_docx(rutas["word"]), _texto_docx(rutas["clave"])
    assert "Evaluación práctica" in hoja and "Pregunta 1 · Clasifica los hechos" in hoja and "Tu respuesta" in hoja
    assert "Transparencia" not in hoja and "Concepto que evalúa" not in hoja      # sin respuestas ni concepto
    assert "Clave de la evaluación" in clave and "Transparencia" in clave and "Errores comunes" in clave
    assert "S1_P1_hechos-ventas.xlsx: El hecho 3 parece de adaptación." in clave  # el problema sembrado del archivo
    from docx import Document
    from docx.shared import Cm
    respuesta = Document(rutas["word"]).tables[-1]                     # la tabla de respuesta de la última pregunta
    assert all(f.height >= Cm(1) for f in respuesta.rows[1:])          # espacio para escribir a mano


@pytest.mark.parametrize("nombre", list(sembrados.SEMBRADOS))
def test_cada_error_sembrado_de_la_evaluacion_se_detecta(sesion, nombre):  # noqa: F811
    cambio, regla, detalle = sembrados.SEMBRADOS[nombre]
    contenido = Evaluacion.model_validate(sembrados.con(cambio))
    _, resultado = evaluacion.armar_y_verificar(sesion, 1, contenido)
    assert any(h.nivel == "FALLA" and h.regla == regla and detalle in h.detalle for h in resultado.hallazgos), \
        _hallazgos(resultado)


def test_la_cantidad_de_preguntas_es_la_de_la_ficha(sesion):  # noqa: F811
    datos = sembrados.limpio()
    datos["preguntas"].pop()
    datos["archivos"].pop()
    hallazgos = evaluacion.controles(Evaluacion.model_validate(datos), evaluacion.preguntas_de_la_ficha(sesion, 1))
    assert ("FALLA", "cantidad de preguntas", "la ficha pide 3 preguntas y la evaluación tiene 2") in [
        (h.nivel, h.regla, h.detalle) for h in hallazgos]


def test_ninguna_respuesta_aparece_en_el_laboratorio_ni_en_la_actividad():
    contenido = Evaluacion.model_validate(sembrados.limpio())
    assert evaluacion.controles(contenido, 3, {"la guía del laboratorio": "otros datos sin relación"}) == []
    # Se busca la fila completa de la respuesta («Hecho 1: Transparencia»), sin tildes, mayúsculas ni signos.
    otros = {"la guía del laboratorio": evaluacion._plano("En el ejemplo, HECHO 1: transparencia.")}
    hallazgos = evaluacion.controles(contenido, 3, otros)
    assert [(h.regla, h.detalle) for h in hallazgos] == [
        ("respuesta fuera de la clave", "la respuesta aparece en la guía del laboratorio")]


def test_una_pregunta_no_usa_un_archivo_del_laboratorio():
    from app.materiales.contenido import ArchivoPractica
    contenido = Evaluacion.model_validate(sembrados.limpio())
    igual = ArchivoPractica.model_validate({**sembrados.LIMPIO["archivos"][0], "nombre": "S1_E2_hechos.xlsx"})
    hallazgos = evaluacion.controles(contenido, 3, del_laboratorio=[igual])
    assert ("FALLA", "archivo del laboratorio") in [(h.nivel, h.regla) for h in hallazgos]
    parecido = ArchivoPractica.model_validate({**sembrados.LIMPIO["archivos"][0], "nombre": "S1_E2_hechos.xlsx",
                                               "hojas": [{"nombre": "Hechos", "encabezados": ["Hecho", "Descripción"],
                                                          "filas": sembrados.LIMPIO["archivos"][0]["hojas"][0]["filas"][:2]
                                                          + [["Hecho 9", "Otro dato."]]}]})
    hallazgos = evaluacion.controles(contenido, 3, del_laboratorio=[parecido])
    assert [(h.nivel, h.regla) for h in hallazgos] == [("AVISO", "datos del laboratorio")]


def test_las_preguntas_de_tamano_muy_distinto_dan_aviso():
    contenido = Evaluacion.model_validate(sembrados.con(sembrados._filas(8)))
    assert [(h.nivel, h.regla) for h in evaluacion.controles(contenido, 3)] == [("AVISO", "tamaño parecido")]


# ---------- Redacción, validación y ejecución (fase 9c), con respuestas simuladas ----------

def _lista_para_generar(carpeta):
    for clave in ("lectura", "laboratorio"):
        _aprobar(carpeta, clave)


async def _todas_coinciden(contenido, **_):
    """La ejecución tiene su propia prueba (abajo): aquí cada pregunta coincide sin llamar a Claude."""
    return [ejecucion.Ejecucion(p.numero, p.titulo, respuesta_ficha="clave", a_mano=True,
                                conversacion=[ejecucion.Turno("guía", "pedido"), ejecucion.Turno("resultado", "tabla")])
            for p in contenido.preguntas]


def _generar(carpeta, *salidas, ejecutar_=_todas_coinciden):
    consulta = consulta_en_secuencia(*salidas)
    material = asyncio.run(evaluacion.generar(carpeta, "scrum", 1, consulta=consulta, ejecutar_=ejecutar_))
    return material, consulta


def test_genera_la_evaluacion_completa_y_su_entrega(sesion):  # noqa: F811
    _lista_para_generar(sesion)
    material, consulta = _generar(sesion, sembrados.limpio(), SIN_ERRORES)
    assert material["estado"] == "verificada", material.get("problemas")
    rutas = evaluacion.archivos(sesion, 1)
    for clave in ("word", "clave", "zip", "excel", "registro", "contenido"):
        assert rutas[clave].exists(), clave
    pedido = consulta.llamadas[0]["prompt"]
    assert "TAREA: prepara la evaluación práctica de la sesión" in pedido and "Material 6 · Evaluación práctica" in pedido
    assert "preguntas: 3, las de la ficha de la sesión" in pedido
    assert material["preguntas"] == 3 and material["ejecucion"]["coinciden"] == [1, 2, 3]
    entrega = json.dumps(material["entrega"], ensure_ascii=False)
    assert "S1_Evaluacion_Clave.docx: la clave de corrección" in entrega
    assert "Resolví 3 preguntas con Claude como un alumno, sin la clave." in entrega
    assert "Problema sembrado en S1_P1_hechos-ventas.xlsx: El hecho 3 parece de adaptación." in entrega


def test_sin_el_laboratorio_aprobado_no_empieza_ni_gasta(sesion):  # noqa: F811
    _aprobar(sesion, "lectura")
    consulta = consulta_en_secuencia()
    with pytest.raises(evaluacion.NoSePuedeEmpezar, match="Aprueba primero el laboratorio"):
        asyncio.run(evaluacion.generar(sesion, "scrum", 1, consulta=consulta))
    assert consulta.llamadas == []


def test_una_falla_de_una_oracion_se_corrige_por_oraciones(sesion):  # noqa: F811
    _lista_para_generar(sesion)
    con_puntaje = sembrados.con(sembrados._puntaje)
    cambio = {"cambios": [{"oracion": "Esta pregunta vale 5 puntos.", "nueva": ""}], "explicaciones": [],
              "datos_nuevos": []}
    material, consulta = _generar(sesion, con_puntaje, cambio, SIN_ERRORES)
    assert material["estado"] == "verificada", material.get("problemas")
    assert material["correcciones"]["verificador"] == 1
    assert "Llena solo «oraciones»" in consulta.llamadas[1]["prompt"]


def test_una_pregunta_demasiado_grande_va_a_correccion_completa(sesion):  # noqa: F811
    _lista_para_generar(sesion)
    corregida = {"evaluacion": sembrados.limpio(), "explicaciones": [], "datos_nuevos": []}
    material, consulta = _generar(sesion, sembrados.con(sembrados._texto_largo), corregida, SIN_ERRORES)
    assert material["estado"] == "verificada", material.get("problemas")
    assert "Devuelve la evaluación completa" in consulta.llamadas[1]["prompt"]
    assert "260 palabras" in consulta.llamadas[1]["prompt"]


def test_una_diferencia_con_la_clave_queda_pendiente_en_la_entrega(sesion):  # noqa: F811
    async def con_diferencia(contenido, **_):
        resultados = await _todas_coinciden(contenido)
        resultados[1].diferencias = [ejecucion.Diferencia("Scrum diario", "Revisión del Sprint", "Otro orden.")]
        return resultados

    _lista_para_generar(sesion)
    material, _ = _generar(sesion, sembrados.limpio(), SIN_ERRORES, ejecutar_=con_diferencia)
    assert material["estado"] == "con pendientes"
    assert ("Ejecución · pregunta 2: la clave dice «Scrum diario» y el alumno respondió «Revisión del Sprint». "
            "Otro orden.") in material["problemas"]
    assert material["ejecucion"]["coinciden"] == [1, 3]
    assert "## Pregunta 2" in evaluacion.archivos(sesion, 1)["registro"].read_text(encoding="utf-8")


def _alumno_y_comparacion(diferencias):
    """El alumno simulado responde con una tabla; la comparación devuelve las diferencias que se le pasan."""
    llamadas, pendientes = [], list(diferencias)

    async def consulta(*, prompt, options):
        llamadas.append({"prompt": prompt, "options": options})
        if options.output_format is None:
            yield AssistantMessage(content=[TextBlock(text="Hecho | Pilar\nHecho 1 | Transparencia\n"
                                                      "Hecho 2 | Adaptación\nHecho 3 | Adaptación")],
                                   model="claude-sonnet-5-5")
            yield resultado_simulado()
        else:
            yield resultado_simulado(structured_output={"diferencias": pendientes.pop(0) if pendientes else []})

    consulta.llamadas = llamadas
    return consulta


def test_cada_pregunta_se_resuelve_como_alumno_sin_la_clave_y_se_compara_con_prueba(sesion):  # noqa: F811
    contenido = Evaluacion.model_validate(sembrados.limpio())
    evaluacion.armar_y_verificar(sesion, 1, contenido)               # crea los archivos reales
    con_prueba = {"esperado": "Inspección", "salio": "Hecho 3 | Adaptación", "explicacion": "Otro pilar."}
    sin_prueba = {"esperado": "Hecho 7 | Valores", "salio": "", "explicacion": "Inventada."}
    consulta = _alumno_y_comparacion([[con_prueba, sin_prueba]])
    resultados = asyncio.run(evaluacion.ejecutar(
        contenido, carpeta_archivos=evaluacion.archivos(sesion, 1)["archivos"],
        respuestas_ficha={1: "Hecho 3: inspección"},
        llamada={"curso": "scrum", "sesion": "S1", "material": "evaluacion", "consulta": consulta}))
    alumno = [l["prompt"] for l in consulta.llamadas if l["options"].output_format is None]
    assert len(alumno) == 3                                             # una conversación nueva por pregunta
    assert "[Archivo: S1_P1_hechos-ventas.xlsx]" in alumno[0] and "Hecho 1 | ____" in alumno[0]
    assert "Transparencia" not in alumno[0] and "Clave" not in alumno[0]    # sin la clave
    comparacion = [l["prompt"] for l in consulta.llamadas if l["options"].output_format is not None][0]
    assert "Respuesta correcta de la ficha: Hecho 3: inspección" in comparacion
    assert [d.esperado for d in resultados[0].diferencias] == ["Inspección"] and resultados[0].descartadas == 1
    assert evaluacion.pendientes_de(resultados[0]) == [
        "Ejecución · pregunta 1: la clave dice «Inspección» y el alumno respondió «Hecho 3 | Adaptación». Otro pilar."]
    assert resultados[1].diferencias == [] and resultados[2].diferencias == []


def test_la_evaluacion_se_aprueba_solo_sin_fallas_y_al_dia(sesion):  # noqa: F811
    _lista_para_generar(sesion)
    with pytest.raises(evaluacion.NoSePuedeAprobar):
        evaluacion.aprobar(sesion, 1)
    _generar(sesion, sembrados.limpio(), SIN_ERRORES)
    assert evaluacion.aprobar(sesion, 1)["estado"] == "aprobada"
    comun.actualizar(sesion, 1, "evaluacion", estado="verificada", desactualizado=True)
    with pytest.raises(evaluacion.NoSePuedeAprobar, match="Genérala de nuevo"):
        evaluacion.aprobar(sesion, 1)
