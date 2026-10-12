"""Controles de las fichas, estilo, portada y modelos en PDF (PLAN.md §0, decisión 27)."""

import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from app import servidor
from app.fichas import almacen, controles, plantillas
from app.materiales import catalogo, estilo, modelos
from app.materiales import docx as generador
from app.materiales.contenido import Lectura
from tests.unidad.test_lectura import LIMPIA, sesion  # noqa: F401  (fixture)

cliente = TestClient(servidor.app)


@pytest.fixture
def curso(entorno):
    cliente.post("/cursos", data={"nombre": "Scrum"})
    return entorno["cursos"] / "scrum"


def _pdf(ruta, color=(0.12, 0.31, 0.62), fuente="helv"):
    """Un PDF de modelo: portada y una página de prosa con texto de color."""
    documento = pymupdf.open()
    portada = documento.new_page()
    portada.insert_text((72, 100), "Curso de prueba", fontname=fuente, fontsize=28, color=color)
    pagina = documento.new_page()
    pagina.insert_text((72, 80), "Titulo del bloque", fontname=fuente, fontsize=16, color=color)
    prosa = ("Revisa cada pedido antes de enviarlo al almacén. Si falta un dato, pídelo al cliente y anótalo en la "
             "ficha del pedido. Así el equipo de despacho trabaja con la información completa y no repite llamadas. ") * 2
    pagina.insert_textbox(pymupdf.Rect(72, 100, 520, 400), prosa, fontname=fuente, fontsize=11, color=(0.1, 0.1, 0.1))
    documento.save(ruta)
    documento.close()


# ---------- Ficha: plantilla y valores por defecto ----------

def test_la_plantilla_tiene_el_campo_estilo():
    texto = plantillas.renderizar("curso", {"visual.estilo": {"valor": "moderno"}})
    assert "- Estilo (profesional, académico, moderno, corporativo, minimalista o cercano): moderno" in texto


def test_la_ficha_nueva_trae_docente_y_contacto_editables(curso):
    ficha = almacen.cargar(curso, "curso")
    assert almacen.valor(ficha, "portada.docente") == "José Fernando Davelouis Mendía"
    assert "www.linkedin.com/in/jdavelouis" in almacen.valor(ficha, "portada.contacto")
    assert "info@futurasoft.net" in almacen.valor(ficha, "portada.contacto")
    cliente.post("/cursos/scrum/ficha", data={"accion": "guardar", "portada.docente": "Otra persona"})
    assert almacen.valor(almacen.cargar(curso, "curso"), "portada.docente") == "Otra persona"


def test_la_pagina_muestra_los_controles(curso):
    pagina = cliente.get("/cursos/scrum/ficha").text
    assert 'enctype="multipart/form-data"' in pagina
    assert '<select id="curso.idioma"' in pagina and "Otro…" in pagina
    assert 'name="paleta"' in pagina and "Azul institucional" in pagina
    assert 'type="color"' in pagina
    assert '<select id="visual.estilo"' in pagina and "Moderno" in pagina
    assert 'type="radio" name="herramientas.neutral"' in pagina
    assert 'data-depende-de="herramientas.neutral"' in pagina
    assert 'type="checkbox" name="materiales.nunca"' in pagina
    assert 'name="modelos.aprobados.archivos" accept=".pdf" multiple' in pagina
    assert 'name="visual.logo.archivo"' in pagina
    assert "José Fernando Davelouis Mendía" in pagina


# ---------- Del formulario a la ficha ----------

def test_los_controles_escriben_el_texto_de_la_ficha(curso):
    cliente.post("/cursos/scrum/ficha", data={
        "accion": "guardar",
        "curso.idioma": controles.OTRO, "curso.idioma.otro": "Quechua",
        "visual.estilo": "académico",
        "visual.principal": "#1F4E79",
        "herramientas.neutral": "no", "herramientas.cual": "Microsoft Copilot en Word 365",
        "herramientas.adjuntos": "no sé",
        "materiales.por_sesion.casillas": "1", "materiales.por_sesion": ["lectura", "diapositivas"],
        "materiales.nunca.casillas": "1", "materiales.nunca": list(controles.NUNCA_POR_OMISION),
        "materiales.nunca.otros": "anexos",
    })
    ficha = almacen.cargar(curso, "curso")
    assert almacen.valor(ficha, "curso.idioma") == "Quechua"
    assert almacen.valor(ficha, "visual.estilo") == "académico"
    assert almacen.valor(ficha, "herramientas.neutral") == "no"
    assert almacen.valor(ficha, "herramientas.adjuntos") == "no sé"
    assert almacen.valor(ficha, "materiales.por_sesion") == "lectura y diapositivas"
    assert almacen.valor(ficha, "materiales.nunca").endswith("lenguaje publicitario y anexos")
    md = almacen.ruta_ficha(curso, "curso", extension=".md").read_text(encoding="utf-8")
    assert "- Estilo (profesional, académico, moderno, corporativo, minimalista o cercano): académico" in md


def test_todas_las_casillas_por_defecto_dejan_el_valor_vacio(curso):
    cliente.post("/cursos/scrum/ficha", data={
        "accion": "guardar", "curso.nombre": "Scrum",
        "materiales.nunca.casillas": "1", "materiales.nunca": list(controles.NUNCA_POR_OMISION),
        "materiales.por_sesion.casillas": "1", "materiales.por_sesion": list(controles.MATERIALES),
    })
    ficha = almacen.cargar(curso, "curso")
    assert almacen.valor(ficha, "materiales.nunca") == ""
    assert almacen.valor(ficha, "materiales.por_sesion") == ""


def test_las_casillas_leen_lo_que_ya_estaba_escrito():
    marcadas, otros = controles.marcadas("materiales.cuales", "Lectura y diapositivas")
    assert marcadas == ["lectura", "diapositivas"] and otros == ""
    assert controles.marcadas("materiales.cuales", "los seis")[0] == list(controles.MATERIALES)
    assert controles.marcadas("materiales.cuales", "")[0] == list(controles.MATERIALES)
    ficha = {"tipo": "sesion", "campos": {"materiales.cuales": {"valor": "lectura y diapositivas"}}}
    assert almacen.materiales_de_la_sesion(ficha) == {"lectura", "diapositivas"}


def test_subir_logo_y_modelos(curso, tmp_path):
    _pdf(tmp_path / "modelo.pdf")
    archivos = [
        ("visual.logo.archivo", ("logo.png", b"\x89PNG-de-prueba", "image/png")),
        ("modelos.aprobados.archivos", ("modelo.pdf", (tmp_path / "modelo.pdf").read_bytes(), "application/pdf")),
        ("modelos.aprobados.archivos", ("no-sirve.docx", b"x", "application/octet-stream")),
    ]
    cliente.post("/cursos/scrum/ficha", files=archivos, data={
        "accion": "guardar", "visual.logo": "", "modelos.aprobados": "", "modelos.aprobados.casillas": "1"})
    ficha = almacen.cargar(curso, "curso")
    assert almacen.valor(ficha, "visual.logo") == "logo.png" and (curso / "logo.png").exists()
    assert almacen.valor(ficha, "modelos.aprobados") == "modelo.pdf"
    assert (curso / "modelos" / "modelo.pdf").exists() and not (curso / "modelos" / "no-sirve.docx").exists()
    cliente.post("/cursos/scrum/ficha", data={"accion": "guardar", "modelos.aprobados": "modelo.pdf",
                                             "modelos.aprobados.casillas": "1", "modelos.aprobados.quitar": "modelo.pdf"})
    assert almacen.valor(almacen.cargar(curso, "curso"), "modelos.aprobados") == ""


# ---------- Estilo en el Word ----------

def _word(tmp_path, nombre_estilo):
    identidad = estilo.Identidad(**estilo.PALETA_SOBRIA, tipografia="Calibri", logo=None,
                                 estilo=catalogo.estilo(nombre_estilo))
    portada = generador.Portada("Scrum Master con IA", "Sesión 1 · Fundamentos", "Lectura",
                                docente="José Fernando Davelouis Mendía",
                                contacto="www.linkedin.com/in/jdavelouis; info@futurasoft.net")
    ruta = generador.generar_lectura(Lectura.model_validate(LIMPIA), identidad, portada, tmp_path / f"{nombre_estilo}.docx")
    return Document(ruta)


def test_profesional_es_el_diseno_de_siempre(tmp_path):
    documento = _word(tmp_path, "")
    assert documento.styles["Title"].font.size.pt == 26
    assert documento.styles["Heading 1"].font.size.pt == 16
    assert documento.styles["Normal"].font.size.pt == 11   # el cuerpo no cambia con el estilo


def test_moderno_lleva_franja_y_titulos_grandes(tmp_path):
    documento = _word(tmp_path, "moderno")
    assert documento.styles["Title"].font.size.pt == 32
    assert documento.styles["Normal"].font.size.pt == 11
    primera = documento.tables[0]
    sombra = primera.rows[0].cells[0]._tc.tcPr.find(qn("w:shd"))
    assert sombra.get(qn("w:fill")) == estilo.PALETA_SOBRIA["principal"]


def test_academico_centra_la_portada_y_subraya_el_titulo_1(tmp_path):
    documento = _word(tmp_path, "Académico")
    titulo = next(p for p in documento.paragraphs if p.style.name == "Title")
    assert titulo.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert documento.styles["Heading 1"].element.pPr.find(qn("w:pBdr")) is not None


def test_la_portada_muestra_docente_y_contacto(tmp_path):
    textos = [p.text for p in _word(tmp_path, "profesional").paragraphs]
    assert "Docente: José Fernando Davelouis Mendía" in textos
    assert "www.linkedin.com/in/jdavelouis" in textos and "info@futurasoft.net" in textos


def test_el_estilo_cambia_el_borde_de_los_recuadros(tmp_path):
    def borde(documento):
        return {t._tbl.tblPr.find(qn("w:tblBorders")).find(qn("w:left")).get(qn("w:sz"))
                for t in documento.tables if t._tbl.tblPr.find(qn("w:tblBorders")).find(qn("w:left")) is not None}
    assert borde(_word(tmp_path, "minimalista")) - {None} == {"12"}   # None: tablas sin borde izquierdo
    assert borde(_word(tmp_path, "moderno")) - {None} == {"36"}


def test_un_estilo_desconocido_vale_profesional():
    assert catalogo.estilo("barroco").nombre == "profesional"
    assert catalogo.estilo("ACADEMICO").nombre == "académico"


# ---------- Modelos en PDF ----------

def _con_modelo(sesion, tmp_path, **pdf):  # noqa: F811
    (sesion / "modelos").mkdir(exist_ok=True)
    _pdf(sesion / "modelos" / "modelo.pdf", **pdf)
    ficha = almacen.cargar(sesion, "curso")
    almacen.aplicar_formulario(ficha, {"modelos.aprobados": "modelo.pdf"})
    almacen.guardar(sesion, "curso", None, ficha)


def test_el_modelo_da_colores_y_tipografia_sin_ia(sesion, tmp_path):  # noqa: F811
    _con_modelo(sesion, tmp_path)
    apariencia = modelos.apariencia(sesion, ("Helvetica", "Calibri"))
    assert apariencia["principal"] == "1F4F9E"
    assert apariencia["tipografia"] == "Helvetica"
    assert apariencia["origen"] == "modelo.pdf"


def test_la_ficha_manda_sobre_el_modelo(sesion, tmp_path):  # noqa: F811
    _con_modelo(sesion, tmp_path)
    identidad = estilo.desde_ficha(sesion)
    assert identidad.principal == "1F4F9E"
    assert any("Tomé del modelo (modelo.pdf)" in a for a in identidad.avisos)
    assert not any("no trae paleta" in a for a in identidad.avisos)
    ficha = almacen.cargar(sesion, "curso")
    almacen.aplicar_formulario(ficha, {"visual.principal": "#6D1A36"})
    almacen.guardar(sesion, "curso", None, ficha)
    assert estilo.desde_ficha(sesion).principal == "6D1A36"


def test_cada_modelo_se_analiza_una_sola_vez(sesion, tmp_path, monkeypatch):  # noqa: F811
    _con_modelo(sesion, tmp_path)
    modelos.apariencia(sesion, ())
    monkeypatch.setattr(modelos, "analizar", lambda ruta: pytest.fail("volvió a analizar el modelo"))
    modelos.apariencia(sesion, ())
    modelos.muestra_de_tono(sesion)


def test_el_redactor_recibe_el_tono_y_el_pasaje_del_modelo(sesion, tmp_path):  # noqa: F811
    from app.materiales import lectura
    _con_modelo(sesion, tmp_path)
    ficha = almacen.cargar(sesion, "curso")
    almacen.aplicar_formulario(ficha, {"visual.estilo": "cercano"})
    almacen.guardar(sesion, "curso", None, ficha)
    pedido = lectura.pedido_de_redaccion(sesion, 1)
    assert "TONO DEL MATERIAL (estilo «cercano»)" in pedido
    assert "Trata al alumno de tú" in pedido
    assert "MODELO DE TONO (pasaje de modelo.pdf" in pedido
    assert "No uses su contenido" in pedido
    assert "Revisa cada pedido antes de enviarlo" in pedido
    assert "Curso de prueba" not in pedido   # la portada del modelo no entra en el pasaje


def test_la_portada_de_las_diapositivas_lleva_franja_y_contacto(tmp_path):
    from pptx import Presentation

    from app.materiales import pptx as diapositivas
    from app.materiales.contenido import Diapositivas
    from tests.unidad.test_diapositivas import LIMPIAS
    identidad = estilo.Identidad(**estilo.PALETA_SOBRIA, tipografia="Calibri", logo=None,
                                 estilo=catalogo.estilo("corporativo"))
    portada = generador.Portada("Scrum Master con IA", "Sesión 1 · Fundamentos", "Diapositivas",
                                docente="José Fernando Davelouis Mendía",
                                contacto="www.linkedin.com/in/jdavelouis; info@futurasoft.net")
    ruta = diapositivas.generar_diapositivas(Diapositivas.model_validate(LIMPIAS), identidad, portada,
                                             tmp_path / "S1_Diapositivas.pptx")
    primera = Presentation(ruta).slides[0]
    textos = [linea for f in primera.shapes if f.has_text_frame for linea in f.text_frame.text.splitlines()]
    assert "www.linkedin.com/in/jdavelouis" in textos and "info@futurasoft.net" in textos
    franjas = [f for f in primera.shapes if f.shape_type == 1 and f.top == 0]   # 1: autoforma
    assert franjas and str(franjas[0].fill.fore_color.rgb) == estilo.PALETA_SOBRIA["principal"]


def test_un_curso_anterior_toma_el_docente_por_defecto(curso):
    ficha = almacen.cargar(curso, "curso")
    del ficha["campos"]["portada.docente"]
    almacen.guardar(curso, "curso", None, ficha)
    assert "portada.docente" not in almacen.ruta_ficha(curso, "curso").read_text(encoding="utf-8")
    assert almacen.valor(almacen.cargar(curso, "curso"), "portada.docente") == "José Fernando Davelouis Mendía"


# ---------- Valores por defecto y propuesta desde las fuentes ----------

def test_tipografia_y_estilo_muestran_su_valor_por_defecto(curso):
    pagina = cliente.get("/cursos/scrum/ficha").text
    assert '<option value="" selected>Calibri (por defecto)</option>' in pagina
    assert '<option value="" selected>Profesional (por defecto): Directo y sobrio. Es el diseño de siempre.</option>' in pagina
    assert pagina.count('value="profesional"') == 0   # la opción por defecto no se repite
    assert pagina.count('value="Calibri"') == 0


def test_la_propuesta_conoce_las_opciones_de_los_controles(curso):
    from app.fichas import flujo
    pedido = flujo.pedido_de_propuesta(curso, "curso", None, almacen.cargar(curso, "curso"))
    assert "curso.idioma: Idioma del material [usa una de estas, tal cual: Español, Inglés, Portugués, Francés" in pedido
    assert "herramientas.adjuntos" in pedido and "[usa una de estas, tal cual: sí, no, no sé]" in pedido
    assert "[usa las que correspondan, separadas por comas: lectura, diapositivas" in pedido


def test_una_ficha_propuesta_se_ve_en_los_controles_y_guardar_no_borra_la_marca(curso):
    """La ficha que se llena sola desde las fuentes (decisión 17) se muestra bien en los controles. Guardar desde la
    página, sin tocar esos controles, conserva «Propuesto» y la fuente."""
    from app.fichas import flujo
    ficha = almacen.cargar(curso, "curso")
    flujo.aplicar_propuesta(curso, ficha, {"campos": [
        {"id": "curso.idioma", "valor": "español", "origen": "deducido", "literal": False},
        {"id": "herramientas.neutral", "valor": "Sí, el material es neutral", "origen": "deducido", "literal": False},
        {"id": "herramientas.adjuntos", "valor": "no sé", "origen": "deducido", "literal": False},
        {"id": "materiales.por_sesion", "valor": "Lectura y diapositivas", "origen": "deducido", "literal": False},
    ], "conflictos": []})
    almacen.guardar(curso, "curso", None, ficha)
    pagina = cliente.get("/cursos/scrum/ficha").text
    assert '<option value="Español" selected>' in pagina
    assert 'name="herramientas.neutral" value="sí" checked' in pagina
    assert 'name="herramientas.adjuntos" value="no sé" checked' in pagina
    assert 'name="materiales.por_sesion" value="lectura" checked' in pagina
    assert 'name="materiales.por_sesion" value="laboratorio" >' in pagina   # sin marcar
    assert pagina.count('<span class="marca">Propuesto</span>') == 4
    # Lo que el navegador envía al pulsar «Guardar» sin tocar nada
    cliente.post("/cursos/scrum/ficha", data={
        "accion": "guardar", "curso.idioma": "Español", "herramientas.neutral": "sí",
        "herramientas.adjuntos": "no sé", "materiales.por_sesion.casillas": "1",
        "materiales.por_sesion": ["lectura", "diapositivas"]})
    campos = almacen.cargar(curso, "curso")["campos"]
    for campo_id in ("curso.idioma", "herramientas.neutral", "herramientas.adjuntos", "materiales.por_sesion"):
        assert campos[campo_id]["estado"] == "propuesto", campo_id
    assert campos["curso.idioma"]["valor"] == "español"
