"""Programa verificador de la skill material-de-clase.

Revisa sin ayuda de IA lo que se puede comprobar de forma mecánica (SKILL.md, "Preparación ·
El programa verificador"). Extrae el texto de los archivos del material, lo separa en
oraciones y crea el archivo de verificación con cuatro hojas, igual al modelo
S1_Laboratorio_Verificacion.xlsx.

Uso:
    python verificar.py --config verificacion.json --salida S1_Lectura_Verificacion.xlsx S1_Lectura.docx [...]

La última línea impresa es:  Oraciones: N · Fallas: F · Avisos: A
Termina con código 0 si no hay fallas, 1 si hay fallas y 2 si la configuración tiene errores.

Formato de verificacion.json
----------------------------
Las rutas son relativas a la carpeta donde está verificacion.json. Solo "sesion" y
"material" son obligatorios.

{
  "sesion": 1,
  "material": "Lectura",

  "fuentes": [                                   # fuentes del curso ya convertidas a texto
    {
      "nombre": "Guía de Scrum, capítulo 1",
      "texto": "../../fuentes_texto/guia.pdf.jsonl",   # un pasaje por línea: {"texto", "ubicacion", ...}
      "norma": true,
      "referencias": {                           # cómo se citan sus cláusulas, controles o láminas
        "patron": "lámina (\\d+)",               # en el material; el grupo 1 es el número
        "indice": "numero_impreso"               # "numero_impreso", "pagina" o "patron"
        # con "indice": "patron", agregar "patron_indice": "^(\\d+(?:\\.\\d+)*)\\s+(.+)$"
        # (grupo 1: número, grupo 2: título)
      }
    }
  ],
  "textos_permitidos": ["../../ficha_del_curso.md", "ficha_de_la_sesion.md"],
                                                 # otros textos donde una cita puede aparecer tal cual
  "datos_fijos": [{"etiqueta": "Fecha de referencia", "valor": "lunes 5 de octubre de 2026"}],
  "vocabulario": [{"concepto": "Manifiesto Ágil", "variantes": ["Agile", "manifiesto agile"]}],
  "nunca_se_incluye": ["objetivos de aprendizaje", "requisitos previos", "glosario"],
  "permitidos": ["notas del profesor", "puntos de historia"],   # frases que no son tiempos ni puntajes
  "carpeta_practica": "materiales/practica",     # o null si el material no tiene archivos de práctica
  "limites": {
    "paginas_max": 6,                            # Word: páginas contadas por el Word instalado
    "diapositivas_min": 14, "diapositivas_max": 18, "notas_en_todas": true,
    "conteos": [{"nombre": "ejercicios", "patron": "^Ejercicio \\d+", "min": 3, "max": 5}],
    "titulos_con": [{"patron": "^Mejora el resultado", "debe_contener": "(opcional)"}]
  },
  "oracion_larga": 25,                           # más palabras que esto es un AVISO
  "cita_min_palabras": 5,                        # las citas más cortas no se buscan en las fuentes
  "separador_decimal": ",",
  "listas": {                                    # por omisión, las de config/ del proyecto
    "lenguaje_ia": "../../../config/lenguaje_ia.txt",
    "relleno": "../../../config/relleno.txt",
    "imprecisas": "../../../config/imprecisas.txt"
  }
}
"""

from __future__ import annotations

import argparse
import ast
import hashlib
import html.parser
import json
import operator
import os
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

RAIZ_PROYECTO = Path(__file__).resolve().parents[4]
CARPETA_LISTAS = RAIZ_PROYECTO / "config"

FALLA = "FALLA"
AVISO = "AVISO"

# ---------- Formato del Excel: copia de S1_Laboratorio_Verificacion.xlsx ----------

HOJAS = {
    "Oraciones": (
        ["N", "Diapositiva o sección", "Parte", "Oración", "Tipo de oración",
         "Pasaje de la fuente", "Fuente", "Veredicto", "Revisor independiente"],
        [5, 26, 12, 60, 16, 60, 22, 13, 22],
    ),
    "Hallazgos": (
        ["Nivel", "Diapositiva", "Oración", "Regla", "Detalle", "Cómo se resolvió"],
        [9, 11, 60, 18, 40, 60],
    ),
    "Datos repetidos": (
        ["Término o cifra", "Apariciones", "Oraciones"],
        [26, 12, 110],
    ),
    "Segunda pasada": (
        ["Bloque", "¿Qué puede hacer el alumno con esto?", "¿Qué dato necesita y dónde está?",
         "¿Qué oración tiene dos lecturas?", "¿Qué decide si el caso no sale como el ejemplo?"],
        [22, 45, 45, 45, 45],
    ),
}
COLOR_ENCABEZADO = "0F4C5C"
LIMITE_CELDA = 32000


class ErrorDeConfiguracion(ValueError):
    """verificacion.json no tiene el formato esperado."""


# ---------- Datos ----------

@dataclass
class Oracion:
    n: int
    archivo: str
    seccion: str
    parte: str
    texto: str
    fuente: str = ""

    @property
    def huella(self) -> str:
        """Cambia si cambia una sola letra. Sirve para validar de nuevo solo lo que cambió."""
        return hashlib.sha256(f"{self.seccion}\n{self.texto}".encode("utf-8")).hexdigest()[:16]


@dataclass
class Hallazgo:
    nivel: str
    seccion: str
    oracion: str
    regla: str
    detalle: str
    n: int | None = None


@dataclass
class Resultado:
    oraciones: list[Oracion] = field(default_factory=list)
    hallazgos: list[Hallazgo] = field(default_factory=list)
    datos_repetidos: list[tuple[str, int, list[str]]] = field(default_factory=list)

    @property
    def fallas(self) -> int:
        return sum(1 for h in self.hallazgos if h.nivel == FALLA)

    @property
    def avisos(self) -> int:
        return sum(1 for h in self.hallazgos if h.nivel == AVISO)

    def resumen(self) -> str:
        return f"Oraciones: {len(self.oraciones)} · Fallas: {self.fallas} · Avisos: {self.avisos}"


# ---------- Texto ----------

def normalizar(texto: str) -> str:
    """Minúsculas y sin tildes, para buscar frases sin distinguir mayúsculas ni tildes."""
    descompuesto = unicodedata.normalize("NFD", texto.lower())
    return "".join(c for c in descompuesto if unicodedata.category(c) != "Mn")


def literal(texto: str) -> str:
    """Normaliza solo espacios, comillas y guiones, para comparar citas tal cual."""
    reemplazos = {"“": '"', "”": '"', "«": '"', "»": '"', "‘": "'", "’": "'",
                  "–": "-", "—": "-", "­": "", " ": " "}
    for viejo, nuevo in reemplazos.items():
        texto = texto.replace(viejo, nuevo)
    return " ".join(texto.split())


def palabras(texto: str) -> int:
    return len(re.findall(r"[\wáéíóúüñÁÉÍÓÚÜÑ]+(?:[.,][\d]+)*", texto))


def patron_frase(frase: str) -> re.Pattern:
    return re.compile(r"(?<!\w)" + re.escape(normalizar(frase.strip())) + r"(?!\w)")


ABREVIATURAS = ("p. ej.", "etc.", "sr.", "sra.", "srta.", "dr.", "dra.", "núm.", "pág.", "págs.",
                "aprox.", "art.", "fig.", "ej.", "vol.", "cap.", "ud.", "uds.", "e. g.", "i. e.", "n.º")
FIN_DE_ORACION = re.compile(r"[.!?…]+[\"'»”)\]]*\s+")


def separar_oraciones(texto: str) -> list[str]:
    """Separa un párrafo en oraciones. Una oración termina en . ! ? o … seguidos de espacio y
    de una mayúscula, un número o un signo de apertura. No corta abreviaturas comunes."""
    texto = " ".join(texto.split())
    oraciones, inicio = [], 0
    for marca in FIN_DE_ORACION.finditer(texto):
        siguiente = texto[marca.end():marca.end() + 1]
        if not siguiente or not (siguiente.isupper() or siguiente.isdigit() or siguiente in "¿¡«\"“(—-•"):
            continue
        cola = texto[inicio:marca.start() + 1].lower()
        if cola.endswith(ABREVIATURAS):
            continue
        oraciones.append(texto[inicio:marca.end()].strip())
        inicio = marca.end()
    resto = texto[inicio:].strip()
    if resto:
        oraciones.append(resto)
    return oraciones


# ---------- Extracción ----------

def _es_titulo(parrafo: Paragraph, nivel_max: int = 1) -> bool:
    nombre = (parrafo.style.name or "") if parrafo.style is not None else ""
    coincidencia = re.match(r"(?:Heading|Título|Titulo)\s*(\d)", nombre)
    return bool(coincidencia) and int(coincidencia.group(1)) <= nivel_max


def _es_algun_titulo(parrafo: Paragraph) -> bool:
    return _es_titulo(parrafo, nivel_max=3)


def extraer_docx(ruta: Path) -> tuple[list[tuple[str, str, str]], list[str]]:
    """Devuelve (sección, parte, párrafo) en el orden del documento, y los títulos."""
    documento = Document(ruta)
    filas, titulos = [], []
    seccion = "inicio"
    vistos_en_encabezado = set()
    for parte_doc in documento.sections:
        for bloque in (parte_doc.header, parte_doc.footer):
            for parrafo in bloque.paragraphs:
                if parrafo.text.strip() and parrafo.text not in vistos_en_encabezado:
                    vistos_en_encabezado.add(parrafo.text)
                    filas.append(("encabezado y pie", "texto", parrafo.text))
    for hijo in documento.element.body.iterchildren():
        if hijo.tag == qn("w:p"):
            parrafo = Paragraph(hijo, documento)
            if _es_titulo(parrafo):
                seccion = parrafo.text.strip() or seccion
            if _es_algun_titulo(parrafo) and parrafo.text.strip():
                titulos.append(parrafo.text.strip())
            filas.append((seccion, "texto", parrafo.text))
        elif hijo.tag == qn("w:tbl"):
            for fila in Table(hijo, documento).rows:
                vistas = []
                for celda in fila.cells:
                    if celda._tc in vistas:  # celdas combinadas se repiten
                        continue
                    vistas.append(celda._tc)
                    for parrafo in celda.paragraphs:
                        filas.append((seccion, "tabla o recuadro", parrafo.text))
    return [f for f in filas if f[2].strip()], titulos


def _figuras(figuras):
    for figura in figuras:
        if figura.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _figuras(figura.shapes)
        else:
            yield figura


def extraer_pptx(ruta: Path) -> tuple[list[tuple[str, str, str]], list[str], dict]:
    presentacion = Presentation(ruta)
    filas, titulos, datos = [], [], {"diapositivas": 0, "sin_notas": []}
    for numero, lamina in enumerate(presentacion.slides, start=1):
        datos["diapositivas"] = numero
        seccion = f"diapositiva {numero}"
        if lamina.shapes.title is not None and lamina.shapes.title.text.strip():
            titulos.append(lamina.shapes.title.text.strip())
        for figura in _figuras(lamina.shapes):
            if figura.has_text_frame:
                for parrafo in figura.text_frame.paragraphs:
                    filas.append((seccion, "texto", "".join(r.text for r in parrafo.runs)))
            elif getattr(figura, "has_table", False) and figura.has_table:
                for fila in figura.table.rows:
                    for celda in fila.cells:
                        filas.append((seccion, "tabla o recuadro", celda.text))
        notas = lamina.notes_slide.notes_text_frame.text if lamina.has_notes_slide else ""
        if notas.strip():
            for parrafo in notas.split("\n"):
                filas.append((seccion, "notas", parrafo))
        else:
            datos["sin_notas"].append(numero)
    return [f for f in filas if f[2].strip()], titulos, datos


def extraer_xlsx(ruta: Path) -> list[tuple[str, str, str]]:
    valores = openpyxl.load_workbook(ruta, data_only=True)
    formulas = openpyxl.load_workbook(ruta, data_only=False)
    filas = []
    for hoja in valores.worksheets:
        for fila_v, fila_f in zip(hoja.iter_rows(), formulas[hoja.title].iter_rows()):
            for celda, celda_f in zip(fila_v, fila_f):
                valor = celda.value if celda.value is not None else celda_f.value
                if valor is not None and str(valor).strip():
                    filas.append((hoja.title, "celda", str(valor)))
    return filas


class _LectorHtml(html.parser.HTMLParser):
    BLOQUES = {"p", "li", "h1", "h2", "h3", "h4", "td", "th", "button", "label", "summary",
               "figcaption", "caption", "dt", "dd", "legend", "option", "title"}

    def __init__(self):
        super().__init__()
        self.filas, self.titulos = [], []
        self.seccion, self.pila, self.texto, self.datos_json = "inicio", [], [], []
        self.en_script = None

    def handle_starttag(self, etiqueta, atributos):
        atributos = dict(atributos)
        if etiqueta in ("script", "style"):
            es_datos = etiqueta == "script" and atributos.get("type") == "application/json"
            self.en_script = "datos" if es_datos else "ignorar"
            return
        if etiqueta in self.BLOQUES:
            self._cerrar()
            self.pila.append(etiqueta)

    def handle_endtag(self, etiqueta):
        if etiqueta in ("script", "style"):
            self.en_script = None
            return
        if etiqueta in self.BLOQUES:
            texto = self._cerrar()
            if etiqueta in ("h1", "h2") and texto:
                self.seccion = texto
            if etiqueta in ("h1", "h2", "h3") and texto:
                self.titulos.append(texto)

    def handle_data(self, datos):
        if self.en_script == "datos":
            self.datos_json.append(datos)
        elif self.en_script is None:
            self.texto.append(datos)

    def _cerrar(self) -> str:
        texto = " ".join("".join(self.texto).split())
        self.texto = []
        if texto:
            parte = "tabla o recuadro" if self.pila and self.pila[-1] in ("td", "th") else "texto"
            self.filas.append((self.seccion, parte, texto))
        return texto


def _textos_json(valor, ruta="datos"):
    if isinstance(valor, str):
        yield ruta, valor
    elif isinstance(valor, list):
        for i, v in enumerate(valor, start=1):
            yield from _textos_json(v, f"{ruta} {i}")
    elif isinstance(valor, dict):
        for k, v in valor.items():
            yield from _textos_json(v, ruta)


def extraer_html(ruta: Path) -> tuple[list[tuple[str, str, str]], list[str]]:
    lector = _LectorHtml()
    lector.feed(ruta.read_text(encoding="utf-8"))
    lector._cerrar()
    filas = list(lector.filas)
    for bloque in lector.datos_json:
        try:
            datos = json.loads(bloque)
        except json.JSONDecodeError:
            continue
        filas.extend(("datos de la página", "texto", t) for _, t in _textos_json(datos))
    return filas, lector.titulos


def extraer_texto_plano(ruta: Path) -> list[tuple[str, str, str]]:
    parrafos = re.split(r"\n\s*\n", ruta.read_text(encoding="utf-8"))
    return [("inicio", "texto", p) for p in parrafos if p.strip()]


# ---------- Configuración ----------

def _leer_lista(ruta: Path) -> list[str]:
    if not ruta.exists():
        raise ErrorDeConfiguracion(f"No existe la lista {ruta}.")
    lineas = ruta.read_text(encoding="utf-8").splitlines()
    return [l.strip() for l in lineas if l.strip() and not l.lstrip().startswith("#")]


@dataclass
class Fuente:
    nombre: str
    pasajes: list[dict]
    texto_continuo: str
    norma: bool
    patron_referencia: re.Pattern | None
    indice: dict[str, str]


@dataclass
class Configuracion:
    sesion: int
    material: str
    carpeta: Path
    fuentes: list[Fuente]
    textos_permitidos: str
    datos_fijos: list[dict]
    vocabulario: list[dict]
    nunca_se_incluye: list[str]
    permitidos: list[str]
    carpeta_practica: Path | None
    limites: dict
    oracion_larga: int
    cita_min_palabras: int
    separador_decimal: str
    lenguaje_ia: list[str]
    relleno: list[str]
    imprecisas: list[str]


def _indice_de_fuente(nombre: str, pasajes: list[dict], referencias: dict) -> dict[str, str]:
    tipo = referencias.get("indice", "numero_impreso")
    indice: dict[str, str] = {}
    if tipo in ("numero_impreso", "pagina"):
        for pasaje in pasajes:
            numero = pasaje.get(tipo)
            if numero is not None and str(numero) not in indice:
                indice[str(numero)] = pasaje["texto"]
    elif tipo == "patron":
        if "patron_indice" not in referencias:
            raise ErrorDeConfiguracion(f"La fuente «{nombre}» usa indice «patron» sin «patron_indice».")
        patron = re.compile(referencias["patron_indice"])
        for pasaje in pasajes:
            coincidencia = patron.match(pasaje["texto"])
            if coincidencia:
                indice.setdefault(coincidencia.group(1), coincidencia.group(2).strip())
    else:
        raise ErrorDeConfiguracion(f"La fuente «{nombre}» tiene un indice desconocido: «{tipo}».")
    return indice


def _ruta(carpeta: Path, relativa: str) -> Path:
    """Ruta sin «..»: en Windows, una ruta de más de 260 caracteres no se puede abrir."""
    return Path(os.path.normpath(carpeta / relativa))


def leer_configuracion(ruta: Path) -> Configuracion:
    try:
        datos = json.loads(ruta.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise ErrorDeConfiguracion(f"No se pudo leer {ruta.name}: {error}") from error
    for clave in ("sesion", "material"):
        if clave not in datos:
            raise ErrorDeConfiguracion(f"Falta «{clave}» en {ruta.name}.")
    carpeta = ruta.resolve().parent

    fuentes = []
    for f in datos.get("fuentes", []):
        if "nombre" not in f or "texto" not in f:
            raise ErrorDeConfiguracion("Cada fuente necesita «nombre» y «texto».")
        ruta_texto = _ruta(carpeta, f["texto"])
        if not ruta_texto.exists():
            raise ErrorDeConfiguracion(f"No existe el texto de la fuente «{f['nombre']}»: {ruta_texto}.")
        pasajes = [json.loads(l) for l in ruta_texto.read_text(encoding="utf-8").splitlines() if l.strip()]
        referencias = f.get("referencias")
        fuentes.append(Fuente(
            nombre=f["nombre"],
            pasajes=pasajes,
            texto_continuo=literal(" ".join(p["texto"] for p in pasajes)),
            norma=bool(f.get("norma", False)),
            patron_referencia=re.compile(referencias["patron"], re.I) if referencias else None,
            indice=_indice_de_fuente(f["nombre"], pasajes, referencias) if referencias else {},
        ))

    permitidos_texto = []
    for relativa in datos.get("textos_permitidos", []):
        ruta_permitida = _ruta(carpeta, relativa)
        if not ruta_permitida.exists():
            raise ErrorDeConfiguracion(f"No existe el texto permitido {ruta_permitida}.")
        permitidos_texto.append(ruta_permitida.read_text(encoding="utf-8"))

    listas = datos.get("listas", {})
    def lista(nombre):
        return _leer_lista(_ruta(carpeta, listas[nombre]) if nombre in listas else CARPETA_LISTAS / f"{nombre}.txt")

    for dato in datos.get("datos_fijos", []):
        if not dato.get("etiqueta") or not dato.get("valor"):
            raise ErrorDeConfiguracion("Cada dato fijo necesita «etiqueta» y «valor».")
    for concepto in datos.get("vocabulario", []):
        if not concepto.get("concepto"):
            raise ErrorDeConfiguracion("Cada concepto del vocabulario necesita «concepto».")

    practica = datos.get("carpeta_practica")
    return Configuracion(
        sesion=int(datos["sesion"]),
        material=str(datos["material"]),
        carpeta=carpeta,
        fuentes=fuentes,
        textos_permitidos=literal(" ".join(permitidos_texto)),
        datos_fijos=datos.get("datos_fijos", []),
        vocabulario=datos.get("vocabulario", []),
        nunca_se_incluye=datos.get("nunca_se_incluye", []),
        permitidos=datos.get("permitidos", ["notas del profesor"]),
        carpeta_practica=_ruta(carpeta, practica) if practica else None,
        limites=datos.get("limites", {}),
        oracion_larga=int(datos.get("oracion_larga", 25)),
        cita_min_palabras=int(datos.get("cita_min_palabras", 5)),
        separador_decimal=datos.get("separador_decimal", ","),
        lenguaje_ia=lista("lenguaje_ia"),
        relleno=lista("relleno"),
        imprecisas=lista("imprecisas"),
    )


# ---------- Revisiones por oración ----------

TIEMPOS_Y_PUNTAJES = [
    (FALLA, "tiempo", re.compile(r"\bminutos?\b|\d+\s*min\b|\bcronometr|\bcuenta regresiva\b|\btemporizador")),
    (FALLA, "tiempo", re.compile(r"\d+\s*(?:segundos?|seg)\b")),
    (AVISO, "tiempo", re.compile(r"\d+\s*(?:horas?|h)\b")),
    (FALLA, "puntaje", re.compile(r"\bpuntajes?\b|\bpuntuacion|\d+\s*(?:puntos?|pts?)\b|\bcalificacion|\bcalificar\b")),
    (FALLA, "puntaje", re.compile(r"\bnotas?\s+(?:final|finales|maxima|minima|aprobatoria|del alumno|de la evaluacion)\b")),
    (AVISO, "porcentaje", re.compile(r"%|\bpor\s*ciento\b|\bporcentajes?\b")),
]
EMOJI = re.compile("[\U0001F000-\U0001FAFF☀-➿\U0001F1E6-\U0001F1FF⭐⭕⌚⌛⏩-⏺️]")
CITA = re.compile(r"«([^»]+)»|“([^”]+)”|\"([^\"]+)\"")
ARCHIVO = re.compile(r"\bS(\d+)_([EP])(\d+)_([A-Za-z0-9-]+)\.([A-Za-z0-9]{2,5})\b")


def _texto_sin_permitidos(normal: str, permitidos: list[str]) -> str:
    for frase in permitidos:
        normal = patron_frase(frase).sub(" ", normal)
    return normal


def _dentro_de_cita(texto: str, inicio: int, fin: int) -> bool:
    return any(c.start() <= inicio and fin <= c.end() for c in CITA.finditer(normalizar(texto)))


UNIDADES = ["cero", "uno", "dos", "tres", "cuatro", "cinco", "seis", "siete", "ocho", "nueve", "diez",
            "once", "doce", "trece", "catorce", "quince", "dieciseis", "diecisiete", "dieciocho",
            "diecinueve", "veinte", "veintiuno", "veintidos", "veintitres", "veinticuatro",
            "veinticinco", "veintiseis", "veintisiete", "veintiocho", "veintinueve"]
DECENAS = {30: "treinta", 40: "cuarenta", 50: "cincuenta", 60: "sesenta", 70: "setenta",
           80: "ochenta", 90: "noventa"}


def _numeros_en_palabras() -> dict[str, int]:
    tabla = {palabra: valor for valor, palabra in enumerate(UNIDADES)}
    tabla["un"] = tabla["una"] = 1
    for decena, palabra in DECENAS.items():
        tabla[palabra] = decena
        for unidad in range(1, 10):
            tabla[f"{palabra} y {UNIDADES[unidad]}"] = decena + unidad
    tabla["cien"] = 100
    return tabla


NUMEROS_EN_PALABRAS = _numeros_en_palabras()
PATRON_NUMERO_EN_PALABRAS = re.compile(
    r"(?<!\w)(" + "|".join(sorted((re.escape(p) for p in NUMEROS_EN_PALABRAS), key=len, reverse=True)) + r")(?!\w)"
)
OPERADORES = {"+": operator.add, "-": operator.sub, "*": operator.mul, "/": operator.truediv}


def _a_numero(texto: str, decimal: str) -> float:
    if decimal == ",":
        texto = texto.replace(".", "").replace(",", ".")
    else:
        texto = texto.replace(",", "")
    return float(texto)


def _evaluar(expresion: str) -> float:
    def valor(nodo):
        if isinstance(nodo, ast.Constant) and isinstance(nodo.value, (int, float)):
            return nodo.value
        if isinstance(nodo, ast.BinOp) and type(nodo.op) in (ast.Add, ast.Sub, ast.Mult, ast.Div):
            simbolo = {ast.Add: "+", ast.Sub: "-", ast.Mult: "*", ast.Div: "/"}[type(nodo.op)]
            return OPERADORES[simbolo](valor(nodo.left), valor(nodo.right))
        raise ValueError("expresión no admitida")
    return valor(ast.parse(expresion, mode="eval").body)


def revisar_operaciones(texto: str, decimal: str) -> list[str]:
    """Busca operaciones escritas como «3 × 4 = 12» o «tres más cuatro es igual a siete»."""
    t = normalizar(texto)
    t = PATRON_NUMERO_EN_PALABRAS.sub(lambda m: str(NUMEROS_EN_PALABRAS[m.group(1)]), t)
    t = re.sub(r"(?<=\d)\s+mas\s+(?=\d)", " + ", t)
    t = re.sub(r"(?<=\d)\s+menos\s+(?=\d)", " - ", t)
    t = re.sub(r"(?<=\d)\s+(?:por|multiplicado por)\s+(?=\d)", " × ", t)
    t = re.sub(r"(?<=\d)\s+(?:entre|dividido entre|dividido por)\s+(?=\d)", " ÷ ", t)
    t = re.sub(r"\s+(?:es igual a|son igual a|da como resultado)\s+", " = ", t)
    numero = r"\d+(?:[.,]\d+)*"
    problemas = []

    porcentaje = re.compile(rf"({numero})\s*(?:%|por ciento)\s+de\s+({numero})\s*(?:=|es|son|da)\s*({numero})")
    for m in porcentaje.finditer(t):
        p, base, escrito = (_a_numero(x, decimal) for x in m.groups())
        if abs(p * base / 100 - escrito) > 0.005 * max(1, abs(escrito)):
            problemas.append(f"«{m.group(0)}»: da {p * base / 100:g}, no {escrito:g}")

    operacion = re.compile(rf"({numero}(?:\s*[+\-−×x*÷/]\s*{numero})+)\s*=\s*({numero})")
    for m in operacion.finditer(t):
        izquierda = m.group(1)
        expresion = re.sub(numero, lambda n: repr(_a_numero(n.group(0), decimal)), izquierda)
        expresion = expresion.replace("×", "*").replace("x", "*").replace("÷", "/").replace("−", "-")
        try:
            calculado = _evaluar(expresion)
        except (ValueError, ZeroDivisionError, SyntaxError):
            continue
        escrito = _a_numero(m.group(2), decimal)
        if abs(calculado - escrito) > 0.005 * max(1, abs(escrito)):
            problemas.append(f"«{m.group(0)}»: da {calculado:g}, no {escrito:g}")
    return problemas


def revisar_oracion(o: Oracion, cfg: Configuracion, hallazgos: list[Hallazgo]) -> None:
    def agregar(nivel, regla, detalle):
        hallazgos.append(Hallazgo(nivel, o.seccion, o.texto, regla, detalle, o.n))

    normal = normalizar(o.texto)
    sin_permitidos = _texto_sin_permitidos(normal, cfg.permitidos)

    for nivel, regla, patron in TIEMPOS_Y_PUNTAJES:
        m = patron.search(sin_permitidos)
        if m:
            agregar(nivel, regla, f"dice «{m.group(0).strip()}»")
    if EMOJI.search(o.texto):
        agregar(FALLA, "emoji", f"tiene «{EMOJI.search(o.texto).group(0)}»")

    for frase in cfg.nunca_se_incluye:
        if patron_frase(frase).search(normal):
            agregar(FALLA, "lo que nunca se incluye", f"dice «{frase}»")

    for concepto in cfg.vocabulario:
        sin_concepto = patron_frase(concepto["concepto"]).sub(" ", normal)
        for variante in concepto.get("variantes", []):
            if patron_frase(variante).search(sin_concepto):
                agregar(FALLA, "variante de un concepto",
                        f"usa «{variante}»; el nombre único es «{concepto['concepto']}»")

    for dato in cfg.datos_fijos:
        patron = re.compile(r"(?<!\w)" + re.escape(normalizar(dato["etiqueta"])) + r"\s*(?:\bes el\b|\bes la\b|\bes\b|:|=)\s*(.+)")
        m = patron.search(normal)
        if m and not m.group(1).startswith(normalizar(dato["valor"])):
            agregar(FALLA, "dato fijo", f"«{dato['etiqueta']}» debe ser «{dato['valor']}»")

    for m in CITA.finditer(o.texto):
        cita = next(g for g in m.groups() if g is not None)
        if palabras(cita) < cfg.cita_min_palabras:
            continue
        buscada = literal(cita).strip(" .,;:")
        if buscada not in cfg.textos_permitidos and not any(buscada in f.texto_continuo for f in cfg.fuentes):
            agregar(FALLA, "cita sin fuente", f"«{cita}» no aparece tal cual en ninguna fuente")

    for fuente in cfg.fuentes:
        if fuente.patron_referencia is None:
            continue
        for m in fuente.patron_referencia.finditer(o.texto):
            numero = m.group(1)
            if numero in fuente.indice:
                o.fuente = f"{fuente.nombre}, {m.group(0)}: {fuente.indice[numero]}"
            elif not any(numero in f.indice for f in cfg.fuentes if f.patron_referencia):
                agregar(FALLA, "cláusula inexistente", f"«{m.group(0)}» no existe en «{fuente.nombre}»")

    for problema in revisar_operaciones(o.texto, cfg.separador_decimal):
        agregar(FALLA, "operación", problema)

    for m in ARCHIVO.finditer(o.texto):
        nombre = m.group(0)
        if int(m.group(1)) != cfg.sesion:
            agregar(FALLA, "nombre de archivo", f"«{nombre}» no es de la sesión {cfg.sesion}")
        if cfg.carpeta_practica is None or not (cfg.carpeta_practica / nombre).exists():
            agregar(FALLA, "archivo inexistente", f"«{nombre}» no está en la carpeta de práctica")

    total = palabras(o.texto)
    if total > cfg.oracion_larga:
        agregar(AVISO, "oración larga", f"{total} palabras")

    for frase in cfg.relleno:
        if patron_frase(frase).search(normal):
            agregar(AVISO, "relleno", f"dice «{frase}»")
    for frase in cfg.imprecisas:
        m = patron_frase(frase).search(normal)
        if m:
            agregar(AVISO, "palabra imprecisa", frase)
    for frase in cfg.lenguaje_ia:
        m = patron_frase(frase).search(normal)
        if m:
            nivel = AVISO if _dentro_de_cita(o.texto, m.start(), m.end()) else FALLA
            agregar(nivel, "lenguaje de IA", f"dice «{frase}»")


# ---------- Revisiones del material completo ----------

def revisar_archivos_de_practica(cfg: Configuracion, oraciones: list[Oracion], hallazgos: list[Hallazgo]) -> None:
    if cfg.carpeta_practica is None or not cfg.carpeta_practica.exists():
        return
    mencionados = {m.group(0) for o in oraciones for m in ARCHIVO.finditer(o.texto)}
    formato = re.compile(rf"^S{cfg.sesion}_[EP]\d+_[a-z0-9]+(?:-[a-z0-9]+)*\.[A-Za-z0-9]{{2,5}}$")
    for ruta in sorted(cfg.carpeta_practica.iterdir()):
        if not ruta.is_file():
            continue
        if not formato.match(ruta.name):
            hallazgos.append(Hallazgo(FALLA, "carpeta de práctica", ruta.name, "nombre de archivo",
                                      f"no sigue el formato S{cfg.sesion}_E[ejercicio]_[descripcion-corta]"))
        if ruta.name not in mencionados:
            hallazgos.append(Hallazgo(FALLA, "carpeta de práctica", ruta.name, "archivo sin mencionar",
                                      "ningún texto del material lo menciona"))


def revisar_limites(cfg: Configuracion, ruta: Path, titulos: list[str], datos_pptx: dict | None,
                    hallazgos: list[Hallazgo], contar_paginas) -> None:
    limites = cfg.limites
    def agregar(nivel, regla, detalle):
        hallazgos.append(Hallazgo(nivel, ruta.name, "", regla, detalle))

    if ruta.suffix.lower() == ".docx" and "paginas_max" in limites:
        try:
            paginas = contar_paginas(ruta)
        except Exception as error:
            agregar(AVISO, "páginas", f"no se pudieron contar las páginas: {error}")
        else:
            if paginas > limites["paginas_max"]:
                agregar(FALLA, "páginas", f"tiene {paginas} páginas; el máximo es {limites['paginas_max']}")

    if datos_pptx is not None:
        total = datos_pptx["diapositivas"]
        if total < limites.get("diapositivas_min", 0) or total > limites.get("diapositivas_max", 10**6):
            agregar(FALLA, "diapositivas", f"tiene {total} diapositivas; deben ser de "
                    f"{limites.get('diapositivas_min', 0)} a {limites.get('diapositivas_max', '—')}")
        if limites.get("notas_en_todas"):
            for numero in datos_pptx["sin_notas"]:
                agregar(FALLA, "notas del profesor", f"la diapositiva {numero} no tiene notas")

    for conteo in limites.get("conteos", []):
        patron = re.compile(conteo["patron"], re.I)
        total = sum(1 for t in titulos if patron.search(t))
        if total < conteo.get("min", 0) or total > conteo.get("max", 10**6):
            agregar(FALLA, conteo["nombre"], f"tiene {total}; deben ser de {conteo.get('min', 0)} a {conteo.get('max', '—')}")

    for regla in limites.get("titulos_con", []):
        patron = re.compile(regla["patron"], re.I)
        for titulo in titulos:
            if patron.search(titulo) and regla["debe_contener"] not in titulo:
                hallazgos.append(Hallazgo(FALLA, ruta.name, titulo, "título",
                                          f"debe contener «{regla['debe_contener']}»"))


def listar_datos_repetidos(cfg: Configuracion, oraciones: list[Oracion]) -> list[tuple[str, int, list[str]]]:
    """Cada concepto, cada variante prohibida y cada cifra, con las oraciones donde aparece."""
    filas = []
    def buscar(termino: str, patron: re.Pattern):
        encontradas = [f"{o.seccion}: {o.texto}" for o in oraciones if patron.search(normalizar(o.texto))]
        filas.append((termino, len(encontradas), encontradas))

    for concepto in cfg.vocabulario:
        for variante in concepto.get("variantes", []):
            buscar(variante, patron_frase(variante))
        buscar(concepto["concepto"], patron_frase(concepto["concepto"]))

    cifras: dict[str, list[str]] = {}
    patron_cifra = re.compile(r"(?<![\w.,])\d+(?:[.,]\d+)*(?![\w])|" + PATRON_NUMERO_EN_PALABRAS.pattern)
    for o in oraciones:
        for m in patron_cifra.finditer(normalizar(o.texto)):
            cifras.setdefault(m.group(0), []).append(f"{o.seccion}: {o.texto}")
    for cifra, lista in cifras.items():
        if cifra in ("un", "una", "uno"):  # artículos, no cifras
            continue
        filas.append((cifra, len(lista), list(dict.fromkeys(lista))))
    return filas


# ---------- Programa ----------

def _contar_paginas_con_word(ruta: Path) -> int:
    sys.path.insert(0, str(RAIZ_PROYECTO))
    from app.office import contar_paginas_word  # Office solo se controla desde app/office.py
    return contar_paginas_word(ruta)


def verificar(ruta_config: Path, archivos: list[Path], contar_paginas=_contar_paginas_con_word) -> Resultado:
    cfg = leer_configuracion(ruta_config)
    resultado = Resultado()
    documentos = []
    for ruta in archivos:
        if not ruta.exists():
            raise ErrorDeConfiguracion(f"No existe el archivo del material {ruta}.")
        extension = ruta.suffix.lower()
        titulos, datos_pptx = [], None
        if extension == ".docx":
            filas, titulos = extraer_docx(ruta)
        elif extension == ".pptx":
            filas, titulos, datos_pptx = extraer_pptx(ruta)
        elif extension in (".xlsx", ".xlsm"):
            filas = extraer_xlsx(ruta)
        elif extension in (".html", ".htm"):
            filas, titulos = extraer_html(ruta)
        elif extension in (".txt", ".md"):
            filas = extraer_texto_plano(ruta)
        elif extension == ".zip":
            continue
        else:
            raise ErrorDeConfiguracion(f"Formato no admitido: {ruta.name}.")
        for seccion, parte, parrafo in filas:
            for texto in separar_oraciones(parrafo):
                resultado.oraciones.append(Oracion(
                    n=len(resultado.oraciones) + 1, archivo=ruta.name,
                    seccion=f"{ruta.name} · {seccion}", parte=parte, texto=texto,
                ))
        documentos.append((ruta, titulos, datos_pptx))

    for oracion in resultado.oraciones:
        revisar_oracion(oracion, cfg, resultado.hallazgos)
    for ruta, titulos, datos_pptx in documentos:
        revisar_limites(cfg, ruta, titulos, datos_pptx, resultado.hallazgos, contar_paginas)
    revisar_archivos_de_practica(cfg, resultado.oraciones, resultado.hallazgos)
    resultado.datos_repetidos = listar_datos_repetidos(cfg, resultado.oraciones)
    return resultado


def _hoja(libro, nombre: str):
    encabezados, anchos = HOJAS[nombre]
    hoja = libro.create_sheet(nombre)
    hoja.append(encabezados)
    for columna, (celda, ancho) in enumerate(zip(hoja[1], anchos), start=1):
        celda.font = Font(name="Arial", bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=COLOR_ENCABEZADO)
        hoja.column_dimensions[get_column_letter(columna)].width = ancho
    hoja.freeze_panes = "A2"
    return hoja


def _agregar_fila(hoja, valores) -> None:
    hoja.append([v[:LIMITE_CELDA] + "…" if isinstance(v, str) and len(v) > LIMITE_CELDA else v for v in valores])
    for celda in hoja[hoja.max_row]:
        celda.font = Font(name="Arial", size=10)
        celda.alignment = Alignment(wrap_text=True, vertical="top")


def escribir_excel(resultado: Resultado, salida: Path) -> Path:
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    oraciones = _hoja(libro, "Oraciones")
    for o in resultado.oraciones:
        _agregar_fila(oraciones, [o.n, o.seccion, o.parte, o.texto, None, None, o.fuente or None, None, None])
    hallazgos = _hoja(libro, "Hallazgos")
    for h in resultado.hallazgos:
        _agregar_fila(hallazgos, [h.nivel, h.seccion, h.oracion, h.regla, h.detalle, None])
    repetidos = _hoja(libro, "Datos repetidos")
    for termino, total, lista in resultado.datos_repetidos:
        _agregar_fila(repetidos, [termino, total, "\n".join(lista) or None])
    _hoja(libro, "Segunda pasada")
    salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(salida)
    return salida


def escribir_json(resultado: Resultado, salida: Path) -> Path:
    """Copia de los datos para AulaLista, con la huella de cada oración."""
    datos = {
        "resumen": resultado.resumen(),
        "oraciones": [{"n": o.n, "huella": o.huella, "archivo": o.archivo, "seccion": o.seccion,
                       "parte": o.parte, "texto": o.texto, "fuente": o.fuente} for o in resultado.oraciones],
        "hallazgos": [vars(h) for h in resultado.hallazgos],
    }
    salida.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
    return salida


def main(argumentos: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Verificador de materiales de clase.")
    parser.add_argument("--config", required=True, type=Path, help="ruta a verificacion.json")
    parser.add_argument("--salida", required=True, type=Path, help="archivo Excel de verificación")
    parser.add_argument("--json", type=Path, help="copia de los resultados en JSON")
    parser.add_argument("archivos", nargs="+", type=Path, help="archivos del material")
    args = parser.parse_args(argumentos)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        resultado = verificar(args.config, args.archivos)
    except ErrorDeConfiguracion as error:
        print(f"ERROR DE CONFIGURACIÓN: {error}")
        return 2
    escribir_excel(resultado, args.salida)
    if args.json:
        escribir_json(resultado, args.json)
    for h in resultado.hallazgos:
        donde = f"[{h.n}] " if h.n else ""
        print(f"{h.nivel} · {h.regla} · {donde}{h.seccion}: {h.detalle}")
    print(resultado.resumen())
    return 1 if resultado.fallas else 0


if __name__ == "__main__":
    sys.exit(main())
