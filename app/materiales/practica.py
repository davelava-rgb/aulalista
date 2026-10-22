"""Archivos de práctica del laboratorio: Excel para datos y Word para textos (SKILL.md, «Material 3»).

El redactor escribe el contenido de cada archivo como datos. Este programa crea el archivo real, con un diseño
sobrio y sin portada, y el .zip con todos. Las figuras de la guía se sacan después de estos archivos reales.
"""

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.materiales import docx as generador_docx
from app.materiales.contenido import ArchivoPractica, Laboratorio
from app.materiales.estilo import Identidad

EXTENSIONES = ("xlsx", "docx")
ANCHO_MAXIMO_COLUMNA = 45
# PLAN.md §0, decisión 60: la hoja con las opciones de cada lista y el mensaje fijo de Excel.
HOJA_DE_OPCIONES = "Opciones"
MENSAJE_DE_LA_LISTA = "Elige una opción de la lista."
# Lo que el alumno copia funciona tal cual: sin espacios por llenar (SKILL.md, «Ejercicios que se ejecutan»).
POR_LLENAR = re.compile(r"\[[^\]\n]{0,40}\]|_{3,}|<[^<>\n]{1,40}>|\bX{3,}\b|\{\{?[^{}\n]{0,40}\}?\}")
# En una fórmula, «<» y «>» son operadores: solo se buscan corchetes vacíos, guiones bajos y «XXX».
POR_LLENAR_EN_FORMULA = re.compile(r"\[\s*\]|_{3,}|\bX{3,}\b")
FIGURAS_POR_EJERCICIO = 1     # PLAN.md §0, decisión 33: una como máximo, y solo si hace falta
# Hace falta cuando un paso pide ubicar algo en el archivo: una columna, una celda, un rango, una fila o un apartado.
# «filas» sola no cuenta: «Revisa que la tabla tenga 10 filas» habla de la respuesta, no del archivo.
UBICA_ALGO = re.compile(r"(?i:\b(?:columnas?|celdas?|rangos?|apartados?|párrafos?|secci(?:ón|ones))\b|\bfilas? \d)"
                        r"|\b[A-Z]{1,2}\d{1,4}:[A-Z]{1,2}\d{1,4}\b")


# Los pasos que un programa sigue tal cual al ejecutar el ejercicio (app/materiales/ejecucion.py).
COPIA_FORMULA = re.compile(r"[Cc]opia la fórmula (\d+) en la celda ([A-Z]{1,3}\d+)")
HACIA_ABAJO = re.compile(r"[Cc]opia la celda ([A-Z]{1,3})(\d+) hacia abajo,? hasta la fila (\d+)")
_NUMERO_DE_FORMULA = re.compile(r"Fórmula (\d+)")


def formulas_por_numero(e) -> dict:
    """Los recuadros de fórmula del ejercicio, también los de la mejora, por su número («Fórmula 2 · …»)."""
    recuadros = [r for r in [*e.copiar, *e.mejora.copiar] if r.tipo == "formula"]
    numeros = {}
    for i, r in enumerate(recuadros, start=1):
        m = _NUMERO_DE_FORMULA.search(r.titulo)
        numeros[int(m.group(1)) if m else i] = r
    return numeros


def textos_del_ejercicio(e) -> list[str]:
    """Todo lo que el alumno lee en un ejercicio, también en «Mejora el resultado»."""
    recuadros = [*e.copiar, *e.mejora.copiar]
    return [e.situacion, *[p.texto for p in e.pasos], *[r.explicacion for r in recuadros],
            *[l for r in recuadros for l in r.lineas], e.mejora.texto, *e.mejora.pasos]


# Nombres de las funciones: la guía los escribe en español; dentro del archivo, Excel los guarda en inglés.
FUNCIONES = {
    "SI": "IF", "Y": "AND", "O": "OR", "NO": "NOT", "SUMA": "SUM", "PROMEDIO": "AVERAGE", "CONTAR": "COUNT",
    "CONTARA": "COUNTA", "CONTAR.SI": "COUNTIF", "CONTAR.SI.CONJUNTO": "COUNTIFS", "SUMAR.SI": "SUMIF",
    "SUMAR.SI.CONJUNTO": "SUMIFS", "PROMEDIO.SI": "AVERAGEIF", "BUSCARV": "VLOOKUP", "BUSCARH": "HLOOKUP",
    "BUSCARX": "XLOOKUP", "INDICE": "INDEX", "COINCIDIR": "MATCH", "SI.ERROR": "IFERROR", "REDONDEAR": "ROUND",
    "MAX": "MAX", "MIN": "MIN", "HOY": "TODAY", "FECHA": "DATE", "DIAS": "DAYS", "TEXTO": "TEXT",
    "CONCATENAR": "CONCATENATE", "IZQUIERDA": "LEFT", "DERECHA": "RIGHT", "EXTRAE": "MID", "LARGO": "LEN",
    "MAYUSC": "UPPER", "ESPACIOS": "TRIM", "SIFECHA": "DATEDIF", "ABS": "ABS", "ENTERO": "INT",
}
_FUNCION = re.compile(r"(?<![A-Za-z0-9_.])([A-ZÁÉÍÓÚ][A-ZÁÉÍÓÚ0-9]*(?:\.[A-ZÁÉÍÓÚ0-9]+)*)\s*\(")


def funciones_usadas(formulas: list[str]) -> list[tuple[str, str]]:
    """Las funciones de las fórmulas de la guía, en español y en inglés, en orden de aparición."""
    vistas = []
    for formula in formulas:
        for m in _FUNCION.finditer(formula.upper()):
            nombre = m.group(1).replace("Í", "I")
            if nombre in FUNCIONES and nombre not in vistas:
                vistas.append(nombre)
    return [(n, FUNCIONES[n]) for n in vistas]


def formula_para_excel(formula: str) -> str:
    """«=SUMA(B2:B9)» → «=SUM(B2:B9)»; los separadores «;» pasan a «,», como los guarda Excel."""
    def cambiar(m):
        nombre = m.group(1).upper().replace("Í", "I")
        return FUNCIONES.get(nombre, m.group(1)) + "("
    partes = re.split(r'("[^"]*")', formula)        # no se toca el texto entre comillas
    return "".join(p if p.startswith('"') else _FUNCION.sub(cambiar, p).replace(";", ",") for p in partes)


def patron_de_nombre(sesion: int, ejercicio: int | None = None) -> re.Pattern:
    numero = str(ejercicio) if ejercicio is not None else r"\d+"
    return re.compile(rf"^S{sesion}_E{numero}_[a-z0-9]+(?:-[a-z0-9]+)*\.(?:{'|'.join(EXTENSIONES)})$")


def numero_de_ejercicio(nombre: str) -> int | None:
    m = re.match(r"^S\d+_E(\d+)_", nombre)
    return int(m.group(1)) if m else None


# ---------- Crear los archivos ----------

def _valor(texto):
    """Un número escrito como texto se guarda como número; un código con ceros a la izquierda queda como texto."""
    if not isinstance(texto, str):
        return texto
    limpio = texto.strip()
    if limpio.startswith("="):
        return formula_para_excel(limpio)
    if re.fullmatch(r"-?\d+", limpio) and not (len(limpio) > 1 and limpio.lstrip("-").startswith("0")):
        return int(limpio)
    if re.fullmatch(r"-?\d+\.\d+", limpio):
        return float(limpio)
    return texto


valor_de_celda = _valor


def _dar_formato(hoja, identidad: Identidad, filas: list[list]) -> None:
    """Encabezados en negrita sobre el fondo de la ficha, anchos según el contenido y la primera fila fija."""
    fuente = identidad.tipografia
    borde = Side(style="thin", color=identidad.principal)
    for celda in hoja[1]:
        celda.font = Font(name=fuente, bold=True, color=identidad.texto)
        celda.fill = PatternFill("solid", fgColor=identidad.fondo)
        celda.border = Border(bottom=borde)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.font = Font(name=fuente, color=identidad.texto)
            celda.alignment = Alignment(vertical="top", wrap_text=isinstance(celda.value, str) and len(celda.value) > 30)
    for i, encabezado in enumerate(filas[0], start=1):
        largo = max([len(str(encabezado))] + [len(str(f[i - 1])) for f in filas[1:]
                                              if i <= len(f) and f[i - 1] is not None])
        hoja.column_dimensions[get_column_letter(i)].width = min(ANCHO_MAXIMO_COLUMNA, max(10, largo + 2))
    hoja.freeze_panes = "A2"


def encabezados_de_opciones(archivo: ArchivoPractica) -> list[tuple[str, str, str]]:
    """(hoja, columna, encabezado en «Opciones») de cada lista. Si dos hojas tienen una lista en una columna del
    mismo nombre, el encabezado lleva la hoja entre paréntesis."""
    listas = [(h.nombre, l.columna) for h in archivo.hojas for l in h.listas]
    repetidas = {c for _, c in listas if sum(1 for _, otra in listas if otra == c) > 1}
    return [(h, c, f"{c} ({h})" if c in repetidas else c) for h, c in listas]


def crear_excel(archivo: ArchivoPractica, identidad: Identidad, destino: Path) -> Path:
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for hoja_ in archivo.hojas:
        hoja = libro.create_sheet(hoja_.nombre)
        hoja.append(hoja_.encabezados)
        for fila in hoja_.filas:
            hoja.append([_valor(v) for v in fila])
        _dar_formato(hoja, identidad, [hoja_.encabezados, *hoja_.filas])
    _agregar_listas(libro, archivo, identidad)
    libro.save(destino)
    return destino


def _agregar_listas(libro, archivo: ArchivoPractica, identidad: Identidad) -> None:
    """PLAN.md §0, decisión 60: las opciones van en una hoja visible, «Opciones», una columna por lista, y cada
    columna con lista solo acepta una de ellas. La celda puede quedar vacía."""
    encabezados = encabezados_de_opciones(archivo)
    if not encabezados:
        return
    por_hoja = {h.nombre: h for h in archivo.hojas}
    opciones = libro.create_sheet(HOJA_DE_OPCIONES)
    columnas = []
    for numero, (nombre_hoja, columna, encabezado) in enumerate(encabezados, start=1):
        hoja_ = por_hoja[nombre_hoja]
        lista = next(l for l in hoja_.listas if l.columna == columna)
        letra = get_column_letter(numero)
        opciones.cell(row=1, column=numero, value=encabezado)
        for fila, opcion in enumerate(lista.opciones, start=2):
            opciones.cell(row=fila, column=numero, value=opcion)
        columnas.append([encabezado, *lista.opciones])
        if not hoja_.filas:
            continue
        destino = get_column_letter(hoja_.encabezados.index(columna) + 1)
        validacion = DataValidation(
            type="list", formula1=f"'{HOJA_DE_OPCIONES}'!${letra}$2:${letra}${len(lista.opciones) + 1}",
            allow_blank=True, showErrorMessage=True, errorStyle="stop", error=MENSAJE_DE_LA_LISTA)
        validacion.add(f"{destino}2:{destino}{len(hoja_.filas) + 1}")
        libro[nombre_hoja].add_data_validation(validacion)
    alto = max(len(c) for c in columnas)
    _dar_formato(opciones, identidad, [[c[i] if i < len(c) else None for c in columnas] for i in range(alto)])


def crear_word(archivo: ArchivoPractica, identidad: Identidad, destino: Path) -> Path:
    documento = Document()
    normal = documento.styles["Normal"]
    normal.font.name = identidad.tipografia
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(identidad.texto)
    for parrafo in archivo.parrafos:
        if parrafo.estilo == "titulo":
            p = documento.add_paragraph(style="Heading 1")
            corrida = p.add_run(parrafo.texto)
            corrida.font.color.rgb = RGBColor.from_string(identidad.principal)
            corrida.font.name = identidad.tipografia
        elif parrafo.estilo == "lista":
            documento.add_paragraph(parrafo.texto, style="List Bullet")
        else:
            documento.add_paragraph(parrafo.texto)
    if archivo.tabla is not None:
        generador_docx.tabla_de_datos(documento, identidad, archivo.tabla)
    documento.save(destino)
    return destino


def crear(contenido: Laboratorio, identidad: Identidad, carpeta: Path) -> list[Path]:
    """Crea todos los archivos de práctica en una carpeta limpia."""
    carpeta.mkdir(parents=True, exist_ok=True)
    for vieja in carpeta.iterdir():
        if vieja.is_file():
            vieja.unlink(missing_ok=True)
    rutas = []
    for archivo in contenido.archivos:
        destino = carpeta / archivo.nombre
        rutas.append(crear_excel(archivo, identidad, destino) if archivo.extension == "xlsx"
                     else crear_word(archivo, identidad, destino))
    return rutas


def empaquetar(rutas: list[Path], destino: Path) -> Path:
    destino.unlink(missing_ok=True)
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as zip_:
        for ruta in sorted(rutas):
            zip_.write(ruta, ruta.name)
    return destino


def nombres_del_zip(ruta: Path) -> list[str]:
    with zipfile.ZipFile(ruta) as zip_:
        return sorted(zip_.namelist())


# ---------- Controles antes de crear los archivos ----------

def _texto_de_archivo(archivo: ArchivoPractica) -> str:
    partes = [p.texto for p in archivo.parrafos]
    for hoja in archivo.hojas:
        partes += hoja.encabezados + [str(v) for fila in hoja.filas for v in fila if v is not None]
    if archivo.tabla:
        partes += archivo.tabla.encabezados + [c for f in archivo.tabla.filas for c in f]
    return "\n".join(partes)


def revisar(contenido: Laboratorio, sesion: int, ejercicios_de_la_ficha: int, fecha_de_referencia: str = "",
            respuestas_a_la_herramienta: dict[int, str] | None = None,
            permitidos: frozenset[str] = frozenset({"prompt", "formula", "plantilla"}),
            origen_de_opciones: "OrigenDeOpciones | None" = None) -> list[str]:
    """Lo que un programa comprueba antes de crear los archivos (SKILL.md, «Material 3» y «Ejercicios que se
    ejecutan»). Cada problema es una FALLA de estructura: se arregla con una corrección completa."""
    problemas = []
    numeros = [e.numero for e in contenido.ejercicios]
    if numeros != list(range(1, len(numeros) + 1)):
        problemas.append(f"FALLA · estructura · ejercicios: van numerados 1, 2, 3…; vienen {numeros}.")
    if len(numeros) != ejercicios_de_la_ficha:
        problemas.append(f"FALLA · estructura · ejercicios: la ficha de la sesión trae {ejercicios_de_la_ficha}; "
                         f"hay {len(numeros)}.")
    for a, b in zip(contenido.ejercicios, contenido.ejercicios[1:]):
        if a.area.strip().lower() == b.area.strip().lower():
            problemas.append(f"FALLA · estructura · ejercicios {a.numero} y {b.numero}: los dos son del área «{a.area}»; "
                             "las áreas rotan.")
    por_nombre = {a.nombre: a for a in contenido.archivos}
    if len(por_nombre) != len(contenido.archivos):
        problemas.append("FALLA · estructura · archivos: hay dos archivos de práctica con el mismo nombre.")
    for archivo in contenido.archivos:
        if not patron_de_nombre(sesion).match(archivo.nombre):
            problemas.append(f"FALLA · nombre de archivo · «{archivo.nombre}»: el formato es "
                             f"S{sesion}_E[ejercicio]_[descripcion-corta].xlsx o .docx, en minúsculas y con guiones.")
        if archivo.extension == "xlsx" and not archivo.hojas:
            problemas.append(f"FALLA · estructura · «{archivo.nombre}»: un Excel necesita al menos una hoja con datos.")
        if archivo.extension == "docx" and not (archivo.parrafos or archivo.tabla):
            problemas.append(f"FALLA · estructura · «{archivo.nombre}»: un Word necesita párrafos o una tabla.")
        if not any(archivo.nombre in e.archivos for e in contenido.ejercicios):
            problemas.append(f"FALLA · estructura · «{archivo.nombre}»: ningún ejercicio lo usa.")
        problemas += _revisar_listas(archivo, origen_de_opciones or OrigenDeOpciones())
    respuestas = respuestas_a_la_herramienta or {}
    for e in contenido.ejercicios:
        for nombre in e.archivos:
            if nombre not in por_nombre:
                problemas.append(f"FALLA · archivo inexistente · ejercicio {e.numero}: «{nombre}» no está en «archivos».")
            elif numero_de_ejercicio(nombre) != e.numero:
                problemas.append(f"FALLA · estructura · ejercicio {e.numero}: usa «{nombre}», que es de otro ejercicio. "
                                 "Cada ejercicio funciona solo, con sus archivos.")
        texto = " ".join(textos_del_ejercicio(e))
        otro = re.search(r"\b(?:resultado|respuesta|salida|archivo)\s+del\s+ejercicio\s+(\d+)", texto, re.I)
        if otro and int(otro.group(1)) != e.numero:
            problemas.append(f"FALLA · estructura · ejercicio {e.numero}: usa el resultado del ejercicio {otro.group(1)}. "
                             "Cada ejercicio funciona solo.")
        problemas += _revisar_herramienta(e, texto, permitidos)
        for recuadro in [*e.copiar, *e.mejora.copiar]:
            if recuadro.tipo == "plantilla":
                continue                      # completar la plantilla es el tema del ejercicio
            patron = POR_LLENAR_EN_FORMULA if recuadro.tipo == "formula" else POR_LLENAR
            for linea in recuadro.lineas:
                m = patron.search(linea)
                if m:
                    problemas.append(f"FALLA · por llenar · ejercicio {e.numero}: «{m.group(0)}» en lo que el alumno "
                                     f"copia («{linea[:80]}»). Lo que copia funciona tal cual.")
        pasos = " ".join([p.texto for p in e.pasos] + list(e.mejora.pasos)).replace("**", "")
        copiadas = {int(m.group(1)) for m in COPIA_FORMULA.finditer(pasos)}
        for numero in sorted(set(formulas_por_numero(e)) - copiadas):
            problemas.append(f"FALLA · ejecución · ejercicio {e.numero}: ningún paso dice dónde se copia la fórmula "
                             f"{numero}. Escribe «Copia la fórmula {numero} en la celda L2» con la celda real.")
        for numero in sorted(copiadas - set(formulas_por_numero(e))):
            problemas.append(f"FALLA · ejecución · ejercicio {e.numero}: un paso copia la fórmula {numero} y no hay un "
                             f"recuadro «Fórmula {numero} · …».")
        con_figura = [paso for paso in e.pasos if paso.figura is not None]
        if len(con_figura) > FIGURAS_POR_EJERCICIO:
            problemas.append(f"FALLA · figura · ejercicio {e.numero}: tiene {len(con_figura)} figuras; va una como "
                             "máximo, y solo si aclara dónde está algo en el archivo. Deja la más necesaria.")
        ubica = next((paso.texto for paso in e.pasos if UBICA_ALGO.search(paso.texto)), None)
        if not con_figura and ubica and e.archivos:
            problemas.append(f"FALLA · figura · ejercicio {e.numero}: el paso «{ubica[:80]}» pide ubicar algo en el "
                             "archivo y el ejercicio no tiene figura. Pon la figura en ese paso.")
        for paso in con_figura:
            problemas += _revisar_figura(e.numero, paso.figura, por_nombre, e.archivos)
        if fecha_de_referencia and re.search(r"\bplazos?\b|\bvenc\w*|\bfechas?\b", texto, re.I):
            de_sus_archivos = " ".join(_texto_de_archivo(por_nombre[n]) for n in e.archivos if n in por_nombre)
            if e.archivos and fecha_de_referencia.lower() not in de_sus_archivos.lower():
                problemas.append(f"FALLA · fecha de referencia · ejercicio {e.numero}: depende de fechas o plazos y su "
                                 f"archivo no trae la fecha de referencia «{fecha_de_referencia}». Escríbela en el archivo.")
        respuesta = (respuestas.get(e.numero) or "").strip()
        if respuesta and respuesta.lower() not in texto.lower():
            problemas.append(f"FALLA · estructura · ejercicio {e.numero}: la ficha dice que la herramienta pregunta algo "
                             f"y que el alumno responde «{respuesta}». Ordena que se detenga a preguntar y escribe esa "
                             "respuesta exacta.")
    return problemas


# ---------- Listas de opciones (PLAN.md §0, decisión 60) ----------

SI_Y_NO = ("Sí", "No")
PALABRAS_DE_UNA_OPCION = 5     # una respuesta más larga es abierta: no lleva lista
# La sesión enseña a crear listas o validación de datos: la lista le haría el trabajo al alumno (decisión 60 f).
ENSENA_LISTAS = re.compile(r"(?i)validaci[oó]n de datos|validar (?:los )?datos|listas? desplegables?"
                           r"|listas? de (?:opciones|validaci[oó]n)")
MENCIONA_LA_LISTA = re.compile(r"(?i)\blista\b")


@dataclass(frozen=True)
class OrigenDeOpciones:
    """Lo que las listas necesitan de las fichas y del material aprobado: el vocabulario de la sesión con sus
    variantes prohibidas, el texto de las dos fichas, donde una opción puede aparecer escrita (decisión 61), los
    procesos, cada uno con sus pasos en orden (decisión 62), y si la sesión enseña listas o validación de datos."""
    vocabulario: tuple[tuple[str, tuple[str, ...]], ...] = ()
    texto_de_las_fichas: str = ""
    ensena_listas: bool = False
    procesos: tuple[tuple[str, ...], ...] = ()


def _revisar_listas(archivo: ArchivoPractica, origen: OrigenDeOpciones) -> list[str]:
    listas = [(hoja, lista) for hoja in archivo.hojas for lista in hoja.listas]
    if not listas:
        return []
    numero = numero_de_ejercicio(archivo.nombre)
    ejercicio = f"ejercicio {numero}, " if numero is not None else ""
    if origen.ensena_listas:
        return [f"FALLA · lista de opciones · {ejercicio}«{archivo.nombre}»: la sesión enseña listas o validación de "
                "datos, y una lista le haría el trabajo al alumno. Quita las listas del archivo."]
    problemas = []
    if any(h.nombre.strip().lower() == HOJA_DE_OPCIONES.lower() for h in archivo.hojas):
        problemas.append(f"FALLA · lista de opciones · {ejercicio}«{archivo.nombre}»: ya tiene una hoja "
                         f"«{HOJA_DE_OPCIONES}» y AulaLista pone ahí las opciones. Cambia el nombre de esa hoja.")
    variantes = {v.lower(): c for c, vs in origen.vocabulario for v in vs}
    for hoja, lista in listas:
        donde = f"{ejercicio}«{archivo.nombre}» · hoja «{hoja.nombre}» · columna «{lista.columna}»"
        if not hoja.filas:
            problemas.append(f"FALLA · lista de opciones · {donde}: la hoja no tiene filas de datos donde elegir.")
        respuestas = [r.strip() for r in lista.respuestas if r.strip()]
        if not respuestas:
            problemas.append(f"FALLA · lista de opciones · {donde}: ninguna fila trae su respuesta. Escribe en "
                             "«respuestas» la respuesta correcta de cada fila.")
        for r in dict.fromkeys(respuestas):
            if len(r.split()) > PALABRAS_DE_UNA_OPCION:
                problemas.append(f"FALLA · lista de opciones · {donde}: la respuesta «{r[:60]}» es abierta. Una "
                                 "respuesta abierta no lleva lista: quita la lista de esa columna.")
            elif r not in lista.opciones:
                problemas.append(f"FALLA · lista de opciones · {donde}: a la lista le falta la respuesta «{r}». Cada "
                                 "respuesta de la columna está entre las opciones, escrita igual.")
        for opcion in lista.opciones:
            if opcion.lower() in variantes:
                problemas.append(f"FALLA · lista de opciones · {donde}: la opción «{opcion}» es una variante "
                                 f"prohibida; el nombre único es «{variantes[opcion.lower()]}».")
        for proceso in origen.procesos:
            pasos = [p.lower() for p in proceso]
            del_proceso = [o for o in lista.opciones if o.lower() in pasos]
            en_orden = sorted(del_proceso, key=lambda o: pasos.index(o.lower()))
            if len(del_proceso) > 1 and del_proceso != en_orden:
                problemas.append(f"FALLA · lista de opciones · {donde}: son pasos de un proceso y van en su orden: "
                                 f"{', '.join(en_orden)}.")
    return problemas


NUMERADA = re.compile(r"^\s*(\d+)\s*[.)]\s+(.+?)\s*$")


def procesos_del_vocabulario(conceptos: list[dict]) -> list[tuple[str, ...]]:
    """Decisión 62 a): los conceptos marcados «(proceso)» seguidos forman un proceso, en el orden del vocabulario."""
    procesos, actual = [], []
    for c in conceptos:
        if c.get("proceso"):
            actual.append(c["concepto"])
        elif actual:
            procesos.append(tuple(actual))
            actual = []
    if actual:
        procesos.append(tuple(actual))
    return [p for p in procesos if len(p) > 1]


def procesos_del_material(lectura: dict | None, diapositivas: dict | None) -> list[tuple[str, ...]]:
    """Decisión 62 b): los pasos de una diapositiva de tipo «proceso» o «pasos», y las líneas numeradas de
    «Aplícalo así», del material ya aprobado de la sesión. Recibe el contenido guardado de cada uno."""
    procesos = []
    for d in (diapositivas or {}).get("diapositivas", []):
        if d.get("tipo") == "proceso" and d.get("grafico"):
            procesos.append(tuple(e["texto"] for e in d["grafico"].get("elementos", [])))
        elif d.get("tipo") == "pasos":
            procesos.append(tuple(d.get("frases", [])))
    lineas = [NUMERADA.match(l) for l in ((lectura or {}).get("aplicalo") or {}).get("plantilla", [])]
    numeradas = [m.group(2).rstrip(".") for m in lineas if m]
    procesos.append(tuple(numeradas))
    return [p for p in procesos if len(p) > 1]


def _aparece(opcion: str, texto: str) -> bool:
    return re.search(r"(?<!\w)" + re.escape(opcion.lower()) + r"(?!\w)", texto) is not None


def avisos_de_listas(contenido: Laboratorio, origen: OrigenDeOpciones | None = None) -> list[tuple[str, str]]:
    """(sección, detalle) de lo que en una lista se corrige o se explica: un AVISO del verificador.
    - Decisión 61: una opción que no aparece escrita en las fichas, en «Antes de empezar» ni en los datos de su
      archivo. «Sí» y «No» valen siempre. Los pasos de la guía no cuentan: ahí escribe el mismo redactor.
    - Decisión 60 e): un ejercicio cuyos pasos no dicen que el alumno elige de la lista."""
    origen = origen or OrigenDeOpciones()
    antes = contenido.antes_de_empezar
    de_todos = " ".join([origen.texto_de_las_fichas, *antes.reglas,
                         *[" | ".join(f) for tabla in antes.tablas for f in [tabla.encabezados, *tabla.filas]]]).lower()
    avisos = []
    for archivo in contenido.archivos:
        datos = " | ".join(" | ".join([*h.encabezados, *[str(v) for f in h.filas for v in f if v is not None]])
                           for h in archivo.hojas).lower()
        numero = numero_de_ejercicio(archivo.nombre)
        seccion = next((f"Ejercicio {e.numero} · {e.titulo}" for e in contenido.ejercicios if e.numero == numero),
                       archivo.nombre)
        for hoja in archivo.hojas:
            for lista in hoja.listas:
                for opcion in lista.opciones:
                    if opcion in SI_Y_NO or _aparece(opcion, de_todos) or _aparece(opcion, datos):
                        continue
                    avisos.append((seccion, f"«{archivo.nombre}» · hoja «{hoja.nombre}» · columna «{lista.columna}»: "
                                            f"la opción «{opcion}» no aparece escrita en las fichas, en «Antes de "
                                            "empezar» ni en los datos del archivo. Escríbela donde la vea el alumno "
                                            "(por ejemplo, en una tabla de criterios) o explica de dónde sale."))
    con_listas = {a.nombre for a in contenido.archivos if any(h.listas for h in a.hojas)}
    for e in contenido.ejercicios:
        archivos = [n for n in e.archivos if n in con_listas]
        pasos = [p.texto for p in e.pasos] + list(e.mejora.pasos)
        if archivos and not any(MENCIONA_LA_LISTA.search(p) for p in pasos):
            avisos.append((f"Ejercicio {e.numero} · {e.titulo}",
                           f"Ejercicio {e.numero}: «{archivos[0]}» tiene una lista de opciones y ningún paso dice que el "
                           "alumno elige su respuesta de la lista de la celda. Dilo en el paso, sin nombrar menús ni "
                           "botones."))
    return avisos


NOMBRE_DEL_RECUADRO = {"prompt": "prompt", "formula": "fórmula", "plantilla": "plantilla"}
_MENCIONA_IA = re.compile(r"\bprompts?\b|\badjunta\b|\basistente de IA\b|\bIA generativa\b", re.I)
_MENCIONA_FORMULA = re.compile(r"\bf[óo]rmulas?\b", re.I)


def _revisar_herramienta(e, texto: str, permitidos: frozenset[str]) -> list[str]:
    """PLAN.md §0, decisión 36: un prompt solo en un curso de IA y una fórmula solo en un curso de Excel. Los
    ejercicios practican el tema de la sesión; la herramienta sale del tema del curso, no de un modelo."""
    problemas = []
    for tipo in sorted({r.tipo for r in [*e.copiar, *e.mejora.copiar]} - permitidos):
        problemas.append(f"FALLA · herramienta · ejercicio {e.numero}: trae un recuadro de {NOMBRE_DEL_RECUADRO[tipo]} y "
                         f"el curso no usa {'IA' if tipo == 'prompt' else 'Excel'} (ficha del curso, «Herramienta de los "
                         "ejercicios»). El alumno hace la acción de la ficha en su archivo de práctica.")
    for tipo, patron in (("prompt", _MENCIONA_IA), ("formula", _MENCIONA_FORMULA)):
        m = patron.search(texto)
        if tipo not in permitidos and m:
            problemas.append(f"FALLA · herramienta · ejercicio {e.numero}: dice «{m.group(0)}» y el curso no usa "
                             f"{'IA' if tipo == 'prompt' else 'Excel'}. Escribe el paso como una acción que el alumno "
                             "hace en su archivo de práctica.")
    return problemas


_CELDA = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)$")


def celda(texto: str) -> tuple[int, int] | None:
    """«C5» → (fila 5, columna 3)."""
    m = _CELDA.match(texto.strip().upper())
    if not m:
        return None
    columna = 0
    for letra in m.group(1):
        columna = columna * 26 + ord(letra) - 64
    return int(m.group(2)), columna


def rango(texto: str) -> tuple[int, int, int, int] | None:
    """«A1:E8» → (fila 1, columna 1, fila 8, columna 5). Una celda sola es un rango de una celda."""
    partes = texto.strip().upper().split(":")
    if len(partes) not in (1, 2):
        return None
    inicio, fin = celda(partes[0]), celda(partes[-1])
    if inicio is None or fin is None:
        return None
    return min(inicio[0], fin[0]), min(inicio[1], fin[1]), max(inicio[0], fin[0]), max(inicio[1], fin[1])


def _revisar_figura(numero: int, figura, por_nombre: dict, archivos_del_ejercicio: list[str]) -> list[str]:
    donde = f"ejercicio {numero}, figura de «{figura.archivo}»"
    if figura.archivo not in archivos_del_ejercicio or figura.archivo not in por_nombre:
        return [f"FALLA · figura · {donde}: la figura sale de un archivo de práctica de ese mismo ejercicio."]
    archivo = por_nombre[figura.archivo]
    problemas = []
    if archivo.extension == "xlsx":
        hojas = {h.nombre: h for h in archivo.hojas}
        hoja = hojas.get(figura.hoja) if figura.hoja else archivo.hojas[0] if archivo.hojas else None
        if hoja is None:
            return [f"FALLA · figura · {donde}: la hoja «{figura.hoja}» no existe en el archivo."]
        caja = rango(figura.rango) if figura.rango else (1, 1, len(hoja.filas) + 1, len(hoja.encabezados))
        if caja is None:
            return [f"FALLA · figura · {donde}: el rango «{figura.rango}» no es válido. Usa la forma «A1:E8»."]
        if caja[2] > len(hoja.filas) + 1 or caja[3] > len(hoja.encabezados):
            problemas.append(f"FALLA · figura · {donde}: el rango «{figura.rango}» se sale de los datos de la hoja.")
        if caja[3] - caja[1] + 1 > 8 or caja[2] - caja[0] + 1 > 16:
            problemas.append(f"FALLA · figura · {donde}: el rango «{figura.rango}» es muy grande para leerse en la "
                             "página: hasta 8 columnas y 16 filas.")
        for marca in figura.marcas:
            lugar = rango(marca.lugar)
            if lugar is None or not (caja[0] <= lugar[0] and lugar[2] <= caja[2] and caja[1] <= lugar[1] and lugar[3] <= caja[3]):
                problemas.append(f"FALLA · figura · {donde}: la marca «{marca.lugar}» no es una celda o un rango dentro "
                                 f"de «{figura.rango or 'la hoja'}».")
    else:
        texto = _texto_de_archivo(archivo).lower()
        for marca in figura.marcas:
            if marca.lugar.strip().lower() not in texto:
                problemas.append(f"FALLA · figura · {donde}: el texto marcado «{marca.lugar[:60]}» no está tal cual en "
                                 "el archivo.")
    return problemas
