"""Generador de las diapositivas de AulaLista, con el diseño común de la skill.

Diapositivas (SKILL.md, «Material 2»): PowerPoint 16:9, una idea por diapositiva, frases cortas,
ejemplos en recuadros legibles a distancia. Sin notas del profesor (decisión 66 a).
Diseño común: la identidad visual de la ficha, tres colores como máximo más los fondos de recuadro, y
tablas con encabezado en el color principal, filas alternas en gris muy claro y sin bordes verticales.
En las diapositivas, los recuadros son tarjetas con encabezado de color, sin borde izquierdo; cada diapositiva
de concepto lleva el nombre de su bloque y una marca de avance (PLAN.md §0, decisión 30).

El generador elige el tamaño de letra que cabe según una estimación. Quien mide de verdad es
PowerPoint (app/office.py, medir_diapositivas): un texto que se sale de su recuadro es una falla.
"""

import math
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Emu, Pt

from app.materiales import estilo
from app.materiales.contenido import Diapositiva, Diapositivas, Grafico, Tabla
from app.materiales.docx import Portada, lineas_de_contacto
from app.materiales.estilo import Identidad

ANCHO, ALTO = Cm(33.867), Cm(19.05)          # 16:9
MARGEN = Cm(1.6)
BLOQUE_ARRIBA, BLOQUE_ALTO = Cm(0.7), Cm(1.0)      # nombre del bloque, arriba del título (decisión 30)
TITULO_ARRIBA, TITULO_ALTO = Cm(1.5), Cm(2.0)
CUERPO_ARRIBA = Cm(3.9)
CUERPO_ALTO = ALTO - CUERPO_ARRIBA - Cm(1.2)
CUERPO_ANCHO = ANCHO - 2 * MARGEN
SEPARACION = Cm(0.8)
ENCABEZADO = Cm(1.3)                          # encabezado de color de las tarjetas

TAMANO_TITULO = 32
TAMANO_CUERPO = 26
TAMANO_MINIMO = 18                            # menos que esto no se lee a distancia
TAMANO_TABLA = 20
ALTO_FILA = Cm(1.25)

DISENO_SOLO_TITULO = 5
DISENO_EN_BLANCO = 6
SIN_ESTILO_NI_CUADRICULA = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color)


def alto_estimado(parrafos: list[str], ancho: Emu, tamano: int) -> Emu:
    """Estimación: letras por renglón según el ancho y renglones según el largo de cada párrafo."""
    letras_por_renglon = max(8, int(Emu(ancho).pt / (tamano * 0.5)))
    renglones = sum(max(1, math.ceil(len(p) / letras_por_renglon)) for p in parrafos)
    return Pt(renglones * tamano * 1.2 + len(parrafos) * tamano * 0.45)


def _cabe(parrafos: list[str], ancho: Emu, alto: Emu, tamano: int) -> bool:
    return alto_estimado(parrafos, ancho, tamano) <= alto


def tamano_que_cabe(parrafos: list[str], ancho: Emu, alto: Emu, base: int = TAMANO_CUERPO) -> int:
    for tamano in range(base, TAMANO_MINIMO - 1, -2):
        if _cabe(parrafos, ancho, alto, tamano):
            return tamano
    return TAMANO_MINIMO


def _vineta(parrafo, numerada: bool, tamano: int) -> None:
    """Viñeta o número real de PowerPoint, con sangría francesa."""
    ppr = parrafo._p.get_or_add_pPr()
    sangria = int(Pt(tamano) * 1.1)
    ppr.set("marL", str(sangria))
    ppr.set("indent", str(-sangria))
    for viejo in ppr.findall(qn("a:buNone")) + ppr.findall(qn("a:buChar")) + ppr.findall(qn("a:buAutoNum")):
        ppr.remove(viejo)
    if numerada:
        vineta = ppr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"})
    else:
        vineta = ppr.makeelement(qn("a:buChar"), {"char": "•"})
    ppr.append(vineta)


def _escribir(marco, parrafos: list[str], identidad: Identidad, tamano: int, *, titulo: str = "",
              color_titulo: str | None = None, vinetas: bool = False, numerada: bool = False) -> None:
    marco.word_wrap = True
    marco.auto_size = MSO_AUTO_SIZE.NONE
    marco.vertical_anchor = MSO_ANCHOR.TOP
    primero = True
    if titulo:
        p = marco.paragraphs[0]
        corrida = p.add_run()
        corrida.text = titulo
        _letra(corrida, identidad, tamano, color_titulo or identidad.texto, negrita=True)
        p.space_after = Pt(tamano * 0.4)
        p.alignment = PP_ALIGN.LEFT
        primero = False
    for texto in parrafos:
        p = marco.paragraphs[0] if primero else marco.add_paragraph()
        primero = False
        corrida = p.add_run()
        corrida.text = texto
        _letra(corrida, identidad, tamano, identidad.texto)
        p.space_after = Pt(tamano * 0.45)
        p.alignment = PP_ALIGN.LEFT
        if vinetas or numerada:
            _vineta(p, numerada, tamano)


def _letra(corrida, identidad: Identidad, tamano: float, color: str, negrita: bool = False) -> None:
    corrida.font.name = identidad.tipografia
    corrida.font.size = Pt(tamano)
    corrida.font.bold = negrita
    corrida.font.color.rgb = _rgb(color)


def _caja(diapositiva, izquierda, arriba, ancho, alto):
    caja = diapositiva.shapes.add_textbox(izquierda, arriba, ancho, alto)
    return caja.text_frame


def _titulo(diapositiva, texto: str, identidad: Identidad, bloque: str = "", avance: tuple[int, int] | None = None) -> None:
    """Título de la diapositiva. Arriba, el nombre del bloque en el color de acento y, a la derecha, un punto por
    bloque con el actual relleno (decisión 30)."""
    if bloque:
        ancho_bloque = CUERPO_ANCHO - (Cm(0.75) * avance[1] if avance else 0)
        marco = _caja(diapositiva, MARGEN, BLOQUE_ARRIBA, ancho_bloque, BLOQUE_ALTO)
        marco.word_wrap = True
        marco.auto_size = MSO_AUTO_SIZE.NONE
        corrida = marco.paragraphs[0].add_run()
        corrida.text = bloque.upper()
        _letra(corrida, identidad, TAMANO_MINIMO, estilo.para_texto(identidad.acento), negrita=True)
    if avance:
        actual, total = avance
        for i in range(total):
            x = ANCHO - MARGEN - (total - i) * Cm(0.75) + Cm(0.2)
            punto = diapositiva.shapes.add_shape(MSO_SHAPE.OVAL, int(x), int(BLOQUE_ARRIBA + Cm(0.3)), int(Cm(0.4)),
                                                 int(Cm(0.4)))
            punto.fill.solid()
            punto.fill.fore_color.rgb = _rgb(identidad.principal if i == actual else identidad.fondo)
            punto.line.color.rgb = _rgb(identidad.principal)
            punto.line.width = Pt(1)
            punto.shadow.inherit = False
    titulo = diapositiva.shapes.title
    titulo.left, titulo.top, titulo.width, titulo.height = MARGEN, TITULO_ARRIBA, CUERPO_ANCHO, TITULO_ALTO
    marco = titulo.text_frame
    marco.clear()
    marco.word_wrap = True
    marco.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = marco.paragraphs[0]
    corrida = p.add_run()
    corrida.text = texto
    tamano = TAMANO_TITULO if _cabe([texto], CUERPO_ANCHO, TITULO_ALTO, TAMANO_TITULO) else 26
    _letra(corrida, identidad, tamano, estilo.para_texto(identidad.principal), negrita=True)
    p.alignment = PP_ALIGN.LEFT


MARGEN_RECUADRO_LADO, MARGEN_RECUADRO_ALTO = Cm(0.6), Cm(0.4)


def medida_de_recuadro(parrafos: list[str], ancho, alto_maximo, titulo: str = "",
                       base: int = TAMANO_CUERPO) -> tuple[int, Emu]:
    """Tamaño de letra y alto de la tarjeta: el alto justo para su texto, sin pasar del espacio disponible."""
    interior_ancho = ancho - 2 * MARGEN_RECUADRO_LADO - Cm(0.8)   # la sangría de las viñetas
    encabezado = ENCABEZADO if titulo else Cm(0.3)
    interior_alto = alto_maximo - encabezado - 2 * MARGEN_RECUADRO_ALTO
    tamano = tamano_que_cabe(parrafos, interior_ancho, interior_alto, base)
    alto = encabezado + alto_estimado(parrafos, interior_ancho, tamano) + 2 * MARGEN_RECUADRO_ALTO + Cm(0.6)
    return tamano, min(alto_maximo, max(Cm(5.5), alto))


def _sin_sombra_ni_linea(figura) -> None:
    figura.line.fill.background()
    figura.shadow.inherit = False


def recuadro(diapositiva, identidad: Identidad, izquierda, arriba, ancho, alto_maximo, parrafos: list[str], *,
             titulo: str = "", color_borde: str | None = None, vinetas: bool = False,
             base: int = TAMANO_CUERPO, medida: tuple[int, Emu] | None = None) -> None:
    """Tarjeta: encabezado del color (con su título, si lo tiene) y cuerpo sobre el fondo de recuadro.
    En las diapositivas reemplaza al recuadro con borde izquierdo de los Word (decisión 30)."""
    color = color_borde or identidad.principal
    tamano, alto = medida or medida_de_recuadro(parrafos, ancho, alto_maximo, titulo, base)
    encabezado = ENCABEZADO if titulo else Cm(0.3)
    cuerpo = diapositiva.shapes.add_shape(MSO_SHAPE.RECTANGLE, izquierda, arriba, ancho, alto)
    cuerpo.fill.solid()
    cuerpo.fill.fore_color.rgb = _rgb(identidad.fondo)
    _sin_sombra_ni_linea(cuerpo)
    banda = diapositiva.shapes.add_shape(MSO_SHAPE.RECTANGLE, izquierda, arriba, ancho, encabezado)
    banda.fill.solid()
    banda.fill.fore_color.rgb = _rgb(color)
    _sin_sombra_ni_linea(banda)
    if titulo:
        marco = banda.text_frame
        marco.margin_left = marco.margin_right = MARGEN_RECUADRO_LADO
        marco.vertical_anchor = MSO_ANCHOR.MIDDLE
        marco.word_wrap = True
        corrida = marco.paragraphs[0].add_run()
        corrida.text = titulo
        claro = estilo.BLANCO if estilo.contraste(color, estilo.BLANCO) >= estilo.CONTRASTE_MINIMO else identidad.texto
        _letra(corrida, identidad, min(tamano, 24), claro, negrita=True)
    texto = diapositiva.shapes.add_textbox(izquierda, arriba + encabezado, ancho, alto - encabezado)
    marco = texto.text_frame
    marco.margin_left = marco.margin_right = MARGEN_RECUADRO_LADO
    marco.margin_top = marco.margin_bottom = MARGEN_RECUADRO_ALTO
    _escribir(marco, parrafos, identidad, tamano, vinetas=vinetas)


def tabla(diapositiva, identidad: Identidad, datos: Tabla, arriba, alto_disponible) -> None:
    """Encabezado en el color principal, filas alternas en gris muy claro y sin bordes verticales."""
    filas, columnas = 1 + len(datos.filas), len(datos.encabezados)
    alto = min(alto_disponible, ALTO_FILA * filas)
    marco = diapositiva.shapes.add_table(filas, columnas, MARGEN, arriba, CUERPO_ANCHO, alto)
    estilo_tabla = marco._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId"))
    if estilo_tabla is not None:
        estilo_tabla.text = SIN_ESTILO_NI_CUADRICULA
    t = marco.table
    t.first_row, t.horz_banding = True, False
    for fila_numero in range(filas):
        valores = datos.encabezados if fila_numero == 0 else datos.filas[fila_numero - 1]
        for columna, texto in enumerate(valores):
            celda = t.cell(fila_numero, columna)
            celda.fill.solid()
            if fila_numero == 0:
                celda.fill.fore_color.rgb = _rgb(identidad.principal)
            else:
                celda.fill.fore_color.rgb = _rgb(estilo.GRIS_FILA_ALTERNA if fila_numero % 2 == 0 else estilo.BLANCO)
            marco_celda = celda.text_frame
            marco_celda.word_wrap = True
            corrida = marco_celda.paragraphs[0].add_run()
            corrida.text = texto
            _letra(corrida, identidad, TAMANO_TABLA, estilo.BLANCO if fila_numero == 0 else identidad.texto,
                   negrita=fila_numero == 0)


# ---------- Cada tipo de diapositiva ----------

def _nueva(presentacion, diseno: int = DISENO_SOLO_TITULO):
    return presentacion.slides.add_slide(presentacion.slide_layouts[diseno])


# ---------- Portada y cierre (PLAN.md §0, decisión 29) ----------

PANEL = 0.62                       # parte de la diapositiva que ocupa el panel del color principal
LINEA_ACENTO = Cm(0.35)


def _sobre(color: str, identidad: Identidad) -> str:
    """Color del texto sobre un relleno: blanco si se lee bien; si no, el color del texto de la ficha."""
    return estilo.BLANCO if estilo.contraste(color, estilo.BLANCO) >= estilo.CONTRASTE_MINIMO else identidad.texto


def _rectangulo(diapositiva, forma, izquierda, arriba, ancho, alto, relleno: str | None, linea: str | None = None,
                grosor: float = 1.5):
    figura = diapositiva.shapes.add_shape(forma, int(izquierda), int(arriba), int(ancho), int(alto))
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        figura.adjustments[0] = 0.12
    if relleno:
        figura.fill.solid()
        figura.fill.fore_color.rgb = _rgb(relleno)
    else:
        figura.fill.background()
    if linea:
        figura.line.color.rgb = _rgb(linea)
        figura.line.width = Pt(grosor)
    else:
        figura.line.fill.background()
    figura.shadow.inherit = False
    return figura


def _lineas(diapositiva, izquierda, arriba, ancho, alto, partes, identidad: Identidad, *, centrado=False,
            anclaje=MSO_ANCHOR.TOP):
    """Una caja de texto con párrafos de distinto tamaño: partes = [(texto, tamaño, color, negrita)]."""
    marco = _caja(diapositiva, int(izquierda), int(arriba), int(ancho), int(alto))
    marco.word_wrap = True
    marco.auto_size = MSO_AUTO_SIZE.NONE
    marco.vertical_anchor = anclaje
    primero = True
    for texto, tamano, color, negrita in partes:
        p = marco.paragraphs[0] if primero else marco.add_paragraph()
        primero = False
        corrida = p.add_run()
        corrida.text = texto
        _letra(corrida, identidad, tamano, color, negrita)
        p.space_after = Pt(tamano * 0.35)
        p.alignment = PP_ALIGN.CENTER if centrado else PP_ALIGN.LEFT
    return marco


def _panel(presentacion, identidad: Identidad) -> tuple:
    """Diapositiva en blanco con el panel del color principal a la izquierda y la línea de acento en su borde."""
    diapositiva = _nueva(presentacion, DISENO_EN_BLANCO)
    ancho_panel = int(ANCHO * PANEL)
    _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, 0, 0, ancho_panel, ALTO, identidad.principal)
    _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, ancho_panel, 0, LINEA_ACENTO, ALTO, identidad.acento)
    if identidad.estilo.franja:   # estilos moderno y corporativo, o la carátula del modelo (decisiones 27 y 28)
        _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, ancho_panel + LINEA_ACENTO, 0,
                    ANCHO - ancho_panel - LINEA_ACENTO, Cm(0.9), identidad.principal)
    return diapositiva, ancho_panel


def _docente_y_contacto(diapositiva, identidad: Identidad, portada: Portada, ancho_panel: int, arriba) -> None:
    """Columna derecha: docente y contacto, en el color del texto, con la etiqueta en el color de acento."""
    izquierda = ancho_panel + LINEA_ACENTO + Cm(1.4)
    ancho = ANCHO - izquierda - Cm(1.2)
    partes = []
    if portada.docente:
        partes += [("Docente", TAMANO_MINIMO, estilo.para_texto(identidad.acento), True),
                   (portada.docente, 22, identidad.texto, True)]
    contacto = lineas_de_contacto(portada.contacto)
    if contacto:
        partes += [("Contacto", TAMANO_MINIMO, estilo.para_texto(identidad.acento), True)]
        partes += [(dato, TAMANO_MINIMO, identidad.texto, False) for dato in contacto]
    if partes:
        _lineas(diapositiva, izquierda, arriba, ancho, ALTO - arriba - Cm(1.2), partes, identidad)


def _logo(diapositiva, identidad: Identidad, ancho_panel: int) -> bool:
    """El logo de la ficha, si lo hay: en el panel si la carátula lo pone a la izquierda; si no, a la derecha."""
    if identidad.logo is None:
        return False
    if identidad.estilo.logo == "izquierda":
        diapositiva.shapes.add_picture(str(identidad.logo), MARGEN, Cm(1.4), height=Cm(2.4))
        return False
    izquierda = ancho_panel + LINEA_ACENTO + Cm(1.4)
    diapositiva.shapes.add_picture(str(identidad.logo), izquierda, Cm(2.2), height=Cm(3.2))
    return True


def _portada(presentacion, identidad: Identidad, portada: Portada) -> None:
    diapositiva, ancho_panel = _panel(presentacion, identidad)
    color = _sobre(identidad.principal, identidad)
    centrado = identidad.estilo.portada_centrada
    logo_a_la_derecha = _logo(diapositiva, identidad, ancho_panel)
    ancho = ancho_panel - 2 * MARGEN
    arriba_titulo = Cm(4.6) if identidad.logo is not None and not logo_a_la_derecha else Cm(3.4)
    titulo_tamano = tamano_que_cabe([portada.curso], ancho, Cm(5.6), base=44)
    _lineas(diapositiva, MARGEN, arriba_titulo, ancho, Cm(5.8), [(portada.curso, titulo_tamano, color, True)],
            identidad, centrado=centrado, anclaje=MSO_ANCHOR.BOTTOM)
    linea_x = MARGEN + (ancho - Cm(4)) // 2 if centrado else MARGEN
    _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, linea_x, arriba_titulo + Cm(6.1), Cm(4), Cm(0.22), identidad.acento)
    _lineas(diapositiva, MARGEN, arriba_titulo + Cm(6.7), ancho, Cm(4.4),
            [(portada.sesion, tamano_que_cabe([portada.sesion], ancho, Cm(4.2), base=28), color, False)],
            identidad, centrado=centrado)
    _docente_y_contacto(diapositiva, identidad, portada, ancho_panel, Cm(7.6) if logo_a_la_derecha else Cm(5.4))


def _cierre(presentacion, identidad: Identidad, portada: Portada, frase: str) -> None:
    """La idea que se llevan, con el docente y el contacto (decisión 29)."""
    diapositiva, ancho_panel = _panel(presentacion, identidad)
    color = _sobre(identidad.principal, identidad)
    ancho = ancho_panel - 2 * MARGEN
    _lineas(diapositiva, MARGEN, Cm(2.2), ancho, Cm(1.4),
            [("La idea que te llevas", TAMANO_MINIMO + 2, color, True)], identidad)
    _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, MARGEN, Cm(3.8), Cm(4), Cm(0.22), identidad.acento)
    _lineas(diapositiva, MARGEN, Cm(4.6), ancho, Cm(11.6),
            [(frase, tamano_que_cabe([frase], ancho, Cm(11.2), base=34), color, True)], identidad,
            anclaje=MSO_ANCHOR.MIDDLE)
    logo_a_la_derecha = _logo(diapositiva, identidad, ancho_panel)
    _docente_y_contacto(diapositiva, identidad, portada, ancho_panel, Cm(7.6) if logo_a_la_derecha else Cm(5.4))


# ---------- Recursos gráficos (decisión 29) ----------

MARGEN_FIGURA = Cm(0.25)


def _en_figura(figura, partes, identidad: Identidad, *, centrado=True, anclaje=MSO_ANCHOR.MIDDLE) -> None:
    """Escribe dentro de una forma: partes = [(texto, tamaño, color, negrita)]."""
    marco = figura.text_frame
    marco.clear()
    marco.word_wrap = True
    marco.auto_size = MSO_AUTO_SIZE.NONE
    marco.vertical_anchor = anclaje
    marco.margin_left = marco.margin_right = marco.margin_top = marco.margin_bottom = MARGEN_FIGURA
    primero = True
    for texto, tamano, color, negrita in partes:
        if not texto:
            continue
        p = marco.paragraphs[0] if primero else marco.add_paragraph()
        primero = False
        corrida = p.add_run()
        corrida.text = texto
        _letra(corrida, identidad, tamano, color, negrita)
        p.space_after = Pt(tamano * 0.25)
        p.alignment = PP_ALIGN.CENTER if centrado else PP_ALIGN.LEFT


def _tamano(textos: list[str], ancho, alto, base: int = 22) -> int:
    textos = [x for x in textos if x]
    return tamano_que_cabe(textos, ancho - 2 * MARGEN_FIGURA, alto - 2 * MARGEN_FIGURA, base) if textos else base


def _textos(e) -> list[str]:
    return [x for x in (e.texto, e.detalle, *e.hijos) if x]


def _tamano_comun(elementos, ancho, alto, base: int = 22) -> int:
    """Un solo tamaño de letra para todas las piezas de un recurso: el que cabe en la más cargada."""
    return min(_tamano(_textos(e), ancho, alto, base) for e in elementos)


def _alto_justo(elementos, ancho, tamano: int, minimo, maximo):
    """El alto que necesita la pieza más cargada, entre un mínimo y el espacio disponible."""
    necesario = max(alto_estimado(_textos(e), ancho - 2 * MARGEN_FIGURA, tamano) for e in elementos)
    return int(min(maximo, max(minimo, necesario + 2 * MARGEN_FIGURA + Cm(0.5))))


def _caja_de_elemento(diapositiva, identidad: Identidad, e, izquierda, arriba, ancho, alto, *, fuerte=False,
                      forma=MSO_SHAPE.ROUNDED_RECTANGLE, hijos=True, tamano: int | None = None) -> None:
    """Una pieza: título en negrita y, debajo, su detalle y sus hijos. Fuerte: relleno del color principal."""
    relleno = identidad.principal if fuerte else identidad.fondo
    figura = _rectangulo(diapositiva, forma, izquierda, arriba, ancho, alto, relleno,
                         None if fuerte else identidad.principal)
    color_titulo = _sobre(relleno, identidad) if fuerte else estilo.para_texto(identidad.principal, identidad.fondo)
    color_texto = _sobre(relleno, identidad) if fuerte else identidad.texto
    resto = [e.detalle, *(e.hijos if hijos else [])]
    tamano = tamano or _tamano([e.texto, *resto], ancho, alto)
    _en_figura(figura, [(e.texto, tamano, color_titulo, True), *[(x, max(TAMANO_MINIMO, tamano - 2), color_texto, False)
                                                                 for x in resto]], identidad)


def _linea(diapositiva, x1, y1, x2, y2, color: str, grosor: float = 2.0, flecha: bool = False):
    conector = diapositiva.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    conector.line.color.rgb = _rgb(color)
    conector.line.width = Pt(grosor)
    conector.shadow.inherit = False
    if flecha:
        ln = conector.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return conector


def _proceso(diapositiva, identidad: Identidad, g: Grafico, izquierda, arriba, ancho, alto) -> None:
    """Pasos en fila, con su número y una flecha entre cada uno."""
    n = len(g.elementos)
    hueco = Cm(1.1)
    caja = (ancho - (n - 1) * hueco) // n
    alto_titulo = Cm(2.6)
    con_detalle = any(e.detalle for e in g.elementos)
    tamano = min(_tamano([e.texto], caja, alto_titulo) for e in g.elementos)
    disponible = alto - Cm(2.0) - alto_titulo - Cm(0.4)
    tamano_detalle = min([_tamano([e.detalle], caja, disponible, 20) for e in g.elementos if e.detalle] or [20])
    alto_detalle = int(min(disponible, max(Cm(2.4), max(
        [alto_estimado([e.detalle], caja - 2 * MARGEN_FIGURA, tamano_detalle) for e in g.elementos if e.detalle] or [0])
        + 2 * MARGEN_FIGURA + Cm(0.5))))
    total = Cm(2.0) + alto_titulo + ((Cm(0.4) + alto_detalle) if con_detalle else 0)
    arriba = arriba + max(0, (alto - total) // 2)
    for i, e in enumerate(g.elementos):
        x = izquierda + i * (caja + hueco)
        circulo = _rectangulo(diapositiva, MSO_SHAPE.OVAL, x + (caja - Cm(1.5)) // 2, arriba, Cm(1.5), Cm(1.5),
                              identidad.acento)
        _en_figura(circulo, [(str(i + 1), TAMANO_MINIMO + 2, _sobre(identidad.acento, identidad), True)], identidad)
        titulo = _rectangulo(diapositiva, MSO_SHAPE.ROUNDED_RECTANGLE, x, arriba + Cm(2.0), caja, alto_titulo,
                             identidad.principal)
        _en_figura(titulo, [(e.texto, tamano, _sobre(identidad.principal, identidad), True)], identidad)
        if i < n - 1:
            _rectangulo(diapositiva, MSO_SHAPE.RIGHT_ARROW, x + caja + Cm(0.15), arriba + Cm(2.0) + alto_titulo // 2 - Cm(0.45),
                        hueco - Cm(0.3), Cm(0.9), identidad.acento)
        if con_detalle and e.detalle:
            detalle = _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, x, arriba + Cm(2.0) + alto_titulo + Cm(0.4), caja,
                                  alto_detalle, identidad.fondo)
            _en_figura(detalle, [(e.detalle, tamano_detalle, identidad.texto, False)], identidad, anclaje=MSO_ANCHOR.TOP)


def _ciclo(diapositiva, identidad: Identidad, g: Grafico, izquierda, arriba, ancho, alto) -> None:
    """Etapas sobre un óvalo, con flechas en el sentido del ciclo."""
    n = len(g.elementos)
    caja_ancho, caja_alto = Cm(7.6), Cm(3.4)
    cx, cy = izquierda + ancho / 2, arriba + alto / 2
    rx, ry = ancho / 2 - caja_ancho / 2, alto / 2 - caja_alto / 2
    _rectangulo(diapositiva, MSO_SHAPE.OVAL, cx - rx, cy - ry, 2 * rx, 2 * ry, None, identidad.principal, 2.5)
    if g.centro:
        centro = _rectangulo(diapositiva, MSO_SHAPE.OVAL, cx - Cm(4), cy - Cm(1.6), Cm(8), Cm(3.2), identidad.fondo)
        _en_figura(centro, [(g.centro, _tamano([g.centro], Cm(7), Cm(3)), estilo.para_texto(identidad.principal,
                                                                                                 identidad.fondo), True)],
                   identidad)
    angulos = [-math.pi / 2 + 2 * math.pi * i / n for i in range(n)]
    for i, angulo in enumerate(angulos):
        medio = angulo + math.pi / n
        mx, my = cx + rx * math.cos(medio), cy + ry * math.sin(medio)
        dx, dy = -rx * math.sin(medio), ry * math.cos(medio)
        flecha = _rectangulo(diapositiva, MSO_SHAPE.ISOSCELES_TRIANGLE, mx - Cm(0.5), my - Cm(0.5), Cm(1.0), Cm(1.0),
                             identidad.acento)
        flecha.rotation = math.degrees(math.atan2(dy, dx)) + 90
    tamano = _tamano_comun(g.elementos, caja_ancho, caja_alto)
    for e, angulo in zip(g.elementos, angulos):
        x, y = cx + rx * math.cos(angulo), cy + ry * math.sin(angulo)
        _caja_de_elemento(diapositiva, identidad, e, x - caja_ancho / 2, y - caja_alto / 2, caja_ancho, caja_alto,
                          fuerte=True, tamano=tamano)


def _organigrama(diapositiva, identidad: Identidad, g: Grafico, izquierda, arriba, ancho, alto) -> None:
    """Un cargo arriba y los que dependen de él en una fila, unidos por líneas."""
    n = len(g.elementos)
    raiz_ancho, raiz_alto = Cm(14), Cm(2.4)
    hueco = Cm(0.6)
    caja = (ancho - (n - 1) * hueco) / n
    disponible = alto - raiz_alto - Cm(1.8)
    tamano = _tamano_comun(g.elementos, caja, disponible)
    alto_caja = _alto_justo(g.elementos, caja, tamano, Cm(3.2), disponible)
    arriba = arriba + max(0, (disponible - alto_caja) // 2)   # el organigrama, centrado en el espacio
    raiz = _rectangulo(diapositiva, MSO_SHAPE.ROUNDED_RECTANGLE, izquierda + (ancho - raiz_ancho) / 2, arriba,
                       raiz_ancho, raiz_alto, identidad.principal)
    _en_figura(raiz, [(g.centro, _tamano([g.centro], raiz_ancho, raiz_alto, 24), _sobre(identidad.principal, identidad),
                       True)], identidad)
    fila = arriba + raiz_alto + Cm(1.8)
    bus = arriba + raiz_alto + Cm(0.9)
    centros = [izquierda + i * (caja + hueco) + caja / 2 for i in range(n)]
    _linea(diapositiva, izquierda + ancho / 2, arriba + raiz_alto, izquierda + ancho / 2, bus, identidad.principal)
    _linea(diapositiva, centros[0], bus, centros[-1], bus, identidad.principal)
    for x in centros:
        _linea(diapositiva, x, bus, x, fila, identidad.principal)
    for e, x in zip(g.elementos, centros):
        _caja_de_elemento(diapositiva, identidad, e, x - caja / 2, fila, caja, alto_caja, tamano=tamano)


def _mapa(diapositiva, identidad: Identidad, g: Grafico, izquierda, arriba, ancho, alto) -> None:
    """Mapa mental: la idea al centro y sus ramas alrededor."""
    n = len(g.elementos)
    caja_ancho, caja_alto = Cm(8.6), Cm(4.0)
    cx, cy = izquierda + ancho / 2, arriba + alto / 2
    rx, ry = ancho / 2 - caja_ancho / 2, alto / 2 - caja_alto / 2
    inicio = -math.pi / 2 if n % 2 else -math.pi / 2 + math.pi / n
    angulos = [inicio + 2 * math.pi * i / n for i in range(n)]
    posiciones = [(cx + rx * math.cos(a), cy + ry * math.sin(a)) for a in angulos]
    for x, y in posiciones:
        _linea(diapositiva, cx, cy, x, y, identidad.acento, 3.0)
    centro = _rectangulo(diapositiva, MSO_SHAPE.OVAL, cx - Cm(4.4), cy - Cm(1.8), Cm(8.8), Cm(3.6), identidad.principal)
    _en_figura(centro, [(g.centro, _tamano([g.centro], Cm(7.6), Cm(3.0), 26), _sobre(identidad.principal, identidad),
                         True)], identidad)
    tamano = _tamano_comun(g.elementos, caja_ancho, caja_alto)
    for e, (x, y) in zip(g.elementos, posiciones):
        _caja_de_elemento(diapositiva, identidad, e, x - caja_ancho / 2, y - caja_alto / 2, caja_ancho, caja_alto,
                          tamano=tamano)


def _matriz(diapositiva, identidad: Identidad, g: Grafico, izquierda, arriba, ancho, alto) -> None:
    """Matriz 2×2 con sus dos ejes. Orden: arriba a la izquierda, arriba a la derecha, abajo a la izquierda,
    abajo a la derecha."""
    eje_lado, eje_abajo, eje_arriba = Cm(0.9), Cm(1.4), Cm(1.1)
    x0, y0 = izquierda + eje_lado, arriba + eje_arriba
    w, h = ancho - eje_lado, alto - eje_abajo - eje_arriba
    hueco = Cm(0.3)
    cw, ch = (w - hueco) / 2, (h - hueco) / 2
    tamano = _tamano_comun(g.elementos, cw, ch)
    for i, e in enumerate(g.elementos):
        fila, columna = divmod(i, 2)
        _caja_de_elemento(diapositiva, identidad, e, x0 + columna * (cw + hueco), y0 + fila * (ch + hueco), cw, ch,
                          fuerte=(i == 1), forma=MSO_SHAPE.RECTANGLE, tamano=tamano)
    _linea(diapositiva, x0 - Cm(0.4), y0 + h, x0 - Cm(0.4), y0, identidad.acento, 2.5, flecha=True)
    _linea(diapositiva, x0, y0 + h + Cm(0.4), x0 + w, y0 + h + Cm(0.4), identidad.acento, 2.5, flecha=True)
    horizontal, vertical = g.ejes
    _lineas(diapositiva, x0, y0 + h + Cm(0.5), w, Cm(1.0), [(horizontal, TAMANO_MINIMO, identidad.texto, True)],
            identidad, centrado=True)
    _lineas(diapositiva, izquierda, arriba, w, eje_arriba, [(vertical, TAMANO_MINIMO, identidad.texto, True)], identidad)


def _tarjetas(diapositiva, identidad: Identidad, g: Grafico, izquierda, arriba, ancho, alto) -> None:
    """Tablero: tarjetas con una franja de color arriba, en una o dos filas."""
    n = len(g.elementos)
    por_fila = n if n <= 3 else (2 if n == 4 else 3)
    filas = math.ceil(n / por_fila)
    hueco = Cm(0.6)
    cw = (ancho - (por_fila - 1) * hueco) / por_fila
    disponible = (alto - (filas - 1) * hueco) / filas
    tamano = _tamano_comun(g.elementos, cw, disponible - Cm(0.35))
    ch = _alto_justo(g.elementos, cw, tamano, Cm(3.6), disponible - Cm(0.35)) + Cm(0.35)
    arriba = arriba + max(0, (alto - filas * ch - (filas - 1) * hueco) // 2)
    for i, e in enumerate(g.elementos):
        fila, columna = divmod(i, por_fila)
        x, y = izquierda + columna * (cw + hueco), arriba + fila * (ch + hueco)
        _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, x, y, cw, Cm(0.35), identidad.acento)
        tarjeta = _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, x, y + Cm(0.35), cw, ch - Cm(0.35), identidad.fondo)
        resto = [e.detalle, *e.hijos]
        _en_figura(tarjeta, [(e.texto, tamano, estilo.para_texto(identidad.principal, identidad.fondo), True),
                             *[(x_, max(TAMANO_MINIMO, tamano - 2), identidad.texto, False) for x_ in resto]],
                   identidad, centrado=False, anclaje=MSO_ANCHOR.TOP)


def _grafica(diapositiva, identidad: Identidad, g: Grafico, izquierda, arriba, ancho, alto) -> None:
    """Gráfica de barras con los colores de la ficha. Las cifras salen de la lectura o de una fuente."""
    datos = CategoryChartData()
    datos.categories = [e.texto for e in g.elementos]
    datos.add_series(g.centro or "Valor", [e.valor for e in g.elementos])
    marco = diapositiva.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, int(izquierda), int(arriba), int(ancho),
                                         int(alto), datos)
    grafica = marco.chart
    grafica.has_legend = False
    grafica.has_title = bool(g.centro)
    if g.centro:
        grafica.chart_title.text_frame.text = g.centro
        _letra(grafica.chart_title.text_frame.paragraphs[0].runs[0], identidad, TAMANO_MINIMO + 2, identidad.texto, True)
    serie = grafica.plots[0].series[0]
    serie.format.fill.solid()
    serie.format.fill.fore_color.rgb = _rgb(identidad.principal)
    grafica.plots[0].gap_width = 80
    grafica.plots[0].has_data_labels = True
    etiquetas = grafica.plots[0].data_labels
    etiquetas.show_value = True
    etiquetas.position = XL_LABEL_POSITION.OUTSIDE_END
    etiquetas.font.size = Pt(TAMANO_MINIMO + 2)
    etiquetas.font.bold = True
    etiquetas.font.color.rgb = _rgb(identidad.texto)
    eje = grafica.category_axis
    eje.tick_labels.font.size = Pt(TAMANO_MINIMO)
    eje.tick_labels.font.color.rgb = _rgb(identidad.texto)
    eje.format.line.color.rgb = _rgb(identidad.texto)
    valores = grafica.value_axis
    valores.visible = False
    valores.has_major_gridlines = False


def _pasos(diapositiva, identidad: Identidad, frases: list[str], izquierda, arriba, ancho, alto) -> None:
    """Pasos numerados: un círculo con el número y el paso al lado."""
    n = len(frases)
    hueco = Cm(0.35)
    fila = min(Cm(2.6), (alto - (n - 1) * hueco) / n)
    diametro = min(Cm(1.6), fila)
    tamano = min(_tamano([f], ancho - diametro - Cm(0.6), fila, 24) for f in frases)
    arriba = arriba + max(0, (alto - n * fila - (n - 1) * hueco) // 2)
    for i, frase in enumerate(frases):
        y = arriba + i * (fila + hueco)
        circulo = _rectangulo(diapositiva, MSO_SHAPE.OVAL, izquierda, y + (fila - diametro) / 2, diametro, diametro,
                              identidad.principal)
        _en_figura(circulo, [(str(i + 1), TAMANO_MINIMO + 2, _sobre(identidad.principal, identidad), True)], identidad)
        caja = _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, izquierda + diametro + Cm(0.4), y,
                           ancho - diametro - Cm(0.4), fila, identidad.fondo)
        _en_figura(caja, [(frase, tamano, identidad.texto, False)], identidad, centrado=False)


DIBUJOS = {"proceso": _proceso, "ciclo": _ciclo, "organigrama": _organigrama, "mapa": _mapa, "matriz": _matriz,
           "tarjetas": _tarjetas, "grafica": _grafica}


def _filas_con_marca(diapositiva, identidad: Identidad, frases: list[str], izquierda, arriba, ancho, alto, *,
                     marca: str = "", color: str | None = None, base: int = 26) -> None:
    """Frases en filas, cada una con una marca de color a la izquierda: un cuadro, un número o un signo."""
    color = color or identidad.acento
    n = len(frases)
    hueco = Cm(0.45)
    fila = min(Cm(2.3), (alto - (n - 1) * hueco) / n)
    lado = Cm(1.4) if marca else Cm(0.45)
    texto_ancho = ancho - lado - Cm(0.7)
    tamano = min(_tamano([f], texto_ancho, fila, base) for f in frases)
    arriba = arriba + max(0, (alto - n * fila - (n - 1) * hueco) // 2)
    for i, frase in enumerate(frases):
        y = arriba + i * (fila + hueco)
        figura = _rectangulo(diapositiva, MSO_SHAPE.OVAL if marca else MSO_SHAPE.RECTANGLE, izquierda,
                             y + (fila - lado) / 2, lado, lado, color)
        if marca:
            _en_figura(figura, [(str(i + 1) if marca == "#" else marca, TAMANO_MINIMO + 2, _sobre(color, identidad), True)],
                       identidad)
        _lineas(diapositiva, izquierda + lado + Cm(0.7), y, texto_ancho, fila, [(frase, tamano, identidad.texto, False)],
                identidad, anclaje=MSO_ANCHOR.MIDDLE)


def _destacada(diapositiva, identidad: Identidad, frase: str, izquierda, arriba, ancho, alto, color: str) -> None:
    """Una sola frase, grande, con una línea de acento arriba: la idea que se quiere fijar."""
    tamano = tamano_que_cabe([frase], ancho - Cm(2), alto - Cm(2.5), base=40)
    alto_texto = min(alto - Cm(1.5), alto_estimado([frase], ancho - Cm(2), tamano) + Cm(0.8))
    y = arriba + max(Cm(0.8), (alto - alto_texto) // 2)            # el bloque, centrado en el espacio
    _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, izquierda, y - Cm(0.7), Cm(3.5), Cm(0.25), identidad.acento)
    _lineas(diapositiva, izquierda, y, ancho - Cm(2), alto_texto, [(frase, tamano, color, True)],
            identidad, anclaje=MSO_ANCHOR.TOP)


def _frases(diapositiva, identidad: Identidad, frases: list[str], izquierda, arriba, ancho, alto) -> None:
    if len(frases) == 1:
        _destacada(diapositiva, identidad, frases[0], izquierda, arriba, ancho, alto,
                   estilo.para_texto(identidad.principal))
    else:
        _filas_con_marca(diapositiva, identidad, frases, izquierda, arriba, ancho, alto)


def _idea_central(presentacion, identidad: Identidad, frases: list[str]) -> None:
    """La idea central sobre el fondo de recuadro, en letra grande."""
    diapositiva = _nueva(presentacion)
    fondo = _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, 0, 0, ANCHO, ALTO, identidad.fondo)
    diapositiva.shapes._spTree.remove(fondo._element)
    diapositiva.shapes._spTree.insert(2, fondo._element)          # detrás del título
    titulo = diapositiva.shapes.title
    titulo.left, titulo.top, titulo.width, titulo.height = MARGEN, Cm(1.6), CUERPO_ANCHO, Cm(1.4)
    marco = titulo.text_frame
    marco.clear()
    corrida = marco.paragraphs[0].add_run()
    corrida.text = "Idea central"
    marco.paragraphs[0].alignment = PP_ALIGN.LEFT
    _letra(corrida, identidad, TAMANO_MINIMO + 4, estilo.para_texto(identidad.acento, identidad.fondo), negrita=True)
    texto = " ".join(frases)
    tamano = tamano_que_cabe([texto], CUERPO_ANCHO - Cm(3), Cm(11), base=40)
    _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, MARGEN, Cm(3.4), Cm(3.5), Cm(0.25), identidad.acento)
    _lineas(diapositiva, MARGEN, Cm(4.2), CUERPO_ANCHO - Cm(3), Cm(12), [(texto, tamano,
            estilo.para_texto(identidad.principal, identidad.fondo), True)], identidad, anclaje=MSO_ANCHOR.MIDDLE)


def _cifras(diapositiva, identidad: Identidad, g: Grafico, izquierda, arriba, ancho, alto) -> None:
    """De 1 a 3 cifras grandes con lo que cuentan (decisión 30). Las cifras salen de la lectura o de una fuente."""
    n = len(g.elementos)
    hueco = Cm(1.0)
    columna = (ancho - (n - 1) * hueco) / n
    alto_tarjeta = min(alto - Cm(1.2), Cm(8.5))
    y = arriba + (alto - alto_tarjeta) // 2
    for i, e in enumerate(g.elementos):
        x = izquierda + i * (columna + hueco)
        tarjeta = _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, x, y, columna, alto_tarjeta, identidad.fondo)
        _rectangulo(diapositiva, MSO_SHAPE.RECTANGLE, x, y, columna, Cm(0.3), identidad.acento)
        partes = [(f"{e.valor:g}", 66, estilo.para_texto(identidad.principal, identidad.fondo), True),
                  (e.texto, 24, identidad.texto, True)]
        if e.detalle:
            partes.append((e.detalle, 20, identidad.texto, False))
        _en_figura(tarjeta, partes, identidad)
    if g.centro:
        _lineas(diapositiva, izquierda, arriba + alto - Cm(0.6), ancho, Cm(0.9),
                [(g.centro, TAMANO_MINIMO, identidad.texto, False)], identidad, centrado=True)


DIBUJOS["cifras"] = _cifras


def _introduccion(diapositiva, identidad: Identidad, frases: list[str], arriba, alto):
    """Una línea de introducción arriba del recurso gráfico. Devuelve el espacio que queda."""
    if not frases:
        return arriba, alto
    alto_frases = Cm(1.6) * len(frases)
    marco = _caja(diapositiva, MARGEN, arriba, CUERPO_ANCHO, alto_frases)
    _escribir(marco, frases, identidad, tamano_que_cabe(frases, CUERPO_ANCHO, alto_frases, 24))
    return arriba + alto_frases + Cm(0.2), alto - alto_frases - Cm(0.2)


def _titulo_de_tabla(diapositiva, identidad: Identidad, titulo: str, arriba, alto):
    """El título de la tabla, arriba de ella, en el color principal. Devuelve el espacio que queda."""
    alto_titulo = Cm(0.9)
    marco = _caja(diapositiva, MARGEN, arriba, CUERPO_ANCHO, alto_titulo)
    marco.margin_top = marco.margin_bottom = 0
    marco.vertical_anchor = MSO_ANCHOR.MIDDLE
    corrida = marco.paragraphs[0].add_run()
    corrida.text = titulo
    _letra(corrida, identidad, TAMANO_MINIMO, estilo.para_texto(identidad.principal), negrita=True)
    return arriba + alto_titulo, alto - alto_titulo


def _de_concepto(presentacion, identidad: Identidad, d: Diapositiva, avance: tuple[int, int]) -> None:
    diapositiva = _nueva(presentacion)
    _titulo(diapositiva, d.titulo, identidad, d.bloque, avance)
    arriba, alto = CUERPO_ARRIBA, CUERPO_ALTO
    if d.tipo == "frases":
        _frases(diapositiva, identidad, d.frases, MARGEN, arriba, CUERPO_ANCHO, alto)
    elif d.tipo in DIBUJOS:
        arriba, alto = _introduccion(diapositiva, identidad, d.frases, arriba, alto)
        DIBUJOS[d.tipo](diapositiva, identidad, d.grafico, MARGEN, arriba, CUERPO_ANCHO, alto)
    elif d.tipo == "pasos":
        _pasos(diapositiva, identidad, d.frases, MARGEN, arriba, CUERPO_ANCHO, alto)
    elif d.tipo == "ejemplo":
        medida = medida_de_recuadro(d.frases, CUERPO_ANCHO, alto, "Ejemplo", base=28)
        recuadro(diapositiva, identidad, MARGEN, arriba + (alto - medida[1]) // 2, CUERPO_ANCHO, alto, d.frases,
                 titulo="Ejemplo", color_borde=identidad.acento, medida=medida)
    elif d.tipo == "columnas":
        arriba, alto = _introduccion(diapositiva, identidad, d.frases, arriba, alto)
        ancho = (CUERPO_ANCHO - SEPARACION) // 2
        # Las dos columnas con la misma letra y el mismo alto, para que se comparen a la vista.
        medidas = [medida_de_recuadro(c.frases, ancho, alto, c.titulo, base=26) for c in d.columnas]
        medida = (min(m[0] for m in medidas), max(m[1] for m in medidas))
        arriba = arriba + (alto - medida[1]) // 2
        for i, columna in enumerate(d.columnas):
            recuadro(diapositiva, identidad, MARGEN + i * (ancho + SEPARACION), arriba, ancho, alto, columna.frases,
                     titulo=columna.titulo, vinetas=True, medida=medida,
                     color_borde=identidad.principal if i == 0 else identidad.acento)
    elif d.tipo == "tabla":
        if d.frases:
            alto_frases = Cm(1.6) * len(d.frases)
            marco = _caja(diapositiva, MARGEN, arriba, CUERPO_ANCHO, alto_frases)
            _escribir(marco, d.frases, identidad, tamano_que_cabe(d.frases, CUERPO_ANCHO, alto_frases, 24))
            arriba, alto = arriba + alto_frases, alto - alto_frases
        if d.tabla.titulo:
            arriba, alto = _titulo_de_tabla(diapositiva, identidad, d.tabla.titulo, arriba, alto)
        tabla(diapositiva, identidad, d.tabla, arriba, alto)


def generar_diapositivas(contenido: Diapositivas, identidad: Identidad, portada: Portada, destino: Path) -> Path:
    presentacion = Presentation()
    presentacion.slide_width, presentacion.slide_height = ANCHO, ALTO
    _portada(presentacion, identidad, portada)
    _idea_central(presentacion, identidad, contenido.idea_central.frases)
    bloques = list(dict.fromkeys(d.bloque for d in contenido.diapositivas))
    for d in contenido.diapositivas:
        _de_concepto(presentacion, identidad, d, (bloques.index(d.bloque), len(bloques)))

    aplicalo = contenido.aplicalo
    diapositiva = _nueva(presentacion)
    _titulo(diapositiva, "Aplícalo así", identidad, "Para tu trabajo")
    arriba, alto = CUERPO_ARRIBA, CUERPO_ALTO
    if aplicalo.frases:
        alto_frases = Cm(1.7) * len(aplicalo.frases)
        marco = _caja(diapositiva, MARGEN, arriba, CUERPO_ANCHO, alto_frases)
        _escribir(marco, aplicalo.frases, identidad, tamano_que_cabe(aplicalo.frases, CUERPO_ANCHO, alto_frases, 24))
        arriba, alto = arriba + alto_frases + Cm(0.3), alto - alto_frases - Cm(0.3)
    _filas_con_marca(diapositiva, identidad, aplicalo.plantilla, MARGEN, arriba, CUERPO_ANCHO, alto, marca="#",
                     color=identidad.principal, base=24)

    diapositiva = _nueva(presentacion)
    _titulo(diapositiva, "Cuidado con", identidad, "Riesgos")
    _filas_con_marca(diapositiva, identidad, contenido.cuidado.riesgos, MARGEN, CUERPO_ARRIBA, CUERPO_ANCHO, CUERPO_ALTO,
                     marca="!", color=identidad.advertencia)

    _cierre(presentacion, identidad, portada, contenido.cierre.frase)

    destino.parent.mkdir(parents=True, exist_ok=True)
    presentacion.save(destino)
    return destino
