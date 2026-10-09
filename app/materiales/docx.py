"""Generador de los Word de AulaLista, con el diseño común de la skill.

Lectura (SKILL.md, «Material 1»): A4, márgenes de 2,5 cm, cuerpo en 11 pt, interlineado 1,15,
6 pt después de cada párrafo, portada limpia, encabezado con el nombre del curso, pie con
número de página y estilos reales de Título 1, 2 y 3. El laboratorio, la guía del profesor y
la evaluación usan este mismo diseño.
"""

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.materiales import estilo
from app.materiales.contenido import Lectura, Tabla
from app.materiales.estilo import Identidad


@dataclass
class Portada:
    curso: str
    sesion: str        # «Sesión 1 · Agilidad y fundamentos de Scrum»
    material: str      # «Lectura», «Laboratorio», «Guía del profesor»…
    docente: str = ""


def _rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color)


def _fuente(estilo_word, identidad: Identidad, tamano: float, color: str, negrita: bool = False) -> None:
    fuente = estilo_word.font
    fuente.name = identidad.tipografia
    fuente.size = Pt(tamano)
    fuente.bold = negrita
    fuente.color.rgb = _rgb(color)
    rpr = estilo_word.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(atributo), identidad.tipografia)


def documento_base(identidad: Identidad, portada: Portada) -> Document:
    documento = Document()
    seccion = documento.sections[0]
    seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
    for margen in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(seccion, margen, Cm(2.5))

    estilos = documento.styles
    _fuente(estilos["Normal"], identidad, 11, identidad.texto)
    formato = estilos["Normal"].paragraph_format
    formato.line_spacing = 1.15
    formato.space_after = Pt(6)
    formato.space_before = Pt(0)
    formato.alignment = WD_ALIGN_PARAGRAPH.LEFT
    titular = estilo.para_texto(identidad.principal)
    for nombre, tamano, antes in (("Heading 1", 16, 18), ("Heading 2", 13, 12), ("Heading 3", 11.5, 10)):
        _fuente(estilos[nombre], identidad, tamano, titular, negrita=True)
        estilos[nombre].paragraph_format.space_before = Pt(antes)
        estilos[nombre].paragraph_format.space_after = Pt(6)
        estilos[nombre].paragraph_format.keep_with_next = True
    _fuente(estilos["Title"], identidad, 26, titular, negrita=True)
    estilos["Title"].paragraph_format.space_after = Pt(12)
    _quitar_borde_del_titulo(estilos["Title"])

    seccion.different_first_page_header_footer = True
    encabezado = seccion.header.paragraphs[0]
    encabezado.text = portada.curso
    encabezado.runs[0].font.size = Pt(9)
    encabezado.runs[0].font.color.rgb = _rgb(identidad.texto)
    _numero_de_pagina(seccion.footer.paragraphs[0], identidad)

    _portada(documento, identidad, portada)
    return documento


def _quitar_borde_del_titulo(estilo_word) -> None:
    ppr = estilo_word.element.get_or_add_pPr()
    for borde in ppr.findall(qn("w:pBdr")):
        ppr.remove(borde)


def _numero_de_pagina(parrafo, identidad: Identidad) -> None:
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    campo = OxmlElement("w:fldSimple")
    campo.set(qn("w:instr"), "PAGE")
    corrida = OxmlElement("w:r")
    texto = OxmlElement("w:t")
    texto.text = "1"
    corrida.append(texto)
    campo.append(corrida)
    parrafo._p.append(campo)


def _portada(documento: Document, identidad: Identidad, portada: Portada) -> None:
    """Portada limpia: logo si existe, curso, sesión y material. Sin imágenes decorativas."""
    for _ in range(4):
        documento.add_paragraph()
    if identidad.logo is not None:
        documento.add_picture(str(identidad.logo), width=Cm(4))
    documento.add_paragraph(portada.curso, style="Title")
    sesion = documento.add_paragraph(portada.sesion)
    sesion.runs[0].font.size = Pt(16)
    material = documento.add_paragraph(portada.material)
    material.runs[0].font.size = Pt(20)
    material.runs[0].font.bold = True
    material.runs[0].font.color.rgb = _rgb(estilo.para_texto(identidad.acento))
    if portada.docente:
        documento.add_paragraph()
        documento.add_paragraph(f"Docente: {portada.docente}")
    documento.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------- Bloques comunes ----------

def _sombrear(celda, color: str) -> None:
    propiedades = celda._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color)
    propiedades.append(sombra)


def _bordes(tabla, **lados) -> None:
    """lados: nombre → (color, grosor en octavos de punto) o None para quitar el borde."""
    propiedades = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elemento = OxmlElement(f"w:{lado}")
        valor = lados.get(lado)
        if valor is None:
            elemento.set(qn("w:val"), "nil")
        else:
            color, grosor = valor
            elemento.set(qn("w:val"), "single")
            elemento.set(qn("w:sz"), str(grosor))
            elemento.set(qn("w:color"), color)
        bordes.append(elemento)
    propiedades.append(bordes)


def _margen_interior(tabla, izquierda: int = 200, derecha: int = 160) -> None:
    """Espacio entre el borde del recuadro y el texto, en vigésimos de punto."""
    margenes = OxmlElement("w:tblCellMar")
    for lado, valor in (("top", 60), ("left", izquierda), ("bottom", 60), ("right", derecha)):
        elemento = OxmlElement(f"w:{lado}")
        elemento.set(qn("w:w"), str(valor))
        elemento.set(qn("w:type"), "dxa")
        margenes.append(elemento)
    tabla._tbl.tblPr.append(margenes)


def _no_partir_fila(fila) -> None:
    propiedades = fila._tr.get_or_add_trPr()
    elemento = OxmlElement("w:cantSplit")
    propiedades.append(elemento)


def recuadro(documento: Document, identidad: Identidad, titulo: str, parrafos: list[str],
             color_borde: str | None = None, lineas_para_copiar: list[str] | None = None) -> None:
    """Recuadro sombreado con borde izquierdo de color. Mismo estilo en todos los materiales."""
    tabla = documento.add_table(rows=1, cols=1)
    tabla.alignment = WD_TABLE_ALIGNMENT.LEFT
    _bordes(tabla, left=(color_borde or identidad.principal, 24))
    _margen_interior(tabla)
    celda = tabla.cell(0, 0)
    _sombrear(celda, identidad.fondo)
    _no_partir_fila(tabla.rows[0])
    primero = celda.paragraphs[0]
    primero.add_run(titulo).bold = True
    primero.runs[0].font.color.rgb = _rgb(estilo.para_texto(color_borde or identidad.principal, identidad.fondo))
    for texto in parrafos:
        celda.add_paragraph(texto)
    for linea in lineas_para_copiar or []:
        celda.add_paragraph(linea)
    documento.add_paragraph()


def tabla_de_datos(documento: Document, identidad: Identidad, datos: Tabla) -> None:
    """Encabezado en el color principal, filas alternas en gris muy claro y sin bordes verticales."""
    if datos.titulo:
        documento.add_paragraph(datos.titulo).runs[0].bold = True
    tabla = documento.add_table(rows=1 + len(datos.filas), cols=len(datos.encabezados))
    _bordes(tabla)
    for i, texto in enumerate(datos.encabezados):
        celda = tabla.cell(0, i)
        _sombrear(celda, identidad.principal)
        corrida = celda.paragraphs[0].add_run(texto)
        corrida.bold = True
        corrida.font.color.rgb = _rgb(estilo.BLANCO)
    for numero, fila in enumerate(datos.filas, start=1):
        _no_partir_fila(tabla.rows[numero])
        for i, texto in enumerate(fila):
            celda = tabla.cell(numero, i)
            if numero % 2 == 0:
                _sombrear(celda, estilo.GRIS_FILA_ALTERNA)
            celda.paragraphs[0].add_run(texto)
    documento.add_paragraph()


# ---------- Lectura ----------

def generar_lectura(contenido: Lectura, identidad: Identidad, portada: Portada, destino: Path) -> Path:
    documento = documento_base(identidad, portada)
    documento.add_heading("Idea central", level=1)
    for texto in contenido.idea_central:
        documento.add_paragraph(texto)
    for bloque in contenido.bloques:
        documento.add_heading(bloque.subtitulo, level=1)
        for texto in bloque.parrafos:
            documento.add_paragraph(texto)
        if bloque.tabla is not None:
            tabla_de_datos(documento, identidad, bloque.tabla)
        recuadro(documento, identidad, bloque.ejemplo.titulo, bloque.ejemplo.parrafos)
    recuadro(documento, identidad, "Aplícalo así", contenido.aplicalo.parrafos,
             lineas_para_copiar=contenido.aplicalo.plantilla)
    recuadro(documento, identidad, "Cuidado con", contenido.cuidado, color_borde=identidad.advertencia)
    destino.parent.mkdir(parents=True, exist_ok=True)
    documento.save(destino)
    return destino
