"""Figuras del laboratorio, hechas con un programa a partir del archivo de práctica real (SKILL.md, «Imágenes»).

- Excel: la cuadrícula del rango, con las letras de columna y los números de fila, como la ve el alumno.
- Word: los párrafos alrededor del texto marcado.
- Las marcas son recuadros de color ámbar, con un número si hay más de una.
- No se dibujan pantallas, menús ni botones de ninguna herramienta.
"""

from dataclasses import dataclass
from pathlib import Path

import openpyxl
from docx import Document
from PIL import Image, ImageDraw, ImageFont

from app.materiales import practica
from app.materiales.contenido import Figura

AMBAR = (230, 149, 0)
GRIS_ENCABEZADO = (242, 242, 242)
GRIS_LINEA = (210, 210, 210)
GRIS_TEXTO = (110, 110, 110)
NEGRO = (31, 41, 51)
BLANCO = (255, 255, 255)
LETRA = 26                      # píxeles; la imagen se pone en el Word a 200 puntos por pulgada
FILA = 46
MARGEN = 14
LETRAS_MAXIMAS_POR_CELDA = 34
PUNTOS_POR_PULGADA = 200
ANCHO_MAXIMO_CM = 16.0          # el ancho del texto en una página A4 con márgenes de 2,5 cm
ALTO_MAXIMO_CM = 12.0
LETRA_MINIMA_PT = 8.0           # menos que esto no se lee sin ampliar
_CARPETA_FUENTES = Path("C:/Windows/Fonts")


@dataclass
class Imagen:
    ruta: Path
    ancho_cm: float
    letra_pt: float              # el tamaño de la letra en la página final


def _fuente(negrita: bool = False, tamano: int = LETRA):
    for nombre in (("calibrib.ttf", "arialbd.ttf") if negrita else ("calibri.ttf", "arial.ttf")):
        ruta = _CARPETA_FUENTES / nombre
        if ruta.exists():
            return ImageFont.truetype(str(ruta), tamano)
    return ImageFont.load_default(tamano)


def _letra_columna(numero: int) -> str:
    letras = ""
    while numero:
        numero, resto = divmod(numero - 1, 26)
        letras = chr(65 + resto) + letras
    return letras


def _texto_de_celda(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float):
        texto = f"{valor:,.2f}".rstrip("0").rstrip(".") if valor % 1 else f"{int(valor):,}"
        return texto.replace(",", " ")
    texto = str(valor)
    return texto if len(texto) <= LETRAS_MAXIMAS_POR_CELDA else texto[:LETRAS_MAXIMAS_POR_CELDA - 1] + "…"


def _numero_de_marca(dibujo, x, y, numero: int) -> None:
    radio = 18
    dibujo.ellipse((x - radio, y - radio, x + radio, y + radio), fill=AMBAR)
    dibujo.text((x, y), str(numero), font=_fuente(True, 24), fill=BLANCO, anchor="mm")


def _medida(imagen: Image.Image, ruta: Path) -> Imagen:
    ancho_cm = min(ANCHO_MAXIMO_CM, imagen.width / PUNTOS_POR_PULGADA * 2.54)
    alto_cm = imagen.height * ancho_cm / imagen.width
    if alto_cm > ALTO_MAXIMO_CM:
        ancho_cm = ancho_cm * ALTO_MAXIMO_CM / alto_cm
    escala = ancho_cm / (imagen.width / PUNTOS_POR_PULGADA * 2.54)
    letra_pt = LETRA / PUNTOS_POR_PULGADA * 72 * escala
    imagen.save(ruta)
    return Imagen(ruta, round(ancho_cm, 2), round(letra_pt, 1))


def de_excel(archivo: Path, figura: Figura, destino: Path) -> Imagen:
    """La cuadrícula del rango del Excel real. Lee los valores que guardó el archivo (con las fórmulas ya
    calculadas por Excel, si las hay)."""
    libro = openpyxl.load_workbook(archivo, data_only=True)
    hoja = libro[figura.hoja] if figura.hoja else libro.worksheets[0]
    caja = practica.rango(figura.rango) if figura.rango else (1, 1, hoja.max_row, hoja.max_column)
    f1, c1, f2, c2 = caja
    fuente, negrita = _fuente(), _fuente(True)
    medidor = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    anchos = []
    for c in range(c1, c2 + 1):
        textos = [_texto_de_celda(hoja.cell(f, c).value) for f in range(f1, f2 + 1)]
        ancho = max([medidor.textlength(t, font=negrita if f1 == 1 else fuente) for t in textos] + [40])
        anchos.append(int(ancho) + 2 * MARGEN)
    ancho_numeros = int(medidor.textlength(str(f2), font=fuente)) + 2 * MARGEN
    ancho = ancho_numeros + sum(anchos) + 1
    alto = FILA * (f2 - f1 + 2) + 1
    imagen = Image.new("RGB", (ancho, alto), BLANCO)
    dibujo = ImageDraw.Draw(imagen)
    dibujo.rectangle((0, 0, ancho, FILA), fill=GRIS_ENCABEZADO)
    dibujo.rectangle((0, 0, ancho_numeros, alto), fill=GRIS_ENCABEZADO)
    xs = [ancho_numeros]
    for a in anchos:
        xs.append(xs[-1] + a)
    for i, c in enumerate(range(c1, c2 + 1)):
        dibujo.text(((xs[i] + xs[i + 1]) / 2, FILA / 2), _letra_columna(c), font=fuente, fill=GRIS_TEXTO, anchor="mm")
    for j, f in enumerate(range(f1, f2 + 1), start=1):
        y = j * FILA
        dibujo.text((ancho_numeros / 2, y + FILA / 2), str(f), font=fuente, fill=GRIS_TEXTO, anchor="mm")
        for i, c in enumerate(range(c1, c2 + 1)):
            valor = hoja.cell(f, c).value
            alineado_derecha = isinstance(valor, (int, float))
            x = xs[i + 1] - MARGEN if alineado_derecha else xs[i] + MARGEN
            dibujo.text((x, y + FILA / 2), _texto_de_celda(valor), font=negrita if f == 1 else fuente, fill=NEGRO,
                        anchor="rm" if alineado_derecha else "lm")
    for x in xs:
        dibujo.line((x, 0, x, alto), fill=GRIS_LINEA, width=1)
    for j in range(f2 - f1 + 3):
        dibujo.line((0, j * FILA, ancho, j * FILA), fill=GRIS_LINEA, width=1)
    varias = len(figura.marcas) > 1
    for k, marca in enumerate(figura.marcas, start=1):
        m1, mc1, m2, mc2 = practica.rango(marca.lugar)
        x1, x2 = xs[mc1 - c1], xs[mc2 - c1 + 1]
        y1, y2 = (m1 - f1 + 1) * FILA, (m2 - f1 + 2) * FILA
        dibujo.rectangle((x1 - 2, y1 - 2, x2 + 2, y2 + 2), outline=AMBAR, width=6)
        if varias or marca.numero:
            _numero_de_marca(dibujo, x1 + 4, y1 + 4, marca.numero or k)
    return _medida(imagen, destino)


def _parrafos_de_word(archivo: Path) -> list[tuple[str, bool]]:
    documento = Document(archivo)
    parrafos = [(p.text, p.style.name.startswith("Heading")) for p in documento.paragraphs if p.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            parrafos.append((" | ".join(c.text for c in fila.cells), False))
    return parrafos


def _partir(texto: str, fuente, ancho: int, medidor) -> list[str]:
    lineas, actual = [], ""
    for palabra in texto.split():
        prueba = f"{actual} {palabra}".strip()
        if medidor.textlength(prueba, font=fuente) <= ancho or not actual:
            actual = prueba
        else:
            lineas.append(actual)
            actual = palabra
    return lineas + ([actual] if actual else [])


def de_word(archivo: Path, figura: Figura, destino: Path) -> Imagen:
    """Los párrafos del Word real alrededor del texto marcado, con el trozo marcado en ámbar."""
    parrafos = _parrafos_de_word(archivo)
    indices = [i for i, (t, _) in enumerate(parrafos) for m in figura.marcas if m.lugar.strip().lower() in t.lower()]
    inicio = max(0, min(indices) - 1) if indices else 0
    fin = min(len(parrafos), (max(indices) + 2) if indices else 6)
    fin = min(fin, inicio + 8)
    ancho_texto = 1000              # así cabe en la página sin achicar la letra
    fuente, negrita = _fuente(), _fuente(True)
    medidor = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    renglones = []                    # (texto, negrita, índice del párrafo)
    for i in range(inicio, fin):
        texto, es_titulo = parrafos[i]
        for linea in _partir(texto, negrita if es_titulo else fuente, ancho_texto, medidor):
            renglones.append((linea, es_titulo, i))
        renglones.append(("", False, i))
    alto_renglon = int(LETRA * 1.5)
    imagen = Image.new("RGB", (ancho_texto + 2 * 40, alto_renglon * len(renglones) + 2 * 30), BLANCO)
    dibujo = ImageDraw.Draw(imagen)
    dibujo.rectangle((0, 0, imagen.width - 1, imagen.height - 1), outline=GRIS_LINEA, width=2)
    posiciones = []
    for j, (linea, es_titulo, _) in enumerate(renglones):
        y = 30 + j * alto_renglon
        dibujo.text((40, y), linea, font=negrita if es_titulo else fuente, fill=NEGRO)
        posiciones.append(y)
    varias = len(figura.marcas) > 1
    for k, marca in enumerate(figura.marcas, start=1):
        buscado = marca.lugar.strip().lower()
        for j, (linea, es_titulo, _) in enumerate(renglones):
            donde = linea.lower().find(buscado[: max(8, len(buscado) // 2)])
            if donde < 0:
                continue
            letra = negrita if es_titulo else fuente
            x1 = 40 + medidor.textlength(linea[:donde], font=letra)
            x2 = 40 + medidor.textlength(linea[:donde + len(buscado)], font=letra)
            x2 = min(x2, imagen.width - 30)
            y = posiciones[j]
            dibujo.rectangle((x1 - 8, y - 6, x2 + 8, y + alto_renglon - 4), outline=AMBAR, width=5)
            if varias or marca.numero:
                _numero_de_marca(dibujo, x1 - 8, y - 6, marca.numero or k)
            break
    return _medida(imagen, destino)


def hacer(archivo: Path, figura: Figura, destino: Path) -> Imagen:
    destino.parent.mkdir(parents=True, exist_ok=True)
    return de_excel(archivo, figura, destino) if archivo.suffix.lower() == ".xlsx" else de_word(archivo, figura, destino)
