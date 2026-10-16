"""Agente diseñador de las diapositivas (PLAN.md §0, decisión 32).

El redactor escribe el contenido y la validación lo aprueba. Después, el diseñador escribe un programa
pptxgenjs que dibuja ese contenido, con íconos y la identidad visual de la ficha. AulaLista:
1. revisa el programa antes de correrlo: sin require, fs, red ni procesos;
2. lo corre con Node con permisos limitados (app/disenador/ejecutar.js);
3. arregla sin IA la letra chica, los colores ajenos y las notas (decisión 41);
4. comprueba el PowerPoint: cantidad de diapositivas, texto que sale tal cual del contenido aprobado,
   letra mínima, colores de la ficha y notas en todas.
La medición con PowerPoint, el verificador y la revisión visual corren después, en diapositivas.py.
"""

import json
import os
import re
import shutil
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

from app.config import RAIZ
from app.materiales import estilo
from app.materiales.contenido import Diapositivas
from app.materiales.docx import Portada, lineas_de_contacto
from app.materiales.estilo import Identidad

LETRA_MINIMA = 12                       # puntos (decisión 32)
PALABRAS_DE_UNA_ETIQUETA = 4            # «EJEMPLO», «1 · Qué es el gobierno de datos», «Cuidado con»
SEGUNDOS_MAXIMOS = 180
EJECUTOR = RAIZ / "app" / "disenador" / "ejecutar.js"
NODE_MODULES = RAIZ / "node_modules"
PROHIBIDO = re.compile(
    r"\brequire\s*\(|\bimport\b|\bprocess\b|child_process|\beval\s*\(|\bFunction\s*\(|globalThis|__dirname|"
    r"__filename|\bfs\.|\bfetch\s*\(|XMLHttpRequest|WebSocket|\bWorker\b|WebAssembly")
# Sin require no se puede cargar ningún módulo de red ni de archivos; fetch y WebSocket son globales en Node 22.

# El programa va por partes: el código común y una función por diapositiva. Un arreglo devuelve solo las partes que
# cambian; reescribir el programa entero en cada arreglo costaba más que el diseño mismo.
ESQUEMA_DISENO = {
    "type": "object",
    "properties": {
        "comun": {"type": "string"},
        "diapositivas": {"type": "array", "items": {
            "type": "object",
            "properties": {"numero": {"type": "integer"}, "codigo": {"type": "string"}},
            "required": ["numero", "codigo"], "additionalProperties": False,
        }},
        "notas": {"type": "string"},
    },
    "required": ["comun", "diapositivas", "notas"],
    "additionalProperties": False,
}


class ErrorDeDiseno(RuntimeError):
    """El programa del diseñador no se pudo correr."""


def disponible() -> bool:
    """Node y los paquetes del diseñador están instalados."""
    return shutil.which("node") is not None and (NODE_MODULES / "pptxgenjs").is_dir() and EJECUTOR.exists()


# ---------- Datos para el programa ----------

def datos(contenido: Diapositivas, identidad: Identidad, portada: Portada, carpeta: Path) -> dict:
    logo = None
    if identidad.logo is not None:
        copia = carpeta / f"logo{identidad.logo.suffix.lower()}"
        shutil.copyfile(identidad.logo, copia)
        logo = str(copia.resolve())
    return {
        "colores": {"principal": identidad.principal, "acento": identidad.acento, "advertencia": identidad.advertencia,
                    "fondo": identidad.fondo, "texto": identidad.texto, "blanco": estilo.BLANCO,
                    "gris": estilo.GRIS_FILA_ALTERNA},
        "tipografia": identidad.tipografia,
        "letra_minima": LETRA_MINIMA,
        "portada": {"curso": portada.curso, "sesion": portada.sesion, "material": portada.material,
                    "docente": portada.docente, "contacto": lineas_de_contacto(portada.contacto)},
        "logo": logo,
        "contenido": contenido.model_dump(exclude={"pasajes", "decisiones"}),
    }


# ---------- Pedidos al agente ----------

GUIA = """ERES EL DISEÑADOR DE LAS DIAPOSITIVAS DE UNA CLASE.
El contenido ya está escrito y aprobado. Tu trabajo es solo el diseño: escribe un programa pptxgenjs que dibuje
ese contenido como una presentación profesional, variada y atractiva.

CÓMO ES EL PROGRAMA
- El programa va por partes. AulaLista las arma así:
    module.exports = async (h) => {
      [comun]
      [diapositiva 1]   // cada una dentro de su propia función async, en orden
      [diapositiva 2]
      ...
    };
- En «comun» va el código que comparten todas las diapositivas: constantes y funciones de ayuda (por ejemplo, el
  pie, la etiqueta del bloque, las tarjetas). Ya tiene «h». No crea diapositivas.
- En «diapositivas», una entrada por diapositiva con su «numero» (1 es la portada) y su «codigo»: el cuerpo de una
  función async sin parámetros que crea esa sola diapositiva con h.pres.addSlide() y usa lo definido en «comun».
- No uses require, import, process, fs, fetch ni nada fuera de «h». Un programa que los use se rechaza.
- «h» trae:
  - h.pres: la presentación de pptxgenjs, ya en 16:9 (10 × 5,625 pulgadas). Crea cada diapositiva con
    h.pres.addSlide() y agrega texto, formas, tablas, gráficas e imágenes con su API.
  - h.colores: principal, acento, advertencia, fondo, texto, blanco y gris, en hexadecimal sin «#». Usa SOLO estos
    colores. Para un tono más suave usa «transparency» en el relleno, no otro color.
  - h.fuente: la tipografía de la ficha. Úsala en todo el texto (fontFace).
  - h.minimo: el tamaño mínimo de letra, en puntos. Ningún texto puede ser más chico.
  - h.portada: curso, sesion, material, docente y contacto (una lista de líneas).
  - h.logo: la ruta del logo para addImage({ path }), o null si no hay logo.
  - h.contenido: el contenido aprobado, con portada, idea_central, diapositivas, aplicalo, cuidado y cierre.
  - await h.icono(nombre, color): un ícono de react-icons como imagen PNG, para addImage({ data, x, y, w, h }).
    El nombre lleva el prefijo de su colección: FaDatabase, FaUsers, LuScale, MdGavel, HiOutlineLightBulb, TbChartBar…
- Gotchas de pptxgenjs: colores sin «#»; isTextBox: true en cada addText; margin: 0 cuando el texto debe alinearse
  con una forma; rectRadius solo en ROUNDED_RECTANGLE; las listas con bullet: true en cada elemento; no compartas
  un mismo objeto de opciones (shadow) entre dos llamadas.

ORDEN Y CANTIDAD
- Una diapositiva por parte del contenido, en este orden y sin agregar ni quitar ninguna: portada, idea central,
  cada elemento de «diapositivas», «Aplícalo así», «Cuidado con» y el cierre.
- Notas: en cada diapositiva, slide.addNotes(...) con «Qué destacar: …» y «Qué preguntar al grupo: …» de las notas de
  esa parte, tal cual.

EL TEXTO
- Todo el texto sale tal cual de h.contenido y de h.portada. Puedes partir una frase en dos líneas, poner en negrita
  o en color algunas de sus palabras (con varios trozos de texto) y elegir dónde va cada una. No cambies palabras,
  no resumas y no agregues afirmaciones.
- Solo puedes agregar etiquetas cortas de {etiqueta} palabras como máximo: «EJEMPLO», «1 · Qué es el gobierno de
  datos», «Primer principio», «vs», números de paso, el pie con el curso y la sesión, el número de diapositiva.
- El «tipo» de cada diapositiva del contenido es una sugerencia: elige la composición que mejor muestre su contenido.

EL DISEÑO
- Varía las composiciones: no repitas la misma dos veces seguidas. Combina, por ejemplo: tarjetas con ícono y título;
  dos columnas con un círculo «vs» entre ellas; un diagrama radial con la idea al centro; una línea o espectro con dos
  extremos; pasos numerados con círculos; un recuadro de cita destacada sobre el color principal; una tabla con
  encabezado de color; un panel lateral con la idea clave; un organigrama; cifras grandes.
- Íconos en casi todas las diapositivas: dentro de un círculo del color principal o junto al título de cada tarjeta.
- Jerarquía clara: arriba, una etiqueta pequeña en mayúsculas con el número y el nombre del bloque, en el color de
  acento; debajo, el título; abajo, un pie discreto con el curso y la sesión, y el número de diapositiva.
- Etiquetas en forma de píldora (ROUNDED_RECTANGLE con el color de acento) para «EJEMPLO», «Primer principio», etc.
- Resalta en negrita y en color las palabras clave de una frase.
- Portada y cierre: fondo del color principal, título grande en blanco, una composición gráfica con íconos a un lado,
  el docente y el contacto. La idea central, sobre el color principal también.
- «Cuidado con» lleva el color de advertencia y un ícono de alerta; «Aplícalo así», pasos numerados.
- Aire: márgenes de 0,4 pulgadas como mínimo y espacio entre los bloques. Ningún texto puede salirse de su recuadro:
  calcula el alto según el largo del texto y el tamaño de letra. Títulos de 22 a 28 puntos; texto de {minimo} a 16.
"""


def sistema(datos_: dict) -> str:
    """La guía y los datos: lo que no cambia entre el diseño y sus arreglos. Va en el mensaje de sistema para que
    el SDK lo guarde en caché y los arreglos no lo paguen completo otra vez (etapa A de la optimización)."""
    return "\n".join([
        GUIA.replace("{etiqueta}", str(PALABRAS_DE_UNA_ETIQUETA)).replace("{minimo}", str(LETRA_MINIMA)),
        "DATOS QUE RECIBE EL PROGRAMA EN «h» (además de h.pres y h.icono):",
        json.dumps(datos_, ensure_ascii=False),
    ])


def escribir_sistema(datos_: dict, carpeta: Path) -> Path:
    """El mensaje de sistema va en un archivo: en Windows, una línea de comandos no admite un texto tan largo."""
    carpeta.mkdir(parents=True, exist_ok=True)
    ruta = carpeta / "diseno_sistema.txt"
    ruta.write_text(sistema(datos_), encoding="utf-8")
    return ruta


def pedido_de_diseno(total: int) -> str:
    return "\n".join([
        "Diseña la presentación según la guía y los datos de tus instrucciones.",
        f"Devuelve «comun» y las {total} diapositivas, de la 1 a la {total}.",
        "En «notas», resume en una o dos frases las decisiones de diseño.",
    ])


def pedido_de_arreglo(partes: "Partes", problemas: list[str]) -> str:
    return "\n".join([
        "Corrige el programa según la guía y los datos de tus instrucciones.",
        "",
        "PROGRAMA ACTUAL, POR PARTES:",
        "// comun", partes.comun,
        *[f"// diapositiva {n}\n{codigo}" for n, codigo in sorted(partes.diapositivas.items())],
        "",
        "PROBLEMAS QUE ENCONTRÓ AULALISTA AL CORRERLO Y REVISARLO. Corrígelos todos.",
        *[f"- {p}" for p in problemas],
        "",
        "CÓMO DEVOLVER EL ARREGLO",
        "- En «diapositivas», devuelve SOLO las diapositivas que cambias, cada una con su número y su código completo.",
        "  Las que no devuelves quedan como están.",
        "- En «comun», devuelve el código común completo solo si lo cambias. Si no lo cambias, déjalo vacío.",
        "- En «notas», di en una frase qué cambiaste.",
    ])


# ---------- El programa por partes ----------

@dataclass
class Partes:
    comun: str = ""
    diapositivas: dict[int, str] = field(default_factory=dict)


def fusionar(partes: Partes | None, respuesta: dict, total: int) -> tuple[Partes, list[str]]:
    """Agrega a las partes las que devolvió el diseñador. Las que no devolvió quedan como estaban.
    Devuelve también lo que falta para armar el programa: una diapositiva sin código."""
    nuevas = Partes(partes.comun, dict(partes.diapositivas)) if partes else Partes()
    if respuesta.get("comun", "").strip():
        nuevas.comun = respuesta["comun"]
    for d in respuesta.get("diapositivas", []):
        if 1 <= d["numero"] <= total and d["codigo"].strip():
            nuevas.diapositivas[d["numero"]] = d["codigo"]
    faltan = [n for n in range(1, total + 1) if n not in nuevas.diapositivas]
    problemas = [f"FALLA · diseño · falta el código de la diapositiva {', '.join(map(str, faltan))}. Devuelve una "
                 "entrada en «diapositivas» por cada una."] if faltan else []
    return nuevas, problemas


def armar(partes: Partes) -> str:
    """El módulo que corre Node: el código común y una función por diapositiva, llamadas en orden."""
    funciones = "\n".join(f"  async () => {{ // diapositiva {n}\n{codigo}\n  }},"
                          for n, codigo in sorted(partes.diapositivas.items()))
    return (f"module.exports = async (h) => {{\n{partes.comun}\n"
            f"const __diapositivas = [\n{funciones}\n];\n"
            "for (const dibujar of __diapositivas) await dibujar();\n};\n")


# ---------- Revisar y correr el programa ----------

def revisar_codigo(codigo: str) -> list[str]:
    """Antes de correrlo: la forma del módulo y nada fuera de «h»."""
    problemas = []
    if "module.exports" not in codigo:
        problemas.append("FALLA · diseño · el programa no exporta su función: module.exports = async (h) => { ... }.")
    prohibidos = sorted({m.group(0).strip() for m in PROHIBIDO.finditer(codigo)})
    if prohibidos:
        problemas.append(f"FALLA · diseño · el programa usa lo que no está permitido: {', '.join(prohibidos)}. "
                         "Usa solo «h».")
    return problemas


@dataclass
class Corrida:
    ok: bool
    error: str = ""
    iconos: tuple[str, ...] = ()


def correr(codigo: str, datos_: dict, carpeta: Path, salida: Path) -> Corrida:
    """Corre el programa con Node y permisos limitados: lee sus librerías y la carpeta del material, y escribe solo ahí."""
    carpeta.mkdir(parents=True, exist_ok=True)
    programa, archivo_datos = carpeta / "diseno.js", carpeta / "diseno_datos.json"
    programa.write_text(codigo, encoding="utf-8")
    archivo_datos.write_text(json.dumps(datos_, ensure_ascii=False), encoding="utf-8")
    salida.unlink(missing_ok=True)
    comando = [
        shutil.which("node") or "node", "--permission", "--allow-addons",   # sharp es un complemento nativo
        f"--allow-fs-read={NODE_MODULES.resolve()}{_TODO}", f"--allow-fs-read={EJECUTOR.parent.resolve()}{_TODO}",
        f"--allow-fs-read={carpeta.resolve()}{_TODO}", f"--allow-fs-write={carpeta.resolve()}{_TODO}",
        f"--allow-fs-write={salida.parent.resolve()}{_TODO}",
        str(EJECUTOR.resolve()), str(programa.resolve()), str(archivo_datos.resolve()), str(salida.resolve()),
    ]
    try:
        proceso = subprocess.run(comando, capture_output=True, text=True, encoding="utf-8", timeout=SEGUNDOS_MAXIMOS,
                                 env={**os.environ, "NODE_PATH": str(NODE_MODULES.resolve())})
    except subprocess.TimeoutExpired:
        return Corrida(False, f"el programa tardó más de {SEGUNDOS_MAXIMOS} segundos")
    if proceso.returncode != 0 or not salida.exists():
        return Corrida(False, _resumir_error(proceso.stderr))
    try:
        resultado = json.loads(proceso.stdout.strip().splitlines()[-1])
    except (ValueError, IndexError):
        resultado = {}
    return Corrida(True, iconos=tuple(resultado.get("iconos", ())))


_TODO = "\\*" if os.name == "nt" else "/*"     # todo lo que está dentro de la carpeta


def _resumir_error(stderr: str) -> str:
    lineas = [l for l in (stderr or "").splitlines() if l.strip() and "SecurityWarning" not in l
              and "--trace-warnings" not in l]
    return " ".join(lineas[:6])[:900] or "el programa terminó con un error sin mensaje"


# ---------- Revisar el PowerPoint ----------

def _normal(texto: str) -> str:
    texto = texto.lower().replace("\u00a0", " ")
    texto = re.sub(r"[«»\"“”'‘’]", "", texto)
    return re.sub(r"\s+", " ", texto).strip(" .,:;·-–—•")


def _textos_del_contenido(contenido: Diapositivas, portada: Portada) -> list[str]:
    textos = [portada.curso, portada.sesion, portada.material, portada.docente, *lineas_de_contacto(portada.contacto),
              *contenido.idea_central.frases, *contenido.aplicalo.frases, *contenido.aplicalo.plantilla,
              *contenido.cuidado.riesgos, contenido.cierre.frase, "Idea central", "Aplícalo así", "Cuidado con"]
    for d in contenido.diapositivas:
        textos += [d.bloque, d.titulo, *d.textos_visibles()]
        if d.grafico:
            textos += [f"{e.valor:g}" for e in d.grafico.elementos if e.valor is not None]
    return [t for t in textos if t]


def _notas_del_contenido(contenido: Diapositivas) -> list[str]:
    partes = [contenido.portada, contenido.idea_central, *contenido.diapositivas, contenido.aplicalo,
              contenido.cuidado, contenido.cierre]
    return [p.notas.texto() for p in partes]


def _figuras(figuras):
    for figura in figuras:
        if figura.shape_type == 6:          # grupo
            yield from _figuras(figura.shapes)
        else:
            yield figura


def _parrafos(figura):
    if figura.has_text_frame:
        for p in figura.text_frame.paragraphs:
            yield p
    elif getattr(figura, "has_table", False) and figura.has_table:
        for fila in figura.table.rows:
            for celda in fila.cells:
                yield from celda.text_frame.paragraphs


def _distancia(a: str, b: str) -> int:
    return sum((int(a[i:i + 2], 16) - int(b[i:i + 2], 16)) ** 2 for i in (0, 2, 4))


def _en_sombra(nodo) -> bool:
    return any(ancestro.tag == qn("a:effectLst") for ancestro in nodo.iterancestors())


def ajustar_pptx(ruta: Path, contenido: Diapositivas, identidad: Identidad) -> list[str]:
    """Arreglos sin IA ni costo, antes de los controles: lo que tiene una sola respuesta correcta no vuelve al
    diseñador. La letra menor que el mínimo sube al mínimo; un color que no es de la ficha pasa al color de la ficha
    más parecido; las notas que no son las del contenido se copian del contenido. Devuelve lo que ajustó."""
    presentacion = Presentation(ruta)
    permitidos = sorted(identidad.colores_permitidos())
    notas_esperadas = _notas_del_contenido(contenido)
    letras, colores, notas = 0, {}, []
    for numero, diapositiva in enumerate(presentacion.slides, start=1):
        for figura in _figuras(diapositiva.shapes):
            for parrafo in _parrafos(figura):
                for corrida in parrafo.runs:
                    if corrida.text.strip() and corrida.font.size is not None and corrida.font.size < Pt(LETRA_MINIMA):
                        corrida.font.size = Pt(LETRA_MINIMA)
                        letras += 1
        for nodo in diapositiva._element.iter(qn("a:srgbClr")):
            valor = (nodo.get("val") or "").upper()
            if re.fullmatch(r"[0-9A-F]{6}", valor) and valor not in permitidos and not _en_sombra(nodo):
                colores[valor] = min(permitidos, key=lambda c: _distancia(valor, c))
                nodo.set("val", colores[valor])
        if numero <= len(notas_esperadas):
            actual = diapositiva.notes_slide.notes_text_frame.text if diapositiva.has_notes_slide else ""
            if _normal(actual) != _normal(notas_esperadas[numero - 1]):
                diapositiva.notes_slide.notes_text_frame.text = notas_esperadas[numero - 1]
                notas.append(numero)
    ajustes = []
    if letras:
        ajustes.append(f"subí a {LETRA_MINIMA} puntos {letras} trozos de texto con letra más chica")
    if colores:
        ajustes.append("cambié colores que no son de la ficha por el más parecido de la ficha: "
                       + ", ".join(f"{a} por {b}" for a, b in sorted(colores.items())))
    if notas:
        ajustes.append(f"copié del contenido las notas de las diapositivas {', '.join(map(str, notas))}")
    if ajustes:
        presentacion.save(ruta)
    return ajustes


def revisar_pptx(ruta: Path, contenido: Diapositivas, portada: Portada, identidad: Identidad) -> list[str]:
    """Cantidad, texto tal cual del contenido, letra mínima, colores de la ficha y notas."""
    presentacion = Presentation(ruta)
    problemas = []
    total = len(presentacion.slides)
    if total != contenido.total:
        problemas.append(f"FALLA · diseño · hay {total} diapositivas; deben ser {contenido.total}: una por parte del "
                         "contenido, sin agregar ni quitar.")
    corpus = " | ".join(_normal(t) for t in _textos_del_contenido(contenido, portada))
    permitidos = identidad.colores_permitidos()
    notas_esperadas = _notas_del_contenido(contenido)
    for numero, diapositiva in enumerate(presentacion.slides, start=1):
        ajenos, chicos, colores = [], set(), set()
        for figura in _figuras(diapositiva.shapes):
            for parrafo in _parrafos(figura):
                texto = "".join(r.text for r in parrafo.runs)
                normal = _normal(texto)
                if normal and len(normal.split()) > PALABRAS_DE_UNA_ETIQUETA and normal not in corpus:
                    ajenos.append(texto.strip())
                for corrida in parrafo.runs:
                    if corrida.text.strip() and corrida.font.size is not None and corrida.font.size < Pt(LETRA_MINIMA):
                        chicos.add(round(corrida.font.size.pt, 1))
        xml = re.sub(r"<a:effectLst>.*?</a:effectLst>", "", diapositiva._element.xml, flags=re.S)   # sombras
        colores |= {c.upper() for c in re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml)} - permitidos
        for texto in ajenos[:3]:
            problemas.append(f"FALLA · diseño · diapositiva {numero}: el texto «{texto[:120]}» no está tal cual en el "
                             "contenido aprobado. Usa el texto del contenido sin cambiar palabras.")
        if chicos:
            problemas.append(f"FALLA · diseño · diapositiva {numero}: hay letra de {', '.join(map(str, sorted(chicos)))} "
                             f"puntos; el mínimo es {LETRA_MINIMA}.")
        if colores:
            problemas.append(f"FALLA · diseño · diapositiva {numero}: usa colores que no son de la ficha: "
                             f"{', '.join(sorted(colores))}. Usa solo h.colores; para un tono suave, «transparency».")
        notas = diapositiva.notes_slide.notes_text_frame.text if diapositiva.has_notes_slide else ""
        if numero <= len(notas_esperadas) and _normal(notas) != _normal(notas_esperadas[numero - 1]):
            problemas.append(f"FALLA · diseño · diapositiva {numero}: sus notas no son las del contenido. Copia tal "
                             "cual «Qué destacar: …» y «Qué preguntar al grupo: …» de esa parte.")
    return problemas
