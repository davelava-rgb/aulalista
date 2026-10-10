"""Completa el Excel de verificación que escribe el verificador (PLAN.md §5.6 y §0, decisión 22).

El Excel de la lectura queda solo con la hoja Hallazgos: los hallazgos actuales del verificador y cada
problema que se envió a corrección, con cómo se resolvió, más los errores de la revisión del contenido.
"""

from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

HOJA = "Hallazgos"


def _estilo(fila) -> None:
    for celda in fila:
        celda.font = Font(name="Arial", size=10)
        celda.alignment = Alignment(wrap_text=True, vertical="top")


def solo_hallazgos(ruta: Path) -> None:
    """Deja solo la hoja Hallazgos: sin tabla de pasaje y veredicto por oración (decisión 22)."""
    libro = openpyxl.load_workbook(ruta)
    for nombre in [n for n in libro.sheetnames if n != HOJA]:
        del libro[nombre]
    libro.save(ruta)


def anotar_reglas(ruta: Path, reglas: tuple[str, ...], nota: str) -> None:
    """«Cómo se resolvió» en la hoja Hallazgos para los avisos de estas reglas que no tienen nota todavía."""
    libro = openpyxl.load_workbook(ruta)
    for fila in libro[HOJA].iter_rows(min_row=2):
        if fila[0].value == "AVISO" and fila[3].value in reglas and not fila[5].value:
            fila[5].value = nota
    libro.save(ruta)


def agregar_historial(ruta: Path, historial: list[dict]) -> None:
    """historial: [{nivel, seccion, oracion, regla, detalle, resolucion}]"""
    if not historial:
        return
    libro = openpyxl.load_workbook(ruta)
    hoja = libro[HOJA]
    for h in historial:
        hoja.append([h["nivel"], h["seccion"], h["oracion"], h["regla"], h["detalle"], h["resolucion"]])
        _estilo(hoja[hoja.max_row])
    libro.save(ruta)
