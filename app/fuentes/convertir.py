"""Conversión de las fuentes del curso a texto, con su página, lámina o celda (SPEC §7 y §9).

Cada fuente se convierte una sola vez: si su huella no cambia, no se vuelve a convertir.
Resultado en cursos/[curso]/fuentes_texto/:
- [archivo].jsonl: un pasaje por línea, con su texto y su ubicación.
- indice.json: estado de cada fuente.
- _pdf/: copia en PDF de cada Word, con las páginas reales que da Word.
"""

import hashlib
import json
import re
from collections.abc import Awaitable, Callable
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pymupdf
from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app import office

# (imagen, tipo MIME, etiqueta) -> líneas de texto
LectorImagen = Callable[[bytes, str, str], Awaitable[list[str]]]

TIPOS_IMAGEN = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}
FORMATOS = {
    ".pdf": "PDF",
    ".docx": "Word",
    ".doc": "Word",
    ".pptx": "PowerPoint",
    ".xlsx": "Excel",
    ".xlsm": "Excel",
    **{extension: "imagen" for extension in TIPOS_IMAGEN},
}
CONVERTIDA = "convertida"
NO_SE_PUEDE_LEER = "no se puede leer"
DPI_PAGINA_ESCANEADA = 150


class ErrorDeLectura(RuntimeError):
    """La fuente no se puede leer. El material no la cita (SPEC §7)."""


@dataclass
class Conversion:
    pasajes: list[dict] = field(default_factory=list)
    revisar: bool = False  # tiene texto leído por IA que el profesor debe revisar
    avisos: list[str] = field(default_factory=list)
    metodo: str = ""

    def agregar(self, texto: str, ubicacion: str, *, origen: str = "texto", **extra) -> None:
        texto = " ".join(str(texto).split())
        if texto:
            self.pasajes.append({"texto": texto, "ubicacion": ubicacion, "origen": origen, **extra})


def formato_admitido(nombre: str) -> bool:
    return Path(nombre).suffix.lower() in FORMATOS


def huella(ruta: Path) -> str:
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


# ---------- PDF ----------

def _bloques_de_texto(pagina) -> list[str]:
    bloques = pagina.get_text("blocks", flags=pymupdf.TEXTFLAGS_TEXT | pymupdf.TEXT_DEHYPHENATE)
    textos = (" ".join(b[4].split()) for b in bloques if b[6] == 0)
    return [t for t in textos if t]


def _numero_impreso(bloques: list[str]) -> str | None:
    """Número de lámina o de página impreso en la propia página, si lo hay."""
    for texto in bloques[:3]:
        if re.fullmatch(r"\d{1,4}", texto):
            return texto
    return None


async def extraer_pdf(ruta: Path, lector: LectorImagen | None, conversion: Conversion | None = None) -> Conversion:
    conversion = conversion or Conversion(metodo="texto del PDF")
    try:
        documento = pymupdf.open(ruta)
    except Exception as error:
        raise ErrorDeLectura(f"El PDF no se puede abrir: {error}") from error
    if documento.needs_pass:
        raise ErrorDeLectura("El PDF tiene contraseña.")
    with documento:
        for indice, pagina in enumerate(documento, start=1):
            bloques = _bloques_de_texto(pagina)
            impreso = _numero_impreso(bloques)
            numero = int(impreso) if impreso is not None else None
            ubicacion = f"página {indice}" + (f" · número impreso {numero}" if numero is not None else "")
            if impreso is not None:
                bloques.remove(impreso)
            for texto in bloques:
                conversion.agregar(texto, ubicacion, pagina=indice, numero_impreso=numero)
            if not bloques and pagina.get_images():
                await _leer_pagina_escaneada(pagina, indice, ruta.name, lector, conversion)
    return conversion


async def _leer_pagina_escaneada(pagina, indice, nombre, lector, conversion) -> None:
    if lector is None:
        conversion.avisos.append(f"Página {indice}: es una imagen y no se leyó.")
        return
    imagen = pagina.get_pixmap(dpi=DPI_PAGINA_ESCANEADA).tobytes("png")
    try:
        lineas = await lector(imagen, "image/png", f"{nombre} · página {indice}")
    except Exception as error:
        conversion.avisos.append(f"Página {indice}: es una imagen y no se pudo leer ({error}).")
        return
    for linea in lineas:
        conversion.agregar(linea, f"página {indice} · leída por IA", origen="ia", pagina=indice)
    conversion.revisar = True


# ---------- Word ----------

async def extraer_word(ruta: Path, lector: LectorImagen | None, carpeta_pdf: Path) -> Conversion:
    """Word pasa a PDF con el Word instalado para conservar las páginas reales."""
    try:
        pdf = office.word_a_pdf(ruta, carpeta_pdf / f"{ruta.name}.pdf")
    except office.ErrorDeOffice as error:
        if ruta.suffix.lower() != ".docx":
            raise ErrorDeLectura(str(error)) from error
        conversion = _extraer_docx_sin_paginas(ruta)
        conversion.avisos.append(f"Word no estuvo disponible ({error}). Las ubicaciones son párrafos, no páginas.")
        return conversion
    return await extraer_pdf(pdf, lector, Conversion(metodo="Word convertido a PDF"))


def _extraer_docx_sin_paginas(ruta: Path) -> Conversion:
    conversion = Conversion(metodo="texto del Word, sin páginas")
    documento = Document(ruta)
    for numero, parrafo in enumerate(documento.paragraphs, start=1):
        conversion.agregar(parrafo.text, f"párrafo {numero}")
    for numero_tabla, tabla in enumerate(documento.tables, start=1):
        for numero_fila, fila in enumerate(tabla.rows, start=1):
            textos = list(dict.fromkeys(celda.text for celda in fila.cells))
            conversion.agregar(" | ".join(textos), f"tabla {numero_tabla}, fila {numero_fila}")
    return conversion


# ---------- PowerPoint ----------

def _figuras(figuras):
    for figura in figuras:
        if figura.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _figuras(figura.shapes)
        else:
            yield figura


async def extraer_pptx(ruta: Path, lector: LectorImagen | None) -> Conversion:
    conversion = Conversion(metodo="texto del PowerPoint")
    try:
        presentacion = Presentation(ruta)
    except Exception as error:
        raise ErrorDeLectura(f"El PowerPoint no se puede abrir: {error}") from error
    for numero, lamina in enumerate(presentacion.slides, start=1):
        ubicacion = f"lámina {numero}"
        antes = len(conversion.pasajes)
        imagenes = []
        for figura in _figuras(lamina.shapes):
            if figura.has_text_frame:
                for parrafo in figura.text_frame.paragraphs:
                    conversion.agregar("".join(r.text for r in parrafo.runs), ubicacion, lamina=numero)
            elif getattr(figura, "has_table", False) and figura.has_table:
                for fila in figura.table.rows:
                    conversion.agregar(" | ".join(c.text for c in fila.cells), ubicacion, lamina=numero)
            elif figura.shape_type == MSO_SHAPE_TYPE.PICTURE:
                imagenes.append(figura.image)
        if len(conversion.pasajes) == antes and imagenes:
            await _leer_imagenes_de_lamina(imagenes, numero, ruta.name, lector, conversion)
        if lamina.has_notes_slide:
            notas = lamina.notes_slide.notes_text_frame
            if notas is not None:
                for parrafo in notas.paragraphs:
                    conversion.agregar(
                        "".join(r.text for r in parrafo.runs), f"lámina {numero} · notas", lamina=numero
                    )
    return conversion


async def _leer_imagenes_de_lamina(imagenes, numero, nombre, lector, conversion) -> None:
    if lector is None:
        conversion.avisos.append(f"Lámina {numero}: solo tiene imágenes y no se leyeron.")
        return
    for imagen in imagenes:
        try:
            lineas = await lector(imagen.blob, imagen.content_type, f"{nombre} · lámina {numero}")
        except Exception as error:
            conversion.avisos.append(f"Lámina {numero}: una imagen no se pudo leer ({error}).")
            continue
        for linea in lineas:
            conversion.agregar(linea, f"lámina {numero} · leída por IA", origen="ia", lamina=numero)
        conversion.revisar = conversion.revisar or bool(lineas)


# ---------- Excel ----------

def _valor_de_celda(valor) -> str:
    if isinstance(valor, datetime) and valor.time() == datetime.min.time():
        valor = valor.date()
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor)


async def extraer_xlsx(ruta: Path, lector: LectorImagen | None) -> Conversion:
    conversion = Conversion(metodo="celdas del Excel")
    try:
        valores = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        formulas = openpyxl.load_workbook(ruta, data_only=False, read_only=True)
    except Exception as error:
        raise ErrorDeLectura(f"El Excel no se puede abrir: {error}") from error
    sin_valor = []
    for hoja in valores.worksheets:
        hoja_formulas = formulas[hoja.title]
        for fila_valores, fila_formulas in zip(hoja.iter_rows(), hoja_formulas.iter_rows()):
            for celda, celda_formula in zip(fila_valores, fila_formulas):
                if celda.value is not None:
                    conversion.agregar(
                        _valor_de_celda(celda.value),
                        f"hoja {hoja.title}, celda {celda.coordinate}",
                        hoja=hoja.title,
                        celda=celda.coordinate,
                    )
                elif isinstance(celda_formula.value, str) and celda_formula.value.startswith("="):
                    sin_valor.append(f"{hoja.title}!{celda_formula.coordinate}")
    valores.close()
    formulas.close()
    if sin_valor:
        conversion.avisos.append(
            "Estas celdas tienen fórmula pero no un valor guardado. Ábrelas en Excel y guarda el archivo: "
            + ", ".join(sin_valor[:20])
            + (" y otras." if len(sin_valor) > 20 else ".")
        )
    return conversion


# ---------- Imagen ----------

async def extraer_imagen(ruta: Path, lector: LectorImagen | None) -> Conversion:
    if lector is None:
        raise ErrorDeLectura("Es una imagen y no hay lector de imágenes.")
    conversion = Conversion(metodo="leída por IA", revisar=True)
    tipo = TIPOS_IMAGEN[ruta.suffix.lower()]
    try:
        lineas = await lector(ruta.read_bytes(), tipo, ruta.name)
    except Exception as error:
        raise ErrorDeLectura(f"La imagen no se pudo leer: {error}") from error
    for numero, linea in enumerate(lineas, start=1):
        conversion.agregar(linea, f"línea {numero} · leída por IA", origen="ia")
    if not conversion.pasajes:
        conversion.avisos.append("La imagen no tiene texto.")
    return conversion


# ---------- Curso completo ----------

def _carpetas(carpeta_curso: Path) -> tuple[Path, Path, Path]:
    salida = carpeta_curso / "fuentes_texto"
    return carpeta_curso / "fuentes", salida, salida / "indice.json"


def leer_indice(carpeta_curso: Path) -> dict:
    _, _, ruta = _carpetas(carpeta_curso)
    return json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else {}


def _guardar_indice(carpeta_curso: Path, indice: dict) -> None:
    _, salida, ruta = _carpetas(carpeta_curso)
    salida.mkdir(parents=True, exist_ok=True)
    ruta.write_text(json.dumps(indice, ensure_ascii=False, indent=2), encoding="utf-8")


def _ruta_pasajes(carpeta_curso: Path, archivo: str) -> Path:
    return _carpetas(carpeta_curso)[1] / f"{archivo}.jsonl"


async def _convertir_archivo(ruta: Path, lector: LectorImagen | None, carpeta_pdf: Path) -> Conversion:
    extension = ruta.suffix.lower()
    if extension == ".pdf":
        return await extraer_pdf(ruta, lector)
    if extension in (".docx", ".doc"):
        return await extraer_word(ruta, lector, carpeta_pdf)
    if extension == ".pptx":
        return await extraer_pptx(ruta, lector)
    if extension in (".xlsx", ".xlsm"):
        return await extraer_xlsx(ruta, lector)
    if extension in TIPOS_IMAGEN:
        return await extraer_imagen(ruta, lector)
    raise ErrorDeLectura("Formato no admitido. Usa PDF, Word, PowerPoint (.pptx), Excel (.xlsx) o imagen.")


async def convertir_curso(carpeta_curso: Path, lector: LectorImagen | None) -> dict:
    """Convierte las fuentes nuevas o cambiadas y devuelve el índice actualizado."""
    fuentes, salida, _ = _carpetas(carpeta_curso)
    indice = leer_indice(carpeta_curso)
    presentes = {
        r.name: r for r in sorted(fuentes.glob("*"))
        if r.is_file() and not r.name.startswith(("~$", "."))
    } if fuentes.exists() else {}

    for nombre in [n for n in indice if n not in presentes]:
        _ruta_pasajes(carpeta_curso, nombre).unlink(missing_ok=True)
        del indice[nombre]

    for nombre, ruta in presentes.items():
        actual = huella(ruta)
        anterior = indice.get(nombre)
        if anterior and anterior["huella"] == actual and anterior["estado"] == CONVERTIDA:
            continue
        destino = _ruta_pasajes(carpeta_curso, nombre)
        entrada = {
            "huella": actual,
            "formato": FORMATOS.get(ruta.suffix.lower(), "desconocido"),
            "fecha": datetime.now().isoformat(timespec="seconds"),
        }
        try:
            conversion = await _convertir_archivo(ruta, lector, salida / "_pdf")
        except ErrorDeLectura as error:
            destino.unlink(missing_ok=True)
            entrada.update(estado=NO_SE_PUEDE_LEER, motivo=str(error), pasajes=0,
                           revisar=False, revisada=False, avisos=[], metodo="")
        else:
            salida.mkdir(parents=True, exist_ok=True)
            with open(destino, "w", encoding="utf-8") as archivo:
                for numero, pasaje in enumerate(conversion.pasajes, start=1):
                    archivo.write(json.dumps({"n": numero, **pasaje}, ensure_ascii=False) + "\n")
            entrada.update(estado=CONVERTIDA, motivo="", pasajes=len(conversion.pasajes),
                           revisar=conversion.revisar, revisada=False,
                           avisos=conversion.avisos, metodo=conversion.metodo)
        indice[nombre] = entrada

    _guardar_indice(carpeta_curso, indice)
    return indice


def leer_pasajes(carpeta_curso: Path, archivo: str) -> list[dict]:
    ruta = _ruta_pasajes(carpeta_curso, archivo)
    if not ruta.exists():
        return []
    return [json.loads(linea) for linea in ruta.read_text(encoding="utf-8").splitlines() if linea.strip()]


def marcar_revisada(carpeta_curso: Path, archivo: str) -> None:
    """El profesor confirma que revisó el texto que leyó la IA."""
    indice = leer_indice(carpeta_curso)
    if archivo not in indice:
        raise KeyError(archivo)
    indice[archivo]["revisada"] = True
    _guardar_indice(carpeta_curso, indice)
